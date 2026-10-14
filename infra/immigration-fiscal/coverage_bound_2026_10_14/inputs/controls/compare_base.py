"""Compare Hispanic April 2020 Census, V2024 blended base, and DA counts.

Run: OPENBLAS_NUM_THREADS=1 uv run --no-project python3
     infra/immigration-fiscal/coverage_bound_2026_10_14/inputs/controls/compare_base.py
MONTH=4.2 selects the April 1 estimates base, not an April monthly estimate.
DA estimates are rounded thousands; their range is not a confidence interval.
"""

from pathlib import Path
import csv
import hashlib
import json
import math
import re
import resource
import time

import openpyxl

HERE = Path(__file__).resolve().parent
LANE_INPUTS = HERE.parent


def require(condition, message):
    if not condition:
        raise SystemExit("[BLOCKED] " + message)


def age_interval(label):
    if label == "Under 5 years":
        return 0, 4
    if label == "85 years and over":
        return 85, 100
    ages = [int(x) for x in re.findall(r"\d+", label)]
    require(1 <= len(ages) <= 2, "unrecognized DHC age label " + label)
    return ages[0], ages[-1]


def main():
    start = time.perf_counter()
    pins = json.loads((HERE / "pins.json").read_text())["sources"]
    require(len({x["file"] for x in pins}) == len(pins), "duplicate controls pin")
    for pin in pins:
        data = (HERE / pin["file"]).read_bytes()
        require(len(data) == pin["size_bytes"], "size mismatch " + pin["file"])
        require(hashlib.sha256(data).hexdigest() == pin["sha256"], "stale pin " + pin["file"])
        if pin["file"].endswith(".json"):
            json.loads(data)
        if pin["file"].endswith(".pdf"):
            require(data.startswith(b"%PDF"), "invalid PDF " + pin["file"])
    with (HERE / "nc-est2024-alldata-r-file01.csv").open(newline="") as stream:
        rows = [r for r in csv.DictReader(stream) if r["YEAR"] == "2020" and r["MONTH"] == "4.2"]
    require(len(rows) == 102, "expected 101 single ages and one all-age base row")
    require(all(r["UNIVERSE"] == "R" for r in rows), "base universe must be resident")
    require(len({int(r["AGE"]) for r in rows}) == 102, "duplicate base age")
    base = {int(r["AGE"]): int(r["H_MALE"]) + int(r["H_FEMALE"]) for r in rows}
    require(set(base) == set(range(101)) | {999}, "missing base ages")
    require(sum(base[a] for a in range(101)) == base[999], "base Hispanic closure")
    response = json.loads((HERE / "dhc-p12h-us.json").read_text())
    require(len(response) == 2 and len(response[0]) == len(response[1]), "expected one DHC US row")
    census = dict(zip(response[0], response[1]))
    require(census["us"] == "1", "DHC geography must be US")
    variables = json.loads((HERE / "dhc-p12h-variables.json").read_text())["variables"]
    cells = []
    for key, metadata in variables.items():
        parts = metadata["label"].split("!!")
        if key.startswith("P12H_") and key.endswith("N") and len(parts) == 4:
            require("HISPANIC OR LATINO" in metadata["concept"], "wrong DHC concept")
            lo, hi = age_interval(parts[-1])
            cells.append((lo, hi, int(census[key])))
    require(len(cells) == 46, "expected 23 DHC age cells per sex")
    require(sum(c for _, _, c in cells) == int(census["P12H_001N"]), "Census Hispanic closure")
    da_pin = json.loads((LANE_INPUTS / "pins.json").read_text())["da_hispanic.xlsx"]
    da_bytes = (LANE_INPUTS / "da_hispanic.xlsx").read_bytes()
    require(len(da_bytes) == da_pin["bytes"], "DA pin size mismatch")
    require(hashlib.sha256(da_bytes).hexdigest() == da_pin["sha256"], "stale DA pin")
    workbook = openpyxl.load_workbook(LANE_INPUTS / "da_hispanic.xlsx", data_only=True)
    sheet = workbook["Table 1C"]
    require("In thousands" in sheet.cell(1, 1).value, "DA unit changed")
    require([sheet.cell(3, i).value for i in range(3, 6)] == ["Low", "Middle", "High"], "DA series changed")
    da = {}
    for row in sheet.iter_rows(min_row=5, max_row=34, values_only=True):
        require(isinstance(row[0], int) and row[0] not in da, "invalid DA both-sex age")
        da[row[0]] = tuple(float(x) * 1000 for x in row[2:5])
    require(set(da) == set(range(30)), "DA must cover ages 0-29")
    bands = [(0, 4), (5, 9), (0, 9), (10, 14), (15, 17), (10, 17), (18, 29), (30, 49), (50, 100), (0, 17), (0, 29), (0, 100)]
    output = []
    for lo, hi in bands:
        chosen = [(a, b, c) for a, b, c in cells if lo <= a and b <= hi]
        require(all(b < lo or a > hi or (lo <= a and b <= hi) for a, b, _ in cells), "band splits DHC cell")
        c = sum(x[2] for x in chosen)
        b = sum(base[a] for a in range(lo, hi + 1))
        require(c > 0 and b > 0, "nonpositive comparison population")
        record = {"age_2020_lo": lo, "age_2020_hi": hi, "census_hispanic": c, "base_hispanic": b,
                  "base_minus_census": b - c, "base_over_census": b / c,
                  "base_change_pct": 100 * (b / c - 1)}
        for i, series in enumerate(("low", "middle", "high")):
            total = sum(da[a][i] for a in range(lo, hi + 1)) if hi < 30 else None
            record["da_" + series] = total
            record["da_residual_factor_" + series] = total / b - 1 if total is not None else None
        require(all(v is None or math.isfinite(v) for v in record.values()), "nonfinite result")
        output.append(record)
    with (HERE / "base_adjustments.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(output[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(output)
    (HERE / "base_adjustments.json").write_text(json.dumps({"reference_date": "2020-04-01", "universe": "resident", "origin": "Hispanic or Latino", "da_rounding": "Each published single-age count is rounded to thousands; sums need not match published rounded totals.", "rows": output}, indent=2) + "\n")
    sex_bands = [(0, 4), (5, 9), (10, 17), (18, 29), (30, 49), (50, 100)]
    sex_output = []
    da_by_sex = {}
    for sex, label in ((1, "Male"), (2, "Female")):
        marker = [i for i in range(1, sheet.max_row + 1) if sheet.cell(i, 1).value == label]
        require(len(marker) == 1, "missing or duplicate DA sex marker")
        values = list(sheet.iter_rows(min_row=marker[0] + 1, max_row=marker[0] + 30, values_only=True))
        require([r[0] for r in values] == list(range(30)), "DA sex ages changed")
        da_by_sex[sex] = {r[0]: tuple(float(x) * 1000 for x in r[2:5]) for r in values}
        sex_base = {int(r["AGE"]): int(r["H_" + label.upper()]) for r in rows}
        require(sum(sex_base[a] for a in range(101)) == sex_base[999], "base sex closure")
        sex_cells = []
        for key, metadata in variables.items():
            parts = metadata["label"].split("!!")
            if key.startswith("P12H_") and key.endswith("N") and len(parts) == 4 and parts[-2] == label + ":":
                lo, hi = age_interval(parts[-1])
                sex_cells.append((lo, hi, int(census[key])))
        total_keys = [key for key, metadata in variables.items() if key.startswith("P12H_") and key.endswith("N") and metadata["label"].split("!!") == [" ", "Total:", label + ":"]]
        require(len(sex_cells) == 23 and len(total_keys) == 1, "DHC sex labels changed")
        require(sum(c for _, _, c in sex_cells) == int(census[total_keys[0]]), "Census sex closure")
        for lo, hi in sex_bands:
            require(all(b < lo or a > hi or (lo <= a and b <= hi) for a, b, _ in sex_cells), "sex band splits DHC cell")
            c = sum(v for a, b, v in sex_cells if lo <= a and b <= hi)
            b = sum(sex_base[a] for a in range(lo, hi + 1))
            require(c > 0 and b > 0, "nonpositive sex population")
            record = {"age_2020_lo": lo, "age_2020_hi": hi, "sex": sex,
                      "census_hispanic": c, "base_hispanic": b,
                      "base_minus_census": b - c, "base_over_census": b / c,
                      "base_change_pct": 100 * (b / c - 1)}
            for i, series in enumerate(("low", "middle", "high")):
                total = sum(da_by_sex[sex][a][i] for a in range(lo, hi + 1)) if hi < 30 else None
                record["da_" + series] = total
                record["da_residual_factor_" + series] = total / b - 1 if total is not None else None
            require(all(v is None or math.isfinite(v) for v in record.values()), "nonfinite sex result")
            sex_output.append(record)
    for age in range(30):
        for series in range(3):
            require(abs(da_by_sex[1][age][series] + da_by_sex[2][age][series] - da[age][series]) <= 1000, "DA sex closure exceeds single-age thousand rounding")
    for lo, hi in sex_bands:
        pair = [r for r in sex_output if (r["age_2020_lo"], r["age_2020_hi"]) == (lo, hi)]
        both = [r for r in output if (r["age_2020_lo"], r["age_2020_hi"]) == (lo, hi)][0]
        require(len(pair) == 2 and {r["sex"] for r in pair} == {1, 2}, "duplicate or missing sex band")
        for field in ("census_hispanic", "base_hispanic"):
            require(sum(r[field] for r in pair) == both[field], "exact sex-to-both-sex closure failed")
    with (HERE / "base_adjustments_sex.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(sex_output[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(sex_output)
    print("PASS: pins, input types, unique ages, exact Census/base closure, finite ratios")
    print("PASS: exact Census/base sex closure; DA single-age rounded sex closure")
    print(f"seconds={time.perf_counter()-start:.3f} maxrss_bytes={resource.getrusage(resource.RUSAGE_SELF).ru_maxrss}")


if __name__ == "__main__":
    main()
