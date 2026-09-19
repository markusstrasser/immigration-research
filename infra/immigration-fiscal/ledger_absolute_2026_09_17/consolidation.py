"""Explicit fiscal consolidation inputs; fail on missing authoritative cells.

Native-First: parse the already held Census/OMB workbooks, with function-specific
ownership. A fiscal transfer is removed only when its final expenditure is priced.
"""
from pathlib import Path
import math
import hashlib

import openpyxl

SOURCES = {
    "outlays_fy2027.xlsx": ("https://www.whitehouse.gov/wp-content/uploads/2026/04/outlays_fy2027.xlsx",
                            "d892f2247e6c1aed68414d3e4168f8b4ab97bcfc7acf82a6a449a3fcb1addb07"),
    "omb_hist12z3_fy2027.xlsx": ("https://www.whitehouse.gov/wp-content/uploads/2026/04/hist12z3_fy2027.xlsx",
                               "43aa30c0d39116909b990a93926540fd173bce4963f26ee3a0d28a6f900d974d"),
}
TRANSPORT_ROWS = [3451, 3481, 3482, 3498, 3506, 3517, 3521, 3528,
                  3530, 3540, 3545, 3546, 3554, 3569, 3572, 3707]
# Physical public housing. Rental assistance, HOME, homeless and tribal programs
# are deliberately not inferred to belong to G: Census may classify them as welfare.
HOUSING_ROWS = [4074, 4079, 4089, 4104, 4107]


def verified_source(cache, name):
    path = Path(cache) / name
    if not path.is_file():
        raise ValueError(f"Missing {path}; run consolidation.py --fetch")
    if hashlib.sha256(path.read_bytes()).hexdigest() != SOURCES[name][1]:
        raise ValueError(f"Authoritative source hash changed: {path}; review the new vintage")
    return path


def federal_programs(cache):
    """Named program-netting scenario, not a recipient/timing-reconciled bridge."""
    path = verified_source(cache, "outlays_fy2027.xlsx")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book.worksheets[0]
    ws.reset_dimensions()
    rows = list(ws.values)
    year = next(i for i, v in enumerate(rows[0]) if str(v).strip() == "2024")
    result = {}
    for name, selected, expected in [("transport", TRANSPORT_ROWS, 67672000000),
                                      ("housing", HOUSING_ROWS, 8659000000)]:
        records = []
        for number in selected:
            row = rows[number - 1]
            if row[11] != "Grant" or (str(row[8])[:2] != "40" if name == "transport" else str(row[8]) != "604"):
                raise ValueError(f"OMB selected program changed: {name}/{number}")
            records.append(dict(row=number, agency=row[1], account=row[4], name=row[5],
                                subfunction=row[8], category=row[10], dollars=numeric(row[year], number) * 1000))
        total = sum(r["dollars"] for r in records)
        if total != expected:
            raise ValueError(f"OMB program sum changed: {name}/{total}")
        result[name] = dict(dollars=total, programs=records)
    justice_rows = [r for r in rows[1:] if len(r) > year and str(r[8]).startswith("75")
                    and str(r[8]) != "753" and r[11] == "Grant" and r[year] is not None]
    result["justice_grants"] = sum(numeric(r[year], "justice grant FY2024") * 1000 for r in justice_rows)
    if result["justice_grants"] != 6264000000:
        raise ValueError("OMB noncorrectional justice-grant total changed")
    book.close()
    path = verified_source(cache, "omb_hist12z3_fy2027.xlsx")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book.worksheets[0]
    ws.reset_dimensions()
    rows = list(ws.values)
    year = next(i for i, v in enumerate(rows[2]) if str(v).strip() == "2024")
    row = rows[385]
    if "Medicaid" not in str(row[0]):
        raise ValueError("OMB Medicaid source row changed")
    result["medicaid_federal"] = numeric(row[year], "Medicaid FY2024") * 1e6
    if result["medicaid_federal"] != 617517000000:
        raise ValueError("OMB Medicaid FY2024 total changed")
    row = rows[386]
    if "Children's Health Insurance" not in str(row[0]):
        raise ValueError("OMB CHIP source row changed")
    result["chip_federal"] = numeric(row[year], "CHIP FY2024") * 1e6
    if result["chip_federal"] != 19449000000:
        raise ValueError("OMB CHIP FY2024 total changed")
    book.close()
    result["limitation"] = "Function match only: federal program totals include some territorial/tribal recipients; G uses 2022 expenditure in 2024 prices. No exact recipient or fiscal-year matching is claimed."
    return result


def numeric(value, context):
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Missing/nonfinite authoritative cell {context}: {value!r}")
    return float(value)


def census_finance(path: Path):
    """2022 Census Table 1. Fees 28..36 correspond only to functions retained in G.

    Omit mixed 'other charges', education and hospitals, and miscellaneous revenue:
    these cannot all be credited against this account's covered service costs.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book["2022_US_WY"].values)
    by_line = {int(r[0]): r for r in rows if isinstance(r[0], (int, float)) and r[0]}
    result = {}
    for col, name in enumerate(rows[8]):
        if col < 2 or not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()
        def amount(line):
            return 1000 * numeric(by_line[line][col], f"{path.name}/{name}/line{line}")
        result[name] = dict(
            fees=sum(amount(i) for i in range(28, 37)),
            fees_by_line={str(i): amount(i) for i in range(28, 37)},
            federal_grants=amount(4), own_source=amount(7), direct_general=amount(66),
            miscellaneous=amount(38), unmapped_other_charges=amount(37))
    book.close()
    if len(result) != 52 or "United States Total" not in result:
        raise ValueError("Census finance requires US plus all 51 state jurisdictions")
    return result


def require_conservation(gross, already_priced, new_netting, remaining, context):
    values = [numeric(x, context) for x in (gross, already_priced, new_netting, remaining)]
    residual = values[0] - sum(values[1:])
    if abs(residual) > max(1.0, abs(gross) * 1e-10):
        raise ValueError(f"Fiscal component conservation failed: {context}: {residual}")
    return residual


if __name__ == "__main__":
    import argparse
    import urllib.request
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fetch", action="store_true")
    args = parser.parse_args()
    cache = Path(__file__).resolve().parent / "_cache"
    cache.mkdir(exist_ok=True)
    if args.fetch:
        for name, (url, expected) in SOURCES.items():
            target = cache / name
            if target.exists():
                verified_source(cache, name)
                continue
            data = urllib.request.urlopen(url).read()
            if hashlib.sha256(data).hexdigest() != expected:
                raise ValueError(f"Unexpected downloaded source vintage: {url}")
            target.write_bytes(data)
    print(federal_programs(cache))
