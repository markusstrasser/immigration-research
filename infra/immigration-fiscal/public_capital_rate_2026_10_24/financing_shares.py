"""Debt-financed and federally financed shares of state and local capital, by function (RESULT revision of 2026-10-24
evening). `rate_model.py` reads the output and sets d per level.

Concepts
- Debt share: long-term debt outstanding / current-cost net stock, the school lane's concept (F-33 debt / BEA K-12
  stock = 21%). In a steady state the stock is carried by its outstanding debt and by past tax- or charge-financed
  investment; the debt-carried part displaces private capital at d_bond, the rest at d_tax. New issuance / capital
  outlay measures how a year's projects are first paid for, not what carries the stock (debt amortizes in 20-30
  years against BEA service lives of 45-60), so it is not used.
- Federal share: federal grants for the function / the function's capital spending, FHWA's highway rule (42.5%).
  A federally financed share f costs the federal margin d_fed; debt carries s of the rest, so
  s = (D/K) / (1 - f) and d = f*d_fed + (1-f)*((1-s)*d_tax + s*d_bond).

Measurement (Census Annual Survey of State and Local Government Finances; BEA Fixed Assets Table 7.1)
- Census has collected debt as one public-purpose code (49U) since 2006 (classification manual, Table 6.1), so debt
  by function is not reported. Two measured routes:
  1. Utilities (water, electric, gas, transit): Census reports interest by utility (I91-I94). Debt = interest / the
     interest rate on the same function's single-function special districts (their 49U and interest).
  2. Sewer, airports, ports: single-function special districts' debt / construction, times the function's national
     construction, divided by the method's calibration. The calibration is the mean over-statement of the same
     estimator against the four functions measured directly (water, electric, transit by route 1; K-12 by F-33).
- Year: FY2022, the Census of Governments (every unit; the 2024 survey is a sample). FY2022 shares are moved to
  FY2024 by the change in the aggregate debt/stock ratio; that predicts K-12's FY2024 share (checked against F-33).
- Lines not measured take the residual: total debt less measured debt, over the unmeasured stock not federally
  financed (state and local highways at FHWA's 42.5%), so the shares add up to the total.

Run (from the repository root):
    uv run --no-project python3 infra/immigration-fiscal/public_capital_rate_2026_10_24/financing_shares.py
Writes derived/financing_shares.json and derived/financing_shares.csv only when every check passes.
"""
import csv
import hashlib
import json
import os
from collections import defaultdict
from pathlib import Path
from statistics import mean

import openpyxl

HERE = Path(__file__).resolve().parent
FISCAL = HERE.parent
CACHE = HERE / "_cache"
OUT = HERE / "derived"
CENSUS = CACHE / "census_dl"
FA = FISCAL / "school_capital_return_2026_09_26" / "_cache" / "fa_Section7All_xls.xlsx"
SCHOOL = FISCAL / "school_capital_return_2026_09_26" / "derived"
SERV = FISCAL / "capital_return_services_2026_09_27" / "derived"
NTD = CACHE / "ntd" / "ntd_funding_by_expense_type.csv"
OMB = CACHE / "omb" / "hist12z3_fy2027.xlsx"
F33_2022 = CENSUS / "elsec22_sumtables.xlsx"
FHWA_HWY_FEDERAL_SHARE = 0.425  # FHWA 23rd C&P Report, Highlights: federal share of highway capital spending, 2014

checks = []


def fail(msg):
    raise SystemExit(f"[BLOCKED] {msg}")


def check(name, ok, detail):
    checks.append({"check": name, "pass": bool(ok), "detail": detail})
    if not ok:
        fail(f"{name}: {detail}")


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------------- BEA Fixed Assets Table 7.1, current-cost net stock, year-end ($mn)
def bea_table(path, sheet):
    rows = list(openpyxl.load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    head = [i for i, r in enumerate(rows) if r and r[0] == "Line"][0]
    years = [int(y) for y in rows[head][3:] if y is not None]
    data, labels = {}, {}
    for r in rows[head + 1:]:
        if isinstance(r[0], int) or (isinstance(r[0], str) and r[0].strip().isdigit()):
            data[int(r[0])] = {y: float(v) for y, v in zip(years, r[3:]) if isinstance(v, (int, float))}
            labels[int(r[0])] = str(r[1]).strip()
    return data, labels


K, KL = bea_table(FA, "FAAt701-A")
LINES = {55: "State and local", 56: "Equipment", 58: "Residential", 62: "Educational", 65: "Transportation",
         66: "Power", 67: "Highways and streets", 68: "Sewer systems", 69: "Water systems",
         73: "Software"}
for ln, text in LINES.items():
    if not KL[ln].startswith(text):
        fail(f"FAAt701 line {ln} is '{KL[ln]}', expected '{text}'")


def avg(line, fy):
    """$bn; the two year-ends that bracket a fiscal year ending mid-year (the case's 2024 average for fy 2024)."""
    return (K[line][fy - 1] + K[line][fy]) / 2000


# Splits of BEA's transportation and highway types: the services lane's enterprise components (2024 averages).
ENT = {r["component"]: float(r["stock_avg2024_bn"]) for r in csv.DictReader(open(SERV / "enterprise_components.csv"))}
SPLIT = {"transit": ENT["ent_transit_sl"] / avg(65, 2024), "airports": ENT["ent_airports_sl"] / avg(65, 2024),
         "ports": ENT["ent_ports_sl"] / avg(65, 2024), "tolls": ENT["ent_tolls_sl"] / avg(67, 2024)}
check("splits_match_the_services_lane",
      abs(SPLIT["transit"] - 0.4229) < 5e-5 and abs(SPLIT["airports"] - 0.4981) < 5e-5 and abs(SPLIT["ports"] - 0.0790) < 5e-5
      and abs(SPLIT["tolls"] - 0.1025) < 5e-5 and abs(sum(SPLIT[k] for k in ("transit", "airports", "ports")) - 1) < 1e-4
      and all(abs(ENT[c] - avg(ln, 2024)) < 1e-6 for c, ln in (("ent_water_sl", 69), ("ent_sewer_sl", 68), ("ent_power_sl", 66))),
      f"transit {SPLIT['transit']:.4f}, airports {SPLIT['airports']:.4f}, ports {SPLIT['ports']:.4f} of S&L transportation; "
      f"tolls {SPLIT['tolls']:.4f} of S&L highways (capital_return_services_2026_09_27 gates); water, sewer, power = FAAt701 l69, l68, l66")
K12KEY = next(r for r in csv.DictReader(open(SCHOOL / "national_k12_capital.csv")) if r["variant"] == "central")


def stocks(fy):
    k12 = (float(K12KEY["key_structures"]) * avg(62, fy)
           + float(K12KEY["key_equipment_software"]) * (avg(56, fy) + avg(73, fy)))
    return {"total": avg(55, fy), "k12": k12, "water": avg(69, fy), "sewer": avg(68, fy), "power": avg(66, fy),
            "transit": avg(65, fy) * SPLIT["transit"], "airports": avg(65, fy) * SPLIT["airports"],
            "ports": avg(65, fy) * SPLIT["ports"], "tolls": avg(67, fy) * SPLIT["tolls"],
            "highways_nontoll": avg(67, fy) * (1 - SPLIT["tolls"])}


STOCK = {fy: stocks(fy) for fy in (2022, 2024)}
check("k12_stock_is_the_school_lanes", abs(STOCK[2024]["k12"] - float(K12KEY["total_avg2024_bn"])) < 0.01,  # keys rounded to 6 decimals in the CSV
      f"K-12 stock 2024 average {STOCK[2024]['k12']:.3f}bn (school lane central {float(K12KEY['total_avg2024_bn']):.3f})")


# ---------------- Census finance files ($bn)
def census_us(fy):
    out = defaultdict(float)
    for line in open(CENSUS / str(fy) / f"{fy % 100}statetypepu.txt"):
        if line[:3] == "001":  # United States, state and local total
            out[line[4:7]] += float(line[8:20]) / 1e6
    return out


def census_units(fy):
    d = CENSUS / str(fy)
    pid = {ln[:12]: (ln[2], ln[136:138]) for ln in open(d / f"Fin_PID_{fy}.txt", encoding="latin-1")}
    dat = [f for f in os.listdir(d) if f.startswith(f"{fy}FinEstDAT")]
    if len(dat) != 1:
        fail(f"one {fy} unit file expected, found {dat}")
    out = defaultdict(lambda: defaultdict(float))
    missing = 0
    for ln in open(d / dat[0], encoding="latin-1"):
        uid, code, amt = ln[:12], ln[12:15], float(ln[15:27]) / 1e6
        if uid not in pid:
            missing += 1
            continue
        t, fn = pid[uid]
        out[f"sd{fn}" if t == "4" else f"type{t}"][code] += amt
        out["all"][code] += amt
    if missing:
        fail(f"{missing} {fy} unit records have no PID entry")
    return out


C = {fy: census_us(fy) for fy in (2022, 2024)}
U = {fy: census_units(fy) for fy in (2022, 2024)}
INT = ("I89", "I91", "I92", "I93", "I94")
check("census_2022_units_add_to_the_us_total",
      all(abs(U[2022]["all"][c] - C[2022][c]) / C[2022][c] < 1e-3 for c in ("49U", "F12", "F80", "F91", "I89", "I91")),
      f"FY2022 unit file 49U {U[2022]['all']['49U']:.2f}bn vs state-by-type {C[2022]['49U']:.2f}bn (Census of Governments: every unit)")
NAT_RATE = {fy: sum(C[fy][c] for c in INT) / C[fy]["49U"] for fy in C}


def district_rate(fy, fn):
    a = U[fy]["sd" + fn]
    return (a["I89"] + a.get("I" + fn, 0.0)) / a["49U"]


def interest_debt(fy, fn, national=False):
    return C[fy]["I" + fn] / (NAT_RATE[fy] if national else district_rate(fy, fn))


def by_construction(fy, fn):
    a = U[fy]["sd" + fn]
    return a["49U"] / a["F" + fn] * C[fy]["F" + fn]


# F-33 FY2022 debt outstanding (Table 13, all school systems, $ thousands)
ws = openpyxl.load_workbook(F33_2022, read_only=True, data_only=True)["13"]
F33_DEBT_2022 = next(next(x for x in r[1:] if isinstance(x, (int, float))) for r in ws.iter_rows(values_only=True)
                     if r[0] and str(r[0]).startswith("Debt outstanding")) / 1e6  # first column: all school systems
F33_DEBT_2024 = next(float(r[1]) for r in csv.reader(open(SCHOOL / "cash_accrual_bridge_national.csv"))
                     if r[0] == "F-33 debt outstanding, end FY2024")
check("f33_debt", 540 < F33_DEBT_2022 < 560 and abs(F33_DEBT_2024 - 581.827914) < 1e-6,
      f"F-33 debt outstanding FY2022 {F33_DEBT_2022:.2f}bn (elsec22 Table 13), FY2024 {F33_DEBT_2024:.2f}bn (school lane)")

FY = 2022
debt = {
    "water": interest_debt(FY, "91"),
    "power": interest_debt(FY, "92") + interest_debt(FY, "93"),
    "transit": interest_debt(FY, "94"),
}
debt_national_rate = {"water": interest_debt(FY, "91", True),
                      "power": interest_debt(FY, "92", True) + interest_debt(FY, "93", True),
                      "transit": interest_debt(FY, "94", True)}
rates = {fn: district_rate(FY, fn) for fn in ("91", "92", "93", "94", "80", "01", "87")}
check("district_interest_rates_near_national", all(0.025 < r < 0.05 for r in rates.values()),
      f"FY2022 interest / debt: national {NAT_RATE[FY]:.4f}; districts " + ", ".join(f"{k} {v:.4f}" for k, v in rates.items())
      + f"; school districts {U[FY]['type5']['I89'] / U[FY]['type5']['49U']:.4f}; F-33 interest/debt "
      f"{20.950958 / F33_DEBT_2022:.4f}")
# Calibration of the district-construction estimator against the four directly measured functions.
calib = {
    "water": by_construction(FY, "91") / interest_debt(FY, "91"),
    "electric": by_construction(FY, "92") / interest_debt(FY, "92"),
    "transit": by_construction(FY, "94") / interest_debt(FY, "94"),
    "k12": U[FY]["type5"]["49U"] / U[FY]["type5"]["F12"] * C[FY]["F12"] / F33_DEBT_2022,
}
CAL = mean(calib.values())
check("district_estimator_calibration", 0.95 < min(calib.values()) and max(calib.values()) < 1.25,
      "district debt/construction x national construction over the direct measure: "
      + ", ".join(f"{k} {v:.3f}" for k, v in calib.items()) + f"; mean {CAL:.3f}")
coverage = {}
for name, fn in (("sewer", "80"), ("airports", "01"), ("ports", "87")):
    debt[name] = by_construction(FY, fn) / CAL
    coverage[name] = U[FY]["sd" + fn]["F" + fn] / C[FY]["F" + fn]
debt_alt = {n: (by_construction(FY, fn) / max(calib.values()), by_construction(FY, fn) / min(calib.values()))
            for n, fn in (("sewer", "80"), ("airports", "01"), ("ports", "87"))}
for n in ("water", "power", "transit"):
    debt_alt[n] = tuple(sorted((debt[n], debt_national_rate[n])))

# FY2022 debt / stock, moved to FY2024 by the aggregate ratio.
S22, S24 = STOCK[2022], STOCK[2024]
agg = {fy: C[fy]["49U"] / STOCK[fy]["total"] for fy in (2022, 2024)}
SCALE = agg[2024] / agg[2022]
k12_22 = F33_DEBT_2022 / S22["k12"]
k12_24 = F33_DEBT_2024 / S24["k12"]
check("fy2022_shares_move_to_fy2024_by_the_aggregate", abs(k12_22 * SCALE - k12_24) < 0.01,
      f"aggregate debt/stock FY2022 {agg[2022]:.4f}, FY2024 {agg[2024]:.4f} (ratio {SCALE:.4f}); K-12 FY2022 {k12_22:.4f} "
      f"x ratio = {k12_22 * SCALE:.4f} vs F-33 FY2024 {k12_24:.4f}")
MEASURED = ("water", "power", "transit", "sewer", "airports", "ports")
dk22 = {n: debt[n] / S22[n] for n in MEASURED}
dk24 = {n: dk22[n] * SCALE for n in MEASURED}
dk24["k12"] = k12_24
dk24_alt = {n: tuple(x / S22[n] * SCALE for x in debt_alt[n]) for n in MEASURED}
dk24_alt["k12"] = (k12_24, k12_24)

# ---------------- federal shares of capital spending
ntd = defaultdict(lambda: defaultdict(float))
for r in csv.DictReader(open(NTD)):
    if r["Fund Expenditure Type"] == "Funds Expended on Capital":
        for c in ("Federal", "Total"):
            ntd[int(r["Report Year"])][c] += float(r[c].replace(",", "") or 0) / 1e9
NTD_SHARE = {y: v["Federal"] / v["Total"] for y, v in sorted(ntd.items())}
rows = list(openpyxl.load_workbook(OMB, data_only=True).worksheets[0].iter_rows(values_only=True))
col = {str(x).strip(): j for j, x in enumerate(rows[2]) if x is not None}


def omb(label, year):
    hits = [r for r in rows if r[0] and str(r[0]).strip() == label]
    if len(hits) != 1:
        fail(f"OMB 12.3 row '{label}' found {len(hits)} times")
    v = hits[0][col[str(year)]]
    return float(v) / 1000 if isinstance(v, (int, float)) else 0.0


AIRPORT_ROWS = ("Grants for airports (Airport and airway trust fund)", "Airport Infrastructure Grants",
                "Airport Terminal Program", "Payment to Grants-in-aid for Airports")
cf = defaultdict(dict)
for r in csv.DictReader(open(SERV / "census_finance.csv")):
    cf[int(r["survey_year"])][r["code"]] = float(r["us_state_and_local_thousands"]) / 1e6
air = {y: sum(omb(x, y) for x in AIRPORT_ROWS) for y in (2017, 2018, 2019, 2024)}
air_share = {y: air[y] / (C[2024]["F01"] if y == 2024 else cf[y]["F01"]) for y in air}
check("airport_grant_share_stable", max(air_share.values()) - min(air_share.values()) < 0.08,
      "FAA airport grants (AIP, its general-fund supplements, AIG, ATP; OMB 12.3 outlays) / Census airport construction: "
      + ", ".join(f"FY{y} {v:.3f}" for y, v in air_share.items()) + "; FY2020-23 omitted: COVID relief paid through AIP")
hwy_census = {fy: C[fy]["B46"] / (C[fy]["F44"] + C[fy]["F45"]) for fy in C}
check("census_highway_aid_reproduces_fhwa", all(abs(v - FHWA_HWY_FEDERAL_SHARE) < 0.03 for v in hwy_census.values()),
      "Census federal highway aid (B46) / highway construction: " + ", ".join(f"FY{y} {v:.3f}" for y, v in hwy_census.items())
      + f" vs FHWA {FHWA_HWY_FEDERAL_SHARE}")
FED = {
    "highways_nontoll": FHWA_HWY_FEDERAL_SHARE,
    "transit": NTD_SHARE[2024],
    "airports": air_share[2024],
    "ports": omb("Port Infrastructure Development Program", 2024) / C[2024]["F87"],
    "sewer": C[2024]["B80"] / C[2024]["F80"],
}
check("transit_federal_share", 0.3 < min(NTD_SHARE.values()) and max(NTD_SHARE.values()) < 0.45,
      "NTD federal share of capital funds expended: " + ", ".join(f"{y} {v:.4f}" for y, v in NTD_SHARE.items()))

# ---------------- debt share of the non-federal stock; residual for the unmeasured lines
MEAS = MEASURED + ("k12",)


def assemble(dk):
    s = {n: dk[n] / (1 - FED.get(n, 0.0)) for n in MEAS}
    d_meas = sum(dk[n] * S24[n] for n in MEAS)
    k_meas = sum(S24[n] for n in MEAS)
    base = S24["total"] - k_meas - FHWA_HWY_FEDERAL_SHARE * S24["highways_nontoll"]
    s["residual"] = (C[2024]["49U"] - d_meas) / base
    return s, d_meas, base


s_central, d_meas, res_base = assemble(dk24)
s_low, _, _ = assemble({n: v[0] for n, v in dk24_alt.items()})
s_high, _, _ = assemble({n: v[1] for n, v in dk24_alt.items()})
total_back = sum(dk24[n] * S24[n] for n in MEAS) + s_central["residual"] * res_base
check("shares_add_to_total_debt", abs(total_back - C[2024]["49U"]) < 1e-6,
      f"measured debt {d_meas:.1f}bn + residual {s_central['residual']:.4f} x {res_base:.1f}bn = Census FY2024 49U "
      f"{C[2024]['49U']:.1f}bn")
old_implied = 0.21 * (S24["total"] - FHWA_HWY_FEDERAL_SHARE * S24["highways_nontoll"])
check("all_shares_in_range", all(0.05 < v < 0.5 for v in s_central.values()),
      ", ".join(f"{k} {v:.4f}" for k, v in s_central.items()))

LEVELS = {
    # level: (federal share, debt share of the non-federal stock, status)
    "state_local": (0.0, s_central["residual"], "residual: total less measured debt over the unmeasured non-federal stock"),
    "state_local_hwy": (FHWA_HWY_FEDERAL_SHARE, s_central["residual"], "FHWA federal share; debt at the residual"),
    "state_local_k12": (0.0, s_central["k12"], "F-33 FY2024 debt / school lane K-12 stock"),
    "state_local_water": (0.0, s_central["water"], "interest-based debt; federal share [GAP]"),
    "state_local_sewer": (FED["sewer"], s_central["sewer"], "district-construction debt; Census federal sewerage aid"),
    "state_local_power": (0.0, s_central["power"], "interest-based debt (electric + gas); no federal capital aid"),
    "state_local_transit": (FED["transit"], s_central["transit"], "interest-based debt; NTD federal share"),
    "state_local_airports": (FED["airports"], s_central["airports"], "district-construction debt; FAA grants (OMB 12.3)"),
    "state_local_ports": (FED["ports"], s_central["ports"], "district-construction debt; PIDP grants (OMB 12.3)"),
}
out = {
    "concepts": {
        "debt_to_stock": "long-term public-purpose debt outstanding (Census 49U) / BEA current-cost net stock, average of "
                         "the year-ends bracketing the fiscal year",
        "federal_share": "federal grants for the function / its capital spending (FHWA's highway rule)",
        "debt_share_of_nonfederal_stock": "debt_to_stock / (1 - federal_share); d = f*d_fed + (1-f)*((1-s)*d_tax + s*d_bond)",
    },
    "levels": {k: {"federal_share": f, "debt_share_of_nonfederal_stock": s, "status": st} for k, (f, s, st) in LEVELS.items()},
    "debt_to_stock_fy2024": dk24,
    "debt_to_stock_fy2022": {**dk22, "k12": k12_22},
    "debt_to_stock_alternatives_fy2024": dk24_alt,
    "debt_share_of_nonfederal_stock_alternatives": {"low": s_low, "high": s_high},
    "aggregate_debt_to_stock": agg, "fy2022_to_fy2024_scale": SCALE,
    "old_rule_implied_debt_fy2024_bn": old_implied, "census_debt_fy2024_bn": C[2024]["49U"],
    "debt_fy2022_bn": debt, "debt_fy2022_national_rate_bn": debt_national_rate,
    "district_interest_rate_fy2022": rates, "national_interest_rate": NAT_RATE,
    "district_estimator_calibration": {**calib, "mean": CAL}, "district_construction_coverage_fy2022": coverage,
    "stocks_bn": {str(fy): v for fy, v in STOCK.items()},
    "federal_shares": FED, "federal_share_checks": {"ntd_transit_by_year": NTD_SHARE, "airports_by_fy": air_share,
                                                    "census_highways_by_fy": hwy_census,
                                                    "census_sewer_fy2022": C[2022]["B80"] / C[2022]["F80"]},
    "residual_base_bn": res_base, "measured_debt_fy2024_bn": d_meas,
    "gaps": ["water: federal capital share (drinking-water SRF capitalization grants) is not coded by Census and EPA's "
             "STAG outlays are not split by fund; no federal blend", "tolls, housing, colleges, health, parks, general "
             "government and other enterprises: no function-level debt measure; they take the residual",
             "housing: federal share of public housing capital (HUD Capital Fund, merged into the Public Housing Fund "
             "from FY2021) is not separable from operating aid; no federal blend",
             "state bonds for school construction are state debt (F-33 omits them): K-12 understated, residual overstated"],
    "inputs_sha256": {str(p.relative_to(FISCAL)): sha(p) for p in sorted(
        [FA, SCHOOL / "national_k12_capital.csv", SCHOOL / "cash_accrual_bridge_national.csv", SERV / "enterprise_components.csv",
         SERV / "census_finance.csv", NTD, OMB, F33_2022]
        + [CENSUS / str(fy) / f for fy in (2022, 2024) for f in sorted(os.listdir(CENSUS / str(fy))) if f.endswith(".txt") and "FinEst" in f or f.endswith("pu.txt") or f.startswith("Fin_PID")])},
    "checks": checks,
}
OUT.mkdir(exist_ok=True)
(OUT / "financing_shares.json").write_text(json.dumps(out, indent=1, sort_keys=False) + "\n")
with open(OUT / "financing_shares.csv", "w", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["level", "federal_share", "debt_share_of_nonfederal_stock", "status"])
    for k, (fs, s, st) in LEVELS.items():
        w.writerow([k, round(fs, 6), round(s, 6), st])
print(json.dumps({k: v for k, v in out.items() if k in ("levels", "debt_to_stock_fy2024", "aggregate_debt_to_stock",
                                                         "federal_shares", "district_estimator_calibration")}, indent=1))
