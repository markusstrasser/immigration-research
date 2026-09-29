"""DHS Yearbook FY2024 nonimmigrant tables: I-94 admissions of Indian citizens.
Table 27 (all classes, FY2015-2024) and NISuppTable1 (class x country of citizenship, FY2024).
Appends rows to context/flows_2015_2025.csv (run after flows_2015_2025.py). I-94 admissions are entries, not persons."""
import csv, hashlib, pathlib
import openpyxl
LANE = pathlib.Path(__file__).resolve().parents[1]
F = LANE.parent / "indian_ledger_2026_09_18/_cache/yearbook_nonimmigrants_fy2024.xlsx"
SHA = hashlib.sha256(F.read_bytes()).hexdigest()
URL = "[UNVERIFIED url] ohss.dhs.gov yearbook nonimmigrants FY2024 xlsx (local name 20260604_ohss_yearbook_nonimmigrants_fy2024.xlsx, same sha256)"
wb = openpyxl.load_workbook(F, read_only=True, data_only=True)
rows = []
rr = [list(r) for r in wb["Table 27"].iter_rows(values_only=True)]
hdr = next(r for r in rr if r and isinstance(r[0], str) and r[0].startswith("Region and country"))
ind = next(r for r in rr if r and r[0] == "India")
for j, y in enumerate(hdr):
    if j and y:
        rows.append(["DHS Yearbook FY2024 Table 27", int(y), "i94_admissions_all_classes", ind[j], "I-94 admissions, citizens of India",
                     str(F.relative_to(LANE.parent)), f"India row: {tuple(ind[:11])}", URL, SHA])
rr = [list(r) for r in wb["NISuppTable1"].iter_rows(values_only=True)]
hdr = next(r for r in rr if r and r[0] == "Class")
ji = hdr.index("India")
WANT = {"H1B", "H1B1", "H4", "F1", "F2", "M1", "J1", "J2", "L1", "L2", "B1", "B2", "B1/B2", "WB", "WT", "O1", "TN"}
for r in rr:
    if r and isinstance(r[0], str) and r[0].strip() in WANT:
        rows.append(["DHS Yearbook FY2024 NISuppTable1", 2024, f"i94_{r[0].strip()}", r[ji], f"I-94 admissions, citizens of India ({r[1]})",
                     str(F.relative_to(LANE.parent)), f"{r[0]} | {r[1]} | India col = {r[ji]}", URL, SHA])
with open(LANE / "context/flows_2015_2025.csv", "a", newline="") as fh:
    csv.writer(fh, lineterminator="\n").writerows(rows)
for r in rows: print(r[1], r[2], r[3], r[4][-40:])
