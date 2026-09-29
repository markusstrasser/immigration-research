"""Tidy India flow series FY2015-FY2025 -> context/flows_2015_2025.csv.

(a) DHS Yearbook Table 10 (LPRs by broad class x country of birth), FY2015-FY2024, from local yearbook files
    (admission_route_2026_09_21/_cache, URLs+sha256 in its acquire_manifest.json; FY2024 from indian_ledger_2026_09_18/_cache).
(a') OHSS special tabulation, LPRs by country x major class incl. employment-based derivatives (spouses/children),
    FY2005-2024 (late_arrival_tail_2026_09_27/_cache). Principals = class total - derivatives.
(b) USCIS H-1B approvals for India-born, from context/h1b_characteristics_fy2003_2025.csv.
(d) CBP encounters, from context/cbp_india_encounters_summary.csv.
"""
import csv, hashlib, json, pathlib, sys
import openpyxl, xlrd
LANE = pathlib.Path(__file__).resolve().parents[1]
FIS = LANE.parent
AR = FIS / "admission_route_2026_09_21/_cache"
MAN = json.load(open(AR / "acquire_manifest.json"))
rows = []
def sha(p): return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def sheet_rows(path, sheet_hint):
    p = str(path)
    if p.endswith(".xls"):
        wb = xlrd.open_workbook(p); sh = wb.sheet_by_index(0)
        return [sh.row_values(i) for i in range(sh.nrows)]
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    names = [s for s in wb.sheetnames if s in (sheet_hint, sheet_hint + "d")]
    name = names[0] if names else wb.sheetnames[-1] if len(wb.sheetnames) <= 2 else None
    if name is None: raise SystemExit(f"[BLOCKED] no {sheet_hint} sheet in {p}: {wb.sheetnames}")
    return [list(r) for r in wb[name].iter_rows(values_only=True)]
T10 = {}
for y in range(2015, 2022):
    d = AR / f"lpr_fy{y}"
    f = next(iter(sorted(d.rglob(f"*table10.xls*")) + sorted(d.rglob(f"*table10d.xls*"))))
    T10[y] = (f, MAN.get(str(y), {}).get("url", ""))
for y in (2022, 2023):
    T10[y] = (AR / f"lpr_fy{y}.xlsx", MAN.get(str(y), {}).get("url", ""))
T10[2024] = (FIS / "indian_ledger_2026_09_18/_cache/yearbook_lpr_fy2024.xlsx", "[UNVERIFIED url] ohss.dhs.gov LPR yearbook FY2024 xlsx")
CLS = ["Total", "Immediate relatives of U.S. citizens", "Family-sponsored preferences", "Employment-based preferences",
       "Diversity", "Refugees and asylees", "Other"]
for y, (f, url) in sorted(T10.items()):
    rr = sheet_rows(f, "Table 10")
    hdr = next(r for r in rr if r and isinstance(r[0], str) and r[0].strip().lower().startswith("region and country of birth"))
    title = next(str(r[0]) for r in rr if r and isinstance(r[0], str) and "BROAD CLASS" in r[0].upper())
    assert str(y) in title, (f, title)
    ind = next(r for r in rr if r and isinstance(r[0], str) and r[0].strip() == "India")
    hn = [str(h).replace("\n", " ").strip() if h else "" for h in hdr]
    for c in CLS:
        j = next(i for i, h in enumerate(hn) if h.lower().startswith(c.lower()[:18]))
        v = ind[j]
        v = int(v) if isinstance(v, float) and v.is_integer() else v
        rows.append(["DHS Yearbook Table 10", y, f"lpr_{c}", v, "persons obtaining LPR, India-born",
                     str(f.relative_to(FIS)), f"{title[:90]} | India row: {tuple(ind[:8])}", url, sha(f)])
# special tabulation with derivatives
SP = FIS / "late_arrival_tail_2026_09_27/_cache/2026_0604_ohss_lpr_by_country_by_major_class_and_deriv_emp-based_fy2005-2024.xlsx"
wb = openpyxl.load_workbook(SP, read_only=True, data_only=True)
for s in wb.sheetnames[1:]:
    rr = [list(r) for r in wb[s].iter_rows(values_only=True)]
    hdr = next(r for r in rr if r and r[0] == "Region and country of birth")
    ind = next(r for r in rr if r and isinstance(r[0], str) and r[0].strip() == "India")
    for j, yy in enumerate(hdr):
        if isinstance(yy, int) and 2015 <= yy <= 2024:
            rows.append(["OHSS special tab LPR by country x major class", yy, "lpr_" + s.replace("Emp- ", "EB ").replace(" ", "_"),
                         ind[j], "persons, rounded to 10", str(SP.relative_to(FIS)), f"sheet '{s}' India row",
                         "https://ohss.dhs.gov/topics/immigration/lawful-permanent-residents/lprs-country-birth-and-major-classes-admission", sha(SP)])
sp = {(r[1], r[2]): r[3] for r in rows if r[0].startswith("OHSS special")}
for yy in range(2015, 2025):
    tot = sum(sp[(yy, f"lpr_EB_{k}_Preference")] for k in ["First", "Second", "Third", "Fourth", "Fifth"])
    der = sum(sp[(yy, f"lpr_EB_{k}_Preference_Deriv")] for k in ["First", "Second", "Third", "Fourth", "Fifth"])
    rows.append(["derived from OHSS special tab", yy, "lpr_EB_all_principals", tot - der, "persons (rounded inputs)", "", f"EB1-5 total {tot} minus derivatives {der}", "", ""])
    rows.append(["derived from OHSS special tab", yy, "lpr_EB_all_derivatives", der, "persons (rounded inputs)", "", f"sum of EB1-5 Deriv sheets", "", ""])
for r in csv.DictReader(open(LANE / "context/h1b_characteristics_fy2003_2025.csv")):
    fy = int(r["fiscal_year"])
    if fy >= 2015 and r["india_approved_n"]:
        a, i = int(r["india_approved_n"]), int(r["india_initial_n"])
        for k, v in [("h1b_approved_all", a), ("h1b_approved_initial", i), ("h1b_approved_continuing", a - i)]:
            rows.append(["USCIS H-1B Characteristics report", fy, k, v, "approved petitions, India-born",
                         f"_cache/context/h1b ({r['india_source']})", r["india_evidence"], "https://www.uscis.gov/tools/reports-and-studies", ""])
for r in csv.DictReader(open(LANE / "context/cbp_india_encounters_summary.csv")):
    for k in ["total", "swb_usbp", "nb_usbp", "ofo"]:
        rows.append(["CBP nationwide encounters", int(r["fiscal_year"]), f"cbp_encounters_{k}", int(r[k]), "encounter events, citizenship India",
                     "context/cbp_india_encounters_summary.csv", "see context_notes item 1", "https://www.cbp.gov/newsroom/stats/nationwide-encounters", ""])
with open(LANE / "context/flows_2015_2025.csv", "w", newline="") as fh:
    w = csv.writer(fh, lineterminator="\n")
    w.writerow(["source", "fiscal_year", "series", "value", "unit", "file", "quoted_line", "url", "file_sha256"])
    w.writerows(rows)
print(len(rows), "rows")
for r in rows:
    if r[2] in ("lpr_Total", "lpr_Employment-based preferences", "lpr_EB_all_principals", "lpr_EB_all_derivatives",
                "lpr_Immediate relatives of U.S. citizens", "lpr_Family-sponsored preferences", "lpr_Refugees and asylees", "lpr_Other", "lpr_Diversity"):
        print(r[1], r[2], r[3])
