"""Business and other NCVS-uncovered crime: the 2024 cost to US residents outside the 44.90M
Mexican-origin lineage (main case v9, key oct16) of offences the NCVS does not see.

    OPENBLAS_NUM_THREADS=1 uv run --no-project python3 infra/immigration-fiscal/business_crime_2026_10_21/business_crime.py
    (add --fetch once to download the raw files into _cache/; normal runs only verify their pins)

Lines (each additive to the case's social rows; see the overlap table in RESULT.md):
    A  retail external theft (shoplifting, ORC, retail burglary/robbery/card fraud by non-employees):
       NRF NRSS 2023 shrink rate x external-theft share x Census MRTS 2024 retail sales, at cost
    B  retail employee theft: the NRSS internal-theft share of the same shrink
    C  police-recorded business losses outside retail stores (NIBRS business, financial, religious
       and other victims; burglary, robbery, larceny, vehicle theft, vandalism, arson, fraud,
       embezzlement, extortion), national by CIUS 2024 counts
    D  violent victimisations of children under 12, outside the family, police-recorded (NIBRS)
    E  violent victimisations at jail/prison locations, victims 12+, police-recorded (NIBRS)
    F  fraud, embezzlement and extortion against individuals, police-recorded losses (NIBRS)
Offender shares are NIBRS police-recorded offender ethnicity in Texas and Arizona 2022-2023
(agencies recording ethnicity for >= 50% of offenders), carried to the nation with the offender-
ethnicity lane's relative-rate translation and to Mexican origin with the victim lane's population
share (0.5975), then to the lineage with the case's crime key. Business losses pass to purchasers:
outsiders bear the forward share (1 - payment share) of group-caused losses, and removal also
returns the group's payment toward others' losses (the live private_security row's rule).

Writes derived/*.csv (LF) and derived/run_log.txt. Every gate stops the run with [BLOCKED].
"""
from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

import bizcrime_stage as S

HERE = Path(__file__).resolve().parent
FISCAL = HERE.parent
CACHE = HERE / "_cache"
OUT = HERE / "derived"
SIB = S.SIB
sys.path.insert(0, str(SIB))
import nibrs_rates as NR  # noqa: E402
from nibrs_stage import OFF_CLASSES  # noqa: E402

VL = FISCAL / "crime_victim_cost_2026_09_23"
ARR = FISCAL / "nibrs_arrests_2026_09_16"
PROP = FISCAL / "sept24_propagation_2026_09_24/derived/oct16/real_costs_totals.json"
SECURITY_KEYS = FISCAL / "social_costs_unpriced_2026_09_28/derived/payment_keys.json"
WAGES = FISCAL / "job_capital_taxes_2026_10_14/derived/wages_by_industry.csv"
MORT = VL / "derived/mortality_arms.csv"
PRICES = VL / "derived/unit_costs_victim_only_2024usd.csv"
SCALING = VL / "derived/scaling_hispanic_to_mexican.csv"
TRANSFER = SIB / "derived/outside_share_by_offence.csv"
CIUS = ARR / "_cache/x/est2024"
G = ["H", "NHW", "NHB", "NHO"]
BIZ = ("B", "F", "R", "O")
SRS = {"burglary": "Burglary", "robbery": "Robbery", "motor_vehicle_theft": "Motor vehicle theft",
       "shoplifting": "Shoplifting", "theft_from_building": "From buildings",
       "theft_coin_machine": "From coin-operated machines", "other_larceny": "Larceny other"}
LOG: list[str] = []

FETCH = {  # cache name -> URL (retrieved 2026-10-21)
    "nrss2023.pdf": "https://info.apprissretail.com/hubfs/Resource-Center/2023.038-national-retail-security-survey-report-2023.pdf",
    "mrtssales92-present.xlsx": "https://www.census.gov/retail/mrts/www/mrtssales92-present.xlsx",
    "arts2022r_gmper.xlsx": "https://www2.census.gov/programs-surveys/arts/tables/2022restated/gmper.xlsx",
}


def say(s: str = "") -> None:
    print(s, flush=True)
    LOG.append(s)


def gate(name: str, ok: bool, detail: str = "") -> None:
    say(f"[gate] {name}: {'PASS' if ok else 'FAIL'}{' — ' + detail if detail else ''}")
    if not ok:
        raise SystemExit(f"[BLOCKED] {name}")


def sha(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def write_csv(name: str, rows: list[dict]) -> None:
    with (OUT / name).open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({k: (f"{v:.6f}" if isinstance(v, float) else v) for k, v in r.items()})


# ---------------------------------------------------------------------------------------------
# 0. Sources and pins
# ---------------------------------------------------------------------------------------------
def fetch() -> None:
    CACHE.mkdir(exist_ok=True)
    ua = "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 Chrome/120 Safari/537.36"
    for name, url in FETCH.items():
        subprocess.run(["curl", "-sSfL", "-A", ua, "-o", str(CACHE / name), url], check=True)
    subprocess.run(["pdftotext", "-layout", str(CACHE / "nrss2023.pdf"), str(CACHE / "nrss2023.txt")], check=True)


def prop_read(prop: dict) -> dict:
    """The fields this lane reads from the pairing's oct16 totals. The pairing carries this lane's row, so a
    file pin would make each lane stale whenever the other is rerun; the pin covers exactly what is read."""
    keys = prop["lineage_social_rows"]["keys"]
    p7 = prop["p7_crime"]
    return dict(case=prop["case"], justice_use_raw_coding_low=keys["low_end"]["justice_use_raw_coding"],
                justice_use_high=keys["high_end"]["justice_use"], recipient_victims=p7["recipient_shares"]["victims"],
                lineage_m=prop["per_member_population_m"]["oct16"],
                victims_lineage={e: p7["lineage_crime_rows"][e]["victims"] for e in ("low", "high")})


def value_digest(values: dict) -> str:
    return hashlib.sha256(json.dumps(values, sort_keys=True).encode()).hexdigest()


def verify_pins() -> dict:
    pins = json.loads((HERE / "inputs/pins.json").read_text())
    for rel, digest in pins["files"].items():
        p = (HERE / rel).resolve()
        gate(f"pin {rel}", p.exists() and sha(p) == digest, digest[:12])
    for rel, digest in pins["read_values"].items():
        p = (HERE / rel).resolve()
        gate(f"read-value pin {rel}", p.exists() and value_digest(prop_read(json.loads(p.read_text()))) == digest,
             digest[:12])
    return pins


# ---------------------------------------------------------------------------------------------
# 1. NIBRS cells and offender shares
# ---------------------------------------------------------------------------------------------
def load_stages() -> dict:
    acs = pd.read_csv(SIB / "derived/acs5_place_county_groups.csv")
    stages = {}
    for st, yr in NR.STATE_YEARS:
        ag = pd.read_pickle(SIB / f"_cache/stage/{st}-{yr}.pkl")["agencies"]
        ag = ag.merge(NR.compositions(st, yr, ag, acs), on="agency_id", how="left")
        d = S.load(SIB / f"_cache/{st}-{yr}.zip", st, yr)
        d["agencies"] = ag
        stages[(st, yr)] = d
    return stages


SPEC_CENTRAL = dict(states=("TX", "AZ"), thresh=0.5, alloc="a", weight="n", arrest_only=False)


def specs() -> list[tuple[str, dict]]:
    c = SPEC_CENTRAL
    return [("central", c),
            ("loss-weighted", {**c, "weight": "l"}),
            ("agencies recording ethnicity >= 80%", {**c, "thresh": 0.8}),
            ("agencies recording ethnicity >= 95%", {**c, "thresh": 0.95}),
            ("incidents with an arrest only", {**c, "arrest_only": True}),
            ("states=TX", {**c, "states": ("TX",)}),
            ("states=AZ", {**c, "states": ("AZ",)}),
            ("states=TX+AZ+CA", {**c, "states": ("TX", "AZ", "CA")}),
            ("alloc=a0 (offence-level proportional)", {**c, "alloc": "a0"}),
            ("bound: unknown offenders all non-Hispanic", {**c, "alloc": "b"}),
            ("bound: unknown offenders all Hispanic", {**c, "alloc": "c"})]


def kept(d: dict, sp: dict) -> pd.DataFrame:
    ag = d["agencies"]
    return ag[ag.comp_source.ne("tribal_excluded") & ag.rec_rate.fillna(0).ge(sp["thresh"])]


def shares(stages: dict, sp: dict, natpop: dict, cell: str, vtypes, retail=None,
           key: str = "group") -> pd.DataFrame:
    """National Hispanic share of offending by `key`, from cells `cell` (pcells, kcells, jcells),
    relative rates per resident of all ages translated to the 2024 national population."""
    cells, pops = [], []
    for (st, yr), d in stages.items():
        if st not in sp["states"]:
            continue
        k = kept(d, sp)
        p = d[cell]
        p = p[p.agency_id.isin(k.agency_id)]
        if cell == "pcells":
            p = p[p.vtype.isin(vtypes)]
            if retail is not None:
                p = p[p.retail.eq(retail)]
            if sp["arrest_only"]:
                p = p[p.arrest]
            cols = [f"{sp['weight']}_{c}" for c in OFF_CLASSES]
        else:
            cols = OFF_CLASSES
        c = p.groupby(key)[cols].sum()
        c.columns = OFF_CLASSES
        c.index.name = "offence"
        c = pd.concat({"vB": c}, names=["vclass"]).swaplevel()
        cells.append(pd.concat({(st, yr): c}, names=["state", "year"]))
        pops.append(k[[f"c_{g}" for g in ["total", "hisp", "nhw", "nhb", "nho"]]].sum())
    C = pd.concat(cells)
    P = pd.concat(pops, axis=1).sum(axis=1)
    A = NR.allocate(C, sp["alloc"]).groupby(level="offence").sum()
    Pg = {"H": P.c_hisp, "NHW": P.c_nhw, "NHB": P.c_nhb, "NHO": P.c_nho}
    nat = natpop["ncest2024_all"]
    rows = []
    for o in A.index:
        r = {g: A.loc[o, g] / Pg[g] for g in G}
        if r["NHW"] <= 0:
            continue
        RR = {g: r[g] / r["NHW"] for g in G}
        rows.append(dict(offence=o, weight_total=float(A.loc[o].sum()), RR_H_NHW=RR["H"],
                         RR_H_all=r["H"] / (A.loc[o].sum() / P.c_total), local_share_H=float(A.loc[o, "H"] / A.loc[o].sum()),
                         local_pop_share_H=float(P.c_hisp / P.c_total),
                         national_share_H=RR["H"] * nat["H"] / sum(RR[g] * nat[g] for g in G),
                         unknown_share=float(C.xs(o, level="offence")[["UW", "UB", "UO", "UU", "NOINFO"]].sum().sum()
                                             / C.xs(o, level="offence").sum().sum())))
    return pd.DataFrame(rows).set_index("offence")


def arrestee_shares(stages: dict, natpop: dict, states=("TX", "AZ")) -> pd.DataFrame:
    """Hispanic share of arrestees for business-victim incidents, translated as offenders are."""
    rows = []
    sp = SPEC_CENTRAL
    agg, pops = [], []
    for (st, yr), d in stages.items():
        if st not in states:
            continue
        k = kept(d, sp)
        a = d["acells"]
        a = a[a.agency_id.isin(k.agency_id) & a.biz_incident]
        cl = a.cls.str
        a = a.assign(g=np.select([a.cls.str.startswith("H"), a.cls.eq("NHW"), a.cls.eq("NHB"), a.cls.eq("NHO")],
                                 ["H", "NHW", "NHB", "NHO"], "U"))
        del cl
        agg.append(a.groupby(["group", "g"]).arrestees.sum())
        pops.append(k[[f"c_{g}" for g in ["total", "hisp", "nhw", "nhb", "nho"]]].sum())
    A = pd.concat(agg).groupby(level=[0, 1]).sum().unstack(fill_value=0)
    P = pd.concat(pops, axis=1).sum(axis=1)
    Pg = {"H": P.c_hisp, "NHW": P.c_nhw, "NHB": P.c_nhb, "NHO": P.c_nho}
    nat = natpop["ncest2024_all"]
    for o in A.index:
        known = A.loc[o, G].sum()
        r = {g: A.loc[o, g] / Pg[g] for g in G}
        if r["NHW"] <= 0:
            continue
        RR = {g: r[g] / r["NHW"] for g in G}
        rows.append(dict(offence=o, arrestees=float(A.loc[o].sum()), eth_unknown_share=float(1 - known / A.loc[o].sum()),
                         local_share_H=float(A.loc[o, "H"] / known), RR_H_NHW=RR["H"],
                         national_share_H=RR["H"] * nat["H"] / sum(RR[g] * nat[g] for g in G)))
    return pd.DataFrame(rows).set_index("offence")


# ---------------------------------------------------------------------------------------------
# 2. National volumes
# ---------------------------------------------------------------------------------------------
def xlsx_rows(path: Path, sheet=None):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return [r for r in ws.iter_rows(values_only=True)]


def mrts_2024() -> dict:
    rows = xlsx_rows(CACHE / "mrtssales92-present.xlsx", "2024")
    hdr = next(r for r in rows if r and "TOTAL" in [str(x).strip() for x in r if x is not None])
    col = [str(x).strip() if x is not None else "" for x in hdr].index("TOTAL")
    out = {}
    adjusted = False
    for r in rows:
        if r[1] and str(r[1]).strip() == "ADJUSTED(2)":
            adjusted = True
        label = str(r[1]).strip() if r[1] else ""
        if not adjusted and label in ("Retail sales, total", "Retail sales, total (excl. motor vehicle and parts dealers)",
                                      "Gasoline stations", "Nonstore retailers", "Motor vehicle and parts dealers"):
            out.setdefault(label, float(r[col]))
    gate("MRTS 2024 not-adjusted annual totals parsed", len(out) == 5 and abs(out["Retail sales, total"] - 7106697) < 1,
         f"retail total ${out.get('Retail sales, total', 0) / 1e6:.4f}tn")
    gate("MRTS: retail total = excl. motor vehicles + motor vehicles",
         abs(out["Retail sales, total"] - out["Retail sales, total (excl. motor vehicle and parts dealers)"]
             - out["Motor vehicle and parts dealers"]) < 2)
    return out


def gross_margin_2022() -> dict:
    rows = xlsx_rows(CACHE / "arts2022r_gmper.xlsx")
    out = {}
    for r in rows:
        lab = str(r[1]) if r[1] else ""
        if lab.startswith("Retail gross margin as a percentage of sales, total (excl. motor"):
            out["excl_mv"] = float(r[2]) / 100
        elif lab.startswith("Retail gross margin as a percentage of sales, total"):
            out["total"] = float(r[2]) / 100
    gate("ARTS 2022 (restated) gross margin parsed", out == {"total": 0.313, "excl_mv": 0.331}, str(out))
    return out


def nrss() -> dict:
    t = (CACHE / "nrss2023.txt").read_text()
    flat = " ".join(t.split())
    for quote in ["average shrink rate in FY 2022 increased to 1.6%, up from 1.4% in FY 2021",
                  "$112.1 billion in losses", "accounted for an average of 36% of total loss",
                  "Internal (employee) theft reported in at 29% of shrink loss",
                  "Process, control failures and errors fell at 27%, with unknown (6%) and other (1%)",
                  "more than three-quarters (78%) of respondents do not include ecommerce goods loss",
                  "over half (57%) do not include supply chain loss or theft"]:
        gate(f"NRSS 2023 text: '{quote[:48]}...'", quote in flat)
    return dict(shrink=0.016, shrink_fy2021=0.014, external=0.36, internal=0.29, process=0.27, unknown=0.06, other=0.01)


def cius_2024() -> dict:
    t7 = xlsx_rows(CIUS / "CIUS_Table_7_Offense_Analysis_United_States_2020-2024.xlsx")
    hdr = next(r for r in t7 if r[0] == "Classification")
    col = list(hdr).index(2024)
    v = {}
    section = None
    for r in t7:
        if r[0]:
            section = str(r[0]).strip()
        lab = str(r[1]).strip() if r[1] else ""
        if section.startswith("Robbery:") and lab == "Total2":
            v["Robbery"] = r[col]
        if section.startswith("Burglary:") and lab == "Total2":
            v["Burglary"] = r[col]
        if section.startswith("Larceny-theft (except") and lab == "Total2":
            v["Larceny total"] = r[col]
        if section == "Larceny-theft by type:" or (section.startswith("Larceny-theft by type") ):
            pass
        if lab in ("Shoplifting", "From buildings", "From coin-operated machines", "Commercial house",
                   "Gas or service station", "Convenience store", "Bank"):
            v[lab] = r[col]
        if section == "Motor vehicle theft":
            v["Motor vehicle theft"] = r[col]
    v["Larceny other"] = v["Larceny total"] - v["Shoplifting"] - v["From buildings"] - v["From coin-operated machines"]
    gate("CIUS 2024 Table 7 parsed", v["Shoplifting"] == 1270097 and v["Motor vehicle theft"] == 880327
         and v["Burglary"] == 779542, f"shoplifting {v['Shoplifting']:,}")
    t1 = xlsx_rows(CIUS / "CIUS_Table_1_Crime_in_the_United_States_by_Volume_and_Rate_per_100000_Inhabitants_2005-2024.xlsx")
    r24 = next(r for r in t1 if r[0] == 2024)
    v["Property crime"], v["Population"] = r24[14], r24[1]
    gate("CIUS 2024 Table 1 property crime", v["Property crime"] == 5986400, f"{v['Property crime']:,}")
    return v


# ---------------------------------------------------------------------------------------------
# 3. The model
# ---------------------------------------------------------------------------------------------
def main() -> None:
    if "--fetch" in sys.argv:
        fetch()
    OUT.mkdir(exist_ok=True)
    pins = verify_pins()
    # case keys
    prop = prop_read(json.loads(PROP.read_text()))
    k_end = {"low": prop["justice_use_raw_coding_low"], "high": prop["justice_use_high"]}
    recipient = prop["recipient_victims"]
    lineage_m = prop["lineage_m"]
    pay = json.loads(SECURITY_KEYS.read_text())
    p_lin, p_union = pay["low"]["lineage_share"], pay["low"]["union_share"]
    gate("payment key equal at both ends", pay["high"]["lineage_share"] == p_lin)
    gate("case keys read (oct16)", prop["case"] == "oct16" and abs(lineage_m - 44.903074) < 1e-5,
         f"k_low {k_end['low']:.6f}, k_high {k_end['high']:.6f}, recipient {recipient:.6f}, lineage {lineage_m:.6f}M, "
         f"payment lineage {p_lin:.6f}")
    victims_v9 = dict(prop["victims_lineage"])
    sc = {r["item"]: float(r["value"]) for r in csv.DictReader(SCALING.open())}
    s_mex = sc["population share, within CPS (central)"]
    s_lo, s_hi = sc["ACS 2020-24 institutional, men 18-39: ratio_generic_kept"], sc["ACS 2020-24 institutional, men 18-39: ratio_generic_reallocated"]
    m_tract = sc["Mexican share of Hispanic neighbours, tract"]
    gate("Hispanic -> Mexican population share is the victim lane's", abs(s_mex - 0.597521) < 1e-6, f"{s_mex:.6f}")

    # murder on the case
    mort = {(r["footing"], r["vsl_standard"]): r for r in csv.DictReader(MORT.open())
            if float(r["real_discount_rate"]) == 0 and r["missing_age_arm"] == "central"}
    murder_share = {"low": float(mort["mixed", "dot_uniform"]["mortality_bn"]) / float(mort["mixed", "dot_uniform"]["full_bn"]),
                    "high": float(mort["custody", "dot_uniform"]["mortality_bn"]) / float(mort["custody", "dot_uniform"]["full_bn"])}
    murder_v9 = {e: victims_v9[e] * murder_share[e] for e in ("low", "high")}
    legacy_share = float(mort["equal", "legacy"]["mortality_bn"]) / float(mort["equal", "legacy"]["full_bn"])
    say(f"murder on v9: ${murder_v9['low']:.3f}–{murder_v9['high']:.3f}bn ({murder_share['low']:.4f} of victims "
        f"${victims_v9['low']:.3f}–{victims_v9['high']:.3f}bn); legacy source-lane share {legacy_share:.4f}")

    # national volumes
    sales = mrts_2024()
    gm = gross_margin_2022()
    nr = nrss()
    cius = cius_2024()
    base = {"central": sales["Retail sales, total (excl. motor vehicle and parts dealers)"] / 1e3,
            "low": (sales["Retail sales, total (excl. motor vehicle and parts dealers)"] - sales["Gasoline stations"]
                    - sales["Nonstore retailers"]) / 1e3,
            "high": sales["Retail sales, total"] / 1e3}                      # $bn
    gate("NRF's $112.1bn is 1.6% of a base near MRTS total retail (NRF applies the rate to all retail sales)",
         abs(112.1 / 0.016 / base["high"] - 1) < 0.08, f"implied 2022 base ${112.1 / 0.016:.0f}bn vs 2024 total ${base['high']:.0f}bn")

    # NIBRS stores unknown age as age_id 103 with age_num '00' (and 104 'NS'): never age 0
    probe = S.age_band(pd.Series(["00", "NS", "05", "NN", "30"]), pd.Series([103, 104, 17, 1, 41]), 12).tolist()
    gate("age_band: '00' (id 103) and 'NS' (id 104) are unknown; id 1 is under 12",
         probe == ["unk", "unk", "lt12", "lt12", "12p"], str(probe))
    stages = load_stages()
    gate("staged victims: age_id 103 rows exist and none sits in the under-12 band",
         all(d["info"].get("violent_victims_unknown_age_in_lt12") == 0 for d in stages.values())
         and sum(d["info"].get("violent_victims_age_id_103", 0) for d in stages.values()) > 0,
         ", ".join(f"{st}-{yr} id103 {d['info'].get('violent_victims_age_id_103')}" for (st, yr), d in sorted(stages.items())))
    natpop = NR.national_pops()

    # ---- offender shares ----
    share_rows = []
    sh = {}
    for name, sp in specs():
        biz_all = shares(stages, sp, natpop, "pcells", BIZ)
        biz_retail = shares(stages, sp, natpop, "pcells", BIZ, retail=True)
        biz_non = shares(stages, sp, natpop, "pcells", BIZ, retail=False)
        ind = shares(stages, sp, natpop, "pcells", ("I",))
        sh[name] = dict(all=biz_all, retail=biz_retail, non=biz_non, ind=ind)
        for block, df in sh[name].items():
            for o, r in df.iterrows():
                share_rows.append(dict(spec=name, block=block, offence=o, **{k: float(v) for k, v in r.items()}))
    write_csv("offender_shares_by_spec.csv", share_rows)
    arr = arrestee_shares(stages, natpop)
    arr_rows = [dict(offence=o, **{k: float(v) for k, v in r.items()}) for o, r in arr.iterrows()]
    write_csv("arrestee_shares_business_incidents.csv", arr_rows)
    cen = sh["central"]
    say("\n-- National Hispanic share of offending, business victims at retail locations (central) --")
    say(cen["retail"][["weight_total", "RR_H_NHW", "RR_H_all", "local_share_H", "local_pop_share_H", "national_share_H",
                       "unknown_share"]].to_string(float_format=lambda x: f"{x:.3f}"))
    say("\n-- arrestees, business-victim incidents --")
    say(arr.to_string(float_format=lambda x: f"{x:.3f}"))

    # spec range for shoplifting (bounds kept apart)
    shop = {n: sh[n]["retail"].loc["shoplifting", "national_share_H"] for n, _ in specs()}
    shop["arrestees (business-victim incidents)"] = arr.loc["shoplifting", "national_share_H"]
    nonbound = {n: v for n, v in shop.items() if not n.startswith("bound")}
    S_shop = {"central": shop["central"], "low": min(nonbound.values()), "high": max(nonbound.values()),
              "bound_low": shop["bound: unknown offenders all non-Hispanic"], "bound_high": shop["bound: unknown offenders all Hispanic"]}
    say(f"\nshoplifting national Hispanic share: central {S_shop['central']:.4f}, specification range "
        f"{S_shop['low']:.4f}–{S_shop['high']:.4f}, allocation bounds {S_shop['bound_low']:.4f}–{S_shop['bound_high']:.4f}")
    nat = natpop["ncest2024_all"]
    pop_share_H = nat["H"] / sum(nat.values())
    say(f"national Hispanic population share, all ages (NC-EST2024): {pop_share_H:.4f}")

    # FBI national arrests: disconfirmation
    fbi = {(r["year"], r["offence"]): r for r in csv.DictReader((ARR / "arrests_by_ethnicity_2023.csv").open())}
    fbi_rows = []
    for o in ["larceny", "burglary", "mvt", "robbery", "fraud", "forgery", "embezzle", "vandalism", "ALL_OFFENCES"]:
        r = fbi["2024", o]
        fbi_rows.append(dict(offence=o, eth_panel_total=float(r["eth_panel_total"]), hispanic=float(r["eth_panel_hispanic"]),
                             hispanic_share=float(r["hispanic_share_of_eth_panel"])))
    write_csv("fbi_2024_adult_arrests_hispanic_share.csv", fbi_rows)

    # ---- A. retail external theft ----
    def lineage_share(S_H, end, ratio=1.0):
        return S_H * s_mex * ratio * (1 + k_end[end])

    def pass_through(L, S_lin):
        gross = L * S_lin * (1 - p_lin)
        reverse = L * p_lin * (1 - S_lin)
        if abs((gross - reverse) - L * (S_lin - p_lin)) > 1e-9:
            raise SystemExit("[BLOCKED] pass-through net is not losses x (offending share - payment share)")
        return gross, gross - reverse

    valuation = {"central": 1 - gm["excl_mv"], "low": 1 - gm["excl_mv"], "high": 1.0}
    ext_share = {"central": nr["external"], "low": nr["external"], "high": nr["external"] / (1 - nr["unknown"] - nr["other"])}
    rate = {"central": nr["shrink"], "low": nr["shrink_fy2021"], "high": nr["shrink"]}
    L_ext = {a: base[a] * rate[a] * ext_share[a] * valuation[a] for a in ("central", "low", "high")}
    # police-recorded floor: CIUS shoplifting x NIBRS mean net loss per business shoplifting victimisation (TX+AZ, all agencies)
    loss, n = 0.0, 0.0
    for (st, yr), d in stages.items():
        if st in ("TX", "AZ"):
            p = d["pcells"]
            p = p[p.group.eq("shoplifting") & p.vtype.isin(BIZ)]
            loss += p.loss_net.sum()
            n += p.n.sum()
    mean_shop = loss / n
    L_shop_police = cius["Shoplifting"] * mean_shop / 1e9
    say(f"\nA. retail external theft at cost: ${L_ext['central']:.2f}bn (low ${L_ext['low']:.2f}bn, high at retail "
        f"${L_ext['high']:.2f}bn); police-recorded shoplifting ${L_shop_police:.3f}bn "
        f"({cius['Shoplifting']:,} offences × ${mean_shop:.0f} mean net loss)")

    # ---- B. retail employee theft ----
    wages = [r for r in csv.DictReader(WAGES.open()) if r["sector"] == "retail" and r["allocation"] == "personal"]
    w_union = sum(float(r["union_wages_bn"]) for r in wages if r["cell"] != "441")
    w_nat = sum(float(r["national_wages_bn"]) for r in wages if r["cell"] != "441")
    worker = {"central": w_union / w_nat, "low": min(float(r["union_share"]) for r in wages if r["cell"] != "441"),
              "high": max(sh["central"]["retail"].loc["embezzlement", "national_share_H"] * s_mex,
                          max(float(r["union_share"]) for r in wages if r["cell"] != "441"))}
    say(f"B. union share of store-retail wages (excl. motor vehicle dealers): {worker['central']:.4f} "
        f"(cells {worker['low']:.4f}–{max(float(r['union_share']) for r in wages if r['cell'] != '441'):.4f}); "
        f"NIBRS retail embezzlement route {sh['central']['retail'].loc['embezzlement', 'national_share_H'] * s_mex:.4f}; "
        f"high arm takes the larger, {worker['high']:.4f}")
    L_int = {a: base[a] * rate[a] * (nr["internal"] if a != "high" else nr["internal"] / (1 - nr["unknown"] - nr["other"]))
             * valuation[a] for a in ("central", "low", "high")}

    # ---- C. police-recorded business losses outside retail stores ----
    def per_year_txaz(f):
        tot = {}
        for (st, yr), d in stages.items():
            if st in ("TX", "AZ"):
                for g, v in f(d).items():
                    tot[g] = tot.get(g, 0.0) + v / 2.0
        return tot
    biz_non_loss = per_year_txaz(lambda d: d["pcells"][d["pcells"].vtype.isin(BIZ) & ~d["pcells"].retail]
                                 .groupby("group").loss_net.sum().to_dict())
    all_n = per_year_txaz(lambda d: d["pcells"].groupby("group").n.sum().to_dict())
    srs_txaz = sum(all_n[g] for g in SRS)
    R_other = cius["Burglary"] + cius["Larceny total"] + cius["Motor vehicle theft"] + cius["Robbery"]
    R_other /= srs_txaz
    nonretail_rows = []
    for g in S.PGROUPS:
        if g not in biz_non_loss:
            continue
        scale = cius[SRS[g]] / all_n[g] if g in SRS else R_other
        Ln = biz_non_loss[g] * scale / 1e9
        S_H = sh["central"]["non"].loc[g, "national_share_H"] if g in sh["central"]["non"].index else sh["central"]["non"].loc["other_larceny", "national_share_H"]
        nonretail_rows.append(dict(group=g, txaz_per_year_loss_m=biz_non_loss[g] / 1e6, national_scale=scale,
                                   national_loss_bn=Ln, national_share_H=float(S_H)))
    write_csv("nonretail_business_losses.csv", nonretail_rows)
    L_non = sum(r["national_loss_bn"] for r in nonretail_rows)
    S_non = sum(r["national_loss_bn"] * r["national_share_H"] for r in nonretail_rows) / L_non
    say(f"C. police-recorded non-retail business losses, national: ${L_non:.3f}bn; loss-weighted Hispanic share {S_non:.4f}")
    cargo = per_year_txaz(lambda d: {"cargo": d["pcells"].loc[d["pcells"].cargo & d["pcells"].vtype.isin(BIZ), "loss_net"].sum()})
    say(f"   (cargo-flagged business losses in TX+AZ NIBRS: ${cargo['cargo'] / 1e6:.2f}M a year; CargoNet 2024 US+Canada $0.455bn)")

    # ---- D/E/F. individuals outside the NCVS frame ----
    prices = {r["offence"]: r for r in csv.DictReader(PRICES.open()) if r["price_set"] == "miller2021"}
    transfer = {r[""]: float(r["nibrs_transferred_national"]) / float(r["nibrs_txaz_observed"]) for r in csv.DictReader(TRANSFER.open())}

    def violent_line(cell, family_rule):
        """National Hispanic-offender victimisations outside the lineage, priced at Miller victim-only."""
        out = []
        sp = SPEC_CENTRAL
        Sh = shares(stages, sp, natpop, cell, None, key="offence")
        # victims of Hispanic offenders by victim class (TX+AZ, kept agencies), allocation a on offenders
        vH, tot = {}, {}
        for (st, yr), d in stages.items():
            if st not in sp["states"]:
                continue
            k = kept(d, sp)
            c = d[cell][d[cell].agency_id.isin(k.agency_id)]
            for r in c.itertuples(index=False):
                h = sum(getattr(r, x) for x in ["HW", "HB", "HO", "HU"])
                fam = getattr(r, "family", False)
                key = (r.offence, r.vclass, fam)
                vH[key] = vH.get(key, 0.0) + h
            # volume from every non-tribal agency (the coverage denominator below)
            ag = d["agencies"]
            call = d[cell][d[cell].agency_id.isin(ag.loc[ag.comp_source.ne("tribal_excluded"), "agency_id"])]
            for o, n_ in call.groupby("offence")[OFF_CLASSES].sum().sum(axis=1).items():
                tot[o] = tot.get(o, 0.0) + n_ / 2.0
        for o in Sh.index:
            if o == "Murder":
                continue                                       # every death is in the case's murder count
            cells_o = {k_: v for k_, v in vH.items() if k_[0] == o}
            known = sum(v for k_, v in cells_o.items() if k_[1] != "vU")
            if known <= 0:
                continue
            outside = 0.0
            for (_, vc, fam), v in cells_o.items():
                if vc == "vU":
                    continue
                if vc == "vH":
                    w = 0.0 if (family_rule and fam) else (1 - m_tract) * transfer.get(o, 1.0)
                else:
                    w = 1.0
                outside += v * w
            out_share = min(outside / known, 1.0)
            # national volume: TX+AZ police-recorded per-year victimisations scaled by the violent-crime ratio
            out.append(dict(offence=o, txaz_per_year=tot[o], national_share_H=float(Sh.loc[o, "national_share_H"]),
                            outside_share=out_share,
                            unit_tangible=float(prices[o]["victim_tangible"]), unit_full=float(prices[o]["full"])))
        return out
    txaz_violent = per_year_txaz(lambda d: {"v": 0.0})
    del txaz_violent
    t5 = xlsx_rows(CIUS / "CIUS_Table_5_Crime_in_the_United_States_by_State_2024.xlsx")
    state, totals5 = None, {}
    for r in t5:
        if r[0] and str(r[0]).strip().isupper():
            state = str(r[0]).strip()
        if r[1] and str(r[1]).strip() == "State Total" and state in ("TEXAS", "ARIZONA"):
            totals5[state] = dict(pop=r[3], violent=r[4], property=r[9])
    gate("CIUS 2024 Table 5 TX and AZ state totals", set(totals5) == {"TEXAS", "ARIZONA"},
         ", ".join(f"{k} {v['pop']:,}" for k, v in totals5.items()))
    # NIBRS coverage of TX+AZ population (all non-tribal agencies), to put police counts on state totals
    cov_pop = sum(d["agencies"].loc[d["agencies"].comp_source.ne("tribal_excluded"), "covered_pop"].sum() / 2.0
                  for (st, yr), d in stages.items() if st in ("TX", "AZ"))
    state_pop = totals5["TEXAS"]["pop"] + totals5["ARIZONA"]["pop"]
    viol_ratio = (cius["Population"] / state_pop) * \
        ((1221345 / cius["Population"]) / ((totals5["TEXAS"]["violent"] + totals5["ARIZONA"]["violent"]) / state_pop))
    say(f"TX+AZ NIBRS covered population {cov_pop / 1e6:.2f}M of {state_pop / 1e6:.2f}M; national/TX+AZ violent scale {viol_ratio:.3f}")

    def price_line(rows, label):
        res = []
        for r in rows:
            nat_v = r["txaz_per_year"] * (state_pop / cov_pop) * viol_ratio
            grp = nat_v * r["national_share_H"] * s_mex
            outv = grp * r["outside_share"] * (1 - recipient)
            res.append(dict(line=label, offence=r["offence"], national_victimisations=nat_v,
                            national_share_H=r["national_share_H"], outside_share=r["outside_share"],
                            union_outside_victimisations=outv, unit_tangible=r["unit_tangible"], unit_full=r["unit_full"],
                            union_tangible_bn=outv * r["unit_tangible"] / 1e9, union_full_bn=outv * r["unit_full"] / 1e9))
        return res
    # Shared table for within_group_harm_2026_10_21: police-recorded violent victimisations of children
    # under 12, by state-year, agency recording flag, offence, victim class and family/household
    # relationship, with raw (unallocated) fractional offender-class weights.
    child_rows = []
    for (st, yr), d in sorted(stages.items()):
        ag = d["agencies"]
        nontribal = set(ag.loc[ag.comp_source.ne("tribal_excluded"), "agency_id"])
        rec50 = set(ag.loc[ag.comp_source.ne("tribal_excluded") & ag.rec_rate.fillna(0).ge(0.5), "agency_id"])
        c = d["kcells"][d["kcells"].agency_id.isin(nontribal)].copy()
        c["agency_records_ethnicity_ge50"] = c.agency_id.isin(rec50)
        c = c.groupby(["agency_records_ethnicity_ge50", "offence", "vclass", "family"], observed=True)[OFF_CLASSES].sum()
        for (r50, o, vc, fam), r in c.iterrows():
            child_rows.append(dict(state=st, year=yr, agency_records_ethnicity_ge50=bool(r50), offence=o,
                                   victim_class=vc, family_or_household=bool(fam), victimisations=float(r.sum()),
                                   off_H=float(r[["HW", "HB", "HO", "HU"]].sum()), off_NHW=float(r["NHW"]),
                                   off_NHB=float(r["NHB"]), off_NHO=float(r["NHO"]), off_NH_race_unknown=float(r["NHU"]),
                                   off_eth_unknown=float(r[["UW", "UB", "UO", "UU"]].sum()), off_no_record=float(r["NOINFO"])))
    write_csv("nibrs_child_victims_under12_by_relationship_ethnicity.csv", child_rows)
    kids = price_line(violent_line("kcells", True), "D children under 12")
    jail = price_line(violent_line("jcells", False), "E jail/prison location 12+")
    write_csv("violent_outside_ncvs.csv", kids + jail)
    D = {"tangible": sum(r["union_tangible_bn"] for r in kids), "full": sum(r["union_full_bn"] for r in kids)}
    E = {"tangible": sum(r["union_tangible_bn"] for r in jail), "full": sum(r["union_full_bn"] for r in jail)}
    say(f"D. children under 12 (non-family, outside): union ${D['tangible']:.3f}bn tangible, ${D['full']:.3f}bn full")
    say(f"E. jail/prison location: union ${E['tangible']:.3f}bn tangible, ${E['full']:.3f}bn full")

    ind_loss = per_year_txaz(lambda d: d["pcells"][d["pcells"].vtype.eq("I")].groupby("group").loss_net.sum().to_dict())
    F_groups = ["fraud", "embezzlement", "extortion"]
    L_F = sum(ind_loss.get(g, 0.0) for g in F_groups) * R_other / 1e9
    S_F = sum(ind_loss.get(g, 0.0) * sh["central"]["ind"].loc[g, "national_share_H"] for g in F_groups) / \
        sum(ind_loss.get(g, 0.0) for g in F_groups)
    out_F = float(next(r for r in csv.DictReader((VL / "derived/outside_victim_shares_central.csv").open())
                       if r["block"] == "nonfatal_violence")["outside_share_estimate"])
    F_union = L_F * S_F * s_mex * out_F * (1 - recipient)
    say(f"F. fraud/embezzlement/extortion against individuals, police-recorded national ${L_F:.3f}bn; "
        f"Hispanic share {S_F:.4f}; union outside ${F_union:.3f}bn")

    # ---- assemble arms ----
    arms = []

    def add(line, arm, end, gross, net, note):
        arms.append(dict(line=line, arm=arm, end=end, gross_forward_bn=gross, net_with_without_bn=net, note=note))

    for end in ("low", "high"):
        # A
        for arm, L, S_H, note in [
            ("central", L_ext["central"], S_shop["central"], "NRSS 1.6% x 36% x store retail excl. motor vehicles, at cost; NIBRS shoplifting share"),
            ("low", L_ext["low"], S_shop["low"], "1.4% (FY2021) x store retail excl. motor vehicles, gas and nonstore, at cost; lowest specification share"),
            ("high", L_ext["high"], S_shop["high"], "1.6% x external incl. pro-rata unknown/other x all retail, at retail value; highest specification share"),
            ("bound_low", L_ext["central"], S_shop["bound_low"], "central volume; every unknown offender non-Hispanic"),
            ("bound_high", L_ext["central"], S_shop["bound_high"], "central volume; every unknown offender Hispanic"),
            ("police_recorded_floor", L_shop_police, S_shop["central"], "CIUS 2024 shoplifting x NIBRS mean net loss; police-recorded only"),
        ]:
            ratio = {"low": s_lo, "high": s_hi}.get(arm, 1.0)
            g_, n_ = pass_through(L, lineage_share(S_H, end, ratio))
            add("A retail external theft", arm, end, g_, n_, note + (f"; Mexican/Hispanic ratio {ratio:.3f}" if ratio != 1 else ""))
        # B
        for arm in ("central", "low", "high"):
            S_lin = worker[arm] * (1 + k_end[end])
            g_, n_ = pass_through(L_int[arm], S_lin)
            add("B retail employee theft", arm, end, g_, n_, f"union retail worker share {worker[arm]:.4f} [ASSUMPTION: theft per worker equal]")
        # C
        for arm, mult, ratio in (("central", 1.0, 1.0), ("low", 1.0, s_lo), ("high", 2.0, s_hi)):
            g_, n_ = pass_through(L_non * mult, lineage_share(S_non, end, ratio))
            add("C non-retail business, police-recorded", arm, end, g_, n_,
                f"unreported multiplier {mult:g} [ASSUMPTION]; Mexican/Hispanic ratio {ratio:.3f}")
        # D, E, F: direct victims, the victims row's boundary
        for line, val in (("D children under 12, non-family (full)", D["full"]), ("E jail/prison location (full)", E["full"]),
                          ("F fraud against individuals, police-recorded", F_union)):
            for arm, ratio in (("central", 1.0), ("low", s_lo), ("high", s_hi)):
                v = val * ratio * (1 + k_end[end])
                add(line, arm, end, v, v, f"union x (1 - recipient {recipient:.4f}) x (1 + crime key); Mexican/Hispanic ratio {ratio:.3f}")
    write_csv("arms.csv", arms)

    # beside, not in the totals: employees' pain and fear in commercial robberies (NCVS covers workers
    # robbed or threatened at work, so part of this is already in the victims row)
    com_rob = cius["Commercial house"] + cius["Gas or service station"] + cius["Convenience store"] + cius["Bank"]
    rob_beside = com_rob * cen["retail"].loc["robbery", "national_share_H"] * s_mex * (1 - worker["central"]) * \
        float(prices["Robbery"]["intangible"]) / 1e9
    say(f"beside: employees' intangible loss in {com_rob:,} commercial robberies, union outside ${rob_beside:.3f}bn "
        f"(overlaps the NCVS victims row; not added)")

    def tot(arm, end, col, lines=None):
        return sum(r[col] for r in arms if r["arm"] == arm and r["end"] == end and (lines is None or r["line"] in lines))
    lines_all = sorted({r["line"] for r in arms})
    summary = []
    for arm in ("central", "low", "high"):
        for end in ("low", "high"):
            g_, n_ = tot(arm, end, "gross_forward_bn"), tot(arm, end, "net_with_without_bn")
            summary.append(dict(arm=arm, end=end, gross_forward_bn=g_, net_with_without_bn=n_,
                                gross_per_member_usd=g_ * 1e3 / lineage_m, net_per_member_usd=n_ * 1e3 / lineage_m,
                                murder_v9_bn=murder_v9[end], gross_over_murder=g_ / murder_v9[end], net_over_murder=n_ / murder_v9[end]))
    write_csv("summary.csv", summary)
    say("\n-- arms ($bn, lineage, outside residents) --")
    say(pd.DataFrame(arms).pivot_table(index=["line"], columns=["arm", "end"], values="gross_forward_bn").round(3).to_string())
    say(pd.DataFrame(summary).round(3).to_string(index=False))
    gate("every arm line present at both ends", len(arms) == 2 * (6 + 3 + 3 + 9))

    # The pairing's row (sept24_propagation_2026_09_24 v9_totals.py) at full precision: the central arm's parts on the
    # union, so the pairing can move each part with the people v9 adds, and the all-offender national volumes its
    # equal-count average comparator needs.
    pairing = dict(rule="central arm; net = sum over A-C of L x (S - p) + D-F; S = S_union x (1 + k), p = payment_share_lineage",
                   payment_share_lineage=p_lin, payment_share_union=p_union, ends={},
                   national_all_offenders_bn=dict(
                       D=sum(r["national_victimisations"] * r["unit_full"] for r in kids) / 1e9,
                       E=sum(r["national_victimisations"] * r["unit_full"] for r in jail) / 1e9, F=L_F))
    for end in ("low", "high"):
        k = k_end[end]
        parts = {"A": (L_ext["central"], lineage_share(S_shop["central"], end) / (1 + k)),
                 "B": (L_int["central"], worker["central"]),
                 "C": (L_non, lineage_share(S_non, end) / (1 + k))}
        direct = {"D": D["full"], "E": E["full"], "F": F_union}
        net = sum(L * (S * (1 + k) - p_lin) for L, S in parts.values()) + sum(direct.values()) * (1 + k)
        gross = sum(L * S * (1 + k) * (1 - p_lin) for L, S in parts.values()) + sum(direct.values()) * (1 + k)
        gate(f"pairing parts reproduce the central totals ({end}, 1e-12)",
             abs(net - tot("central", end, "net_with_without_bn")) < 1e-12
             and abs(gross - tot("central", end, "gross_forward_bn")) < 1e-12, f"net {net:.9f}, gross {gross:.9f}")
        pairing["ends"][end] = dict(crime_key=k, net_bn=net, gross_bn=gross,
                                    loss_lines={x: dict(L_bn=L, S_union=S) for x, (L, S) in parts.items()},
                                    direct_union_bn=direct)

    # bounds row for A on the total
    extra = dict(pairing=pairing,murder_v9_low_bn=murder_v9["low"], murder_v9_high_bn=murder_v9["high"], murder_share_v9=murder_share["low"],
                 murder_share_source_lane_legacy=legacy_share, victims_v9_low_bn=victims_v9["low"], victims_v9_high_bn=victims_v9["high"],
                 lineage_m=lineage_m, payment_share_lineage=p_lin, payment_share_union=p_union, k_low=k_end["low"],
                 k_high=k_end["high"], recipient_share=recipient, s_mex=s_mex, shoplifting_share=S_shop,
                 national_hispanic_pop_share_all_ages=pop_share_H, L_ext_bn=L_ext, L_int_bn=L_int, L_non_bn=L_non,
                 L_shop_police_bn=L_shop_police, mean_shop_loss_usd=mean_shop, worker_share=worker, sales_base_bn=base,
                 gross_margin=gm, nrss=nr, lines=lines_all, commercial_robbery_employee_intangible_union_bn=rob_beside,
                 pins=pins["files"], read_value_pins=pins["read_values"])
    (OUT / "summary.json").write_text(json.dumps(extra, indent=1, sort_keys=True, default=float) + "\n")
    (OUT / "run_log.txt").write_text("\n".join(LOG) + "\n")


if __name__ == "__main__":
    main()
