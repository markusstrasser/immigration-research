"""Build the New York City school x year panel for the newcomer school-shock lane.

Inputs are the pinned raw pulls in `_cache/nyc/` (URLs in reads/nyc_sources.md):
  - NYC Public Schools demographic snapshots (school tab): 2017-18 to 2021-22 (NYC Open Data c7ru-d68s) and
    2021-22 to 2025-26 (InfoHub). Enrollment is the October 31 audited register; ELL status is as of June 30.
  - School grade 3-8 ELA and math results 2018-2026 (InfoHub), tabs "All" and "ELL".
  - School Allocation Memorandum tables: Project Open Arms (FY2023 SAM 65, Wayback copy of 2022-10-31),
    Increase in Students in Temporary Housing (FY2024 SAM 90), Title III Immigrant (FY2023-FY2026) and
    register relief (FY2022, FY2023, FY2025, FY2026).
  - NYSED report-card tables (Expenditures per Pupil; teacher counts), extracted by extract_nysed_src.py.
  - Class size reports (K-5 or K-8 average class size by grade and program; pupil-teacher ratio).

Outputs (derived/): nyc_schools.csv, nyc_scores.csv, nyc_class_size.csv, nyc_audit.json.
`year` is the spring of the school year (2023 = 2022-23) and FY2023 is school year 2022-23.

Run from the repository root:
    OPENBLAS_NUM_THREADS=1 uv run --no-project python3 \
        infra/immigration-fiscal/newcomer_school_shock_2026_09_28/build_nyc.py
"""

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

LANE = Path(__file__).resolve().parent
CACHE = LANE / "_cache" / "nyc"
OUT = LANE / "derived"

PINS = {
    "demographic-snapshot-2021-22-to-2025-26-public.xlsx": "d7759afaf3e1cd821a512f84a23fae6b70aa818f116ee2fc56e18a7bcfcb69f2",
    "demographic-snapshot-2017-18-to-2021-22-opendata-c7ru-d68s.csv": "26106be957979e78c953090cc4d1904e8684e6c7099002bfcb70f14b57975545",
    "school-ela-results-public.xlsx": "5a8c419e64091c65b06a01d2694d9bd11bcd45f424fe2756ffa47b8f02b4ce18",
    "school-math-results-public.xlsx": "fc11957633937e324a7456ca80cce1b075ea379d85c4c7153d9c477f6f255a6a",
    "sams/FY2022_SAM086_T01.xlsx": "a4480447c9dad47f40c604100e357a8f7a34d6b43e8210387fc3f40888cd224a",
    "sams/FY2023_SAM065_T01_wayback20221031.xlsx": "76f3379e06f596847766f51bd2fe40007b3b1db70d0a46a2f48f7b1e6fa86112",
    "sams/FY2023_SAM084_T01.xlsx": "2de1e4243518580067772bc6925a7bca5a60d150e8f4847f2b57ce3213438c9b",
    "sams/FY2023_SAM085_T01.xlsx": "25e3eb7af2568bd4ede29594dc975242834feeab8db6d0c97e0e40829c0e39b1",
    "sams/FY2024_SAM073_T01.xlsx": "e9310e53aeb77422e347a3aceb6c90be9cf05ae7d8be8beddff917cfced713b0",
    "sams/FY2024_SAM090_T01.xlsx": "8bac2860d2a807d3bc8f5977e9f4983bb68a35174de0596dd4442d26ff96cacb",
    "sams/FY2025_SAM081_t01.xlsx": "cec59bc9912d73d322f4fa65517c956955698f276a3869a6dfc3a7045f2ccc75",
    "sams/FY2025_SAM086_T01.xlsx": "9ec1e8ff8e0b30615d9685d999681f0035b7bef6074426ecdeef3b6d191cd2d0",
    "sams/FY2026_SAM073_T01.xlsx": "1c36ac6f03fff8c6c6fc6792e1f984e82541b689ec18d88adf4de684c18244d4",
    "sams/FY2026_SAM085_T01.xlsx": "f6d32910f8f486a1ddfa37c044208e83fa52e30e82165bba1662b36800005e44",
    "class_size/opendata-2017-18-school-class-size-2ic6-4hrf.csv": "b2c56191ce76f4bc4c84b61352c72335d22e25d72ad181d980af36e001eed1ff",
    "class_size/november2018_avg_classsize_school.xlsx": "0c15e64a8f8d5a110a88f6638d8d6486425ae767e98b32bae9b94180be197dce",
    "class_size/february2019_avg_classsize_school.xlsx": "017e3b3f19b6e6a065959bb3940ea0ebec093aa63bdc56e56e092f5130b10c97",
    "class_size/november2019_avg_classsize_school.xlsx": "50da078b3bb79b0ca4aed2948ca0626b322cd1537b1121d55dd4a633a49dc5dd",
    "class_size/february2020_avg_classsize_school.xlsx": "c73761ee685620adfb80b2c89d315ab6a24c8dbf0ca583994a666260ee9ca840",
    "class_size/opendata-2021-22-avg-class-size-school-sgr7-hhwp.csv": "c5ce5dad9f96e5730b2ec28371e9189d1dd7875b45fd3d45707df99001d44ec7",
    "class_size/prelim2022_avg_classsize_school.xlsx": "63cd47c7d935477719abdc993cbf2ef7c45876b4ce622372c7f5b7d8b192ee26",
    "class_size/updated2023_avg_classsize_schl.xlsx": "c10123ea906b80bbb1b8b9da6be5df92a4c9d157556a048d70d43470b2f8cf13",
    "class_size/preliminary-2023-24-class-size---school.xlsx": "011794a58d871724173c1d1008840bfb60584952b3a48831667aa62c428523e1",
    "class_size/updated-2023-24-class-size---school.xlsx": "b53bfd1fe63691828323493ab116acc1f9e163e3e85b583ce3524a6172505e9b",
    "class_size/june-2023-24-class-size---school.xlsx": "0eddb8c4a31d2e16974e5597224fd402a9ccd58c6d12e18ff5a3e0f127efd168",
    "class_size/november-2024-25-class-size---school.xlsx": "e7e6701cad3d2d29b6d39a9168e933e720015985625ae895e32595173d046a43",
    "class_size/february-2024-25-class-size---school.xlsx": "679368aee33aeb78f98b1a4534d97a24d313c57d6cf2bc171599b6e57f89f092",
    "class_size/june-2024-25-class-size---school.xlsx": "d4ffe75ef5ea284811948418a2998beec1834452530ad36dddcfb58e467ddd0c",
    "class_size/november-2025-26-class-size---school.xlsx": "f173073e07ce6bac0edee4b7338d980e31f1f81d9cdcf9d99a3e07b11a387687",
    "class_size/february-2025-26-class-size---school.xlsx": "5dd906c727cfa940fb1c4538469bf570c02b7e580d9caedd00a2039e985dbd4a",
    "class_size/june-2025-26-class-size---school.xlsx": "73bb22598a60a57e7dfdf17c5885baa8a84829606c12ce2528b165827afadd0f",
    "nysed_src/tables/SRC2019_expenditures.csv": "1a13986b885e0f89d54253dc04fa0cef53e31a4140a5e2c5a820667a039a1699",
    "nysed_src/tables/SRC2019_teachers.csv": "fb83fa06bbc26dd231f240f05f99eae3d0016f816e0fdf88bf6f0cad3bf9890f",
    "nysed_src/tables/SRC2021_expenditures.csv": "dfd227087eccf47c14f9a46936069eb300fac352e273eab843538acb53562827",
    "nysed_src/tables/SRC2021_teachers.csv": "e4a62f9efbf1af5fa62ebd80dcbdadb785889cc137769ca73a031b5e69eec8a7",
    "nysed_src/tables/SRC2022_expenditures.csv": "a7bf0fcb7f93851481a6b12c63b396ce2248a1163a5814cc2ea224d02cf68b50",
    "nysed_src/tables/SRC2022_teachers.csv": "28652712d28dbbdb5668a7c98095aa51fd69190dfd4535225a26d24156fbc457",
    "nysed_src/tables/SRC2023_expenditures.csv": "68de553a28fe386a9a52c3240e87ef1a153b5efd6b67959b0774d125ec077dcd",
    "nysed_src/tables/SRC2023_teachers.csv": "8702cffd00713d7e1311f0666a6af99dca5339a3d2ded5b6ccaaa7fda5cf3f33",
    "nysed_src/tables/SRC2024_expenditures.csv": "991d5970b5b98e36be21d7bbb611b7b9476efb125fb6252073f664986ce5fedb",
    "nysed_src/tables/SRC2024_teachers.csv": "696093eba85c2b99e3f4b1383b7b7419c48fd05e4741e0487ddbfe9e1b59a523",
    "nysed_src/tables/SRC2025_expenditures.csv": "131dfbdcd2fc6a90703286646ca571acfeeecddd9fe6742f38d9b43ac4d89e90",
    "nysed_src/tables/SRC2025_teachers.csv": "c934d62ce631218c93b2fad5edc9c99b914889841a599d77612f9a96d2ef7da7",
}

DBN_RE = re.compile(r"^\d\d[MXKQR]\d\d\d$")
BORO = {"31": "M", "32": "X", "33": "K", "34": "Q", "35": "R"}
YEARS = range(2018, 2027)

# class-size reports: (file, spring year, report timing)
CLASS_SIZE = [
    ("class_size/opendata-2017-18-school-class-size-2ic6-4hrf.csv", 2018, "opendata"),
    ("class_size/november2018_avg_classsize_school.xlsx", 2019, "nov"),
    ("class_size/february2019_avg_classsize_school.xlsx", 2019, "feb"),
    ("class_size/november2019_avg_classsize_school.xlsx", 2020, "nov"),
    ("class_size/february2020_avg_classsize_school.xlsx", 2020, "feb"),
    ("class_size/opendata-2021-22-avg-class-size-school-sgr7-hhwp.csv", 2022, "opendata"),
    ("class_size/prelim2022_avg_classsize_school.xlsx", 2023, "nov"),
    ("class_size/updated2023_avg_classsize_schl.xlsx", 2023, "feb"),
    ("class_size/preliminary-2023-24-class-size---school.xlsx", 2024, "nov"),
    ("class_size/updated-2023-24-class-size---school.xlsx", 2024, "feb"),
    ("class_size/june-2023-24-class-size---school.xlsx", 2024, "jun"),
    ("class_size/november-2024-25-class-size---school.xlsx", 2025, "nov"),
    ("class_size/february-2024-25-class-size---school.xlsx", 2025, "feb"),
    ("class_size/june-2024-25-class-size---school.xlsx", 2025, "jun"),
    ("class_size/november-2025-26-class-size---school.xlsx", 2026, "nov"),
    ("class_size/february-2025-26-class-size---school.xlsx", 2026, "feb"),
    ("class_size/june-2025-26-class-size---school.xlsx", 2026, "jun"),
]
K5_GRADES = {"K", "0K", "1", "2", "3", "4", "5"}
REG_PROGRAMS = {"Gen Ed", "ICT", "G&T", "ICT & G&T"}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def verify():
    for rel, want in PINS.items():
        p = CACHE / rel
        if not p.exists():
            raise SystemExit(f"[BLOCKED] missing source {p}")
        got = sha256(p)
        if got != want:
            raise SystemExit(f"[BLOCKED] stale source {rel}: {got} != {want}")


def num(v):
    """Published cell -> float or None ('s', blanks and '<15' style markers are None)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    if s.endswith("%"):
        try:
            return float(s[:-1]) / 100.0
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, int) or (isinstance(v, float) and v.is_integer() and abs(v) < 1e15):
        return str(int(v))
    return f"{v:.6f}".rstrip("0").rstrip(".")


def sheet_rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        # some class-size workbooks carry title rows above the header; the header is the row that starts "DBN"
        for n, first in enumerate(it):
            if first and first[0] is not None and str(first[0]).strip() in ("DBN", "Grade", "Year", "Borough"):
                break
            if n > 15:
                raise SystemExit(f"[BLOCKED] {path.name}/{sheet}: no header row in the first 16 rows")
        header = [str(h).strip() if h is not None else "" for h in first]
        rows = [dict(zip(header, r)) for r in it if any(c is not None for c in r)]
    finally:
        wb.close()
    return header, rows


def spring(label):
    """'2022-23' -> 2023."""
    m = re.fullmatch(r"(\d{4})-(\d{2})", str(label).strip())
    if not m:
        raise SystemExit(f"[BLOCKED] unexpected school-year label {label!r}")
    return int(m.group(1)) + 1


def load_demographics(audit):
    schools = {}
    header, rows = sheet_rows(CACHE / "demographic-snapshot-2021-22-to-2025-26-public.xlsx", "School")
    for need in ("DBN", "Year", "Total Enrollment", "# English Language Learners", "# Students with Disabilities",
                 "# Poverty", "Economic Need Index"):
        if need not in header:
            raise SystemExit(f"[BLOCKED] demographic snapshot lacks {need}")
    new = {}
    for r in rows:
        new[(r["DBN"], spring(r["Year"]))] = r
    with open(CACHE / "demographic-snapshot-2017-18-to-2021-22-opendata-c7ru-d68s.csv", newline="") as f:
        old = {(r["DBN"], spring(r["Year"])): r for r in csv.DictReader(f)}
    # 2021-22 appears in both releases; keep the newer one and record agreement
    both = [k for k in old if k[1] == 2022 and k in new]
    agree = sum(1 for k in both if num(old[k]["# English Language Learners"]) == num(new[k]["# English Language Learners"])
                and num(old[k]["Total Enrollment"]) == num(new[k]["Total Enrollment"]))
    audit["demographics_2022_overlap"] = {"schools_in_both": len(both), "enrollment_and_ell_identical": agree}
    for src, table in (("opendata_c7ru", old), ("infohub_2021_26", new)):
        for (dbn, yr), r in table.items():
            if src == "opendata_c7ru" and (dbn, yr) in new:
                continue
            g = {f"enroll_g{k}": num(r.get(f"Grade {k}")) for k in range(3, 9)}
            schools[(dbn, yr)] = {
                "school_name": str(r.get("School Name") or "").strip(),
                "enroll_total": num(r["Total Enrollment"]),
                **g,
                "ell_n": num(r["# English Language Learners"]),
                "swd_n": num(r["# Students with Disabilities"]),
                "poverty_n": num(r["# Poverty"]),
                "eni": num(r["Economic Need Index"]),
                "demog_source": src,
            }
    return schools


def load_scores(audit):
    out = []
    for subj, fname, tabs in (("ela", "school-ela-results-public.xlsx", ("ELA - All", "ELA - ELL")),
                              ("math", "school-math-results-public.xlsx", ("Math - All", "Math - ELL"))):
        for tab in tabs:
            header, rows = sheet_rows(CACHE / fname, tab)
            for need in ("DBN", "Grade", "Year", "Category", "Number Tested", "Mean Scale Score", "# Level 3+4"):
                if need not in header:
                    raise SystemExit(f"[BLOCKED] {fname}/{tab} lacks {need}")
            for r in rows:
                dbn = str(r["DBN"]).strip()
                if not DBN_RE.match(dbn):
                    continue
                grade = str(r["Grade"]).strip()
                grade = "3-8" if grade == "All Grades" else grade
                mss = num(r["Mean Scale Score"])
                out.append({
                    "dbn": dbn, "year": int(r["Year"]), "subject": subj, "grade": grade,
                    "category": str(r["Category"]).strip(),
                    "n_tested": num(r["Number Tested"]), "mean_scale_score": mss,
                    "n_level34": num(r["# Level 3+4"]), "pct_level34": num(r["% Level 3+4"]),
                    "suppressed": int(str(r["Mean Scale Score"]).strip() == "s"),
                })
    audit["scores_rows"] = len(out)
    audit["scores_by_category"] = dict(sorted(Counter(r["category"] for r in out).items()))
    audit["scores_suppressed_by_category"] = dict(sorted(Counter(r["category"] for r in out if r["suppressed"]).items()))
    return out


def sam_table(rel, amount_col, dbn_col=2):
    wb = openpyxl.load_workbook(CACHE / rel, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        vals = {}
        total = None
        for r in ws.iter_rows(values_only=True):
            cells = list(r)
            if len(cells) <= amount_col:
                continue
            key = str(cells[dbn_col]).strip() if cells[dbn_col] is not None else ""
            if DBN_RE.match(key):
                v = num(cells[amount_col])
                if v is None:
                    continue
                if key in vals:
                    raise SystemExit(f"[BLOCKED] duplicate DBN {key} in {rel}")
                vals[key] = v
            elif any(str(c).strip().lower() == "total" for c in cells if c is not None):
                total = num(cells[amount_col])
    finally:
        wb.close()
    s = sum(vals.values())
    if total is None or abs(s - total) > 1.0:
        raise SystemExit(f"[BLOCKED] {rel}: school rows sum to {s}, table total {total}")
    return vals, total


def load_sams(audit):
    spec = {
        # name: (file, amount column index, spring year)
        "sam65_open_arms_alloc": ("sams/FY2023_SAM065_T01_wayback20221031.xlsx", 4, 2023),
        "sam90_sth_increase_alloc": ("sams/FY2024_SAM090_T01.xlsx", 4, 2024),
    }
    out = defaultdict(dict)
    for name in ("sam65_open_arms_alloc", "sam90_sth_increase_alloc"):
        rel, col, yr = spec[name]
        vals, total = sam_table(rel, col)
        for dbn, v in vals.items():
            out[(dbn, yr)][name] = v
        audit[f"{name}_{yr}"] = {"schools": len(vals), "total": total}
    for rel, col, yr in (("sams/FY2023_SAM084_T01.xlsx", 4, 2023), ("sams/FY2024_SAM073_T01.xlsx", 4, 2024),
                         ("sams/FY2025_SAM081_t01.xlsx", 4, 2025), ("sams/FY2026_SAM073_T01.xlsx", 4, 2026)):
        vals, total = sam_table(rel, col)
        for dbn, v in vals.items():
            out[(dbn, yr)]["title3_immigrant_alloc"] = v
        audit[f"title3_immigrant_{yr}"] = {"schools": len(vals), "total": total}
    for rel, col, yr in (("sams/FY2022_SAM086_T01.xlsx", 4, 2022), ("sams/FY2023_SAM085_T01.xlsx", 4, 2023),
                         ("sams/FY2025_SAM086_T01.xlsx", 4, 2025), ("sams/FY2026_SAM085_T01.xlsx", 4, 2026)):
        vals, total = sam_table(rel, col)
        for dbn, v in vals.items():
            out[(dbn, yr)]["register_relief_alloc"] = v
        audit[f"register_relief_{yr}"] = {"schools": len(vals), "total": total}
    return out


def entity_to_dbn(ent):
    """NYSED ENTITY_CD 'CC DD 0001 T SSS' -> DBN 'DD' + borough + 'SSS'.

    CC is the county (31-35), DD the community district, 0001 marks a DOE school (charters carry 0086 and are
    dropped), T is 0 for elementary/middle and 1 for high schools, SSS the school number.
    """
    if len(ent) != 12 or ent[:2] not in BORO or ent[4:8] != "0001" or ent[9:] == "000":
        return None
    return ent[2:4] + BORO[ent[:2]] + ent[9:12]


def load_nysed(audit):
    """Each school year from the latest release that carries it."""
    out = defaultdict(dict)
    chosen = {}
    for kind in ("expenditures", "teachers"):
        by_year = {}
        for rel in sorted(k for k in PINS if k.endswith(f"_{kind}.csv")):
            release = int(re.search(r"SRC(\d{4})", rel).group(1))
            with open(CACHE / rel, newline="") as f:
                for r in csv.DictReader(f):
                    dbn = entity_to_dbn(r["ENTITY_CD"])
                    if dbn is None:
                        continue
                    yr = int(r["YEAR"])
                    prev = by_year.get((dbn, yr))
                    if prev is not None and prev[1]["ENTITY_CD"] != r["ENTITY_CD"]:
                        raise SystemExit(f"[BLOCKED] {dbn} {yr}: two entity codes {prev[1]['ENTITY_CD']} {r['ENTITY_CD']}")
                    if prev is None or prev[0] < release:
                        by_year[(dbn, yr)] = (release, r)
        for (dbn, yr), (release, r) in by_year.items():
            chosen[(kind, yr)] = max(chosen.get((kind, yr), 0), release)
            if kind == "expenditures":
                out[(dbn, yr)].update({
                    "nysed_pupils": num(r["PUPIL_COUNT_TOT"]),
                    "exp_federal": num(r["FEDERAL_EXP"]), "exp_state_local": num(r["STATE_LOCAL_EXP"]),
                    "exp_total": num(r["FED_STATE_LOCAL_EXP"]), "ppe_total": num(r["PER_FED_STATE_LOCAL_EXP"]),
                    "ppe_federal": num(r["PER_FEDERAL_EXP"]), "ppe_state_local": num(r["PER_STATE_LOCAL_EXP"]),
                    "exp_release": release,
                })
            else:
                out[(dbn, yr)].update({"num_teach": num(r["NUM_TEACH"]),
                                       "num_teach_inexp": num(r.get("NUM_TEACH_INEXP")),
                                       "teach_release": release})
    audit["nysed_release_by_year"] = {f"{k}_{y}": v for (k, y), v in sorted(chosen.items())}
    return out


def class_size_rows(rel):
    """Yield (dbn, grade_level, program, students, classes) for K-5 regular classes and the PTR map."""
    p = CACHE / rel
    k5, ptr = [], {}
    if rel.endswith(".csv"):
        with open(p, newline="") as f:
            rows = list(csv.DictReader(f))
    else:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        names = wb.sheetnames
        wb.close()
        sheet = next((s for s in names if s.startswith("K-8 Avg") or s.startswith("K-5 Average")), None)
        if sheet is None:
            raise SystemExit(f"[BLOCKED] {rel}: no K-8/K-5 average tab in {names}")
        _, rows = sheet_rows(p, sheet)
        if "PTR" in names:
            _, prow = sheet_rows(p, "PTR")
            for r in prow:
                dbn = str(r.get("DBN") or "").strip()
                if DBN_RE.match(dbn):
                    ptr[dbn] = num(r.get("School Pupil-Teacher Ratio"))
    for r in rows:
        dbn = str(r.get("DBN") or "").strip()
        grade = str(r.get("Grade Level") or "").strip()
        grade = grade.lstrip("0") if grade.isdigit() else grade  # 2024-25 on: '01'..'05'
        prog = str(r.get("Program Type") or "").strip()
        if not DBN_RE.match(dbn) or grade not in K5_GRADES or prog not in REG_PROGRAMS:
            continue
        st, cl = num(r.get("Number of Students")), num(r.get("Number of Classes"))
        if st is None or cl is None or cl <= 0:
            continue
        k5.append((dbn, grade, prog, st, cl))
    return k5, ptr


def load_class_size(audit):
    per = []
    progs = Counter()
    for rel, yr, timing in CLASS_SIZE:
        k5, ptr = class_size_rows(rel)
        agg = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
        for dbn, grade, prog, st, cl in k5:
            progs[prog] += 1
            a = agg[dbn]
            a[0] += st
            a[1] += cl
            if grade in {"3", "4", "5"}:
                a[2] += st
                a[3] += cl
        for dbn in sorted(set(agg) | set(ptr)):
            a = agg.get(dbn, [0.0, 0.0, 0.0, 0.0])
            per.append({"dbn": dbn, "year": yr, "timing": timing, "source_file": rel,
                        "k5_students": a[0] or None, "k5_classes": a[1] or None,
                        "k5_avg_class": (a[0] / a[1]) if a[1] else None,
                        "g35_avg_class": (a[2] / a[3]) if a[3] else None,
                        "ptr": ptr.get(dbn)})
        audit.setdefault("class_size_files", {})[rel] = {"k5_rows": len(k5), "schools_k5": len(agg),
                                                         "schools_ptr": len(ptr)}
    audit["class_size_programs"] = dict(sorted(progs.items()))
    return per


def write_csv(path, header, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([fmt(r.get(h)) for h in header])


def main():
    verify()
    audit = {"inputs": dict(sorted(PINS.items()))}
    demo = load_demographics(audit)
    scores = load_scores(audit)
    sams = load_sams(audit)
    nysed = load_nysed(audit)
    cls = load_class_size(audit)

    feb = {(r["dbn"], r["year"]): r for r in cls if r["timing"] in ("feb", "opendata")}
    keys = sorted(set(k for k in demo if k[1] in YEARS or k[1] == 2018) | set(sams) |
                  set(k for k in nysed if k[1] in YEARS))
    rows = []
    for dbn, yr in keys:
        if yr < 2018 or yr > 2026:
            continue
        d, s, n, c = demo.get((dbn, yr), {}), sams.get((dbn, yr), {}), nysed.get((dbn, yr), {}), feb.get((dbn, yr), {})
        rows.append({"dbn": dbn, "district": int(dbn[:2]), "boro": dbn[2], "year": yr, **d, **s, **n,
                     "k5_avg_class": c.get("k5_avg_class"), "g35_avg_class": c.get("g35_avg_class"),
                     "ptr": c.get("ptr"), "class_size_timing": c.get("timing")})
    # join rates, the gate that the NYSED entity-code mapping is right
    demo_dbns = {k[0] for k in demo}
    for label, keyset in (("nysed", set(nysed)), ("sams", set(sams))):
        dbns = {k[0] for k in keyset}
        audit[f"{label}_dbns_matched_to_snapshot"] = {"dbns": len(dbns), "matched": len(dbns & demo_dbns)}
    m = audit["nysed_dbns_matched_to_snapshot"]
    if m["matched"] < 0.9 * m["dbns"]:
        raise SystemExit(f"[BLOCKED] NYSED entity codes match only {m} snapshot DBNs")

    header = ["dbn", "school_name", "district", "boro", "year", "demog_source", "enroll_total",
              "enroll_g3", "enroll_g4", "enroll_g5", "enroll_g6", "enroll_g7", "enroll_g8", "ell_n", "swd_n",
              "poverty_n", "eni", "sam65_open_arms_alloc", "sam90_sth_increase_alloc", "title3_immigrant_alloc",
              "register_relief_alloc", "nysed_pupils", "exp_total", "exp_federal", "exp_state_local", "ppe_total",
              "ppe_federal", "ppe_state_local", "exp_release", "num_teach", "num_teach_inexp", "teach_release",
              "k5_avg_class", "g35_avg_class", "ptr", "class_size_timing"]
    write_csv(OUT / "nyc_schools.csv", header, rows)
    scores.sort(key=lambda r: (r["dbn"], r["year"], r["subject"], r["grade"], r["category"]))
    write_csv(OUT / "nyc_scores.csv", ["dbn", "year", "subject", "grade", "category", "n_tested",
                                       "mean_scale_score", "n_level34", "pct_level34", "suppressed"], scores)
    cls.sort(key=lambda r: (r["dbn"], r["year"], r["timing"], r["source_file"]))
    write_csv(OUT / "nyc_class_size.csv", ["dbn", "year", "timing", "source_file", "k5_students", "k5_classes",
                                           "k5_avg_class", "g35_avg_class", "ptr"], cls)
    audit["schools_rows_by_year"] = dict(sorted(Counter(str(r["year"]) for r in rows).items()))
    (OUT / "nyc_audit.json").write_text(json.dumps(audit, indent=1, sort_keys=True) + "\n")
    print(f"nyc_schools {len(rows)} rows; nyc_scores {len(scores)}; nyc_class_size {len(cls)}")


if __name__ == "__main__":
    main()
