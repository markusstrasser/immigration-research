"""Parse CMS's CY2019-2023 Medicaid LTSS tables by race/ethnicity and language (T-MSIS TAF).

Source: Mathematica for CMS, "Medicaid Long Term Services and Supports Annual Expenditures and
Users: Calendar Year <year> TAF Data", workbooks C1 (users) and C2 (expenditures). 2023 comes from
the zip pinned by health_admin_2026_09_20; 2019-2022 from Wayback copies (fetch_sources.py,
SOURCE_PINS.json). Hashes are checked. Reads every "ByRaceEthn", "ByLanguage" and "ByAge" sheet, national
and state rows, plus the DQ Measures sheet; nothing is recomputed. "DS" (suppressed) and "NC"
(not calculated) cells become blank with a flag.
Writes derived/taf_race_ethn.csv, derived/taf_language.csv, derived/taf_age.csv, derived/taf_dq.csv.
Run from the repo root:
  OPENBLAS_NUM_THREADS=1 uv run --no-project --with openpyxl python3 \
      infra/immigration-fiscal/ltss_share_2026_09_23/taf_tables.py
"""
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
ZIP = HERE.parent / "health_admin_2026_09_20/raw/ltss2023_tables.zip"
ZIP_SHA = "2e817d2e94bf50aa94a1673d82371dbb0690d1136db95578a096f90d8b09d367"
WAYBACK = HERE / "_cache/wayback"
SOURCES = {2023: ZIP, 2022: WAYBACK / "ltss-expenditures-user-data-2022.zip",
           2021: WAYBACK / "ltss-expenditures-user-data-2019-2021.zip",
           2020: WAYBACK / "ltss-expenditures-user-data-2019-2021.zip",
           2019: WAYBACK / "ltss-expenditures-user-data-2019-2021.zip"}
GROUPS = {"Asian and Pacific Islander, non-Hispanic": "api_nh",
          "American Indian and Alaska Native, non-Hispanic": "aian_nh", "Black, non-Hispanic": "black_nh",
          "White, non-Hispanic": "white_nh", "Multiracial, non-Hispanic": "multi_nh",
          "Hispanic, any race": "hispanic", "Race/ethnicity unknown": "unknown",
          "Race and ethnicity Unknown": "unknown",
          "English": "english", "Spanish": "spanish", "Other language": "other_language",
          "Age group 0-20": "age_0_20", "Age group 21-44": "age_21_44", "Age group 45-64": "age_45_64",
          "Age group 65 and older": "age_65plus", "Age group unknown": "age_unknown"}


def sheets(book, suffix):
    for name in book.sheetnames:
        m = re.match(r"C\.(\d)\.(\d+) (\w+?)By" + suffix + r"$", name)
        if m:
            yield name, m.group(3)


def parse(book, name, category, measure):
    rows = list(book[name].values)
    top = [i for i, r in enumerate(rows[:3]) if r and r[0] == "State"]  # 2023 has a title row
    if len(top) != 1:
        raise ValueError(f"{name}: unexpected header")
    header = rows[top[0]]
    out = []
    for r in rows[top[0] + 1:]:
        if r[0] is None or any("end of worksheet" in str(v).lower() for v in r if v is not None):
            continue
        total = r[1]
        for j in range(2, len(header), 2):
            label = re.sub(r" \(total\)$", "", header[j])
            code = GROUPS.get(label) or (re.sub(r"\W+", "_", "age_" + label[10:].replace(" and older", "plus"))
                                         if label.startswith("Age group ") else None)
            if code is None or not header[j + 1].startswith(label + " (% of"):
                raise ValueError(f"{name}: column {header[j]!r}")
            out.append(dict(sheet=name, category=category, measure=measure, state=r[0], total=total,
                            group=code, value=r[j], pct=r[j + 1]))
    return out


def main():
    with ZIP.open("rb") as f:
        if hashlib.file_digest(f, "sha256").hexdigest() != ZIP_SHA:
            raise SystemExit("[BLOCKED] pinned LTSS workbook changed")
    pins = json.loads((HERE / "SOURCE_PINS.json").read_text())
    for path in set(SOURCES.values()) - {ZIP}:
        if hashlib.sha256(path.read_bytes()).hexdigest() != pins[path.name]["sha256"]:
            raise SystemExit(f"[BLOCKED] {path.name} changed")
    race, lang, ages, dqs = [], [], [], []
    for year, path in SOURCES.items():
        z = zipfile.ZipFile(path)
        for member, measure in [(f"C1_LTSSUsrChar_{year}.xlsx", "users"), (f"C2_LTSSExpChar_{year}.xlsx", "expenditures")]:
            book = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True, data_only=True)
            for name, cat in sheets(book, "RaceEthn"):
                race += [dict(year=year, **r) for r in parse(book, name, cat, measure)]
            for name, cat in sheets(book, "Language"):
                lang += [dict(year=year, **r) for r in parse(book, name, cat, measure)]
            for name, cat in sheets(book, "Age"):
                ages += [dict(year=year, **r) for r in parse(book, name, cat, measure)]
            if measure == "expenditures" and "DQ Measures" in book.sheetnames:
                rows = [r[:5] for r in list(book["DQ Measures"].values)[2:]]
                dq = pd.DataFrame(rows, columns=["state", "inst_ffs", "inst_mc", "hcbs_ffs", "hcbs_mc"])
                dqs.append(dq.dropna(subset=["inst_ffs"]).assign(year=year))
    dq = pd.concat(dqs, ignore_index=True)
    race, lang, ages = pd.DataFrame(race), pd.DataFrame(lang), pd.DataFrame(ages)
    for frame in (race, lang, ages):
        # "DS" = data suppressed (small user counts); "NC" = not calculated (whole state-years).
        frame["suppressed"] = frame[["total", "value", "pct"]].eq("DS").any(axis=1)
        frame["not_calculated"] = frame[["total", "value", "pct"]].eq("NC").any(axis=1)
        for c in ("total", "value", "pct"):
            frame[c] = pd.to_numeric(frame[c].where(~frame[c].isin(["DS", "NC"])))
        if frame.loc[frame.measure.eq("expenditures"), "suppressed"].any():
            raise SystemExit("[BLOCKED] suppressed expenditure cell")
        # Each table's group cells should sum to its total within the stated approximation.
        full = frame[~frame.groupby(["year", "sheet", "state"]).suppressed.transform("any")].dropna(subset=["total"])
        chk = full.groupby(["year", "sheet", "state"]).agg(total=("total", "first"), s=("value", "sum"))
        gap = ((chk.s - chk.total).abs() / chk.total.where(chk.total > 0)).max()
        print(f"  ✓ max relative gap between group sum and table total (unsuppressed tables): {gap:.2e}")
    out = HERE / "derived"
    out.mkdir(exist_ok=True)
    race.to_csv(out / "taf_race_ethn.csv", index=False)
    lang.to_csv(out / "taf_language.csv", index=False)
    ages.to_csv(out / "taf_age.csv", index=False)
    dq.to_csv(out / "taf_dq.csv", index=False)
    nat = race[race.state.eq("National") & race.group.eq("hispanic") & race.category.isin(
        ["LTSS", "INST", "NF", "ICF", "MHF", "HCBS"])]
    pd.set_option("display.width", 200)
    print(nat.pivot_table(index=["measure", "category"], columns="year", values="pct").round(2).to_string())
    print((nat.pivot_table(index=["measure", "category"], columns="year", values="total") / 1e9).round(3).to_string())
    nc = race[race.not_calculated].groupby("year").state.unique()
    print("not calculated:", {y: list(v) for y, v in nc.items()})


if __name__ == "__main__":
    main()
