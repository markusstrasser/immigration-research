"""INEGI Censo 2020 ISCED × 5-year age vs ACS Mexico-born at the same ages.

Origin: cuestionario básico tabulado 13, national Total, 15+ by quinquenal age,
grado promedio and ISCED counts.
https://www.inegi.org.mx/contenidos/programas/ccpv/2020/tabulados/cpv2020_b_eum_07_educacion.xlsx
(fetched 2026-09-22).

LTHS ≈ ISCED 0–2 (less than primary + primary + lower secondary).
BA+ ≈ ISCED 6–8 (bachelor-equivalent tertiary, master's, doctorate);
short-cycle tertiary is not BA+.

ACS: staged 1-year PUMS, Mexico-born (POBP=303), ages 25–54, modern SCHL
(<=15 LTHS, >=21 BA+). 2019 is the last 1-year ACS before the missing 2020 file;
2021 is the first after. Stock vs 0–5 years since arrival (YOEP >= YEAR-5).
Levels are same-age stocks, not return-migration corrected. Child arrivals sit
in the stock; the 0–5 cut is the selection-relevant object.
"""
import os
from openpyxl import load_workbook
import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "../../.."))
XLSX = os.path.join(
    ROOT,
    "sources/immigration-fiscal/data/external/stage3/inegi/cpv_educacion/"
    "cpv2020_b_eum_07_educacion.xlsx",
)
CACHE = os.path.join(HERE, "_cache")
OUT = os.path.join(HERE, "derived")

BANDS = [
    ("25-29 años", 25, 29, 1991, 1995),
    ("30-34 años", 30, 34, 1986, 1990),
    ("35-39 años", 35, 39, 1981, 1985),
    ("40-44 años", 40, 44, 1976, 1980),
    ("45-49 años", 45, 49, 1971, 1975),
    ("50-54 años", 50, 54, 1966, 1970),
]


def inegi_age():
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["13"]
    rows = []
    total_mean = None
    for row in ws.iter_rows(min_row=9, values_only=True):
        loc, sex, age, pop = row[0], row[1], row[2], row[3]
        if loc != "Estados Unidos Mexicanos" or sex != "Total":
            continue
        less_pri, pri, lo_sec, hi_sec, short, tert, mast, doc, ns, mean_y = row[4:14]
        if age == "Total":
            total_mean = mean_y
            continue
        lths = less_pri + pri + lo_sec
        ba = tert + mast + doc
        specified = pop - ns
        rows.append({
            "age_band": age, "pop": pop, "n_unspecified": ns,
            "mean_years": mean_y,
            "sh_lths": lths / specified, "sh_ba_plus": ba / specified,
            "sh_upper_secondary": hi_sec / specified,
            "sh_short_cycle": short / specified,
        })
    wb.close()
    if total_mean is None or abs(total_mean - 9.7) > 0.05:
        raise SystemExit(f"15+ mean years {total_mean} != Cuéntame 9.7")
    d = pd.DataFrame(rows)
    meta = pd.DataFrame(
        BANDS, columns=["age_band", "age_lo", "age_hi", "born_lo", "born_hi"])
    return d.merge(meta, on="age_band")


def acs_age(year, recent_only):
    p = os.path.join(CACHE, f"acs_{year}.parquet")
    d = pd.read_parquet(p)
    d = d[(d.POBP == 303) & (d.AGEP >= 25) & (d.AGEP <= 54)].copy()
    d["ysm"] = d.YEAR - d.YOEP
    if recent_only:
        d = d[(d.ysm >= 0) & (d.ysm <= 5)]
    d["edu4"] = np.where(
        d.SCHL <= 15, "lths",
        np.where(d.SCHL <= 17, "hs",
                 np.where(d.SCHL <= 20, "somecoll", "ba_plus")))
    tag = "recent" if recent_only else "stock"
    out = []
    for lab, lo, hi, *_ in BANDS:
        g = d[(d.AGEP >= lo) & (d.AGEP <= hi)]
        w = g.PWGTP.sum()
        out.append({
            "age_band": lab,
            f"n_{tag}_{year}": len(g),
            f"wpop_{tag}_{year}": w,
            f"sh_lths_{tag}_{year}": (
                g.loc[g.edu4 == "lths", "PWGTP"].sum() / w if w else np.nan),
            f"sh_ba_{tag}_{year}": (
                g.loc[g.edu4 == "ba_plus", "PWGTP"].sum() / w if w else np.nan),
        })
    return pd.DataFrame(out)


def main():
    o = inegi_age()
    o.to_csv(os.path.join(OUT, "inegi_2020_isced_by_age.csv"), index=False)
    m = o.copy()
    for year in (2019, 2021):
        for recent in (False, True):
            m = m.merge(acs_age(year, recent), on="age_band")
    m["lths_recent2019_minus_mx"] = m.sh_lths_recent_2019 - m.sh_lths
    m["ba_recent2019_minus_mx"] = m.sh_ba_recent_2019 - m.sh_ba_plus
    m["lths_stock2019_minus_mx"] = m.sh_lths_stock_2019 - m.sh_lths
    m["ba_stock2019_minus_mx"] = m.sh_ba_stock_2019 - m.sh_ba_plus
    m.to_csv(os.path.join(OUT, "origin_age_vs_acs.csv"), index=False)
    cols = [
        "age_band", "born_lo", "born_hi", "mean_years",
        "sh_lths", "sh_lths_recent_2019", "lths_recent2019_minus_mx",
        "sh_ba_plus", "sh_ba_recent_2019", "ba_recent2019_minus_mx",
        "n_recent_2019", "sh_lths_stock_2019", "sh_ba_stock_2019",
        "sh_lths_recent_2021", "sh_ba_recent_2021",
    ]
    print(m[cols].to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    print("\n25-54 recent 2019 vs origin, weighted by origin pop:")
    w = m["pop"]
    print(
        "  LTHS mx {:.3f}  US recent {:.3f}  gap {:+.3f}".format(
            np.average(m.sh_lths, weights=w),
            np.average(m.sh_lths_recent_2019, weights=w),
            np.average(m.lths_recent2019_minus_mx, weights=w),
        ))
    print(
        "  BA+  mx {:.3f}  US recent {:.3f}  gap {:+.3f}".format(
            np.average(m.sh_ba_plus, weights=w),
            np.average(m.sh_ba_recent_2019, weights=w),
            np.average(m.ba_recent2019_minus_mx, weights=w),
        ))


if __name__ == "__main__":
    main()
