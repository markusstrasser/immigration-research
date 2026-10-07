"""Public-employee retirement costs inside the account's service lines: what is legacy, what is accrual.

The account's service lines are BEA T3.17 consumption by function (full_account_spending_2026_09_20/builder.py,
partition()). This script measures, from pinned primary files, the three pieces the brief asks about:

  1. Pensions. BEA books defined-benefit employer compensation at the accrual normal cost and removes the cash excess
     (amortization of unfunded liabilities) as a negative imputed contribution (T7.24 line 6, T7.23 line 8); NIPA T3.19
     line 29 removes Census's excess cash employer contributions from S&L expenditure. Interest on the legacy sits in
     domestic_interest, held at response 0. So no pension amortization is on a responding line: traced, not priced.
  2. Retiree health (OPEB), state and local. Census classifies pay-as-you-go payments to former employees and group
     premiums as current operations (Classification Manual 2006, ch. 5, 5.3.4); BEA builds S&L consumption from Census
     total expenditure and T3.19 has no OPEB line, so pay-go sits on the service lines [INFERENCE]. The correction swaps
     the pay-go (legacy: response 0) for the GASB 75 normal cost (accrual of today's service: the line's response).
     States, FY2019: Pew 2023 Appendix B (normal cost, employer contributions). All S&L: times mu, the ratio of all-S&L
     net OPEB liability to the state plans' (Reason FY2019 / Pew FY2019 central; CRR / Pew 2016 low). 2024: times the
     growth of employer group-health contributions (T7.8 line 17, 2024 / 2019).
  3. Retiree health, federal civilian. BEA FAQ 553: federal health contributions cover current employees and retirees.
     FR FY2025 Note 13, FY2024 column: civilian normal cost against benefits paid; swapped the same way.
  4. Retiree health, military. Its accrual (MERHCF normal cost, T7.8 line 16) and the military treatment facilities sit
     in defense (response 0), but care bought from civilian providers for retirees and their dependents is in T3.12
     line 26 (footnote 7), the account's other_federal_benefits, keyed all_cash at response 1. It is pay-go for past
     service: to response 0, no normal cost added (that is defense's). Measured: MERHCF's FY2024 purchased care, $9.7bn
     (DoD MERHCF Audited Financial Report FY2024); pre-65 retirees: FR military benefits paid less MERHCF's total, times
     MERHCF's purchased share at the central [ASSUMPTION], none at the low end, all at the high end.

Each amount is spread over the account's lines by the pension legacy lane's payroll mix (function_mix.csv, oct05);
S&L highways fold into economic_affairs_services (their parent line), enterprises are reported and left out.

Writes derived/national.csv (arm x plan x line), derived/inputs.json, derived/pension_trace.csv and
derived/school_price_beside.csv. Run from the repository root:
  OPENBLAS_NUM_THREADS=1 uv run --no-project python3 infra/immigration-fiscal/public_pension_legacy_2026_10_07/opeb_national.py
"""

import csv
import hashlib
import html
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
FISCAL = HERE.parent
ROOT = FISCAL.parent.parent
OUT = HERE / "derived"
STAGE = ROOT / "sources/immigration-fiscal/data/external/stage3"
LEGACY = FISCAL / "pension_legacy_2026_09_30"

PINS = {
    STAGE / "pew/opeb_2023/dostateshaveenoughsavedforretireehealthcarebenefits.pdf":
        "9b7c47c77a355f1188aa167f34e98344bfee232299e36b8384cd1561b113ac6a",
    STAGE / "crr/slp48/slp_48.pdf": "9d2e61bbf3f4f869d4f15d151ec87212699bea5b8531002a7c7aa3fcf868db61",
    STAGE / "reason/opeb_2021/survey.html": "5dc6897aa1f63091eb44f62bc314eedfa44260c6e9b792f8bb579403800ec75d",
    STAGE / "dod/merhcf_afr_fy2024/2024_AFR_MERHCF.pdf": "89da584ad87ca890e6b236bc3fa60d2a026d4aee0801392bb7c771f53f008843",
    LEGACY / "_cache/fr2025_note13.pdf": "deed540cf26fc8d8d77a6468019b1781c02ebca70a0f267318c3ca899de3cd55",
    LEGACY / "_cache/Section3All_xls.xlsx": "69b5c7aefb38675324887ce31d6feb4fcde7c903ab952db7328da0813096615e",
    LEGACY / "_cache/Section7All_xls.xlsx": "ce107c8ce92393613c0abae38156afcfaef8ab571eedf0302dc09ca44738b9ef",
    LEGACY / "derived/oct05/function_mix.csv": None,      # tracked; hash recorded in inputs.json
    FISCAL / "school_cost_where_enrolled_2026_09_24/_cache/elsec24_sumtables.xlsx":
        "848133174f299957653bbc7070633cf3ff23e4242695ccb628893fb8f6571154",
    FISCAL / "school_cost_where_enrolled_2026_09_24/derived/account_pupils_by_state.csv": None,
}

# The account's lines a payroll share can land on; highways and enterprises are handled below.
LINES = ["general_public_services", "defense", "public_order_safety", "economic_affairs_services",
         "housing_community_services", "health_services", "recreation_culture", "education_services",
         "income_security_services"]

# Arms: mu = all-S&L over state-plan scale, rho = normal cost over pay-go, fed = federal civilian and military in,
# mil = the military retiree purchased care taken out of other_federal_benefits.
C = dict(mu="central", rho="pew_2019", fed=True, swap=True, mil="central", sl=True, role="arm")
ARMS = {
    "central":            C,
    "mu_low":             {**C, "mu": "low"},
    "mu_high":            {**C, "mu": "high"},
    "rho_one":            {**C, "rho": "one"},
    "rho_high":           {**C, "rho": "high"},
    "military_low":       {**C, "mil": "low"},
    "military_high":      {**C, "mil": "high"},
    "no_federal":         {**C, "fed": False},
    "state_local_only":   {**C, "fed": False},
    "federal_only":       {**C, "sl": False},
    "paygo_removed_only": {**C, "swap": False, "role": "beside"},
}
del ARMS["state_local_only"]              # the same as no_federal


def fail(msg):
    print(f"[BLOCKED] {msg}")
    sys.exit(1)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def check_pins():
    out = {}
    for path, want in PINS.items():
        if not path.exists():
            fail(f"missing source {path.relative_to(ROOT)}")
        got = sha(path)
        if want and got != want:
            fail(f"stale source {path.relative_to(ROOT)}: {got}")
        out[str(path.relative_to(ROOT))] = got
    return out


def pdf_text(path, layout=True):
    args = ["pdftotext"] + (["-layout"] if layout else []) + [str(path), "-"]
    return subprocess.run(args, check=True, capture_output=True, text=True).stdout


def bea_cells(book, sheet, lines, year):
    rows = list(book[sheet].values)
    header = [r for r in rows if r[0] == "Line"]
    if len(header) != 1 or rows[1][0] != "[Millions of dollars]":
        fail(f"{sheet}: units or header changed")
    cols = [i for i, x in enumerate(header[0]) if str(x) == str(year)]
    if len(cols) != 1:
        fail(f"{sheet}: no single {year} column")
    table = {int(r[0]): (r[1].strip(), r[cols[0]]) for r in rows if str(r[0]).isdigit()}
    return {n: table[n] for n in lines}


def money(s):
    return float(s.replace("$", "").replace(",", ""))


# ----------------------------------------------------------------------------------------------- sources
def pew_states():
    """Pew 2023 Appendix B, FY2019, thousands of dollars: one row per state with a GASB 75 plan (NE, SD: N/A)."""
    text = pdf_text(STAGE / "pew/opeb_2023/dostateshaveenoughsavedforretireehealthcarebenefits.pdf")
    d = r"(-?\$[\d,]+)"
    pat = re.compile(rf"^\s*([A-Z]{{2}})\s+{d}\s+([\d.]+)%\s+{d}\s+{d}\s+{d}\s+{d}\s+{d}\s+{d}\s+(\d+)%\s+{d}\s*$")
    rows = []
    for line in text.splitlines():
        m = pat.match(line)
        if m:
            g = m.groups()
            rows.append(dict(state=g[0], nol=money(g[1]), interest=money(g[3]), normal=money(g[4]),
                             employee=money(g[6]), benchmark=money(g[7]), employer=money(g[8]), net_amort=money(g[10])))
    if len(rows) != 48 or sum("N/A" in l and re.match(r"^\s*(NE|SD)\s", l) is not None for l in text.splitlines()) != 2:
        fail(f"Pew Appendix B: {len(rows)} state rows, expected 48 plus NE and SD at N/A")
    tot = {k: sum(r[k] for r in rows) / 1e6 for k in ["nol", "interest", "normal", "employee", "benchmark", "employer", "net_amort"]}
    # The brief's own aggregate: "fell short of the net amortization benchmark by $30 billion" (2019).
    if not 29.5 <= -tot["net_amort"] <= 30.5 or abs(tot["benchmark"] - tot["employer"] + tot["net_amort"]) > 0.01:
        fail(f"Pew Appendix B sums disagree with the text's $30 billion: {tot}")
    nat = re.search(r"National\s+\$([\d,]+)\s+\$([\d,]+)\s+9\.2%", text)
    y2016 = re.search(r"In 2016, states reported \$(\d+) billion in liabilities and \$(\d+) billion in assets, with a funding gap of \$(\d+) billion", " ".join(text.split()))
    if not nat or not y2016:
        fail("Pew national 2019 row or 2016 sentence not found")
    tot["assets_2019"] = money(nat.group(1)) / 1e6
    tot["liabilities_2019"] = money(nat.group(2)) / 1e6
    tot["gap_2019"] = tot["liabilities_2019"] - tot["assets_2019"]
    tot["gap_2016"] = float(y2016.group(3))
    return rows, tot


def crr_total():
    text = " ".join(pdf_text(STAGE / "crr/slp48/slp_48.pdf", layout=False).split())   # two columns
    m = re.search(r"producing a total of \$(\d+) billion", text)
    parts = re.search(r"counties are responsible for \$(\d+) billion, cities are responsible for \$(\d+) billion, and school districts for \$(\d+) billion", text)
    if not m or not parts:
        fail("CRR SLP 48: total or local parts not found")
    return dict(total_bn=float(m.group(1)), counties_bn=float(parts.group(1)), cities_bn=float(parts.group(2)),
                school_districts_bn=float(parts.group(3)))


def reason_total():
    raw = (STAGE / "reason/opeb_2021/survey.html").read_text(encoding="utf-8", errors="ignore")
    text = " ".join(html.unescape(re.sub(r"<[^>]+>", " ", raw)).split())
    m = re.search(r"Totals \$([\d,]+) ([\d,]+) \$([\d,]+)", text)
    if not m or "end of fiscal year 2019" not in text:
        fail("Reason survey: totals row or FY2019 statement not found")
    return dict(total_bn=money(m.group(1)) / 1e9, population=money(m.group(2)))


def fr_federal():
    """FR FY2025 Note 13, FY2024 columns ($bn): post-retirement health and pension normal cost, benefits paid."""
    text = pdf_text(LEGACY / "_cache/fr2025_note13.pdf")
    i = text.find("Change in Post-Retirement Health Benefits as of September 30, 2025, and 2024")
    j = text.find("Change in Pension Benefits as of September 30, 2025, and 2024")
    if i < 0 or j < 0:
        fail("FR Note 13 tables not found")
    num = r"\(?([\d.,]+)\)?"
    def row(block, label):
        m = re.search(rf"{label}\s+{num}\s+{num}\s+{num}\s+{num}\s+{num}\s+{num}", block)
        if not m:
            fail(f"FR Note 13: no row {label}")
        v = [float(x.replace(",", "")) for x in m.groups()]
        return dict(civilian=v[1], military=v[3], total=v[5])      # the 2024 columns
    opeb, pens = text[i:i + 2500], text[j:j + 2500]
    out = dict(opeb_normal=row(opeb, "Normal costs"), opeb_paid=row(opeb, "Less benefits paid"),
               opeb_interest=row(opeb, "Interest on liability"),
               pension_normal=row(pens, "Normal costs"), pension_paid=row(pens, "Less benefits paid"))
    for k in out:
        if abs(out[k]["civilian"] + out[k]["military"] - out[k]["total"]) > 0.15:
            fail(f"FR Note 13 {k}: civilian + military != total")
    if (out["opeb_normal"]["civilian"], out["opeb_paid"]["civilian"]) != (20.4, 16.6):
        fail(f"FR Note 13 civilian OPEB FY2024 changed: {out['opeb_normal']}, {out['opeb_paid']}")
    return out


def merhcf():
    text = " ".join(pdf_text(STAGE / "dod/merhcf_afr_fy2024/2024_AFR_MERHCF.pdf", layout=False).split())
    m = re.search(r"In FY 2024 and FY 2023 respectively, the Fund authorized approximately \$([\d.]+) billion and \$[\d.]+ billion "
                  r"in total health care services, \$([\d.]+) billion and \$[\d.]+ billion to civilian providers for purchased care", text)
    if not m:
        fail("MERHCF AFR FY2024: the authorized-care sentence not found")
    return dict(total_fy2024_bn=float(m.group(1)), purchased_fy2024_bn=float(m.group(2)))


def bea():
    s7 = openpyxl.load_workbook(LEGACY / "_cache/Section7All_xls.xlsx", read_only=True, data_only=True)
    s3 = openpyxl.load_workbook(LEGACY / "_cache/Section3All_xls.xlsx", read_only=True, data_only=True)
    t78_24 = bea_cells(s7, "T70800-A", [10, 17], 2024)
    t78_19 = bea_cells(s7, "T70800-A", [10, 17], 2019)
    t724 = bea_cells(s7, "T72400-A", [5, 6], 2024)
    t723 = bea_cells(s7, "T72300-A", [5, 8], 2024)
    t319 = bea_cells(s3, "T31900-A", [26, 28, 29, 31, 32, 47], 2023)
    s7.close(); s3.close()
    labels = {("T70800-A", 17): "Private group health insurance", ("T72400-A", 5): "Actual employer contributions",
              ("T72400-A", 6): "Imputed employer contributions", ("T72300-A", 5): "Actual employer contributions",
              ("T72300-A", 8): "Imputed employer contributions",
              ("T31900-A", 29): "Contributions from general government employers to own defined-benefit pension plans"}
    cells = {("T70800-A", 17): t78_24[17], ("T72400-A", 5): t724[5], ("T72400-A", 6): t724[6], ("T72300-A", 5): t723[5],
             ("T72300-A", 8): t723[8], ("T31900-A", 29): t319[29]}
    for k, want in labels.items():
        if not cells[k][0].startswith(want):
            fail(f"BEA {k}: label {cells[k][0]!r}")
    return dict(group_health_2024=t78_24[17][1] / 1e3, group_health_2019=t78_19[17][1] / 1e3,
                sl_retirement_supplement_2024=t78_24[10][1] / 1e3,
                sl_actual_employer_2024=t724[5][1] / 1e3, sl_imputed_employer_2024=t724[6][1] / 1e3,
                fed_actual_employer_2024=t723[5][1] / 1e3, fed_imputed_employer_2024=t723[8][1] / 1e3,
                t319_fy2023={str(n): dict(label=l, bn=v / 1e3) for n, (l, v) in t319.items()})


def function_mix():
    mix = {}
    with (LEGACY / "derived/oct05/function_mix.csv").open() as f:
        for r in csv.DictReader(f):
            mix.setdefault(r["plan"], {})[r["line"]] = float(r["weight"])
    for plan, w in mix.items():
        if abs(sum(w.values()) - 1) > 1e-5:
            fail(f"function mix {plan} does not sum to 1")
    return mix


def school_price_beside():
    """How much stripping legacy from the per-pupil state prices could move the school key (a bound, not an edit).

    The key prices each pupil at the state's ASSF current spending per pupil, which carries cash benefits (amortization
    and OPEB pay-go included). With legacy a fraction lam of benefits in every state, the group's cost-weighted price
    relative to the nation moves by (1 - lam b_g) / (1 - lam b_n), b the cost-weighted benefit share."""
    book = openpyxl.load_workbook(FISCAL / "school_cost_where_enrolled_2026_09_24/_cache/elsec24_sumtables.xlsx",
                                  read_only=True, data_only=True)
    t6 = {}
    for r in book["6"].values:
        if r[0] and isinstance(r[2], (int, float)):
            t6[re.sub(r"\.+$", "", r[0].strip())] = (r[2], r[4])
    book.close()
    us_total, us_benefits = t6["United States"]
    with (FISCAL / "school_cost_where_enrolled_2026_09_24/derived/account_pupils_by_state.csv").open() as f:
        states = list(csv.DictReader(f))
    if len(states) != 51 or any(s["state"] not in t6 for s in states):
        fail("school beside: state rows do not match ASSF Table 6")
    out = []
    for alloc in ["personal", "shared"]:
        g = n = gb = nb = 0.0
        for s in states:
            tot, ben = t6[s["state"]]
            g += float(s[f"target_school_{alloc}_bn"]); gb += float(s[f"target_school_{alloc}_bn"]) * ben / tot
            n += float(s[f"national_school_{alloc}_bn"]); nb += float(s[f"national_school_{alloc}_bn"]) * ben / tot
        for lam in [0.15, 0.25, 0.35]:
            out.append(dict(allocation=alloc, legacy_fraction_of_benefits=lam, benefit_share_group=gb / g,
                            benefit_share_national=nb / n, benefit_share_us_table6=us_benefits / us_total,
                            key_factor=(1 - lam * gb / g) / (1 - lam * nb / n)))
    return out


# ----------------------------------------------------------------------------------------------- build
def main():
    pins = check_pins()
    pew_rows, pew = pew_states()
    crr, reason, fr, b, mix, mer = crr_total(), reason_total(), fr_federal(), bea(), function_mix(), merhcf()

    mu = dict(central=reason["total_bn"] / pew["gap_2019"],          # same year, both net of assets
              low=crr["total_bn"] / pew["gap_2016"])                 # CRR (FY2013-14 data) over Pew 2016
    mu["high"] = mu["central"] + (mu["central"] - mu["low"])         # [ASSUMPTION] symmetric about the central
    rho = dict(pew_2019=pew["normal"] / pew["employer"], one=1.0,
               high=fr["opeb_normal"]["military"] / fr["opeb_paid"]["military"])   # benefits paid parse as positive
    growth = b["group_health_2024"] / b["group_health_2019"]
    sl_paygo_states_2024 = pew["employer"] * growth
    pre65 = fr["opeb_paid"]["military"] - mer["total_fy2024_bn"]          # FR benefits paid less MERHCF's care
    if not 0 < pre65 < fr["opeb_paid"]["military"]:
        fail(f"military pre-65 residual {pre65}")
    military = dict(low=mer["purchased_fy2024_bn"],
                    central=mer["purchased_fy2024_bn"] + pre65 * mer["purchased_fy2024_bn"] / mer["total_fy2024_bn"],
                    high=mer["purchased_fy2024_bn"] + pre65)

    rows = []
    for arm, a in ARMS.items():
        paygo_sl = sl_paygo_states_2024 * mu[a["mu"]] if a["sl"] else 0.0
        normal_sl = paygo_sl * rho[a["rho"]] if a["swap"] else 0.0
        paygo_fed = fr["opeb_paid"]["civilian"] if a["fed"] else 0.0
        normal_fed = fr["opeb_normal"]["civilian"] if (a["fed"] and a["swap"]) else 0.0
        mil = military[a["mil"]] if a["fed"] else 0.0
        rows.append(dict(arm=arm, role=a["role"], plan="federal_military_retiree_purchased_care", line="other_federal_benefits",
                         mix_weight=1.0, paygo_bn=-mil, normal_bn=0.0, delta_bn=-mil, on_engine_line=True))
        for plan, p, nc in [("state_local", paygo_sl, normal_sl), ("federal_civilian", paygo_fed, normal_fed)]:
            w = mix[plan]
            folded = {l: w.get(l, 0.0) for l in LINES}
            folded["economic_affairs_services"] += w.get("highways", 0.0)
            for line in LINES + ["enterprises"]:
                share = w.get("enterprises", 0.0) if line == "enterprises" else folded[line]
                rows.append(dict(arm=arm, role=a["role"], plan=plan, line=line, mix_weight=share,
                                 paygo_bn=-p * share, normal_bn=nc * share, delta_bn=(nc - p) * share,
                                 on_engine_line=line != "enterprises"))

    OUT.mkdir(exist_ok=True)
    cols = list(rows[0])
    def fmt(v):
        return f"{v:.9f}" if isinstance(v, float) else str(v)
    with (OUT / "national.csv").open("w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cols)
        for r in rows:
            w.writerow([fmt(r[c]) for c in cols])

    trace = [
        ("sl_db_employer_cash_2024", b["sl_actual_employer_2024"], "BEA T7.24 line 5"),
        ("sl_db_imputed_2024", b["sl_imputed_employer_2024"], "BEA T7.24 line 6 (negative: cash above normal cost)"),
        ("sl_db_employer_accrual_2024", b["sl_actual_employer_2024"] + b["sl_imputed_employer_2024"], "T7.24 lines 5+6 (in compensation)"),
        ("sl_db_amortization_share_of_cash", -b["sl_imputed_employer_2024"] / b["sl_actual_employer_2024"], "-(line 6) / line 5"),
        ("fed_db_employer_cash_2024", b["fed_actual_employer_2024"], "BEA T7.23 line 5"),
        ("fed_db_imputed_2024", b["fed_imputed_employer_2024"], "BEA T7.23 line 8"),
        ("fed_db_amortization_share_of_cash", -b["fed_imputed_employer_2024"] / b["fed_actual_employer_2024"], "-(line 8) / line 5"),
        ("t319_census_total_expenditures_fy2023", b["t319_fy2023"]["26"]["bn"], "BEA T3.19 line 26"),
        ("t319_retirement_plan_coverage_fy2023", b["t319_fy2023"]["28"]["bn"], "BEA T3.19 line 28"),
        ("t319_excess_employer_contributions_fy2023", b["t319_fy2023"]["29"]["bn"], "BEA T3.19 line 29 (removed from Census expenditure)"),
        ("t319_imputed_interest_fy2023", b["t319_fy2023"]["32"]["bn"], "BEA T3.19 line 32"),
        ("fed_pension_normal_cost_fy2024_fr", fr["pension_normal"]["total"], "FR Note 13 (Treasury-rate basis, employee share included)"),
    ]
    with (OUT / "pension_trace.csv").open("w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["item", "value", "source"])
        for k, v, s in trace:
            w.writerow([k, f"{v:.6f}", s])

    beside = school_price_beside()
    with (OUT / "school_price_beside.csv").open("w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(list(beside[0]))
        for r in beside:
            w.writerow([fmt(r[c]) for c in r])

    summary = {}
    for arm in ARMS:
        rs = [r for r in rows if r["arm"] == arm]
        summary[arm] = {k: sum(r[k] for r in rs if r["on_engine_line"]) for k in ["paygo_bn", "normal_bn", "delta_bn"]}
        summary[arm]["enterprises_delta_bn"] = sum(r["delta_bn"] for r in rs if not r["on_engine_line"])
    (OUT / "inputs.json").write_text(json.dumps(dict(
        sources=pins, pew_states_fy2019_bn=pew, crr=crr, reason=reason, fr_fy2024_bn=fr, bea=b,
        merhcf_fy2024_bn=mer, military_retiree_purchased_care_bn=military, military_pre65_benefits_paid_bn=pre65,
        mu=mu, rho=rho, group_health_growth_2019_2024=growth, sl_paygo_states_2024_bn=sl_paygo_states_2024,
        arms=ARMS, national_by_arm_bn=summary), indent=1, sort_keys=True) + "\n")

    print(f"[pensions] S&L DB cash {b['sl_actual_employer_2024']:.3f}bn, accrual in compensation "
          f"{b['sl_actual_employer_2024'] + b['sl_imputed_employer_2024']:.3f}bn; T3.19 line 29 FY2023 "
          f"{b['t319_fy2023']['29']['bn']:.3f}bn removed from Census expenditure")
    print(f"[OPEB] states FY2019: normal {pew['normal']:.3f}, employer {pew['employer']:.3f}, rho {rho['pew_2019']:.4f}; "
          f"mu {mu['low']:.3f}/{mu['central']:.3f}/{mu['high']:.3f}; growth {growth:.4f}")
    for arm, s in summary.items():
        print(f"  {arm:20s} pay-go {s['paygo_bn']:8.3f}  normal {s['normal_bn']:8.3f}  delta {s['delta_bn']:+8.3f}bn "
              f"(enterprises {s['enterprises_delta_bn']:+.3f}, off the engine)")


if __name__ == "__main__":
    main()
