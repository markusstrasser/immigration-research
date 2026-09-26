"""School capital return lane (2026-09-26).

The main case charges the Mexican-origin group's pupils at full average cost, from BEA
government consumption. That consumption includes consumption of fixed capital (CFC) but "the
use of depreciation assumes a zero net return on these assets" (NIPA Table 3.10.5, note 2).
This script measures the missing return on the K-12 capital behind the group's pupils, compares
it with the depreciation already inside the school line, checks the account for interest that
would double count it, and sets it beside the ledger's cash-basis item K.

    uv run --no-project python3 infra/immigration-fiscal/school_capital_return_2026_09_26/capital_return.py

Reads `_cache/` (staged by `acquire.py`, sha256-checked against `derived/sources.csv`), the
pinned NIPA Section 3 workbook, and read-only facts from sibling lanes (hashes recorded in
`derived/summary.json`). Writes `derived/*.csv` and `derived/summary.json`, prints the results,
and exits 1 when a gate fails.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

LANE = Path(__file__).resolve().parent
ROOT = LANE.parents[2]
CACHE = LANE / "_cache"
DERIVED = LANE / "derived"
FISCAL = ROOT / "infra" / "immigration-fiscal"
NIPA3 = ROOT / "sources" / "immigration-fiscal" / "data" / "external" / "bea_nipa" / "Section3All_xls.xlsx"
EMBED = FISCAL / "school_cost_where_enrolled_2026_09_24" / "derived" / "account_embedded_price.json"
PUPILS = FISCAL / "school_cost_where_enrolled_2026_09_24" / "derived" / "account_pupils_by_state.csv"
ROC = FISCAL / "school_cost_where_enrolled_2026_09_24" / "derived" / "r_other_concepts.csv"
MAIN = FISCAL / "main_case_schools_full_2026_09_26" / "derived" / "summary.json"
MODEL = FISCAL / "assumption_explorer_2026_09_21" / "derived" / "model.json"
ENGINE = FISCAL / "assumption_explorer_2026_09_21" / "engine.js"
PACKAGES = [FISCAL / "main_case_2026_09_24" / "package.cjs", FISCAL / "main_case_2026_09_26" / "package.cjs",
            FISCAL / "main_case_schools_full_2026_09_26" / "package.cjs"]
DEBT_RESULT = FISCAL / "debt_legacy_2026_09_23" / "RESULT.md"
LEDGER_ITEMS = FISCAL / "ledger_absolute_2026_09_17" / "derived" / "items_by_group.csv"
LEDGER_PARAMS = FISCAL / "ledger_absolute_2026_09_17" / "params" / "params.json"
LEDGER_SCRIPT = FISCAL / "ledger_absolute_2026_09_17" / "absolute_ledger.py"
LEDGER_EXT = FISCAL / "gen_ledger_extension_2026_09_16" / "extend_ledger.py"

YEAR = 2024
BRIEF_S = 0.17480600215120404
RATES = {"A-4 2023 (2%)": 0.02, "A-4 2003 (3%)": 0.03}
A4_2003_PRIVATE = 0.07          # A-4 2003's before-tax return to private capital; linear, reported only
STRUCT_KD_RANGE = (30.0, 80.0)  # net stock / depreciation for structures: geometric rates 1.25-3.3%
CFC_TOL = 0.005                 # FA depreciation vs NIPA CFC where both exist
PIM_TOL = 0.10                  # perpetual inventory vs BEA's published 2024 stock
VINTAGE_SPLIT = 1993            # first year of Census VIP detail
EDU_LINE = 62                   # FAAt70x line: state and local structures, Educational

GATES: list[tuple[str, bool, str]] = []


def gate(name: str, ok: bool, detail: str) -> None:
    GATES.append((name, bool(ok), detail))


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(name: str, header: list[str], rows: list[list]) -> None:
    with (DERIVED / name).open("w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([round(x, 6) if isinstance(x, float) else x for x in r])


# ---------------------------------------------------------------- readers
def check_cache() -> dict:
    reg = {}
    with (DERIVED / "sources.csv").open() as f:
        for r in csv.DictReader(f):
            p = CACHE / r["file"]
            if not p.exists():
                raise SystemExit(f"[BLOCKED] missing {p}; run acquire.py")
            if sha(p) != r["sha256"]:
                raise SystemExit(f"[BLOCKED] {r['file']} changed since acquire.py recorded it")
            reg[r["file"]] = r
    return reg


def bea_table(path: Path, sheet: str) -> tuple[dict, dict, list[str]]:
    """{line: {year: $mn}}, {line: label}, header text rows of one BEA annual sheet."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    head = [i for i, r in enumerate(rows) if r and r[0] == "Line"][0]
    years = [int(y) for y in rows[head][3:] if y is not None]
    data, labels = {}, {}
    for r in rows[head + 1:]:
        if isinstance(r[0], int) or (isinstance(r[0], str) and r[0].strip().isdigit()):
            line = int(r[0])
            labels[line] = str(r[1]).strip()
            data[line] = {y: float(v) for y, v in zip(years, r[3:]) if isinstance(v, (int, float))}
    return data, labels, [str(r[0]) for r in rows[:head] if r and r[0]]


def label_is(labels: dict, line: int, text: str, where: str) -> None:
    if not labels[line].startswith(text):
        raise SystemExit(f"[BLOCKED] {where} line {line} is '{labels[line]}', expected '{text}'")


def vip_series() -> dict:
    """Census VIP, state and local educational construction, $mn by year and category."""
    want = ("Educational", "Primary/secondary", "Higher education", "Other educational")
    out: dict = {}

    def take(name, years, vals):
        name = name.replace("\xa0", " ").strip()
        if name in want:
            for y, v in zip(years, vals):
                out.setdefault(y, {})[name] = float(v)

    for f in ("c30_stateha__Annual.csv", "c30_stateha1__Annual.csv"):
        rows = list(csv.reader((CACHE / f).open()))
        src = rows[0][0].split("=")[1]
        if src != sha(CACHE / f.replace("__Annual.csv", ".xls")):
            raise SystemExit(f"[BLOCKED] {f} is stale; rerun acquire.py")
        hdr = [r for r in rows if r and r[0].startswith("Type of")][0]
        years = [int(float(x)) for x in hdr[1:] if x]
        for r in rows:
            if r:
                take(r[0], years, r[1:1 + len(years)])
    for f in ("c30_stateha2.xlsx", "c30_state.xlsx"):
        rows = list(openpyxl.load_workbook(CACHE / f, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
        hdr = [r for r in rows if r and r[0] and str(r[0]).startswith("Type of")][0]
        years = [int(x) for x in hdr[1:] if x is not None]
        for r in rows:
            if r and r[0]:
                take(str(r[0]), years, r[1:1 + len(years)])
    return out


def f33_state_table(path: Path, sheet: str, cols: dict) -> dict:
    """F-33 summary table rows by state name: {name: {field: $thousand or count}}."""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    out = {}
    for r in rows:
        if r and r[0] and isinstance(r[2], (int, float)):
            name = re.sub(r"\.+$", "", str(r[0]).strip()).strip()
            # "(N)" marks an amount Census does not publish; it stays None and fails loudly if used
            out[name] = {k: (float(r[i]) if isinstance(r[i], (int, float)) else None) for k, i in cols.items()}
    return out


def f33_2019_us(sheet: str, cols: dict) -> dict:
    rows = list(csv.reader((CACHE / f"elsec19_sumtables__{sheet}.csv").open()))
    if rows[0][0].split("=")[1] != sha(CACHE / "elsec19_sumtables.xls"):
        raise SystemExit("[BLOCKED] elsec19 extracts are stale; rerun acquire.py")
    us = [r for r in rows if r and r[0].strip().startswith("United States")][0]
    return {k: float(us[i]) for k, i in cols.items()}


# ---------------------------------------------------------------- main
def main() -> int:
    reg = check_cache()
    DERIVED.mkdir(exist_ok=True)
    fa = CACHE / "fa_Section7All_xls.xlsx"
    K, klab, khead = bea_table(fa, "FAAt701-A")       # current-cost net stock, yearend
    D, dlab, _ = bea_table(fa, "FAAt703-A")           # current-cost depreciation
    INV, ilab, _ = bea_table(fa, "FAAt705-A")         # investment
    Q, qlab, _ = bea_table(fa, "FAAt706-A")           # investment quantity index
    QK, _, _ = bea_table(fa, "FAAt702-A")             # net stock quantity index
    AGE, _, _ = bea_table(fa, "FAAt707-A")            # average age
    for tab, lab in (("FAAt701", klab), ("FAAt703", dlab), ("FAAt705", ilab), ("FAAt706", qlab)):
        for line, text in ((1, "Government fixed assets"), (2, "Equipment"), (3, "Structures"), (18, "Intellectual"),
                           (21, "Federal"), (43, "Health care"), (44, "Educational"), (45, "Public safety"),
                           (55, "State and local"), (56, "Equipment"), (57, "Structures"), (58, "Residential"),
                           (61, "Health care"), (62, "Educational"), (63, "Public safety"), (68, "Sewer systems"),
                           (69, "Water systems"), (72, "Intellectual"), (73, "Software"), (74, "Research")):
            label_is(lab, line, text, tab)
    t317, l317, h317 = bea_table(NIPA3, "T31700-A")
    t3155, l3155, _ = bea_table(NIPA3, "T31505-A")
    t3105, l3105, h3105 = bea_table(NIPA3, "T31005-A")
    t705, l705, _ = bea_table(CACHE / "nipa_Section7All_xls.xlsx", "T70500-A")
    t510, l510, _ = bea_table(CACHE / "nipa_Section5All_xls.xlsx", "T51000-A")
    for tab, lab, checks in (("T31700", l317, ((9, "Education"), (19, "Education"), (21, "State and local"), (30, "Education"),
                                                (132, "Education"))),
                             ("T31505", l3155, ((109, "Education"), (110, "Elementary and secondary"), (111, "Higher"))),
                             ("T31005", l3105, ((47, "State and local consumption"), (51, "Consumption of general government fixed capital"))),
                             ("T70500", l705, ((22, "General government"), (24, "State and local"), (25, "Government enterprises"),
                                               (27, "State and local"))),
                             ("T51000", l510, ((39, "Government"), (40, "Structures"), (41, "Equipment"), (42, "Intellectual")))):
        for line, text in checks:
            label_is(lab, line, text, tab)
    note2 = [h for h in _footnotes(NIPA3, "T31005-A") if h.startswith("2.")][0]
    gate("nipa_cfc_zero_net_return_note", "assumes a zero net return on these assets" in note2, note2[-90:])
    gate("reused_sl_cfc_share", t3105[51][YEAR] == 312559 and t3105[47][YEAR] == 2550362,
         f"T3.10.5 l51/l47 {YEAR}: {t3105[51][YEAR]:.0f}/{t3105[47][YEAR]:.0f} = {t3105[51][YEAR] / t3105[47][YEAR]:.4f}")
    for line, text in ((26, "Health (net)"), (27, "Gross expenditures")):
        label_is(l317, line, text, "T31700")

    bn = lambda v: v / 1000.0  # BEA and VIP $mn -> $bn

    # ---------------- reused facts
    emb = json.loads(EMBED.read_text())
    s = emb["target_pupils"] / emb["national_pupils"]
    gate("pupil_share_matches_brief", abs(s - BRIEF_S) < 1e-12, f"s={s!r}")
    main_sum = json.loads(MAIN.read_text())
    school_line = main_sum["school"]["school_line_at_full_cost"]          # [low end, high end] of the main case
    # Each end is read at its own specification (main_case.cjs since f1e4f5b); its school share pairs with it.
    end_specs = main_sum["school"]["end_specifications"]["average"]
    end_shares = [{e[end]["share"] for e in end_specs} for end in ("low_end", "high_end")]
    gate("end_specification_shares_agree_across_methods", all(len(x) == 1 for x in end_shares), f"{end_shares}")
    end_share = (min(end_shares[0]), min(end_shares[1]))
    main_case = main_sum["main_case"]
    model = json.loads(MODEL.read_text())
    fractions = sorted({p["school_share"] for p in model["service"]["profiles"] if p["school_share"] > 0})
    school_fraction = (fractions[0], fractions[-1])
    gate("school_fraction_bounds", abs(school_fraction[0] - 0.715) < 0.001 and abs(school_fraction[1] - 0.865) < 0.001,
         f"{school_fraction}")
    lines = {l["id"]: l for l in model["spending"]["lines"]}
    edu_cons_account = lines["education_services"]["national_bn"]
    gate("account_education_line_is_nipa_3_17_line_9", abs(edu_cons_account - bn(t317[9][YEAR])) < 1e-6,
         f"model {edu_cons_account} vs T3.17 l9 {bn(t317[9][YEAR])}")

    # ---------------- 1. K-12 structures: key by vintage
    L = EDU_LINE
    k_edu = {y: bn(K[L][y]) for y in (YEAR - 1, YEAR)}
    d_edu = bn(D[L][YEAR])
    k_edu_avg = (k_edu[YEAR - 1] + k_edu[YEAR]) / 2
    delta = d_edu / k_edu_avg
    p_inv = INV[L][YEAR] / Q[L][YEAR]                      # $mn per index point, 2024 investment prices
    years = sorted(y for y in Q[L] if y <= YEAR)
    wraw = {t: p_inv * Q[L][t] * (1 - delta) ** (YEAR - t) * (1 - delta / 2) for t in years}
    pim_2024 = sum(wraw.values())
    pim_1993 = sum(p_inv * Q[L][t] * (1 - delta) ** (1993 - t) * (1 - delta / 2) for t in years if t <= 1993)
    gate("pim_reproduces_bea_stock_2024", abs(pim_2024 / K[L][YEAR] - 1) < PIM_TOL,
         f"PIM {bn(pim_2024):.1f} vs BEA {k_edu[YEAR]:.1f} (ratio {pim_2024 / K[L][YEAR]:.4f}) at delta {delta:.5f}")
    gate("pim_reproduces_bea_real_growth_1993_2024",
         abs((pim_2024 / pim_1993) / (QK[L][YEAR] / QK[L][1993]) - 1) < 0.05,
         f"PIM x{pim_2024 / pim_1993:.3f} vs FAAt702 x{QK[L][YEAR] / QK[L][1993]:.3f}")
    w = {t: v / pim_2024 for t, v in wraw.items()}
    pre_weight = sum(v for t, v in w.items() if t < VINTAGE_SPLIT)

    vip = vip_series()
    vip_share = {y: v["Primary/secondary"] / v["Educational"] for y, v in vip.items() if y <= YEAR}
    vip_share_ph = {y: v["Primary/secondary"] / (v["Primary/secondary"] + v["Higher education"]) for y, v in vip.items() if y <= YEAR}
    gate("vip_covers_1993_2024", all(y in vip_share for y in range(VINTAGE_SPLIT, YEAR + 1)), f"{min(vip_share)}-{max(vip_share)}")
    ratios = [vip[y]["Educational"] / INV[L][y] for y in range(VINTAGE_SPLIT, YEAR + 1)]
    gate("vip_educational_tracks_bea_investment", 0.8 <= min(ratios) and max(ratios) <= 1.3,
         f"VIP/BEA educational investment {min(ratios):.3f}-{max(ratios):.3f}, 1993-2024")

    # pre-1993 vintages: NCES Digest 236.10 K-12 capital outlay x structures fraction / BEA investment
    nces = _nces_capital_outlay()
    f33 = f33_state_table(CACHE / "elsec24_sumtables.xlsx", "9",
                          {"capital_outlay": 2, "construction": 3, "land_existing": 4, "equip_instr": 5,
                           "equip_other": 6, "interest": 7})
    f33_debt = f33_state_table(CACHE / "elsec24_sumtables.xlsx", "10", {"debt": 2, "debt_long": 3})
    f33_enr = f33_state_table(CACHE / "elsec24_sumtables.xlsx", "19", {"enrollment": 4})
    us24 = f33["United States"]
    us19 = f33_2019_us("9", {"capital_outlay": 2, "construction": 3, "land_existing": 4, "equip_instr": 5,
                             "equip_other": 6, "interest": 7})
    f_struct = (us24["construction"] / us24["capital_outlay"] + us19["construction"] / us19["capital_outlay"]) / 2
    implied = {}
    for y, outlay in nces.items():
        if y - 1 in INV[L] and y in INV[L]:
            implied[y] = f_struct * outlay / ((INV[L][y - 1] + INV[L][y]) / 2)
    overlap = {y: implied[y] - (vip[y - 1]["Primary/secondary"] + vip[y]["Primary/secondary"])
               / (vip[y - 1]["Educational"] + vip[y]["Educational"])
               for y in implied if y - 1 in vip and y in vip_share}
    gate("nces_method_tracks_vip_in_overlap", max(abs(v) for v in overlap.values()) < 0.07,
         f"NCES-implied minus VIP, {len(overlap)} school years: {min(overlap.values()):+.3f} to {max(overlap.values()):+.3f}")
    anchors = {y: implied[y] for y in (1950, 1960, 1970, 1980)}
    # 1989-90 is set aside: NCES outlay exceeds BEA's educational investment (implied 0.91), and in
    # 1993-94, the first VIP years, VIP runs 22-25% above BEA's series, so the pairing is off there.
    anchor_years = sorted(anchors)

    def pre_share_central(t: int) -> float:
        pts = [(y, anchors[y]) for y in anchor_years] + [(VINTAGE_SPLIT, vip_share[VINTAGE_SPLIT])]
        if t <= pts[0][0]:
            return pts[0][1]
        for (y0, v0), (y1, v1) in zip(pts, pts[1:]):
            if y0 <= t <= y1:
                return v0 + (v1 - v0) * (t - y0) / (y1 - y0)
        raise ValueError(t)

    pre_lo = min(list(anchors.values()) + [vip_share[VINTAGE_SPLIT]])
    pre_hi = max(list(anchors.values()) + [vip_share[VINTAGE_SPLIT]])
    post = sum(w[t] * vip_share[t] for t in years if t >= VINTAGE_SPLIT)
    key = {
        "central": post + sum(w[t] * pre_share_central(t) for t in years if t < VINTAGE_SPLIT),
        "low": post + pre_lo * pre_weight,
        "high": post + pre_hi * pre_weight,
        "high_with_1990_anchor": post + max(pre_hi, implied[1990]) * pre_weight,
    }
    post_w = sum(w[t] for t in years if t >= VINTAGE_SPLIT)
    key_alt = {
        "vip_post1993_vintages_only": post / post_w,
        "vip_post1993_k12_over_k12_plus_higher": sum(w[t] * vip_share_ph[t] for t in years if t >= VINTAGE_SPLIT) / post_w,
        "vip_2024_only": vip_share[YEAR],
        "vip_1993_2024_min": min(vip_share[y] for y in range(VINTAGE_SPLIT, YEAR + 1)),
        "vip_1993_2024_max": max(vip_share[y] for y in range(VINTAGE_SPLIT, YEAR + 1)),
        "cog_fy2022_capital_outlay": None,
        "nipa_3_15_5_gross_k12_share": t3155[110][YEAR] / t3155[109][YEAR],
    }
    cog = _cog_education_capital()
    key_alt["cog_fy2022_capital_outlay"] = cog["elsec"] / cog["education"]
    gate("k12_key_plausible", 0.5 < key["low"] <= key["central"] <= key["high"] < 0.85,
         f"low {key['low']:.4f} central {key['central']:.4f} high {key['high']:.4f}")

    key_rows = []
    for t in years:
        v = vip.get(t, {})
        key_rows.append([t, INV[L][t], Q[L][t], w[t], v.get("Educational"), v.get("Primary/secondary"),
                         v.get("Higher education"), v.get("Other educational"), vip_share.get(t),
                         vip_share[t] if t >= VINTAGE_SPLIT else pre_share_central(t),
                         vip_share[t] if t >= VINTAGE_SPLIT else pre_lo,
                         vip_share[t] if t >= VINTAGE_SPLIT else pre_hi,
                         nces.get(t), implied.get(t)])
    write_csv("k12_share_key.csv",
              ["year", "bea_sl_educational_investment_mn", "bea_investment_quantity_index", "pim_weight_in_2024_stock",
               "vip_sl_educational_mn", "vip_primary_secondary_mn", "vip_higher_education_mn", "vip_other_educational_mn",
               "vip_k12_share", "k12_share_central", "k12_share_low", "k12_share_high",
               "nces_k12_capital_outlay_mn_school_year_ending", "nces_implied_k12_share"], key_rows)

    # ---------------- 2. K-12 equipment and software, and the NIPA-basis depreciation
    fa_sl_dep = bn(D[55][YEAR])
    nipa_sl_cfc = bn(t705[24][YEAR] + t705[27][YEAR])
    eq_gap_sl = fa_sl_dep - nipa_sl_cfc
    eq_gap_gov = bn(D[2][YEAR] - t510[41][YEAR])
    gate("fa_vs_nipa_structures_cfc_2024", abs(D[3][YEAR] / t510[40][YEAR] - 1) < CFC_TOL,
         f"FA 7.3 l3 {bn(D[3][YEAR]):.3f} vs NIPA 5.10 l40 {bn(t510[40][YEAR]):.3f}")
    gate("fa_vs_nipa_ipp_cfc_2024", abs(D[18][YEAR] / t510[42][YEAR] - 1) < CFC_TOL,
         f"FA 7.3 l18 {bn(D[18][YEAR]):.3f} vs NIPA 5.10 l42 {bn(t510[42][YEAR]):.3f}")
    for y in (2022, 2023):
        gate(f"fa_vs_nipa_state_local_cfc_{y}",
             abs(D[55][y] / (t705[24][y] + t705[27][y]) - 1) < CFC_TOL,
             f"FA 7.3 l55 {bn(D[55][y]):.3f} vs NIPA 7.5 l24+l27 {bn(t705[24][y] + t705[27][y]):.3f}")
    gate("state_local_2024_cfc_gap_sits_in_equipment", abs(eq_gap_sl - eq_gap_gov) < 0.5,
         f"S&L FA-NIPA {eq_gap_sl:.3f}bn; government equipment FA-NIPA {eq_gap_gov:.3f}bn")
    gate("fa_vs_nipa_gov_closing_stock_2024", abs(K[1][YEAR] / t510[74][YEAR] - 1) < 1e-6,
         f"FA 7.1 l1 {bn(K[1][YEAR]):.3f} vs NIPA 5.10 l74 closing {bn(t510[74][YEAR]):.3f}")
    d_eq_nipa = bn(D[56][YEAR]) - eq_gap_sl
    d_sw = bn(D[73][YEAR])
    d_rd = bn(D[74][YEAR])
    eqsw = {y: bn(K[56][y] + K[73][y]) for y in (YEAR - 1, YEAR)}
    eqsw_avg = (eqsw[YEAR - 1] + eqsw[YEAR]) / 2
    us24_equip = us24["equip_instr"] + us24["equip_other"]
    us19_equip = us19["equip_instr"] + us19["equip_other"]
    k_eq = {"fy2024": (us24_equip / 1e3) / (INV[56][YEAR] + INV[73][YEAR]),
            "fy2019": (us19_equip / 1e3) / (INV[56][2019] + INV[73][2019])}
    k_eq["central"] = (k_eq["fy2024"] + k_eq["fy2019"]) / 2
    k_eq["low"], k_eq["high"] = min(k_eq["fy2024"], k_eq["fy2019"]), max(k_eq["fy2024"], k_eq["fy2019"])

    # education's share of state and local non-structure investment (NIPA 3.17 l132 less FA structures)
    edu_nonstruct = bn(t317[132][YEAR] - INV[L][YEAR])
    e_rd_in_edu = (edu_nonstruct - bn(INV[74][YEAR])) / bn(INV[56][YEAR] + INV[73][YEAR])
    e_spread = edu_nonstruct / bn(INV[56][YEAR] + INV[73][YEAR] + INV[74][YEAR])
    gate("education_structures_inside_nipa_education_investment", INV[L][YEAR] < t317[132][YEAR],
         f"FA educational structures {bn(INV[L][YEAR]):.3f} < NIPA S&L education investment {bn(t317[132][YEAR]):.3f}")
    gate("k12_equipment_share_below_education_share", k_eq["high"] < e_rd_in_edu,
         f"K-12 {k_eq['low']:.3f}-{k_eq['high']:.3f} of S&L equipment+software investment; all education {e_rd_in_edu:.3f}")

    # ---------------- 3. national K-12 capital and its depreciation
    comp = []
    for v in ("low", "central", "high"):
        ks = key[v]
        ke = k_eq[v]
        comp.append(dict(variant=v, key_structures=ks, key_equipment=ke,
                         struct_2023=ks * k_edu[YEAR - 1], struct_2024=ks * k_edu[YEAR], struct_avg=ks * k_edu_avg,
                         eq_2023=ke * eqsw[YEAR - 1], eq_2024=ke * eqsw[YEAR], eq_avg=ke * eqsw_avg,
                         dep_struct=ks * d_edu, dep_eq=ke * (d_eq_nipa + d_sw)))
    for c in comp:
        c["total_avg"] = c["struct_avg"] + c["eq_avg"]
        c["total_2024"] = c["struct_2024"] + c["eq_2024"]
        c["dep_total"] = c["dep_struct"] + c["dep_eq"]
    kd = k_edu_avg / d_edu
    gate("net_stock_over_cfc_structures_plausible", STRUCT_KD_RANGE[0] <= kd <= STRUCT_KD_RANGE[1],
         f"S&L educational structures: average stock / depreciation = {kd:.1f} (range {STRUCT_KD_RANGE})")
    kd_all = comp[1]["total_avg"] / comp[1]["dep_total"]
    gate("net_stock_over_cfc_k12_total_plausible", 20 <= kd_all <= STRUCT_KD_RANGE[1], f"K-12 all assets {kd_all:.1f}")
    write_csv("national_k12_capital.csv",
              ["variant", "key_structures", "key_equipment_software", "structures_end2023_bn", "structures_end2024_bn",
               "structures_avg2024_bn", "equipment_software_end2023_bn", "equipment_software_end2024_bn",
               "equipment_software_avg2024_bn", "total_avg2024_bn", "total_end2024_bn", "depreciation_structures_2024_bn",
               "depreciation_equipment_software_2024_nipa_basis_bn", "depreciation_total_2024_bn"],
              [[c["variant"], c["key_structures"], c["key_equipment"], c["struct_2023"], c["struct_2024"], c["struct_avg"],
                c["eq_2023"], c["eq_2024"], c["eq_avg"], c["total_avg"], c["total_2024"], c["dep_struct"], c["dep_eq"],
                c["dep_total"]] for c in comp])

    # ---------------- 4. the return, national and for the group
    ret_rows = []
    ret = {}
    for label, r in RATES.items():
        for c in comp:
            for basis, stock in (("avg2024", c["total_avg"]), ("end2024", c["total_2024"])):
                nat = r * stock
                ret[(label, c["variant"], basis)] = (nat, s * nat)
                ret_rows.append([label, r, c["variant"], basis, stock, nat, s * nat])
    for c in comp:
        nat = A4_2003_PRIVATE * c["total_avg"]
        ret_rows.append(["A-4 2003 private capital (7%), reported only", A4_2003_PRIVATE, c["variant"], "avg2024",
                         c["total_avg"], nat, s * nat])
    tips = _tips_2024()
    for tenor in ("20 YR", "30 YR"):
        r = tips[tenor] / 100
        nat = r * comp[1]["total_avg"]
        ret_rows.append([f"Treasury par real yield {tenor}, 2024 mean (comparator)", r, "central", "avg2024",
                         comp[1]["total_avg"], nat, s * nat])
    write_csv("capital_return.csv", ["rate_label", "real_rate", "key_variant", "stock_basis", "k12_capital_bn",
                                     "national_return_bn", "group_return_bn"], ret_rows)
    g2 = ret[("A-4 2023 (2%)", "central", "avg2024")][1]
    g3 = ret[("A-4 2003 (3%)", "central", "avg2024")][1]
    g2r = (ret[("A-4 2023 (2%)", "low", "avg2024")][1], ret[("A-4 2023 (2%)", "high", "avg2024")][1])
    g3r = (ret[("A-4 2003 (3%)", "low", "avg2024")][1], ret[("A-4 2003 (3%)", "high", "avg2024")][1])

    # land: absent from BEA; conversion factor only
    land_per10 = {label: 0.10 * r * comp[1]["struct_avg"] * s for label, r in RATES.items()}

    # ---------------- 5. depreciation already inside the school line
    d_edu_fed = bn(D[44][YEAR])
    edu_cfc = {
        "structures_only": d_edu + d_edu_fed,
        "all_assets_rd_in_education": d_edu + d_edu_fed + d_rd + e_rd_in_edu * (d_eq_nipa + d_sw),
        "all_assets_rd_spread": d_edu + d_edu_fed + e_spread * (d_eq_nipa + d_sw + d_rd),
    }
    # pairing check: school line / school fraction must land on the education line's group amount
    edu_amounts = [lines["education_services"]["keys"]["education_mix"][a]["target_bn"] for a in ("personal", "shared")]
    gate("end_shares_are_the_fraction_bounds", set(end_share) == set(school_fraction), f"{end_share} vs {school_fraction}")
    implied_edu = (school_line[0] / end_share[0], school_line[1] / end_share[1])
    gate("school_line_pairs_with_fraction", all(0.95 * min(edu_amounts) <= x <= 1.05 * max(edu_amounts) for x in implied_edu),
         f"low end {school_line[0]:.1f}/{end_share[0]:.4f} and high end {school_line[1]:.1f}/{end_share[1]:.4f} imply "
         f"education lines {implied_edu[0]:.1f}, {implied_edu[1]:.1f} against group amounts "
         f"{min(edu_amounts):.1f}-{max(edu_amounts):.1f}")
    dep_rows, dep = [], {}
    for name, cfc in edu_cfc.items():
        for end, frac, line in (("low_end", end_share[0], school_line[0]), ("high_end", end_share[1], school_line[1])):
            brief = cfc * frac * s
            via_line = cfc / edu_cons_account * line
            dep[(name, end)] = (brief, via_line)
            dep_rows.append([f"brief formula: education CFC ({name}) x school fraction x s", end, cfc, frac, brief,
                             g2, g3, g2 / brief, g3 / brief])
            dep_rows.append([f"CFC share of education consumption ({name}) x school line at full cost", end, cfc, line,
                             via_line, g2, g3, g2 / via_line, g3 / via_line])
    direct = s * comp[1]["dep_total"]
    dep_rows.append(["direct: K-12 depreciation (central key) x s", "both", comp[1]["dep_total"], s, direct,
                     g2, g3, g2 / direct, g3 / direct])
    head = "all_assets_rd_in_education"
    brief_dep = (dep[(head, "low_end")][0], dep[(head, "high_end")][0])
    ratio2 = (g2 / max(brief_dep), g2 / min(brief_dep))
    ratio3 = (g3 / max(brief_dep), g3 / min(brief_dep))
    # the account's own attribution: school fraction x ALL education capital, the split it applies to CFC
    edu_capital = (k_edu_avg + (bn(K[44][YEAR - 1]) + bn(K[44][YEAR])) / 2 + (bn(K[74][YEAR - 1]) + bn(K[74][YEAR])) / 2
                   + e_rd_in_edu * eqsw_avg)
    # [low end, high end], each at its own specification's school share, like the brief-formula depreciation
    prop = {label: tuple(r * edu_capital * frac * s for frac in end_share) for label, r in RATES.items()}
    prop_ratio = {label: r * edu_capital / edu_cfc[head] for label, r in RATES.items()}
    lab2, lab3 = list(RATES)
    for i, end in enumerate(("low_end", "high_end")):
        dep_rows.append([f"proportional split: return on school fraction x all education capital (${edu_capital:.1f}bn) x s,"
                         " over the brief-formula depreciation", end, edu_cfc[head], end_share[i],
                         dep[(head, end)][0], prop[lab2][i], prop[lab3][i], prop_ratio[lab2], prop_ratio[lab3]])
    write_csv("depreciation_comparison.csv",
              ["measure", "main_case_end", "national_cfc_bn", "multiplier", "group_depreciation_bn",
               "group_return_2pct_bn", "group_return_3pct_bn", "ratio_2pct", "ratio_3pct"], dep_rows)

    # ---------------- 6. double count: interest in the account
    engine = ENGINE.read_text()
    default_zero = re.search(r"interest_response:\s*0\s*,", engine) is not None
    routed = 'case "interest": return state.interest_response;' in engine
    pkg_text = "\n".join(p.read_text() for p in PACKAGES)
    overrides_interest = re.search(r"domestic_interest|interest_response", pkg_text) is not None
    starts_default = "Engine.defaultState(m)" in PACKAGES[0].read_text()
    di = lines["domestic_interest"]
    gate("interest_row_held_at_zero_in_main_case",
         default_zero and routed and not overrides_interest and starts_default and di["response_class"] == "interest",
         f"engine default 0: {default_zero}; interest class routed: {routed}; packages override: {overrides_interest}")
    debt = re.sub(r"\s+", " ", DEBT_RESULT.read_text())
    debt_sl = "State-local interest ($274.6bn) sits in the account's interest row" in debt
    debt_nolegacy = "no-legacy framing for state-local gaps" in debt
    gate("debt_legacy_is_federal_only", debt_sl and debt_nolegacy, "debt_legacy RESULT: state-local interest stays in the interest row")
    pop_share = di["keys"]["population"]["personal"]["target_bn"] / di["national_bn"]
    school_interest = us24["interest"] / 1e6
    overlap_if_interest_on = pop_share * school_interest

    # ---------------- 7. comparator: ledger item K (cash basis)
    items = list(csv.DictReader(LEDGER_ITEMS.open()))
    k_union = [r for r in items if r["item"] == "K" and r["arm"] == "central" and r["group"] == "mexican_observed_total"][0]
    k_bn = -float(k_union["total_bn"])
    k_pop = float(k_union["total_bn"]) * 1e9 / float(k_union["per_person"])
    params = json.loads(LEDGER_PARAMS.read_text())["k12"]
    gate("ledger_k_uses_same_f33_totals",
         abs(params["f33_capital_outlay_us"]["value"] - us24["capital_outlay"]) < 1
         and abs(params["f33_interest_on_school_debt_us"]["value"] - us24["interest"]) < 1,
         f"ledger params {params['f33_capital_outlay_us']['value']}+{params['f33_interest_on_school_debt_us']['value']}")
    pupils = list(csv.DictReader(PUPILS.open()))
    cash_out = cash_int = debt_g = eq_pp = 0.0
    tp_total = 0.0
    state_rows = []
    for r in pupils:
        name, tp = r["state"], float(r["target_pupils"])
        f, e, dd = f33[name], f33_enr[name]["enrollment"], f33_debt[name]
        cash_out += tp * f["capital_outlay"] * 1e3 / e
        cash_int += tp * f["interest"] * 1e3 / e
        debt_g += tp * dd["debt"] * 1e3 / e
        tp_total += tp
        state_rows.append([name, tp, f["capital_outlay"] * 1e3 / e, f["interest"] * 1e3 / e, dd["debt"] * 1e3 / e])
    gate("pupils_by_state_sum_to_account", abs(tp_total - emb["target_pupils"]) < 1, f"{tp_total:.0f} vs {emb['target_pupils']:.0f}")
    us_enr = f33_enr["United States"]["enrollment"]
    nat_pp = {"capital_outlay": us24["capital_outlay"] * 1e3 / us_enr, "interest": us24["interest"] * 1e3 / us_enr,
              "debt": f33_debt["United States"]["debt"] * 1e3 / us_enr}
    grp_pp = {"capital_outlay": cash_out / tp_total, "interest": cash_int / tp_total, "debt": debt_g / tp_total}
    state_rows.sort(key=lambda r: -r[1])
    write_csv("capital_per_pupil_where_enrolled.csv",
              ["state", "group_pupils", "capital_outlay_per_pupil_fy2024", "interest_per_pupil_fy2024", "debt_per_pupil_end_fy2024"],
              [["United States (all pupils)", us_enr, nat_pp["capital_outlay"], nat_pp["interest"], nat_pp["debt"]],
               ["Group pupils, weighted by state", tp_total, grp_pp["capital_outlay"], grp_pp["interest"], grp_pp["debt"]]]
              + state_rows)
    nat_cash = (us24["capital_outlay"] + us24["interest"]) / 1e6
    r_cash = (nat_cash - comp[1]["dep_total"]) / comp[1]["total_avg"]      # real rate at which accrual = cash
    ledger_pupils = k_bn * 1e9 / (grp_pp["capital_outlay"] + grp_pp["interest"])
    # item K's pupil base: children 5-17 x the NATIVE ACS pupil ratio for every group (absolute_ledger.py)
    ext_src = LEDGER_EXT.read_text()
    ratio_of = lambda name: (lambda m: float(m.group(1)) / float(m.group(2)))(
        re.search(rf"^{name} = ([0-9.]+) / ([0-9.]+)", ext_src, flags=re.M))
    native_ratio, mexico_ratio = ratio_of("PUPIL_RATIO_NATIVE_ACS"), ratio_of("PUPIL_RATIO_MEXICO_BORN_ACS")
    k_uses_native = re.search(r"cap_person = children \* .*ext\.PUPIL_RATIO_NATIVE_ACS", LEDGER_SCRIPT.read_text()) is not None
    gate("ledger_k_pupil_ratio_read", k_uses_native and 0.7 < native_ratio < mexico_ratio < 1,
         f"item K multiplies children 5-17 by the native ratio {native_ratio:.4f}; Mexico-born ratio {mexico_ratio:.4f}")
    roc = {(r["concept"], r["weighting"]): float(r["R"]) for r in csv.DictReader(ROC.open())}
    district_R = {"capital_outlay": roc[("pp_capout", "mexican_district_share")],
                  "interest": roc[("pp_interest", "mexican_district_share")]}
    acc_nat = {label: comp[1]["dep_total"] + r * comp[1]["total_avg"] for label, r in RATES.items()}
    acc_grp = {label: (min(brief_dep), max(brief_dep), min(brief_dep) + s * r * comp[1]["total_avg"],
                       max(brief_dep) + s * r * comp[1]["total_avg"]) for label, r in RATES.items()}
    cmp_rows = [
        ["cash", "ledger item K, Mexican-origin union (capital outlay + interest, F-33 FY2024 by state)", k_bn, k_pop],
        ["cash", "same F-33 per-pupil amounts on the account's 8.49m pupils, by state", (cash_out + cash_int) / 1e9, tp_total],
        ["cash", "  of which capital outlay", cash_out / 1e9, tp_total],
        ["cash", "  of which interest on school debt", cash_int / 1e9, tp_total],
        ["cash", "national F-33 FY2024 x s", s * nat_cash, s * us_enr],
    ]
    for label in RATES:
        lo_d, hi_d, lo, hi = acc_grp[label]
        cmp_rows.append(["accrual", f"depreciation in school line (brief formula) + return at {label}, low", lo, emb["target_pupils"]])
        cmp_rows.append(["accrual", f"depreciation in school line (brief formula) + return at {label}, high", hi, emb["target_pupils"]])
        cmp_rows.append(["accrual", f"direct K-12 depreciation x s + return at {label}", direct + s * RATES[label] * comp[1]["total_avg"],
                         emb["target_pupils"]])
    write_csv("comparator_item_k.csv", ["basis", "measure", "group_bn", "population_or_pupils"], cmp_rows)
    bridge = [
        ["F-33 FY2024 capital outlay", us24["capital_outlay"] / 1e6],
        ["  construction", us24["construction"] / 1e6],
        ["  land and existing structures", us24["land_existing"] / 1e6],
        ["  equipment", us24_equip / 1e6],
        ["F-33 FY2024 interest on school debt", us24["interest"] / 1e6],
        ["F-33 debt outstanding, end FY2024", f33_debt["United States"]["debt"] / 1e6],
        ["BEA-basis K-12 depreciation 2024 (central key)", comp[1]["dep_total"]],
        ["BEA K-12 capital, average 2024 (central key)", comp[1]["total_avg"]],
        ["return at 2% on it", 0.02 * comp[1]["total_avg"]],
        ["return at 3% on it", 0.03 * comp[1]["total_avg"]],
        ["BEA K-12 structures investment 2024 (central key x FAAt705 l62)", key["central"] * bn(INV[L][YEAR])],
        ["real growth of S&L educational structures stock 2024 (FAAt702 l62)", QK[L][YEAR] / QK[L][YEAR - 1] - 1],
        ["school debt as share of K-12 capital", (f33_debt["United States"]["debt"] / 1e6) / comp[1]["total_avg"]],
        ["interest / average debt proxy (end-FY2024 debt)", us24["interest"] / f33_debt["United States"]["debt"]],
    ]
    write_csv("cash_accrual_bridge_national.csv", ["item", "value_bn_or_ratio"], bridge)

    # ---------------- 8. secondary: other services the main case charges in full (structures by TYPE)
    other = []
    services = [
        ("public_order_safety", "Public order and safety", [(63, "S&L Public safety"), (45, "Federal nondefense Public safety")],
         ["population", "use"], "type maps to function; courthouses may sit in Office"),
        ("health_services", "Health", [(61, "S&L Health care"), (43, "Federal nondefense Health care")], ["health_other"],
         f"the account's line is net of hospital sales (S&L gross ${bn(t317[27][YEAR]):.1f}bn, net ${bn(t317[26][YEAR]):.1f}bn,"
         " NIPA 3.17 l26-27); fee recovery not netted here"),
        ("housing_community_services", "Housing and community services",
         [(58, "S&L Residential"), (68, "S&L Sewer systems"), (69, "S&L Water systems")], ["population"],
         "largely fee-financed utilities and rented housing; the account's line is net of sales"),
        ("income_security_services", "Income security", [], ["cash_assistance"], "no structure type: offices sit in Office [GAP]"),
    ]
    for sid, label, parts, keys, note in services:
        stock = sum((bn(K[l][YEAR - 1]) + bn(K[l][YEAR])) / 2 for l, _ in parts)
        dep_ = sum(bn(D[l][YEAR]) for l, _ in parts)
        for kname in keys:
            share = lines[sid]["keys"][kname]["personal"]["target_bn"] / lines[sid]["national_bn"]
            other.append([label, " + ".join(p for _, p in parts) or "none", stock, dep_, 0.02 * stock, 0.03 * stock, kname,
                          share, share * 0.02 * stock, share * 0.03 * stock, note])
    write_csv("other_services.csv", ["function", "bea_structure_types", "stock_avg2024_bn", "depreciation_2024_bn",
                                     "national_return_2pct_bn", "national_return_3pct_bn", "group_key", "group_share",
                                     "group_return_2pct_bn", "group_return_3pct_bn", "note"], other)

    # ---------------- 9. rate evidence in the primary documents
    a4_23 = (CACHE / "omb_a4_2023.txt").read_text()
    a4_03 = (CACHE / "omb_a4_2003.txt").read_text()
    m25 = (CACHE / "omb_m25_15.txt").read_text()
    gate("a4_2023_states_2_percent", "time preference of 2.0 percent per year" in a4_23, "A-4 (2023) p.76-77")
    gate("a4_2003_states_3_and_7_percent", "using both 3 percent and 7 percent" in a4_03
         and "average before-tax rate of return to private capital" in a4_03,
         "A-4 (2003), Discount Rates, 2. Real Discount Rates of 3 Percent and 7 Percent")
    gate("m25_15_revokes_2023_reinstates_2003", "revoke Office of Management and Budget Circular A-4 of November 9, 2023" in re.sub(r"\s+", " ", m25)
         and "reinstate" in m25, "OMB M-25-15, 12 Feb 2025")

    # ---------------- outputs
    write_csv("gates.csv", ["gate", "pass", "detail"], [[n, "PASS" if ok else "FAIL", d] for n, ok, d in GATES])
    summary = {
        "year": YEAR,
        "pupil_share_s": s,
        "school_line_at_full_cost_bn": school_line,
        "main_case_bn": main_case,
        "school_fraction_bounds": school_fraction,
        "k12_structures_key": key, "k12_structures_key_alternatives": key_alt,
        "k12_equipment_software_key": k_eq,
        "pre_1993_weight_in_2024_stock": pre_weight, "delta_educational_structures": delta,
        "nces_pre1993_anchors": anchors, "nces_1990_anchor_set_aside": implied[1990],
        "national_k12_capital_bn": {c["variant"]: {"avg2024": c["total_avg"], "end2024": c["total_2024"],
                                                   "structures_avg2024": c["struct_avg"], "depreciation_2024": c["dep_total"]}
                                    for c in comp},
        "sl_educational_structures_bn": {"end2023": k_edu[YEAR - 1], "end2024": k_edu[YEAR], "depreciation_2024": d_edu,
                                         "average_age_end2024": AGE[L][YEAR]},
        "group_return_bn": {"2pct": {"central": g2, "low": g2r[0], "high": g2r[1]},
                            "3pct": {"central": g3, "low": g3r[0], "high": g3r[1]},
                            "7pct_reported_only": s * A4_2003_PRIVATE * comp[1]["total_avg"],
                            "treasury_real_2024_mean_pct": tips},
        "land_return_per_10pct_land_to_structure_ratio_bn": land_per10,
        "context": {
            "group_return_with_1990_anchor_bn": {label: s * r * (key["high_with_1990_anchor"] * k_edu_avg + comp[1]["eq_avg"])
                                                 for label, r in RATES.items()},
            "k12_capital_per_account_pupil_usd": comp[1]["total_avg"] * 1e9 / emb["national_pupils"],
            "group_k12_capital_bn": s * comp[1]["total_avg"],
            "structures_share_of_k12_capital": comp[1]["struct_avg"] / comp[1]["total_avg"],
            "education_cfc_share_of_education_line": {k: v / edu_cons_account for k, v in edu_cfc.items()},
            "return_share_of_school_line": {label: [s * r * comp[1]["total_avg"] / x for x in school_line]
                                            for label, r in RATES.items()},
            "main_case_if_added_bn": {label: [x + s * r * comp[1]["total_avg"] for x in main_case]
                                      for label, r in RATES.items()},
            "national_cash_per_pupil_usd": nat_pp["capital_outlay"] + nat_pp["interest"],
            "state_mix_uplift_on_cash": (cash_out + cash_int) / (tp_total * (nat_pp["capital_outlay"] + nat_pp["interest"])),
        },
        "education_cfc_bn": edu_cfc,
        "depreciation_in_school_line_bn": {f"{n}|{e}": v for (n, e), v in dep.items()},
        "depreciation_direct_k12_group_bn": direct,
        "ratio_return_to_depreciation_brief_formula": {"2pct": ratio2, "3pct": ratio3},
        "ratio_return_to_depreciation_direct_k12": {"2pct": g2 / direct, "3pct": g3 / direct},
        "proportional_split": {"education_capital_avg2024_bn": edu_capital,
                               "group_return_bn": {k: list(v) for k, v in prop.items()},
                               "ratio_to_brief_depreciation": prop_ratio},
        "education_share_of_sl_equipment_software_investment": {"rd_in_education": e_rd_in_edu, "rd_spread": e_spread},
        "sl_cfc_2024_fa_minus_nipa_bn": {"state_local_all_assets": eq_gap_sl, "government_equipment": eq_gap_gov},
        "interest": {"main_case_interest_response": 0, "domestic_interest_national_bn": di["national_bn"],
                     "f33_school_interest_bn": school_interest, "population_share": pop_share,
                     "overlap_if_interest_charged_bn": overlap_if_interest_on},
        "item_k": {"ledger_union_bn": k_bn, "ledger_union_population": k_pop,
                   "account_pupils_cash_bn": (cash_out + cash_int) / 1e9,
                   "account_pupils_capital_outlay_bn": cash_out / 1e9, "account_pupils_interest_bn": cash_int / 1e9,
                   "accrual_group_bn": {k: list(v) for k, v in acc_grp.items()}, "national_cash_bn": nat_cash,
                   "national_accrual_bn": acc_nat, "real_rate_equating_national_cash_and_accrual": r_cash,
                   "ledger_pupil_equivalents_at_group_cash_per_pupil": ledger_pupils,
                   "ledger_pupil_ratio_native": native_ratio, "ledger_pupil_ratio_mexico_born": mexico_ratio,
                   "ledger_k_at_mexico_born_ratio_bn": k_bn * mexico_ratio / native_ratio,
                   "accrual_group_direct_bn": {label: direct + s * r * comp[1]["total_avg"] for label, r in RATES.items()}},
        "per_pupil": {"national": nat_pp, "group_state_weighted": grp_pp,
                      "school_cost_lane_district_weighted_R": district_R},
        "inputs": {str(p.relative_to(ROOT)): sha(p) for p in [NIPA3, EMBED, PUPILS, ROC, MAIN, MODEL, ENGINE, *PACKAGES,
                                                             DEBT_RESULT, LEDGER_ITEMS, LEDGER_PARAMS, LEDGER_SCRIPT,
                                                             LEDGER_EXT]},
        "gates_failed": [n for n, ok, _ in GATES if not ok],
    }
    (DERIVED / "summary.json").write_text(json.dumps(summary, indent=1, sort_keys=True) + "\n")

    # ---------------- print
    print(f"K-12 capital, national, avg 2024: ${comp[1]['total_avg']:.1f}bn (key {key['low']:.3f}-{key['high']:.3f}, "
          f"central {key['central']:.3f}); depreciation ${comp[1]['dep_total']:.1f}bn")
    print(f"Group return (s={s:.4f}): 2% ${g2:.2f}bn ({g2r[0]:.2f}-{g2r[1]:.2f}); 3% ${g3:.2f}bn ({g3r[0]:.2f}-{g3r[1]:.2f})")
    print(f"Depreciation in school line (brief formula, all assets): ${min(brief_dep):.2f}-{max(brief_dep):.2f}bn; "
          f"ratio 2% {ratio2[0]:.2f}-{ratio2[1]:.2f}, 3% {ratio3[0]:.2f}-{ratio3[1]:.2f}")
    print(f"Direct K-12 depreciation x s: ${direct:.2f}bn; ratio 2% {g2 / direct:.2f}, 3% {g3 / direct:.2f}")
    print(f"Interest row response in main case: 0; overlap only if charged: ${overlap_if_interest_on:.2f}bn")
    print(f"Item K (cash, union): ${k_bn:.2f}bn; same per pupil on account pupils ${(cash_out + cash_int) / 1e9:.2f}bn; "
          f"accrual 2% ${acc_grp['A-4 2023 (2%)'][2]:.2f}-{acc_grp['A-4 2023 (2%)'][3]:.2f}bn, "
          f"3% ${acc_grp['A-4 2003 (3%)'][2]:.2f}-{acc_grp['A-4 2003 (3%)'][3]:.2f}bn")
    for n, ok, d in GATES:
        print(f"  {'PASS' if ok else 'FAIL'}  {n}: {d}")
    return 0 if all(ok for _, ok, _ in GATES) else 1


def _footnotes(path: Path, sheet: str) -> list[str]:
    rows = list(openpyxl.load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    return [str(r[0]).strip() for r in rows if r and isinstance(r[0], str) and re.match(r"^\d+\.\s", r[0].strip())]


def _nces_capital_outlay() -> dict:
    """NCES Digest 2023 Table 236.10, capital outlay, $mn, keyed by the calendar year a school year ends."""
    rows = list(openpyxl.load_workbook(CACHE / "nces_d23_tabn236.10.xlsx", read_only=True, data_only=True)
                .worksheets[0].iter_rows(values_only=True))
    hdr = rows[2]
    if hdr[9] != "Capital outlay":
        raise SystemExit(f"[BLOCKED] NCES 236.10 column 10 is {hdr[9]!r}")
    out = {}
    for r in rows:
        if r[0] and isinstance(r[0], str) and r[0][:4].isdigit() and isinstance(r[9], (int, float)):
            end = int(r[0][:4]) + 1
            out.setdefault(end, r[9] / 1e3)       # first block is unadjusted dollars (thousands)
    return out


def _cog_education_capital() -> dict:
    rows = list(openpyxl.load_workbook(CACHE / "cog_22slsstab1.xlsx", read_only=True, data_only=True)
                .worksheets[0].iter_rows(values_only=True))
    by = {r[0]: r for r in rows if isinstance(r[0], int)}
    for line, text in ((69, "Education"), (70, "Capital outlay"), (71, "Higher education"), (72, "Capital outlay"),
                       (73, "Elementary & secondary"), (74, "Capital outlay")):
        if str(by[line][1]).strip() != text:
            raise SystemExit(f"[BLOCKED] CoG line {line} is {by[line][1]!r}")
    return {"education": float(by[70][2]), "higher": float(by[72][2]), "elsec": float(by[74][2])}


def _tips_2024() -> dict:
    rows = list(csv.DictReader((CACHE / "treasury_real_yield_2024.csv").open()))
    out = {}
    for tenor in ("10 YR", "20 YR", "30 YR"):
        vals = [float(r[tenor]) for r in rows if r.get(tenor)]
        out[tenor] = sum(vals) / len(vals)
    out["days"] = len(rows)
    return out


if __name__ == "__main__":
    sys.exit(main())
