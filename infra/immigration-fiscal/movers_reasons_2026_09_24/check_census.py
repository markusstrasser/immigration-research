#!/usr/bin/env python3
"""Check the IPUMS coding against the Census Bureau's own files and published tables.

1. Census public ASEC microdata (api.census.gov, movers only) against IPUMS, 2014 (3/8 file),
   2015 and 2019-2025: weighted counts of interstate movers and of California leavers by main
   reason, after mapping the two Census NXTRES code schemes onto IPUMS WHYMOVE.
2. The allocation flags only the Census file carries: whole-supplement imputation (FL_665 = 0)
   and reasons allocated from the hot-deck matrix (I_NXTRES = 5), among US-born adult
   California leavers.
3. Gate G2: national interstate movers against Census CPS Table A-1 (published, thousands), and
   California out-migrants against the ACS state-to-state migration tables (published, with MOE).

Writes derived/gate_census_check.csv, derived/gate_movers_count.csv and
derived/census_allocation_flags.csv. Exits non-zero if the IPUMS-Census microdata comparison or
the CPS Table A-1 comparison fails; the ACS level comparison is recorded, not gated, because the
two surveys measure interstate migration at different levels (see RESULT.md).

Run from the repository root (xlrd reads the pre-2022 .xls tables):
  OPENBLAS_NUM_THREADS=1 uv run --no-project --with xlrd python3 infra/immigration-fiscal/movers_reasons_2026_09_24/check_census.py
"""
import csv
import json
import math
import sys
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
CACHE = HERE / "_cache"
DERIVED = HERE / "derived"

# Census NXTRES -> IPUMS WHYMOVE. ASEC 2019 and earlier: 19 codes; ASEC 2020 on: 20 codes with
# "relationship with unmarried partner" inserted as 4 and "better neighborhood/less crime" at 12.
OLD = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: 7, 8: 8, 9: 9, 10: 10, 11: 11, 12: 12, 13: 19, 14: 13,
       15: 14, 16: 15, 17: 16, 18: 18, 19: 17}
NEW = {1: 1, 2: 2, 3: 3, 4: 20, 5: 4, 6: 5, 7: 6, 8: 7, 9: 8, 10: 9, 11: 10, 12: 11, 13: 12, 14: 19,
       15: 13, 16: 14, 17: 15, 18: 16, 19: 18, 20: 17}


def census_movers(year: int) -> pd.DataFrame:
    rows = json.loads((CACHE / "cps_api" / f"asec_{year}_movers.json").read_text())
    df = pd.DataFrame(rows[1:], columns=rows[0]).apply(pd.to_numeric)
    df["WHYMOVE"] = df["NXTRES"].map(OLD if year <= 2019 else NEW)
    if df.loc[df.NXTRES > 0, "WHYMOVE"].isna().any():
        sys.exit(f"[FAILED] unmapped NXTRES codes in {year}: {sorted(df.loc[df.WHYMOVE.isna(), 'NXTRES'].unique())}")
    return df


def compare_microdata(con) -> list[dict]:
    out = []
    for y in [2014, 2015] + list(range(2019, 2026)):
        c = census_movers(y)
        c = c[(c.MIG_ST.between(1, 56)) & (c.MIG_ST != c.GESTFIPS)]
        # the API's 2014 file is the 3/8 file (HFLAG = 1); compare it with that file, and undo the
        # build's split-file weight factor below
        hsel = "and HFLAG = 1" if y == 2014 else ""
        ip = con.execute(f"""select WHYMOVE, sum(case when MIGSTA1 = 6 then wt else 0 end) as ca,
                                    sum(wt) as allint, count(*) filter (where MIGSTA1 = 6) as n_ca, count(*) as n_all
                             from movers where YEAR = ? and MIGRATE1 = 5 {hsel} group by 1""", [y]).df().set_index("WHYMOVE")
        for label, sel_c, col, ncol in (("interstate movers", c.MIG_ST > 0, "allint", "n_all"),
                                        ("California leavers", c.MIG_ST == 6, "ca", "n_ca")):
            cz = c[sel_c].groupby("WHYMOVE").MARSUPWT.sum()
            cn = c[sel_c].groupby("WHYMOVE").size()
            iz, inn = ip[col], ip[ncol]
            if y == 2014:
                iz = iz / F2014_HFLAG1
            codes = sorted(set(cz.index) | set(iz.index))
            n_diff = max(abs(int(cn.get(k, 0)) - int(inn.get(k, 0))) for k in codes)
            s_diff = max(abs(float(cz.get(k, 0)) / cz.sum() - float(iz.get(k, 0)) / iz.sum()) for k in codes)
            w_diff = max(abs(float(cz.get(k, 0)) - float(iz.get(k, 0))) for k in codes)
            tot_c, tot_i = float(cz.sum()), float(iz.sum())
            # records and reason codes must agree exactly; weights may differ only where the two files
            # carry different weight vintages (IPUMS holds the 2020-based reweights for 2020-2021)
            out.append({"year": y, "group": label, "records_census": int(cn.sum()), "records_ipums": int(inn.sum()),
                        "max_code_count_diff": n_diff, "census_weighted": round(tot_c), "ipums_weighted": round(tot_i),
                        "max_abs_code_diff_weighted": round(w_diff, 1), "max_share_diff_pp": round(100 * s_diff, 3),
                        "codes": len(codes), "pass": n_diff == 0 and int(cn.sum()) == int(inn.sum()) and s_diff < 0.01})
    return out


F2014_HFLAG1 = None


def allocation_flags() -> list[dict]:
    out = []
    frames = []
    for y in list(range(2019, 2026)) + ["2019-2025"]:
        if y == "2019-2025":
            d = pd.concat(frames)
        else:
            c = census_movers(y)
            sel = (c.MIG_ST == 6) & (c.GESTFIPS != 6) & (c.PENATVTY == 57) & (c.A_AGE >= 18)
            d = c[sel]
            frames.append(d)
        w = d.MARSUPWT
        nb = d.WHYMOVE == 11
        clean = (d.FL_665 != 0) & (d.I_NXTRES != 5)
        out.append({"year": y, "n": len(d), "weighted": round(w.sum()),
                    "share_whole_supplement_imputed": round(float(w[d.FL_665 == 0].sum() / w.sum()), 4),
                    "share_reason_allocated_matrix": round(float(w[d.I_NXTRES == 5].sum() / w.sum()), 4),
                    "share_reason_from_household_member": round(float(w[d.I_NXTRES.between(1, 4)].sum() / w.sum()), 4),
                    "nbhd_share_all": round(float(w[nb].sum() / w.sum()), 4),
                    "nbhd_share_excl_imputed": round(float(w[nb & clean].sum() / w[clean].sum()), 4),
                    "n_excl_imputed": int(clean.sum())})
    return out


def acs_s2s_ca_out(year: int) -> tuple[float, float, float, float]:
    """California out-migrants (persons 1+) and the national interstate total, with MOEs, from the
    published ACS state-to-state tables. 2024: the table's own state summary ("tool_input" sheet,
    Out-Movers Estimate and MOE). 2010-2023: the sum of the California column over destination rows,
    MOE by the root sum of squares. 2005-2009 use a multi-block layout and are not parsed."""
    # match the table year, not a revision date ("..._2022_T13_updated_2024_06_27.xlsx")
    files = [f for f in sorted((CACHE / "acs_s2s").glob("*.xls*"))
             if f.name.lower().startswith("state_to_state_migration")
             and f.name.lower().split("table_")[1][:4] == str(year)]
    if not files or year < 2010:
        return (math.nan,) * 4
    if year >= 2024:
        t = pd.read_excel(files[0], sheet_name="tool_input", header=0)
        row = t[t["State"] == "California"].iloc[0]
        states = t["State"].dropna()
        us_in = pd.to_numeric(t.loc[states.index, "In-Movers Estimate"], errors="coerce").sum()
        return float(row["Out-Movers Estimate"]), float(row["Out-Movers MOE"]), float(us_in), math.nan
    df = pd.read_excel(files[0], header=None, sheet_name=0)
    # locate the header row that names origin states and the "Estimate"/"MOE" row beneath it
    hdr_r = next(r for r in range(15) if (df.iloc[r].astype(str).str.strip() == "California").any())
    ca_col = [c for c in range(df.shape[1]) if str(df.iat[hdr_r, c]).strip() == "California"][0]
    lab = df.iloc[:, 0].astype(str).str.strip().str.replace(r"\d+$", "", regex=True)
    states = [s for s in lab.unique() if s not in ("nan", "", "United States", "Puerto Rico")]
    est = moe = 0.0
    for i in df.index[lab.isin(states)]:
        if lab[i] == "California":
            continue
        v, m = df.iat[i, ca_col], df.iat[i, ca_col + 1]
        if isinstance(v, (int, float)) and not pd.isna(v):
            est += float(v)
            moe += float(m) ** 2 if isinstance(m, (int, float)) and not pd.isna(m) else 0.0
    us = df.index[lab == "United States"][0]
    tot_col = [c for c in range(df.shape[1]) if str(df.iat[hdr_r, c]).strip() == "Total"]
    tot_col = [c for c in tot_col if c > 5][0] if tot_col else None
    us_int = float(df.iat[us, tot_col]) if tot_col is not None else math.nan
    us_moe = float(df.iat[us, tot_col + 1]) if tot_col is not None else math.nan
    n_states = int((lab.isin(states)).sum())
    if n_states < 50:
        sys.exit(f"[FAILED] ACS {year} table parsed {n_states} state rows")
    return est, math.sqrt(moe), us_int, us_moe


def cps_table_a1() -> dict[int, list[tuple[str, float]]]:
    """Survey year -> [(row label, movers from a different state)] in persons. Census prints
    several rows for years with revised population controls (2001, 2010, 2011, 2020, 2021)."""
    wb = openpyxl.load_workbook(CACHE / "census_cps_tables" / "hst_mig_a_1.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[int, list] = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] is not None and len(row) > 8 and isinstance(row[8], (int, float)):
            lab = str(row[0])
            yr = int(lab[:4]) if lab[:4].isdigit() else None
            if yr and isinstance(row[1], (int, float)) and row[1] > 1000:  # numbers block, not percentages
                out.setdefault(yr, []).append((lab, float(row[8]) * 1000))
    return out


def main() -> None:
    global F2014_HFLAG1
    DERIVED.mkdir(exist_ok=True)
    con = duckdb.connect()
    con.execute(f"create view movers as select * from read_parquet('{CACHE / 'build' / 'movers.parquet'}')")
    # the build scaled 2014 weights by each file's share of all 2014 person records; recover it
    with open(DERIVED / "gate_extract.csv") as fh:
        for r in csv.DictReader(fh):
            if r["gate"] == "2014 split files":
                F2014_HFLAG1 = float(r["value"].split("HFLAG1 ")[1])
    micro = compare_microdata(con)
    flags = allocation_flags()

    a1 = cps_table_a1()
    ip_nat = dict(con.execute("select YEAR, sum(wt) from movers where MIGRATE1 = 5 group by 1").fetchall())
    ip_ca = dict(con.execute("select YEAR, sum(wt) from movers where MIGRATE1 = 5 and MIGSTA1 = 6 group by 1").fetchall())
    # replicate SE for the California count, 2005+
    rep = con.execute(f"""select m.YEAR, sum(m.wt) as est, {', '.join(f'sum(r.r{i})' for i in range(1, 161))}
                          from movers m join read_parquet('{CACHE / 'build' / 'repwt.parquet'}') r using (YEAR, SERIAL, PERNUM)
                          where m.MIGRATE1 = 5 and m.MIGSTA1 = 6 group by 1""").fetchall()
    ca_se = {r[0]: math.sqrt(4 / 160 * sum((x - r[1]) ** 2 for x in r[2:])) for r in rep}
    count_rows = []
    for y in sorted(ip_nat):
        best = min(a1.get(y, []), key=lambda t: abs(ip_nat[y] - t[1]), default=None)
        row = {"asec_year": y, "cps_interstate_ipums": round(ip_nat[y]),
               "cps_table_a1": best[1] if best else None, "a1_row": best[0] if best else None,
               "a1_diff_pct": round(100 * (ip_nat[y] - best[1]) / best[1], 3) if best else None,
               # published in thousands; weight vintages differ by up to a few tenths of a percent
               "a1_match": (abs(ip_nat[y] - best[1]) <= max(500, 0.005 * best[1]) if best else None),
               "cps_ca_out": round(ip_ca.get(y, 0)), "cps_ca_out_se": round(ca_se[y]) if y in ca_se else None}
        for lag, key in ((1, "acs_prev"), (0, "acs_same")):
            est, moe, us_int, us_moe = acs_s2s_ca_out(y - lag)
            row[f"{key}_year"] = y - lag
            row[f"{key}_ca_out"] = None if math.isnan(est) else round(est)
            row[f"{key}_ca_out_moe90"] = None if math.isnan(moe) else round(moe)
            row[f"{key}_us_interstate"] = None if math.isnan(us_int) else round(us_int)
        count_rows.append(row)
    for r in count_rows:
        a, b = r["acs_prev_ca_out"], r["acs_same_ca_out"]
        if a and b and r["cps_ca_out_se"]:
            mean = (a + b) / 2
            moe_acs = math.sqrt(r["acs_prev_ca_out_moe90"] ** 2 + r["acs_same_ca_out_moe90"] ** 2) / 2
            moe_cps = 1.645 * r["cps_ca_out_se"]
            r["acs_mean_ca_out"] = round(mean)
            r["diff_cps_minus_acs"] = round(r["cps_ca_out"] - mean)
            r["combined_moe90"] = round(math.sqrt(moe_acs ** 2 + moe_cps ** 2))
            r["within_margin"] = abs(r["cps_ca_out"] - mean) <= r["combined_moe90"]
            r["ratio_ca_cps_to_acs"] = round(r["cps_ca_out"] / mean, 3)
            if r["acs_prev_us_interstate"] and r["acs_same_us_interstate"]:
                r["ratio_us_cps_to_acs"] = round(r["cps_interstate_ipums"] / ((r["acs_prev_us_interstate"] + r["acs_same_us_interstate"]) / 2), 3)

    def write(name, rows):
        keys = []
        for r in rows:
            keys += [k for k in r if k not in keys]
        with open(DERIVED / name, "w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=keys, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)

    write("gate_census_check.csv", micro)
    write("census_allocation_flags.csv", flags)
    write("gate_movers_count.csv", count_rows)
    ok_micro = all(r["pass"] for r in micro)
    ok_a1 = all(r["a1_match"] for r in count_rows if r["a1_match"] is not None)
    print(f"{'PASS' if ok_micro else 'FAIL'} IPUMS vs Census microdata, {len(micro)} year-group rows; "
          f"records and codes identical; max weighted share diff {max(r['max_share_diff_pp'] for r in micro)} pp (weight vintages 2020-21)")
    print(f"{'PASS' if ok_a1 else 'FAIL'} national interstate movers vs CPS Table A-1, "
          f"{sum(r['a1_match'] is not None for r in count_rows)} years")
    for r in count_rows:
        if "within_margin" in r:
            print(f"  ASEC {r['asec_year']}: CPS CA out {r['cps_ca_out']:,} (SE {r['cps_ca_out_se']:,}) vs ACS mean "
                  f"{r['acs_mean_ca_out']:,}; ratio {r['ratio_ca_cps_to_acs']} (US ratio {r.get('ratio_us_cps_to_acs')}); "
                  f"within margin {r['within_margin']}")
    if not (ok_micro and ok_a1):
        sys.exit(1)


if __name__ == "__main__":
    main()
