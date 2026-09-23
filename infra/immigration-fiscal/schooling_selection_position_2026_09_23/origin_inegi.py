"""Mexico's schooling distribution by census year, sex and five-year age, from INEGI tabulations.

Builds, for the 2000 and 2010 censuses and the 2020 census (full-count cuestionario basico
tables), the national population aged 15+ by sex x age group x attainment interval on the shared
scale in levels.py. Within-level "no especificado" grades are allocated pro rata to the level's
specified grades; people whose level is unspecified are dropped (shares are of the specified).

Sources (pinned in sources.json):
  2000  CPyV2000_NAL_Caracteristicas_educativas.pdf, Educacion 5, 7, 8 (parts 1-2), 9
  2010  Basico 07_08B, 07_10B, 07_11B, 07_12B (grades), 07_14B (anchor)
  2020  cpv2020_b_eum_07_educacion.xlsx sheet 11 (levels and grades), sheet 13 (ISCED anchor)

Anchors run before anything is written: every parsed row satisfies its table's population
identity; cross-table level totals agree; mean years on the shared scale are compared with INEGI's
published grado promedio by age; the 2020 ISCED 0-2 share by age is compared with sheet 13.

Output: derived/origin_levels.csv (census, sex, age_lo, age_hi, lvl_lo, lvl_hi, count),
        derived/origin_five.csv (five-category shares), derived/origin_anchors.csv.

Needs xlrd (not in the project environment) and poppler's pdftotext on PATH:
  OPENBLAS_NUM_THREADS=1 uv run --no-project --with xlrd python3 infra/immigration-fiscal/schooling_selection_position_2026_09_23/origin_inegi.py
"""
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import xlrd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from levels import FIVE, FIVE_NAMES, YEARS  # noqa: E402

ROOT = HERE.parents[2]
INEGI = HERE / "_cache" / "inegi"
OUT = HERE / "derived"
X2020 = ROOT / "sources/immigration-fiscal/data/external/stage3/inegi/cpv_educacion/cpv2020_b_eum_07_educacion.xlsx"
SEX = {"Total": "T", "Hombres": "H", "Mujeres": "M", "TOTAL": "T", "HOMBRES": "H", "MUJERES": "M"}


def age_bounds(label: str) -> tuple[int, int]:
    s = label.replace("años", "").replace("AÑOS", "").strip()
    if "y m" in s.lower():
        return int(re.match(r"\d+", s).group()), 99
    m = re.match(r"(\d+)\s*-\s*(\d+)", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    a = int(re.match(r"\d+", s).group())
    return a, a


def prorate(parts: dict, ne: float) -> dict:
    """Allocate a within-level unspecified count to the level's specified parts."""
    tot = sum(parts.values())
    if tot <= 0:
        return parts
    return {k: v + ne * v / tot for k, v in parts.items()}


# ---------------------------------------------------------------- 2020
def census_2020() -> tuple[pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(X2020, read_only=True, data_only=True)
    rows, isced = [], []
    for r in wb["11"].iter_rows(min_row=11, values_only=True):
        if r[0] != "Estados Unidos Mexicanos" or r[3] != "Total" or r[2] in ("Total", None):
            continue
        lo, hi = age_bounds(r[2])
        if lo < 15:
            continue
        v = [0 if x is None else x for x in r]
        comp = [5, 6, 8, 9, 10, 12, 13, 14, 15, 17, 18, 19, 21, 22, 23, 24, 25, 27, 28, 29, 30, 31]
        if abs(v[4] - sum(v[i] for i in comp)) > 0:
            raise SystemExit(f"[FAILED] 2020 sheet 11 identity {r[1]} {r[2]}")
        sx = SEX[r[1]]
        pri = prorate({(1, 5): v[8], (6, 6): v[9]}, v[10])
        sec = prorate({(7, 8): v[12], (9, 9): v[13]}, v[14])
        tsec = prorate({(10, 11): v[17], (13, 13): v[18]}, v[19])
        prep = prorate({(10, 11): v[21], (13, 13): v[22]}, v[23])
        lic = prorate({(14, 14): v[27], (16, 16): v[28]}, v[29])
        cells = {(0, 0): v[5] + v[6]}
        for d in (pri, sec):
            for k, x in d.items():
                cells[k] = cells.get(k, 0) + x
        cells["tp"] = v[15]  # tecnico con primaria, no grades in 2020; split below
        for d in (tsec, prep):
            for k, x in d.items():
                cells[k] = cells.get(k, 0) + x
        cells[(13, 13)] += v[24]  # normal basica
        cells[(14, 14)] = lic[(14, 14)]            # licenciatura 1-3 grados
        cells[(15, 15)] = v[25]                    # tecnico con preparatoria (short cycle)
        cells[(16, 16)] = lic[(16, 16)] + v[30]    # licenciatura 4+ grados, posgrado
        rows.append((sx, lo, hi, cells))
    for r in wb["13"].iter_rows(min_row=9, values_only=True):
        if r[0] == "Estados Unidos Mexicanos" and r[2] != "Total":
            lo, hi = age_bounds(r[2])
            less_pri, pri, lo_sec, hi_sec, short, tert, mast, doc, ns, mean_y = r[4:14]
            isced.append({"sex": SEX[r[1]], "age_lo": lo, "age_hi": hi,
                          "isced02_share": (less_pri + pri + lo_sec) / (r[3] - ns),
                          "isced58_share": (short + tert + mast + doc) / (r[3] - ns),
                          "inegi_mean_years": mean_y})
    wb.close()
    return rows, pd.DataFrame(isced)


# ---------------------------------------------------------------- 2010
def x2010(name: str) -> dict:
    sh = xlrd.open_workbook(str(INEGI / "cpv2010" / name)).sheet_by_index(0)
    out = {}
    for r in range(sh.nrows):
        if sh.cell_value(r, 0) != "Estados Unidos Mexicanos":
            continue
        age = sh.cell_value(r, 2)
        if age == "Total":
            continue
        vals = [sh.cell_value(r, c) for c in range(3, sh.ncols)]
        out[(SEX[sh.cell_value(r, 1)], age)] = [0.0 if v == "" else float(v) for v in vals]
    return out


def group_single_ages(tab: dict, groups: list[tuple[int, int]]) -> dict:
    """Aggregate single-year rows (e.g. '17 años') into five-year groups."""
    out = {}
    for (sx, age), vals in tab.items():
        lo, hi = age_bounds(age)
        for g in groups:
            if g[0] <= lo and hi <= g[1]:
                key = (sx, g)
                out[key] = [a + b for a, b in zip(out.get(key, [0.0] * len(vals)), vals)]
    return out


def census_2010() -> tuple[list, pd.DataFrame]:
    e8, e10, e11, e12, e14 = (x2010(f) for f in ("07_08B_ESTATAL.xls", "07_10B_ESTATAL.xls", "07_11B_ESTATAL.xls",
                                                   "07_12B_ESTATAL.xls", "07_14B_ESTATAL.xls"))
    groups = sorted({age_bounds(a) for (_, a) in e14})
    e8, e10, e11, e12, e14 = (group_single_ages(t, groups) for t in (e8, e10, e11, e12, e14))
    rows, anchors = [], []
    for (sx, g), v8 in sorted(e8.items()):
        if g[0] < 15:
            continue
        v10, v11, v14 = e10[(sx, g)], e11[(sx, g)], e14[(sx, g)]
        v12 = e12.get((sx, g), [0.0] * 16)
        # identities: population = sum of components in each table
        for nm, v in (("08", v8), ("10", v10), ("11", v11), ("12", v12)):
            if v[0] and abs(v[0] - sum(v[1:])) > 0.5:
                raise SystemExit(f"[FAILED] 2010 Ed {nm} identity {sx} {g}: {v[0]} vs {sum(v[1:])}")
        # cross-table: Ed 8 secundaria+TP+posbasica == Ed 10 detail + posbasica
        if abs((v8[10] + v8[11]) - sum(v10[4:12])) > 0.5 or abs(v8[12] - v10[12]) > 0.5:
            raise SystemExit(f"[FAILED] 2010 Ed8 vs Ed10 {sx} {g}")
        cells = {(0, 0): v8[1] + v8[2]}
        pri = prorate({(i, i): v8[2 + i] for i in range(1, 7)}, v8[9])
        cells.update(pri)
        sec = prorate({(7, 7): v10[4], (8, 8): v10[5], (9, 9): v10[6]}, v10[7])
        tp = prorate({(7, 7): v10[8], (8, 8): v10[9], (9, 9): v10[10]}, v10[11])
        for k in sec:
            cells[k] = sec[k] + tp[k]
        tsec = prorate({(10, 10): v11[2], (11, 11): v11[3], (13, 13): v11[4] + v11[5]}, v11[6])
        prep = prorate({(10, 10): v11[7], (11, 11): v11[8], (13, 13): v11[9]}, v11[10])
        for k in tsec:
            cells[k] = tsec[k] + prep[k]
        cells[(13, 13)] += v11[11]  # normal basica (no grades in 2010)
        sup_total = v11[12]
        if g[0] >= 20 and abs(sup_total - sum(v12[2:15])) > 0.5:
            raise SystemExit(f"[FAILED] 2010 Ed11 superior vs Ed12 {sx} {g}")
        tpr = v12[2] + v12[3] + v12[4] + v12[5]
        prof = prorate({(14, 14): v12[6] + v12[7] + v12[8], (16, 16): v12[9] + v12[10] + v12[11]}, v12[12])
        top = {(14, 14): prof[(14, 14)], (15, 15): tpr, (16, 16): prof[(16, 16)] + v12[13] + v12[14]}
        if g[0] < 20:  # Ed 12 covers 18+ only; scale its split to the Ed 11 superior count
            s = sum(top.values())
            top = {k: (x / s * sup_total if s else 0.0) for k, x in top.items()}
        cells.update(top)
        rows.append((sx, g[0], g[1], cells))
        anchors.append({"sex": sx, "age_lo": g[0], "age_hi": g[1], "inegi_mean_years": v14[-1],
                        "inegi_media_superior": v14[8] / (v14[0] - v14[10]),
                        "inegi_superior": v14[9] / (v14[0] - v14[10])})
    return rows, pd.DataFrame(anchors)


# ---------------------------------------------------------------- 2000
def pdf_text() -> list[str]:
    txt = INEGI / "cpv2000" / "CPyV2000_NAL_Caracteristicas_educativas.txt"
    if not txt.exists():
        subprocess.run(["pdftotext", "-layout", str(INEGI / "cpv2000" / "CPyV2000_NAL_Caracteristicas_educativas.pdf"),
                        str(txt)], check=True)
    return txt.read_text().splitlines()


def splits(tokens: list[tuple[str, int, int]], ncols: int) -> list[list[int]]:
    """All ways to read digit-group tokens as ncols numbers. Tokens separated by 2+ spaces are
    always different numbers; a single space may be a thousands separator (then the next group has
    exactly three digits) or a column gap."""
    res = []

    def rec(i: int, acc: list[int]):
        if len(acc) > ncols:
            return
        if i == len(tokens):
            if len(acc) == ncols:
                res.append(acc)
            return
        s = tokens[i][0]
        if len(s) > 3 and s != "0":
            return
        num, j = s, i + 1
        rec(j, acc + [int(num)])
        while j < len(tokens) and tokens[j][1] - tokens[j - 1][2] == 1 and len(tokens[j][0]) == 3:
            num += tokens[j][0]
            j += 1
            rec(j, acc + [int(num)])

    rec(0, [])
    return res


def parse_block(lines: list[str], start: int, ncols: int, identity) -> dict:
    """Parse the national block that follows `start`; returns {(sex, age_label): numbers}."""
    i = start
    while not lines[i].lstrip().startswith("ESTADOS UNIDOS MEXICANOS"):
        i += 1
    sex, out = "T", {}
    for line in lines[i:]:
        if re.match(r"^\s*0?1 AGUASCALIENTES", line):
            break
        m = re.match(r"^\s*(ESTADOS UNIDOS MEXICANOS|HOMBRES|MUJERES|\d+\s*-\s*\d+\s+AÑOS|\d+\s+AÑOS|65 Y MÁS AÑOS)(.*)$", line)
        if not m:
            continue
        label, rest = m.group(1), m.group(2)
        if label in ("HOMBRES", "MUJERES"):
            sex = SEX[label]
        off = m.start(2)
        toks = [(t.group(), t.start() + off, t.end() + off) for t in re.finditer(r"\d+", rest)]
        cands = [c for c in splits(toks, ncols) if identity(c)]
        if len(cands) != 1:
            raise SystemExit(f"[FAILED] 2000 parse {label!r}: {len(cands)} readings of {rest.strip()!r}")
        key = "Total" if label in ("ESTADOS UNIDOS MEXICANOS", "HOMBRES", "MUJERES") else label
        out[(sex, key)] = cands[0]
    return out


def census_2000() -> tuple[list, pd.DataFrame]:
    L = pdf_text()
    first = {k: next(i for i, s in enumerate(L) if re.search(rf"EDUCACIÓN {k}\s*$", s)) for k in (5, 7, 8, 9)}
    p2 = next(i for i in range(first[8], len(L)) if "PARTE 2" in L[i])
    full = lambda c: c[0] == sum(c[1:])  # noqa: E731
    e5 = parse_block(L, first[5], 11, full)
    e7 = parse_block(L, first[7], 12, full)
    e9 = parse_block(L, first[9], 11, full)
    e8a = parse_block(L, first[8], 11, lambda c: c[0] >= sum(c[1:]))
    e8b = parse_block(L, p2, 8, lambda c: c[0] >= sum(c[1:]))
    groups = sorted({age_bounds(a) for (_, a) in e5 if a != "Total"})
    e5, e7, e9, e8a, e8b = (group_single_ages({k: v for k, v in t.items() if k[1] != "Total"}, groups)
                            for t in (e5, e7, e9, e8a, e8b))
    rows = []
    for (sx, g), v5 in sorted(e5.items()):
        if g[0] < 15:
            continue
        v7, a, b = e7[(sx, g)], e8a[(sx, g)], e8b[(sx, g)]
        v9 = e9.get((sx, g), [0] * 11)
        if a[0] != sum(a[1:]) + sum(b[1:]) or a[0] != b[0]:
            raise SystemExit(f"[FAILED] 2000 Ed 8 identity {sx} {g}")
        if v5[9] != sum(v7[2:11]):  # posprimaria == secundaria + TP + media superior y superior
            raise SystemExit(f"[FAILED] 2000 Ed5 posprimaria vs Ed7 {sx} {g}")
        if v7[10] != sum(a[2:11]) + sum(b[1:7]):  # media sup. y superior == Ed 8 detail
            raise SystemExit(f"[FAILED] 2000 Ed7 vs Ed8 {sx} {g}")
        if g[0] >= 20 and b[6] != sum(v9[2:10]):
            raise SystemExit(f"[FAILED] 2000 Ed8 superior vs Ed9 {sx} {g}")
        cells = {(0, 0): v5[1]}
        cells.update(prorate({(i, i): v5[1 + i] for i in range(1, 7)}, v5[8]))
        sec = prorate({(7, 7): v7[2], (8, 8): v7[3], (9, 9): v7[4]}, v7[5])
        tp = prorate({(7, 7): v7[6], (8, 8): v7[7], (9, 9): v7[8]}, v7[9])
        for k in sec:
            cells[k] = sec[k] + tp[k]
        ts = prorate({(10, 10): a[2], (11, 11): a[3], (13, 13): a[4] + a[5]}, a[6])
        pr = prorate({(10, 10): a[7], (11, 11): a[8], (13, 13): a[9]}, a[10])
        nb = prorate({(10, 10): b[1], (11, 11): b[2], (13, 13): b[3] + b[4]}, b[5])
        for k in ts:
            cells[k] = ts[k] + pr[k] + nb[k]
        # 2000 "profesional" includes carrera tecnica con preparatoria (Ed 9 note 1), so its first
        # three grades cannot be split into tertiary years vs short-cycle credential.
        prof = prorate({(14, 15): v9[2] + v9[3] + v9[4], (16, 16): v9[5] + v9[6] + v9[7]}, v9[8])
        top = {(14, 15): prof[(14, 15)], (16, 16): prof[(16, 16)] + v9[9]}
        s = sum(top.values())
        top = {k: (x / s * b[6] if s else 0.0) for k, x in top.items()}  # Ed 9 is 18+; scale to Ed 8 superior
        cells.update(top)
        rows.append((sx, g[0], g[1], cells))
    return rows, pd.DataFrame()


# ---------------------------------------------------------------- assemble + anchors
def tp_split_2010() -> dict:
    """Grade split of tecnico con primaria by sex from 2010 Ed 10 (15+), used for 2020 (no grades)."""
    e10 = x2010("07_10B_ESTATAL.xls")
    out = {}
    for sx in "THM":
        tot = [0.0] * 3
        for (s, age), v in e10.items():
            if s == sx and age_bounds(age)[0] >= 15:
                for j in range(3):
                    tot[j] += v[8 + j]
        out[sx] = {(7, 7): tot[0] / sum(tot), (8, 8): tot[1] / sum(tot), (9, 9): tot[2] / sum(tot)}
    return out


def to_long(census: int, rows: list, tp: dict | None = None) -> pd.DataFrame:
    recs = []
    for sx, lo, hi, cells in rows:
        cells = dict(cells)
        if "tp" in cells:
            x = cells.pop("tp")
            for k, f in tp[sx].items():
                if k in cells:
                    cells[k] += x * f
                else:  # 2020 has (7,8) and (9,9) intervals, not single grades 7 and 8
                    key = (7, 8) if k[0] in (7, 8) else k
                    cells[key] = cells.get(key, 0) + x * f
        for (a, b), n in cells.items():
            recs.append({"census": census, "sex": sx, "age_lo": lo, "age_hi": hi, "lvl_lo": a, "lvl_hi": b,
                         "count": float(n)})
    return pd.DataFrame(recs)


def five_shares(d: pd.DataFrame) -> pd.DataFrame:
    out = []
    for key, g in d.groupby(["census", "sex", "age_lo", "age_hi"]):
        tot = g["count"].sum()
        rec = dict(zip(["census", "sex", "age_lo", "age_hi"], key))
        rec["n"] = tot
        for name, (lo, hi) in FIVE.items():
            inside = g[(g.lvl_lo >= lo) & (g.lvl_hi <= hi)]["count"].sum()
            straddle = g[~((g.lvl_lo >= lo) & (g.lvl_hi <= hi)) & (g.lvl_hi >= lo) & (g.lvl_lo <= hi)]
            if len(straddle):
                raise SystemExit(f"[FAILED] {key}: interval straddles {name}")
            rec[name] = inside / tot
        yrs = sum(r["count"] * np.mean(YEARS[int(r.lvl_lo):int(r.lvl_hi) + 1]) for _, r in g.iterrows())
        rec["mean_years_scale"] = yrs / tot
        out.append(rec)
    return pd.DataFrame(out)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    r20, isced20 = census_2020()
    r10, anc10 = census_2010()
    r00, _ = census_2000()
    tp = tp_split_2010()
    d = pd.concat([to_long(2000, r00), to_long(2010, r10), to_long(2020, r20, tp)], ignore_index=True)
    f = five_shares(d)
    # anchors
    a10 = f[f.census == 2010].merge(anc10, on=["sex", "age_lo", "age_hi"])
    a10["ms_scale"] = [
        d[(d.census == 2010) & (d.sex == r.sex) & (d.age_lo == r.age_lo) & (d.lvl_lo >= 10) & (d.lvl_hi <= 13)]["count"].sum() / r.n
        for r in a10.itertuples()]
    a20 = f[f.census == 2020].merge(isced20, on=["sex", "age_lo", "age_hi"])
    a20["c123"] = a20.c1_none_primary_incomplete + a20.c2_primary + a20.c3_lower_secondary
    a20["c5_credential"] = [
        d[(d.census == 2020) & (d.sex == r.sex) & (d.age_lo == r.age_lo) & (d.lvl_lo >= 15)]["count"].sum() / r.n
        for r in a20.itertuples()]
    anchors = pd.concat([
        a10.assign(anchor="2010 grado promedio vs scale mean; Ed 14 media superior and superior shares")[
            ["anchor", "census", "sex", "age_lo", "age_hi", "mean_years_scale", "inegi_mean_years", "ms_scale",
             "inegi_media_superior", "c5_tertiary", "inegi_superior"]],
        a20.assign(anchor="2020 sheet 13: ISCED 0-2 vs C1-C3; grado promedio vs scale mean")[
            ["anchor", "census", "sex", "age_lo", "age_hi", "mean_years_scale", "inegi_mean_years", "c123",
             "isced02_share", "c5_tertiary", "c5_credential", "isced58_share"]]], ignore_index=True)
    d.to_csv(OUT / "origin_levels.csv", index=False)
    f.to_csv(OUT / "origin_five.csv", index=False, float_format="%.6f")
    anchors.to_csv(OUT / "origin_anchors.csv", index=False, float_format="%.4f")
    pd.set_option("display.width", 220)
    print(f[f.sex == "T"].to_string(index=False, float_format=lambda x: f"{x:.3f}"))
    print(anchors.to_string(index=False, float_format=lambda x: f"{x:.3f}"))


if __name__ == "__main__":
    main()
