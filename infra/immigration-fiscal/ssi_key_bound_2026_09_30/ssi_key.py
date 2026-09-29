#!/usr/bin/env python3
"""An administrative key for the SSI line, beside the adopted case's CPS-reported key.

The adopted case (main_case_2026_09_29) keys the SSI line by CPS ASEC 2025 reported receipt (SSI_VAL) on audit
row 4's weights. No administrative source splits SSI by ethnicity (admin_benefit_keys_2026_09_24/RESULT.md, SSI
[BLOCKED]). This script builds the key SSA's published totals give when the group's take-up equals everyone
else's within each state and age group:

  D(s, a)  federal SSI paid in calendar 2024 in state s to age group a (under 18, 18-64, 65 or older). The state
           total is SSA's Annual Statistical Supplement 2025 Table 7.B7 (federal SSI, calendar 2024). It is split
           over the three ages by December 2024 federal payments: SSI Annual Statistical Report 2024 Tables 10 x 11
           (federally administered recipients x average payment) less the state's December supplementation
           (Supplement Table 7.B3, number x average), which is spread over ages at the national supplement's age
           mix (ASR Table 5; California holds 95% of the supplement).
  key      each civilian in cell (s, a) carries D(s, a) / N(s, a), N the cell's civilians on audit row 4's
           weights; the shared allocation splits each SPM unit's sum equally among its members, as the case's
           key does (cps_imputation_keys_2026_09_23/common.py spending_vectors).

Variants beside it bound the key's own bias (none is adopted):
  admin_state_only           federal SSI by state, no age split.
  admin_with_supplement      federally administered totals (federal SSI plus state supplementation; 7.B7 total)
                             by state x age: the key of a line that also paid the supplements.
  admin_unauthorized_barred  the adopted stack's state-aware status flag (combine_onbooks_lane.state_aware_flag)
                             carries no SSI; equal take-up among everyone else in the cell.
  admin_citizens_only        no noncitizen carries SSI (the 1996 bar applied to every noncitizen); equal take-up
                             among citizens in the cell.
  hybrid_cps_within_cell     SSA's federal totals by state x {under 65, 65 or older}; within each cell, the group's
                             share of CPS-reported SSI (the case's vector). It keeps the reported take-up and
                             replaces the CPS's state and age distribution. CPS income items exist only for persons
                             15+, so children's SSI sits on a parent's record and "under 65" pools them.

Writes (derived/):
  ssa_cells.csv    state x age: the parsed SSA inputs and the key dollars D (thousands of dollars)
  key_cells.csv    state x age: D's share, row 4's civilians, the group's population share, CPS-reported SSI
                   and the group's share of it, noncitizen and flagged shares
  shares.csv       variant x allocation: the group's key share, its replicate SE, and its difference from the case
  by_age.csv       national age groups: SSA's dollars, CPS-reported dollars, the group's shares
  by_state.csv     states: SSA's and the CPS key's state shares, the group's shares under both keys
  cps_checks.csv   CPS-reported SSI against SSA's totals; imputed and noncitizen parts of the group's reports
  audit.json       sources and hashes, constants, gates

Gates (exit with [BLOCKED]): the SSA files match the back-test lane's pins; every table holds the 51 areas; the
parsed tables add up (Table 10's categories and ages to its totals, the areas to "All areas", 7.B7's parts to its
totals, the December cells to 7.B3's federal payments within 0.5%); the key's state totals equal 7.B7's federal
SSI (and, for the supplement variant, 7.B7's totals); row 4 reproduces the back-test's union (39,712,493), household
fraction and flag counts; the case's shares reproduce the back-test's stored SSI key shares and their SEs.

Run from the repository root (about a minute; openpyxl for the SSA workbooks):
  OPENBLAS_NUM_THREADS=1 uv run --no-project --with openpyxl python3 infra/immigration-fiscal/ssi_key_bound_2026_09_30/ssi_key.py
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import sys
from pathlib import Path

sys.dont_write_bytecode = True  # read-only imports from other lanes

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
FISCAL = ROOT / "infra/immigration-fiscal"
sys.path.insert(0, str(FISCAL / "cps_imputation_keys_2026_09_23"))
import combine_onbooks_lane as L  # noqa: E402
import combine_status as cs  # noqa: E402
import common as lane  # noqa: E402  the CPS lane's frame, masks and sharing rule

BACKTEST = FISCAL / "backtest_admin_totals_2026_09_28"
ASR = BACKTEST / "_cache/ssa_ssi_asr24.xlsx"
SUPPLEMENT = BACKTEST / "_cache/ssa_supplement2025_7b.xlsx"
PINS = BACKTEST / "derived/sources.json"
STORED_SCALARS = BACKTEST / "derived/prediction_scalars.csv"
STORED_AUDIT = BACKTEST / "derived/prediction_audit.json"
PRIMARY = "asec2025_adopted"  # the back-test's frame for the adopted stack, audit row 4
UNION_PERSONS = 39_712_493.331187886  # audit row 4's union (the lead's 39,712,493; prediction_audit.json union_m)
AGES = ["under_18", "18_64", "65_plus"]
ALLOCATIONS = ["personal", "shared"]
REPS = 161
NAMES = {"Alabama": 1, "Alaska": 2, "Arizona": 4, "Arkansas": 5, "California": 6, "Colorado": 8, "Connecticut": 9,
         "Delaware": 10, "District of Columbia": 11, "Florida": 12, "Georgia": 13, "Hawaii": 15, "Idaho": 16,
         "Illinois": 17, "Indiana": 18, "Iowa": 19, "Kansas": 20, "Kentucky": 21, "Louisiana": 22, "Maine": 23,
         "Maryland": 24, "Massachusetts": 25, "Michigan": 26, "Minnesota": 27, "Mississippi": 28, "Missouri": 29,
         "Montana": 30, "Nebraska": 31, "Nevada": 32, "New Hampshire": 33, "New Jersey": 34, "New Mexico": 35,
         "New York": 36, "North Carolina": 37, "North Dakota": 38, "Ohio": 39, "Oklahoma": 40, "Oregon": 41,
         "Pennsylvania": 42, "Rhode Island": 44, "South Carolina": 45, "South Dakota": 46, "Tennessee": 47,
         "Texas": 48, "Utah": 49, "Vermont": 50, "Virginia": 51, "Washington": 53, "West Virginia": 54,
         "Wisconsin": 55, "Wyoming": 56}
FIPS = sorted(NAMES.values())
ABBR = dict(zip(FIPS, ["AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "DC", "FL", "GA", "HI", "ID", "IL", "IN", "IA",
                       "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ", "NM",
                       "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA",
                       "WV", "WI", "WY"]))
OTHER_AREAS = ("All areas", "Northern Mariana Islands")


def sha(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def sdr(values) -> np.ndarray:
    """Successive-difference replicate standard error along the last axis; index 0 is the full sample."""
    values = np.asarray(values, float)
    return np.sqrt(4 / 160 * np.square(values[..., 1:] - values[..., :1]).sum(axis=-1))


def num(x) -> float:
    """A table cell: a number, SSA's '. . .' (not applicable: 0), or a string with a footnote letter or commas."""
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if s.replace(" ", "") == "...":
        return 0.0
    if s[:2] in ("a ", "b "):
        s = s[2:]
    return float(s.replace(",", ""))


def area_rows(ws, width: int) -> tuple[dict[int, np.ndarray], dict[str, np.ndarray]]:
    """The 51 areas' rows (by FIPS) and the other areas' rows of a state table, first `width` values each."""
    states, other = {}, {}
    for row in ws.iter_rows(values_only=True):
        filled = [v for v in (row or ()) if v is not None]  # labels sit in the first, second or third column (indent)
        if not filled:
            continue
        name = str(filled[0]).strip()
        if name not in NAMES and name not in OTHER_AREAS:
            continue
        values = np.array([num(v) for v in filled[1:]][:width])
        if len(values) != width:
            raise SystemExit(f"[BLOCKED] {ws.title}: {name} has {len(values)} values, not {width}")
        target = states if name in NAMES else other
        key = NAMES.get(name, name)
        if key in target:
            raise SystemExit(f"[BLOCKED] {ws.title}: {name} appears twice")
        target[key] = values
    if sorted(states) != FIPS or sorted(other) != sorted(OTHER_AREAS):
        raise SystemExit(f"[BLOCKED] {ws.title}: not the 51 areas plus {OTHER_AREAS}")
    return states, other


def table5_supplement(ws) -> np.ndarray:
    """ASR Table 5: December 2024 state supplementation payments by age (thousands), from the total-payments block."""
    block = False
    for row in ws.iter_rows(values_only=True):
        filled = [v for v in (row or ()) if v is not None]
        first = str(filled[0]).strip() if filled else ""
        if first.startswith("Total payments"):
            block = True
        elif first.startswith("Average monthly payment"):
            block = False
        elif block and first == "State supplementation":
            values = [num(v) for v in filled[1:]]
            if len(values) != 7 or abs(sum(values[4:]) - values[0]) > 1 or abs(sum(values[1:4]) - values[0]) > 1:
                raise SystemExit(f"[BLOCKED] ASR Table 5: the supplementation row does not add up: {values}")
            return np.array(values[4:])
    raise SystemExit("[BLOCKED] ASR Table 5: no state supplementation row in the total-payments block")


def ssa_inputs(audit: dict) -> pd.DataFrame:
    pins = json.loads(PINS.read_text())
    for path in (ASR, SUPPLEMENT):
        pin = pins[path.name]
        if sha(path) != pin["sha256"]:
            raise SystemExit(f"[BLOCKED] {path.name} does not match the back-test lane's pin")
        audit["sources"][path.name] = {"path": str(path.relative_to(ROOT)), "sha256": pin["sha256"],
                                       "url": pin["url"], "route": pin["route"], "retrieved": pin["retrieved"]}
    asr = openpyxl.load_workbook(ASR, read_only=True, data_only=True)
    sup = openpyxl.load_workbook(SUPPLEMENT, read_only=True, data_only=True)
    t10, t10o = area_rows(asr["Table 10"], 7)   # total, aged, blind, disabled, under 18, 18-64, 65+
    t11, _ = area_rows(asr["Table 11"], 7)      # average monthly payment, same columns
    b7, b7o = area_rows(sup["7.B7"], 3)         # total, federal SSI, federally administered supplementation ($k)
    b3, _ = area_rows(sup["7.B3"], 4)           # federal: number, average; supplementation: number, average
    supp_age = table5_supplement(asr["Table 5"])
    gaps = {}
    # Table 10 adds up: categories and ages to the total, the areas to "All areas".
    for f, v in t10.items():
        if v[1:4].sum() != v[0] or v[4:].sum() != v[0]:
            raise SystemExit(f"[BLOCKED] Table 10: {ABBR[f]}'s categories or ages do not add to its total")
    if np.any(sum(t10.values()) + t10o["Northern Mariana Islands"] != t10o["All areas"]):
        raise SystemExit("[BLOCKED] Table 10: the areas do not add to All areas")
    # 7.B7: federal SSI plus supplementation equals the total in every area ($1k rounding), the areas "All areas".
    gaps["b7_parts_max_k"] = max(abs(v[1] + v[2] - v[0]) for v in b7.values())
    gaps["b7_areas_max_k"] = float(np.abs(sum(b7.values()) + b7o["Northern Mariana Islands"] - b7o["All areas"]).max())
    if gaps["b7_parts_max_k"] > 1 or gaps["b7_areas_max_k"] > 5:
        raise SystemExit(f"[BLOCKED] 7.B7 does not add up: {gaps}")
    rows = []
    for f in FIPS:
        recipients, average = t10[f][4:], t11[f][4:]
        dec_total = recipients * average / 1000                          # $k, federally administered
        supp_dec = b3[f][2] * b3[f][3] / 1000                            # $k, December supplementation
        supp = supp_dec * supp_age / supp_age.sum()
        federal_dec = dec_total - supp
        if np.any(federal_dec <= 0):
            raise SystemExit(f"[BLOCKED] {ABBR[f]}: a non-positive December federal cell {federal_dec}")
        # The December cells against the state's own totals: Table 11's all-ages average and 7.B3's federal payments.
        gaps.setdefault("dec_total_vs_table11_max", 0.0)
        gaps["dec_total_vs_table11_max"] = max(gaps["dec_total_vs_table11_max"],
                                               abs(dec_total.sum() / (t10[f][0] * t11[f][0] / 1000) - 1))
        gaps.setdefault("dec_federal_vs_7b3_max", 0.0)
        gaps["dec_federal_vs_7b3_max"] = max(gaps["dec_federal_vs_7b3_max"],
                                             abs(federal_dec.sum() / (b3[f][0] * b3[f][1] / 1000) - 1))
        federal_cy, total_cy = b7[f][1], b7[f][0]
        key = federal_dec * federal_cy / federal_dec.sum()
        with_supp = dec_total * total_cy / dec_total.sum()
        for i, age in enumerate(AGES):
            rows.append(dict(state=ABBR[f], fips=f, age=age, recipients_dec2024=recipients[i],
                             avg_payment_dec2024=average[i], fa_dollars_dec2024_k=dec_total[i],
                             supplement_dec2024_k=supp[i], federal_dec2024_k=federal_dec[i],
                             federal_cy2024_state_k=federal_cy, total_cy2024_state_k=total_cy,
                             key_federal_cy2024_k=key[i], key_with_supplement_cy2024_k=with_supp[i]))
    if gaps["dec_total_vs_table11_max"] > 0.005 or gaps["dec_federal_vs_7b3_max"] > 0.005:
        raise SystemExit(f"[BLOCKED] the December cells miss the state totals by more than 0.5%: {gaps}")
    cells = pd.DataFrame(rows)
    # The gate the lead named: the key's state totals equal SSA's published totals.
    by_state = cells.groupby("fips")[["key_federal_cy2024_k", "key_with_supplement_cy2024_k"]].sum()
    gaps["key_state_vs_7b7_federal_max_k"] = float(max(abs(by_state.key_federal_cy2024_k[f] - b7[f][1]) for f in FIPS))
    gaps["key_state_vs_7b7_total_max_k"] = float(max(abs(by_state.key_with_supplement_cy2024_k[f] - b7[f][0])
                                                     for f in FIPS))
    if gaps["key_state_vs_7b7_federal_max_k"] > 1e-6 or gaps["key_state_vs_7b7_total_max_k"] > 1e-6:
        raise SystemExit(f"[BLOCKED] the key's state totals are not 7.B7's: {gaps}")
    audit["ssa"] = {"federal_cy2024_51_areas_k": float(sum(v[1] for v in b7.values())),
                    "total_cy2024_51_areas_k": float(sum(v[0] for v in b7.values())),
                    "supplementation_cy2024_51_areas_k": float(sum(v[2] for v in b7.values())),
                    "all_areas_7b7_k": b7o["All areas"].tolist(),
                    "supplement_dec2024_by_age_k_table5": supp_age.tolist(), "gates": gaps}
    print(f"  ✓ SSA: 51 areas in Tables 10, 11, 7.B3 and 7.B7; the key's state totals equal 7.B7's federal SSI "
          f"(max gap {gaps['key_state_vs_7b7_federal_max_k']:.1e} $k; total {audit['ssa']['federal_cy2024_51_areas_k'] / 1e6:.6f} $bn)")
    return cells


def quiet_flag(run) -> np.ndarray:
    """A status flag from the lane's rules; impute() prints its Medicaid caveat, which must still be there."""
    with contextlib.redirect_stderr(io.StringIO()) as note:
        flag = run()
    if "[DEGRADED]" not in note.getvalue():
        raise SystemExit("[BLOCKED] impute_status no longer reports its Medicaid caveat; recheck the rule")
    return flag


class Frame:
    """The CPS lane's frame on audit row 4's weights, with the key's cells."""

    def __init__(self, audit: dict):
        d = lane.load_frame()
        self.civ, self.union = lane.masks(d)
        self.index = lane.spm_index(d)
        W = d[lane.REPS].to_numpy(float)
        self.hf = float(W[self.civ, 0].sum() / lane.RESIDENT)
        arms, info = L.weight_arms(d, W, L.acs_cells())
        self.W = arms["row4"]
        del arms, W
        s_in, hh = cs.status_inputs()
        self.flag = quiet_flag(lambda: L.state_aware_flag(s_in, hh, d, L.STATUS_BLIND_2024))
        self.age = d.A_AGE.to_numpy()
        self.state = pd.Categorical(d.GESTFIPS, categories=FIPS).codes.astype(np.int64)
        if (self.state < 0).any():
            raise SystemExit("[BLOCKED] a person outside the 51 areas")
        self.age3 = np.digitize(self.age, [18, 65])
        self.cell = self.state * 3 + self.age3
        self.cell2 = self.state * 2 + (self.age >= 65)
        self.noncitizen = d.PRCITSHP.eq(5).to_numpy()
        self.ssi = d.SSI_VAL.to_numpy(float)
        self.imputed = lane.person_status(d)["ssi"]
        self.vec = {"personal": self.ssi, "shared": lane.unit_equal(self.ssi, self.index)}
        w0 = self.W[:, 0]
        stored = json.loads(STORED_AUDIT.read_text())
        checks = {"union_persons": (float(w0[self.union].sum()), UNION_PERSONS, 1e-3),
                  "household_fraction": (self.hf, stored["household_fraction"], 1e-12),
                  "flagged_m": (float(w0[self.civ & self.flag].sum() / 1e6), stored[PRIMARY]["flagged_m"], 1e-9),
                  "flagged_union_m": (float(w0[self.union & self.flag].sum() / 1e6),
                                      stored[PRIMARY]["flagged_union_m"], 1e-9)}
        for name, (got, want, tol) in checks.items():
            if abs(got - want) > tol:
                raise SystemExit(f"[BLOCKED] row 4 does not reproduce the back-test's {name}: {got} vs {want}")
        young = self.civ & (self.age < 15) & (self.ssi != 0)
        if young.any():
            raise SystemExit("[BLOCKED] the CPS records SSI for persons under 15; the pooled hybrid cell needs rechecking")
        audit["frame"] = {"persons": int(len(d)), "row4_factors": {k: v for k, v in info.items()},
                          **{k: v[0] for k, v in checks.items()}}
        print(f"  ✓ row 4: union {checks['union_persons'][0]:,.0f} persons, household fraction {self.hf:.10f}, "
              f"flagged {checks['flagged_m'][0]:.6f}M (union {checks['flagged_union_m'][0]:.6f}M)")

    def share(self, v: np.ndarray) -> np.ndarray:
        """The union's share of a key vector, by replicate (161)."""
        return (v[self.union] @ self.W[self.union]) / (v[self.civ] @ self.W[self.civ])

    def cell_sums(self, x, mask: np.ndarray, cell: np.ndarray, k: int) -> np.ndarray:
        """Weighted sums of x over the rows in mask, by cell and replicate (k x 161)."""
        x = np.broadcast_to(np.asarray(x, float), mask.shape)[mask]
        return np.stack([np.bincount(cell[mask], weights=x * self.W[mask, r], minlength=k) for r in range(REPS)],
                        axis=1)

    def admin(self, dollars: np.ndarray, cell: np.ndarray, eligible: np.ndarray) -> dict[str, np.ndarray]:
        """Equal take-up: each eligible civilian in a cell carries the cell's dollars over its eligible civilians."""
        out = {a: np.empty(REPS) for a in ALLOCATIONS}
        m = self.civ & eligible
        for r in range(REPS):
            w = self.W[:, r]
            n = np.bincount(cell[m], weights=w[m], minlength=len(dollars))
            if np.any((dollars > 0) & (n <= 0)):
                raise SystemExit("[BLOCKED] a cell with dollars and no eligible civilians")
            v = np.where(m, dollars[cell] / np.where(n > 0, n, 1.0)[cell], 0.0)
            for a, x in (("personal", v), ("shared", lane.unit_equal(v, self.index))):
                out[a][r] = (x[self.union] @ w[self.union]) / (x[self.civ] @ w[self.civ])
        return out

    def hybrid(self, dollars: np.ndarray, cell: np.ndarray, audit: dict) -> dict[str, np.ndarray]:
        """SSA's dollars by cell, split within each cell by the case's vector (the group's share of reported SSI).
        A cell without reported SSI in a replicate takes the group's population share there (counted in audit)."""
        out, empty = {}, {}
        k = len(dollars)
        for a in ALLOCATIONS:
            v = self.vec[a]
            union, civ, pop_u, pop_c = (self.cell_sums(x, mask, cell, k) for x, mask in (
                (v, self.union), (v, self.civ), (1.0, self.union), (1.0, self.civ)))
            g = np.where(civ > 0, union / np.where(civ > 0, civ, 1.0), pop_u / pop_c)
            out[a] = dollars @ g / dollars.sum()
            none = civ[:, 0] <= 0
            empty[a] = {"cells": int(none.sum()), "dollars_share": float(dollars[none].sum() / dollars.sum())}
        audit["hybrid_cells_without_reports"] = empty
        return out


def main() -> None:
    out_dir = HERE / "derived"
    out_dir.mkdir(exist_ok=True)
    audit = {"sources": {}}
    cells = ssa_inputs(audit)
    fr = Frame(audit)
    w0 = fr.W[:, 0]

    # Cell dollars in the frame's cell order (state index x 3 + age).
    order = {(f, a): i for i, (f, a) in enumerate((f, a) for f in FIPS for a in AGES)}
    cells["cell"] = [order[(f, a)] for f, a in zip(cells.fips, cells.age)]
    cells = cells.sort_values("cell").reset_index(drop=True)
    d3 = cells.key_federal_cy2024_k.to_numpy()
    d3_supp = cells.key_with_supplement_cy2024_k.to_numpy()
    d_state = d3.reshape(51, 3).sum(axis=1)
    d2 = np.column_stack([d3.reshape(51, 3)[:, :2].sum(axis=1), d3.reshape(51, 3)[:, 2]]).ravel()
    everyone = np.ones(len(w0), bool)

    shares = {"case_cps": {a: fr.share(fr.vec[a]) for a in ALLOCATIONS},
              "admin_state_age": fr.admin(d3, fr.cell, everyone),
              "admin_state_only": fr.admin(d_state, fr.state.astype(np.int64), everyone),
              "admin_with_supplement": fr.admin(d3_supp, fr.cell, everyone),
              "admin_unauthorized_barred": fr.admin(d3, fr.cell, ~fr.flag),
              "admin_citizens_only": fr.admin(d3, fr.cell, ~fr.noncitizen),
              "hybrid_cps_within_cell": fr.hybrid(d2, fr.cell2, audit)}

    # The case's shares against the back-test's stored ones.
    stored = pd.read_csv(STORED_SCALARS)
    stored = stored[(stored.frame == PRIMARY) & (stored.check == "group_key_share")].set_index("quantity")
    worst = {"share": 0.0, "se": 0.0}
    for a in ALLOCATIONS:
        s = shares["case_cps"][a]
        worst["share"] = max(worst["share"], abs(s[0] - stored.value[f"ssi_{a}"]))
        worst["se"] = max(worst["se"], abs(sdr(s) - stored.se[f"ssi_{a}"]))
    if worst["share"] > 1e-10 or worst["se"] > 1e-10:
        raise SystemExit(f"[BLOCKED] the case's SSI key shares are not the back-test's stored ones: {worst}")
    audit["gate_case_shares_max_abs"] = worst
    print(f"  ✓ the case's SSI key shares reproduce the back-test's stored ones (|share| {worst['share']:.1e}, "
          f"|SE| {worst['se']:.1e})")

    rows = []
    for variant, by in shares.items():
        for a in ALLOCATIONS:
            s, c = by[a], shares["case_cps"][a]
            rows.append(dict(variant=variant, allocation=a, share=s[0], se=float(sdr(s)), diff_vs_case=s[0] - c[0],
                             diff_se=float(sdr(s - c))))
    table = pd.DataFrame(rows)

    # State x age cells: the key, the population and the CPS reports on row 4 (point estimate).
    def by_cell(x, mask, k=153, cell=fr.cell):
        return np.bincount(cell[mask], weights=(x * w0)[mask], minlength=k)
    pop_c, pop_u = by_cell(1.0, fr.civ), by_cell(1.0, fr.union)
    rep_c, rep_u = by_cell(fr.ssi, fr.civ), by_cell(fr.ssi, fr.union)
    key = cells[["state", "fips", "age", "key_federal_cy2024_k"]].copy()
    key["key_share"] = d3 / d3.sum()
    key["civilians_row4"] = pop_c
    key["group_pop_share"] = pop_u / pop_c
    key["noncitizen_share"] = by_cell(1.0, fr.civ & fr.noncitizen) / pop_c
    key["group_noncitizen_share"] = np.divide(by_cell(1.0, fr.union & fr.noncitizen), pop_u,
                                              out=np.full(153, np.nan), where=pop_u > 0)  # no group member: blank
    key["flagged_share"] = by_cell(1.0, fr.civ & fr.flag) / pop_c
    key["cps_ssi_k"] = rep_c / 1000
    key["cps_ssi_share"] = rep_c / rep_c.sum()
    key["group_share_of_cps_ssi"] = np.divide(rep_u, rep_c, out=np.full(153, np.nan), where=rep_c > 0)
    key["cps_ssi_persons_sample"] = np.bincount(fr.cell[fr.civ & (fr.ssi > 0)], minlength=153)

    # National ages.
    ages = []
    for i, age in enumerate(AGES):
        m = key.age == age
        admin_group = (key.key_federal_cy2024_k[m] * key.group_pop_share[m]).sum() / key.key_federal_cy2024_k[m].sum()
        ages.append(dict(age=age, ssa_federal_cy2024_k=key.key_federal_cy2024_k[m].sum(),
                         ssa_share=key.key_share[m].sum(), ssa_with_supplement_k=d3_supp[m.to_numpy()].sum(),
                         cps_ssi_k=key.cps_ssi_k[m].sum(), cps_share=key.cps_ssi_share[m].sum(),
                         group_pop_share=pop_u[m.to_numpy()].sum() / pop_c[m.to_numpy()].sum(),
                         group_share_of_cps_ssi=rep_u[m.to_numpy()].sum() / rep_c[m.to_numpy()].sum(),
                         group_share_admin_key=admin_group))
    ages = pd.DataFrame(ages)

    # States: both keys' state shares and the group's shares of each state's dollars.
    st = []
    for j, f in enumerate(FIPS):
        m = (key.fips == f).to_numpy()
        dk = key.key_federal_cy2024_k[m]
        st.append(dict(state=ABBR[f], fips=f, ssa_federal_share=dk.sum() / d3.sum(),
                       cps_key_share=rep_c[m].sum() / rep_c.sum(),
                       group_pop_share=pop_u[m].sum() / pop_c[m].sum(),
                       group_share_of_state_cps_ssi=rep_u[m].sum() / rep_c[m].sum() if rep_c[m].sum() > 0 else np.nan,
                       group_share_of_state_admin=(dk * key.group_pop_share[m]).sum() / dk.sum(),
                       group_admin_share_of_national=(dk * key.group_pop_share[m]).sum() / d3.sum(),
                       group_case_share_of_national=rep_u[m].sum() / rep_c.sum()))
    st = pd.DataFrame(st).sort_values("group_admin_share_of_national", ascending=False)

    # The CPS's reports against SSA's totals, and what the group's reports rest on.
    ssa_total = audit["ssa"]["total_cy2024_51_areas_k"] * 1e3
    ssa_65 = d3_supp[(cells.age == "65_plus").to_numpy()].sum() * 1e3
    cps_total, cps_65 = float(fr.ssi[fr.civ] @ w0[fr.civ]), float(fr.ssi[fr.civ & (fr.age >= 65)] @ w0[fr.civ & (fr.age >= 65)])
    union_ssi = fr.ssi * fr.union
    checks = [
        ("cps_reported_ssi_bn", cps_total / 1e9, "row 4 civilians, SSI_VAL (income 2024)"),
        ("ssa_federally_administered_cy2024_bn", ssa_total / 1e9, "7.B7 total, 51 areas (respondents report SSI with any supplement)"),
        ("reporting_ratio_all", cps_total / ssa_total, "CPS over SSA, all ages"),
        ("reporting_ratio_under_65", (cps_total - cps_65) / (ssa_total - ssa_65), "CPS 15-64 (children's SSI on parents' records) over SSA under 65"),
        ("reporting_ratio_65_plus", cps_65 / ssa_65, "CPS 65+ over SSA 65+ (SSA's age split from December payments)"),
        ("group_ssi_imputed_share", float((union_ssi * fr.imputed) @ w0 / (union_ssi @ w0)), "share of the group's reported SSI dollars that the CPS imputed (lane rule: SSI flags or whole-supplement imputation)"),
        ("others_ssi_imputed_share", float((fr.ssi * (fr.civ & ~fr.union) * fr.imputed) @ w0 / ((fr.ssi * (fr.civ & ~fr.union)) @ w0)), "the same for civilians outside the group"),
        ("group_ssi_on_noncitizens_share", float((union_ssi * fr.noncitizen) @ w0 / (union_ssi @ w0)), "share of the group's reported SSI dollars on noncitizens"),
        ("group_ssi_on_flagged_share", float((union_ssi * fr.flag) @ w0 / (union_ssi @ w0)), "share of the group's reported SSI dollars on the state-aware flag (imputed unauthorized)"),
        ("group_noncitizen_pop_share", float(w0[fr.union & fr.noncitizen].sum() / w0[fr.union].sum()), "noncitizens' share of the group"),
        ("group_flagged_pop_share", float(w0[fr.union & fr.flag].sum() / w0[fr.union].sum()), "flagged share of the group"),
        ("group_ssi_recipients_sample", float((fr.union & (fr.ssi > 0)).sum()), "unweighted group persons reporting SSI"),
        ("cps_ssi_recipients_sample", float((fr.civ & (fr.ssi > 0)).sum()), "unweighted civilians reporting SSI"),
    ]
    checks = pd.DataFrame(checks, columns=["quantity", "value", "note"])

    fmt = dict(index=False, float_format="%.10g", lineterminator="\n")
    cells.drop(columns="cell").to_csv(out_dir / "ssa_cells.csv", **fmt)
    key.to_csv(out_dir / "key_cells.csv", **fmt)
    table.to_csv(out_dir / "shares.csv", **fmt)
    ages.to_csv(out_dir / "by_age.csv", **fmt)
    st.to_csv(out_dir / "by_state.csv", **fmt)
    checks.to_csv(out_dir / "cps_checks.csv", **fmt)
    audit["household_fraction"] = fr.hf
    audit["line"] = {"id": "ssi", "key": "ssi", "note": "price_arm.cjs reads the national total from the adopted payload model"}
    (out_dir / "audit.json").write_text(json.dumps(audit, indent=1, sort_keys=True) + "\n")
    with pd.option_context("display.width", 200, "display.max_columns", 20):
        print(table.to_string(index=False))
        print(ages.to_string(index=False))
        print(checks[["quantity", "value"]].to_string(index=False))


if __name__ == "__main__":
    main()
