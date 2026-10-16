#!/usr/bin/env python3
"""Independent raw-source audit of the FY2022 administration scale association.

Run with --data-root MAIN_REPO --out-dir OUTPUT. Requires numpy, scipy, openpyxl.
The Census function amounts include capital; no function-specific current-only
administration series is available in this Table 1. BEA current expenditures are
a different accounting construct and are used solely for the declared mix rule.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
from scipy.stats import t


SOURCES = {
    "census_finance": ("infra/immigration-fiscal/assumption_explorer_2026_09_21/_cache/slf2022.xlsx",
                       "4dd123c5ad520e33d9692f622d389a020e2d3a7e6e782763663854dd84643b9b"),
    "population": ("sources/immigration-fiscal/data/external/census_popest_2024/NST-EST2024-ALLDATA.csv",
                   "b8b50bd345d8a5f33d5c8a28c95aeaa0a50971356f000d9b48276a65fd9022a7"),
    "bea": ("sources/immigration-fiscal/data/external/bea_nipa/Section3All_xls.xlsx",
            "69b5c7aefb38675324887ce31d6feb4fcde7c903ab952db7328da0813096615e"),
}
FUNCTIONS = {
    106: ("financial_administration", "Financial administration"),
    107: ("judicial_and_legal", "Judicial and legal"),
    108: ("general_public_buildings", "General public buildings"),
    109: ("other_administration", "Other governmental administration"),
}
METHOD = {
    "purpose": "Descriptive reconstruction of one administration-population response link; not causal identification of population removal.",
    "source_units": "Census FY2022 direct expenditures by function, thousands of dollars; Census July2022 population, persons; BEA 2024 function amounts, millions of dollars.",
    "baskets": "Full Census lines106-109; same basket excluding line107 judiciary; financial administration line106 alone. The four function totals do not separate current operations from capital outlay.",
    "models": "Log amount on intercept and log population. All50 states, lower48, and 50states+DC. Equal-state and population weights. No trimming or winsorization. Every fit retained.",
    "uncertainty": "Independent WLS sandwich HC3 with t(N-2) intervals. Census source is a census, not a sample. Intervals describe regression/model uncertainty and omit causal/transport uncertainty.",
    "composition": "BEA Table3.16 lines3+4+6 total; federal components44+45+47. Residual = national minus federal component by component; this is net of grants, not the literal published state/local sum. Low=(residual_total*admin_beta + federal_tax_financial*financial_beta)/total; federal executive/legislative and federal other are fixed at zero response. High=admin_beta for whole pool. This is an explicit assumed allocation rule, not a measured causal elasticity. Applied to both admin baskets in every sample/weight specification.",
    "accounting_basis": "BEA3.16 current components are kept distinct from BEA3.15.5 consumption+gross investment. No finite-removal transform is computed here; parent audit owns full-case sensitivity.",
    "validation": "Pin hashes, validate workbook titles/header blocks/line labels/unique states, population and finance national sums, state+local component identities, BEA function and federal+state/local-grant identities, and independent implementation tests.",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n")


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def numeric(value, location: str) -> float:
    # Census explicitly defines '-' as zero/rounds to zero. No other string,
    # empty cell, suppression or not-applicable marker is silently made zero.
    if isinstance(value, str) and value.strip() == "-":
        return 0.0
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"[BLOCKED] missing/nonnumeric amount {location}: {value!r}")
    if not math.isfinite(float(value)):
        raise ValueError(f"[BLOCKED] nonfinite amount {location}")
    return float(value)


def assert_rounding(name: str, residual: float, tolerance: float, checks: list[dict]) -> None:
    checks.append({"identity": name, "residual": residual, "tolerance": tolerance})
    if abs(residual) > tolerance:
        raise ValueError(f"[BLOCKED] {name}: residual {residual} exceeds {tolerance}")


def parse_population(path: Path) -> tuple[dict, list[dict]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        required = {"SUMLEV", "STATE", "NAME", "POPESTIMATE2022"}
        if not required <= set(reader.fieldnames or []):
            raise ValueError("[BLOCKED] population fields missing")
        rows = list(reader)
    national = [r for r in rows if r["SUMLEV"] == "010"]
    if len(national) != 1 or national[0]["NAME"] != "United States":
        raise ValueError("[BLOCKED] ambiguous national population")
    states = [r for r in rows if r["SUMLEV"] == "040" and r["STATE"].zfill(2) != "72"]
    if len(states) != 51 or len({r["STATE"] for r in states}) != 51 or len({r["NAME"] for r in states}) != 51:
        raise ValueError("[BLOCKED] expected 50 unique states and DC")
    result = {r["NAME"]: {"state_fips": r["STATE"].zfill(2), "population_2022": int(r["POPESTIMATE2022"])} for r in states}
    if any(r["population_2022"] <= 0 for r in result.values()):
        raise ValueError("[BLOCKED] nonpositive population")
    checks = []
    assert_rounding("51 state/DC population sum minus United States", sum(r["population_2022"] for r in result.values()) - int(national[0]["POPESTIMATE2022"]), 0, checks)
    return result, checks


def parse_finance(path: Path, populations: dict) -> tuple[list[dict], dict, list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["2022_US_WY"]:
        raise ValueError("[BLOCKED] changed Census workbook sheets")
    sheet = workbook.active
    cells = list(sheet.iter_rows(values_only=True))
    if "State and Local Government Finances" not in cells[0][1] or "2022" not in cells[0][1] or "thousands" not in cells[1][1]:
        raise ValueError("[BLOCKED] changed Census title/units")
    line_rows = {}
    for index, row in enumerate(cells):
        if row[0] in FUNCTIONS:
            if row[0] in line_rows or row[1].strip() != FUNCTIONS[row[0]][1]:
                raise ValueError(f"[BLOCKED] duplicate or changed Census line {row[0]}")
            line_rows[row[0]] = index
    if set(line_rows) != set(FUNCTIONS):
        raise ValueError("[BLOCKED] missing function rows")
    geography_columns = {str(value): col for col, value in enumerate(cells[8]) if col >= 2 and value is not None}
    if set(geography_columns) != set(populations) | {"United States Total"} or len(geography_columns) != 52:
        raise ValueError("[BLOCKED] Census geographies differ from population frame")
    records, mapping, checks = [], [], []
    national = {}
    for name, col in geography_columns.items():
        if cells[9][col].strip() != "State & local" or cells[11][col] != "amount1" or cells[12][col] != 1:
            raise ValueError(f"[BLOCKED] {name} is not a combined amount column")
        if cells[9][col + 2].strip() != "State" or cells[9][col + 3].strip() != "Local":
            raise ValueError(f"[BLOCKED] state/local columns shifted for {name}")
        values = {}
        for line, (key, label) in FUNCTIONS.items():
            row = line_rows[line]
            locations = [f"{get_column_letter(c + 1)}{row + 1}" for c in (col, col + 2, col + 3)]
            combined, state, local = [numeric(cells[row][c], f"{name}:{locations[j]}") for j, c in enumerate((col, col + 2, col + 3))]
            assert_rounding(f"{name} {key}: combined minus state minus local", combined - state - local, 2, checks)
            values[key] = combined
            mapping.append({"geography": name, "function": key, "source_line": line, "source_label": label,
                            "excel_row": row + 1, "combined_cell": locations[0], "state_cell": locations[1], "local_cell": locations[2]})
        values["administration_with_judiciary"] = sum(values.values())
        values["administration_without_judiciary"] = values["administration_with_judiciary"] - values["judicial_and_legal"]
        if name == "United States Total":
            national = values
        else:
            if any(value <= 0 for value in values.values()):
                raise ValueError(f"[BLOCKED] nonpositive function expenditure: {name}")
            records.append({"state": name, **populations[name], **values})
    for key, amount in national.items():
        assert_rounding(f"51 states/DC sum minus national {key}", sum(r[key] for r in records) - amount, 52, checks)
    notes = [str(row[1]) for row in cells if len(row) > 1 and isinstance(row[1], str) and
             any(word in row[1] for word in ("zero or rounds", "intergovernmental", "not subject to sampling", "Response Rates"))]
    return records, {"sheet": sheet.title, "mappings": mapping, "national_thousand_dollars": national,
                     "notes": notes, "current_only_admin_available": False,
                     "definition": "Lines106-109 are direct expenditure by function, including current operations and capital. Table1 does not report separate administration capital outlay."}, checks


def fit_hc3(y: np.ndarray, x: np.ndarray, weight: np.ndarray) -> dict:
    y, x, weight = [np.asarray(a, dtype=float) for a in (y, x, weight)]
    if len(y) < 4 or not (len(y) == len(x) == len(weight)) or not all(np.isfinite(a).all() for a in (y, x, weight)) or (weight <= 0).any():
        raise ValueError("invalid regression inputs")
    weight = weight / weight.mean()
    design = np.column_stack([np.ones(len(x)), x])
    xw = design * np.sqrt(weight)[:, None]
    yw = y * np.sqrt(weight)
    bread = np.linalg.inv(xw.T @ xw)
    beta = bread @ xw.T @ yw
    residual = yw - xw @ beta
    leverage = np.einsum("ij,jk,ik->i", xw, bread, xw)
    if np.max(leverage) >= 1 - 1e-10:
        raise ValueError("singular HC3 leverage")
    adjusted = residual / (1 - leverage)
    covariance = bread @ (xw.T @ (xw * adjusted[:, None] ** 2)) @ bread
    se = math.sqrt(max(float(covariance[1, 1]), 0))
    critical = float(t.ppf(0.975, len(x) - 2))
    return {"n": len(x), "intercept": float(beta[0]), "beta": float(beta[1]),
            "hc3_se": se, "ci95_low": float(beta[1]) - critical * se,
            "ci95_high": float(beta[1]) + critical * se, "t_df": len(x) - 2,
            "maximum_leverage": float(leverage.max())}


def extract_bea(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mappings, checks = [], []

    def extract(sheet_name: str, line: int, label: str) -> float:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if "Millions of dollars" not in str(rows[1][0]):
            raise ValueError("[BLOCKED] BEA units changed")
        columns = [i for i, value in enumerate(rows[7]) if str(value) == "2024"]
        matches = [(i, r) for i, r in enumerate(rows) if str(r[0]) == str(line)]
        if len(columns) != 1 or len(matches) != 1 or not str(matches[0][1][1]).strip().startswith(label):
            raise ValueError(f"[BLOCKED] BEA mapping changed {sheet_name}:{line}")
        row_number, row = matches[0]
        cell = f"{get_column_letter(columns[0] + 1)}{row_number + 1}"
        value = numeric(row[columns[0]], f"{sheet_name}!{cell}")
        mappings.append({"sheet": sheet_name, "line": line, "label": row[1].strip(), "series": row[2],
                         "year": 2024, "cell": cell, "million_dollars": value})
        return value

    current, federal, state_local = {}, {}, {}
    for key, label, all_line, fed_line, sl_line in [
        ("executive_legislative", "Executive and legislative", 3, 44, 82),
        ("tax_financial", "Tax collection and financial management", 4, 45, 83),
        ("other", "Other", 6, 47, 85),
    ]:
        current[key] = extract("T31600-A", all_line, label)
        federal[key] = extract("T31600-A", fed_line, label)
        state_local[key] = extract("T31600-A", sl_line, label)
    gps = extract("T31600-A", 2, "General public service")
    interest = extract("T31600-A", 5, "Interest payments")
    grant = extract("T31700-A", 58, "General public service")
    assert_rounding("BEA GPS noninterest components plus interest minus GPS", sum(current.values()) + interest - gps, 3, checks)
    assert_rounding("BEA federal plus literal SL minus grants minus government noninterest", sum(federal.values()) + sum(state_local.values()) - grant - sum(current.values()), 3, checks)
    residual = {key: current[key] - federal[key] for key in current}
    cgi = {key: extract("T31505-A", line, "General public service") for key, line in
           [("national", 2), ("federal", 42), ("state_local", 79)]}
    assert_rounding("BEA consumption/investment federal plus SL minus national", cgi["federal"] + cgi["state_local"] - cgi["national"], 2, checks)
    return {"units": "million dollars", "year": 2024, "current_national_components": current,
            "current_federal_components": federal, "current_published_state_local_components": state_local,
            "current_national_minus_federal_components": residual, "current_noninterest_component_total": sum(current.values()),
            "current_gps_interest": interest, "current_gps_total": gps,
            "current_gps_federal_grants_to_state_local": grant, "consumption_and_gross_investment_gps": cgi,
            "checks": checks, "mappings": mappings,
            "note": "National-minus-federal residual is net of federal grants. Current is not consumption plus gross investment. Low/high mix is a scenario assumption, not an identified response."}


def estimate_all(records: list[dict], bea: dict) -> tuple[list[dict], list[dict]]:
    fits, composition = [], []
    for sample, selected in [("all50", [r for r in records if r["state_fips"] != "11"]),
                             ("lower48", [r for r in records if r["state_fips"] not in {"02", "11", "15"}]),
                             ("all50_plus_dc", records)]:
        pop = np.array([r["population_2022"] for r in selected])
        for weighting, weight in [("equal_state", np.ones(len(selected))), ("population", pop)]:
            slopes = {}
            for basket in ["administration_with_judiciary", "administration_without_judiciary", "financial_administration"]:
                fit = {"sample": sample, "weight": weighting, "basket": basket,
                       **fit_hc3(np.log([r[basket] for r in selected]), np.log(pop), weight)}
                fits.append(fit)
                slopes[basket] = fit["beta"]
            for basket in ["administration_with_judiciary", "administration_without_judiciary"]:
                admin, financial = slopes[basket], slopes["financial_administration"]
                total = bea["current_noninterest_component_total"]
                residual = sum(bea["current_national_minus_federal_components"].values())
                federal_tax = bea["current_federal_components"]["tax_financial"]
                composition.append({"sample": sample, "weight": weighting, "admin_basket": basket,
                                    "admin_beta": admin, "finance_beta": financial,
                                    "national_minus_federal_share": residual / total,
                                    "federal_tax_financial_share": federal_tax / total,
                                    "federal_fixed_share": 1 - (residual + federal_tax) / total,
                                    "low_assumed_mix": (residual * admin + federal_tax * financial) / total,
                                    "high_assumed_mix": admin})
    return fits, composition


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    dump_json(args.out_dir / "finance_method.json", METHOD)
    receipts = []
    for name, (relative, pin) in SOURCES.items():
        actual = sha256(args.data_root / relative)
        if actual != pin:
            raise ValueError(f"[BLOCKED] changed source hash {name}")
        receipts.append({"source": name, "path": relative, "sha256": actual})
    populations, pop_checks = parse_population(args.data_root / SOURCES["population"][0])
    records, mapping, finance_checks = parse_finance(args.data_root / SOURCES["census_finance"][0], populations)
    bea = extract_bea(args.data_root / SOURCES["bea"][0])
    fits, mix = estimate_all(records, bea)
    write_csv(args.out_dir / "finance_states.csv", records)
    write_csv(args.out_dir / "finance_fits.csv", fits)
    write_csv(args.out_dir / "finance_composition.csv", mix)
    dump_json(args.out_dir / "finance_mapping_checks.json", {"census": mapping, "population_checks": pop_checks, "finance_checks": finance_checks, "bea": bea})
    output_hashes = {p.name: sha256(p) for p in sorted(args.out_dir.glob("finance_*")) if p.name != "finance_sources.json"}
    dump_json(args.out_dir / "finance_sources.json", {"source_files": receipts, "script_sha256": sha256(Path(__file__)), "outputs_sha256": output_hashes})
    print(f"Wrote {len(fits)} fits and {len(mix)} explicit mix scenarios; all source checks passed.")
    for fit in fits:
        if fit["sample"] == "all50":
            print(f"{fit['weight']} {fit['basket']}: {fit['beta']:.6f} [{fit['ci95_low']:.6f}, {fit['ci95_high']:.6f}]")
    for row in mix:
        if row["sample"] == "all50":
            print(f"mix {row['weight']} {row['admin_basket']}: {row['low_assumed_mix']:.6f}–{row['high_assumed_mix']:.6f}")


if __name__ == "__main__":
    main()
