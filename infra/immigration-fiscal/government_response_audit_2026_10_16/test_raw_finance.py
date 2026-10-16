"""Meaningful alignment/fail-loud and independent inference checks."""
import numpy as np
import openpyxl
import pytest
import statsmodels.api as sm

from raw_finance import FUNCTIONS, fit_hc3, numeric, parse_finance


def test_missing_amount_is_not_zero():
    assert numeric("-", "test") == 0
    for value in [None, "(X)", ".....", "", float("nan")]:
        with pytest.raises(ValueError):
            numeric(value, "test")


def test_hc3_matches_independent_statsmodels_wls():
    rng = np.random.default_rng(28)
    x = rng.normal(size=50)
    weight = np.exp(x)
    y = 1.2 + 0.8 * x + rng.normal(size=50) * np.exp(x / 2)
    actual = fit_hc3(y, x, weight)
    expected = sm.WLS(y, sm.add_constant(x), weights=weight).fit(cov_type="HC3", use_t=True)
    assert actual["beta"] == pytest.approx(expected.params[1], rel=1e-11)
    assert actual["hc3_se"] == pytest.approx(expected.bse[1], rel=1e-11)
    assert actual["ci95_low"] == pytest.approx(expected.conf_int()[1, 0], rel=1e-11)
    assert actual["ci95_high"] == pytest.approx(expected.conf_int()[1, 1], rel=1e-11)


def test_known_slope_and_currency_invariance():
    x = np.arange(10, dtype=float)
    y = 2 + 0.83 * x
    actual = fit_hc3(y, x, np.ones(10))
    changed_units = fit_hc3(y + np.log(1000), x, np.ones(10))
    assert actual["beta"] == pytest.approx(0.83)
    assert changed_units["beta"] == pytest.approx(actual["beta"])


def test_finance_amount_columns_and_line_labels(tmp_path):
    path = tmp_path / "finance.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "2022_US_WY"
    sheet.cell(1, 2, "State and Local Government Finances 2022")
    sheet.cell(2, 2, "Dollar amounts in thousands")
    populations = {f"state{i}": {"state_fips": str(i).zfill(2), "population_2022": 100 + i} for i in range(51)}
    for index, name in enumerate(["United States Total", *populations]):
        col = 3 + 5 * index
        sheet.cell(9, col, name)
        sheet.cell(10, col, "State & local")
        sheet.cell(10, col + 2, "State")
        sheet.cell(10, col + 3, "Local")
        sheet.cell(12, col, "amount1")
        sheet.cell(13, col, 1)
        multiplier = 51 if index == 0 else 1
        for line, (_, label) in FUNCTIONS.items():
            row = line + 30
            sheet.cell(row, 1, line)
            sheet.cell(row, 2, label)
            sheet.cell(row, col, 12 * multiplier)
            sheet.cell(row, col + 1, 999999)  # CV must never become the amount.
            sheet.cell(row, col + 2, 5 * multiplier)
            sheet.cell(row, col + 3, 7 * multiplier)
    book.save(path)
    records, mapping, _ = parse_finance(path, populations)
    assert records[0]["financial_administration"] == 12
    assert records[0]["administration_with_judiciary"] == 48
    assert mapping["mappings"][0]["excel_row"] == 136
    sheet.cell(136, 2, "Wrong function")
    book.save(path)
    with pytest.raises(ValueError, match="changed Census line"):
        parse_finance(path, populations)
