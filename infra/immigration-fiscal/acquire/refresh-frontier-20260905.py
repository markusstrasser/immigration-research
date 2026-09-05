#!/usr/bin/env python3
"""Bounded 2026-09-05 data refresh; raw files are never overwritten.

Run: uv run --no-project --with openpyxl python3 infra/immigration-fiscal/acquire/refresh-frontier-20260905.py
Native-First: curl performs HTTP probes/downloads; this task-specific file records provenance and audits heterogeneous source formats, using existing acquisition conventions without changing warehouses.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import subprocess
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
STAGE = ROOT / ".scratch/frontier-20260905/datasets"
LIMIT = 25_000_000
VERA_TREE_SHA = "a6bf48e2627323f01827d52776f0d08023c410ba"
manifest = []
errors = []
previous_records = {}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fetch(name, url, dataset, license_note, kind="data"):
    target = STAGE / name
    reused = target.exists()
    target.parent.mkdir(parents=True, exist_ok=True)
    probe = STAGE / "probes" / (name.replace("/", "_") + ".headers")
    probe.parent.mkdir(exist_ok=True)
    if not target.exists():
        with probe.open("wb") as out:
            subprocess.run(["curl", "-sS", "-I", "-L", "--max-time", "20", url], stdout=out, check=False)
        lengths = re.findall(r"content-length:\s*(\d+)", probe.read_text(errors="replace"), re.I)
        if lengths and int(lengths[-1]) > LIMIT:
            raise ValueError(f"Probe exceeds {LIMIT} bytes: {url}")
        part = target.with_suffix(target.suffix + ".part")
        subprocess.run(["curl", "-sS", "--fail", "-L", "--max-time", "40", "--max-filesize", str(LIMIT), "-o", str(part), url], check=True)
        if part.stat().st_size > LIMIT:
            raise ValueError(f"Downloaded file exceeds bound: {part}")
        part.rename(target)
    digest = sha(target)
    previous = previous_records.get(name)
    if reused and previous and digest != previous["sha256"]:
        raise ValueError(f"Existing source hash differs from manifest; file left unchanged: {target}")
    now = datetime.now(timezone.utc).isoformat()
    acquired_at = previous.get("acquired_at") if reused and previous else (None if reused else now)
    manifest.append(dict(dataset=dataset, file=str(target.relative_to(STAGE)), url=url, kind=kind,
                         acquired_at=acquired_at, verified_at=now, reused=reused,
                         bytes=target.stat().st_size, sha256=digest, license=license_note))
    return target


def safe_fetch(*args, **kwargs):
    try:
        return fetch(*args, **kwargs)
    except Exception as exc:
        errors.append(dict(file=args[0], error=str(exc)))
        print(f"[DEGRADED] {args[0]}: {exc}", flush=True)
        return None


def csv_audit(path, skip=0, header=None):
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig", errors="replace"))))
    if header is None:
        header, rows = rows[skip], rows[skip + 1:]
    else:
        rows = rows[skip:]
    rows = [r for r in rows if r and any(c.strip() for c in r)]
    return dict(columns=header, rows=len(rows), widths=dict(Counter(map(len, rows))),
                missing_by_column={header[i]: sum(i >= len(r) or r[i].strip() in ("", "NA", "(NA)", "(D)", ".") for r in rows) for i in range(len(header))},
                first_rows=rows[:3], last_rows=rows[-3:]), rows


def load_pinned_tree():
    path = fetch("probes/vera-tree.json",
                 f"https://api.github.com/repos/vera-institute/ice-fytd-stats/git/trees/{VERA_TREE_SHA}?recursive=1",
                 "ICE_DETENTION_VERA_ARCHIVE", "GitHub repository metadata", "provenance")
    tree = json.loads(path.read_text())
    if tree.get("sha") != VERA_TREE_SHA or tree.get("truncated") is not False:
        raise ValueError("GitHub tree must match the pinned complete snapshot; existing file left unchanged")
    return tree


def principal_checks(labor, bea):
    """Descriptive checks only; preserve source vintages and program hierarchy."""
    labor_index = {(r["nativity"], r["metric"], r["date"]): r for r in labor}
    comparisons = []
    for nativity in ("native_born", "foreign_born"):
        for month in range(1, 9):
            record = dict(nativity=nativity, month=f"{month:02}")
            for metric in ("population", "employed", "employment_population_ratio", "unemployment_rate"):
                for year in (2025, 2026):
                    record[f"{metric}_{year}"] = float(labor_index[nativity, metric, f"{year}-{month:02}"]["value"])
                record[f"{metric}_change"] = round(record[f"{metric}_2026"] - record[f"{metric}_2025"], 6)
            comparisons.append(record)
    with (STAGE / "bls/matched_calendar_months.csv").open("w") as out:
        writer = csv.DictWriter(out, fieldnames=list(comparisons[0]))
        writer.writeheader(); writer.writerows(comparisons)

    ice_path = STAGE / "ice/2026-07-20_FY26_detentionStats07202026.xlsx"
    wb = openpyxl.load_workbook(ice_path, read_only=True, data_only=True)
    ws = wb["Detention FY26"]
    assert "FY2026" in ws["A36"].value and ws["K37"].value == "Jun"
    assert ws["A38"].value == "Total" and ws["A71"].value == "Release to Remove"
    cutoff = wb["Footnotes"]["B60"].value
    assert "07/11/2026" in cutoff and "Final Book Out" in cutoff
    raw_total, raw_removed = ws["K38"].value, ws["K71"].value
    wb.close()
    ice_csv = STAGE / "ice/monthly_detention_summary.csv"
    processed = [r for r in csv.DictReader(ice_csv.open()) if r["year_month"] == "2026_06"]
    assert len(processed) == 1
    row = processed[0]
    assert row["fiscal_year"] == "2026" and row["table_date"] == "2026-07-11"
    assert raw_total == int(row["final_bookouts"])
    assert raw_removed == int(row["final_bookouts_release_to_remove"])
    percent = round(100 * raw_removed / raw_total, 2)
    assert percent == float(row["percent_of_final_bookouts_as_removals"])
    ice = dict(month="2026-06", raw_file=str(ice_path), raw_sha256=sha(ice_path),
               processed_sha256=sha(ice_csv), sheet="Detention FY26", month_header="K37=Jun",
               total_cell="K38", removed_cell="K71", total=raw_total, removed=raw_removed,
               other_bookouts=raw_total-raw_removed, removed_percent=percent,
               processed_table_date=row["table_date"], processed_file_date=row["file_date"],
               raw_footnote_cell="Footnotes!B60", raw_footnote=cutoff, passed=True)

    # Hierarchy follows the principal CSV's indented descriptions; footnote 14
    # says the 5000 refundable-credit block is already included above.
    hierarchy = {1000: [2000, 3000, 4000], 2000: [2100, 2200, 2300, 2400, 2500, 2600, 2700],
                 2100: [2110, 2120], 2120: [2121, 2122, 2123], 2200: [2210, 2220, 2230],
                 2220: [2221, 2222], 2300: [2310, 2320, 2330, 2340], 2340: [2341, 2342],
                 2400: [2410, 2420], 2420: [2421, 2422, 2423, 2424],
                 2500: [2510, 2520, 2530, 2540], 3000: [3100, 3200, 3300],
                 5000: [5100, 5200, 5300]}
    with bea.open(encoding="utf-8-sig") as source:
        bea_rows = [r for r in csv.DictReader(source) if r.get("LineCode") and r["LineCode"].isdigit()]
    bea_index = {(r["GeoFIPS"].strip().strip('"'), int(r["LineCode"])): r for r in bea_rows}
    assert len(bea_index) == len(bea_rows)
    geographies = sorted({key[0] for key in bea_index})
    checks = []
    for geo in geographies:
        for parent, children in hierarchy.items():
            total = int(bea_index[geo, parent]["2024"])
            components = [int(bea_index[geo, child]["2024"]) for child in children]
            checks.append(dict(geofips=geo, parent=parent, children=children, total=total,
                               components=components, difference=total-sum(components)))
    bea_check = dict(year=2024, unit="thousands of nominal dollars", source=str(bea), sha256=sha(bea),
                     rows=len(bea_rows), geographies=len(geographies), hierarchy=hierarchy,
                     checks=len(checks), nonzero_differences=[r for r in checks if r["difference"] != 0],
                     us=[r for r in checks if r["geofips"] == "00000"],
                     refundable_credits_already_in_total=int(bea_index["00000", 5000]["2024"]))
    result = dict(bls_matched_months=comparisons, ice_completed_month_anchor=ice, bea_hierarchy=bea_check)
    (STAGE / "principal_checks.json").write_text(json.dumps(result, indent=2))
    return result


def main():
    STAGE.mkdir(parents=True, exist_ok=True)
    manifest.clear()
    errors.clear()
    previous_records.clear()
    manifest_path = STAGE / "manifest.json"
    if manifest_path.exists():
        previous_records.update({r["file"]: r for r in json.loads(manifest_path.read_text())["files"]})
    tree = load_pinned_tree()
    audits = {}
    labor = []
    metrics = {"000": ("population", "thousand persons"), "020": ("employed", "thousand persons"),
               "023": ("employment_population_ratio", "percent"), "040": ("unemployment_rate", "percent")}
    for nativity, suffix in (("foreign_born", "73395"), ("native_born", "73413")):
        for prefix, (metric, unit) in metrics.items():
            sid = f"LNU{prefix}{suffix}"
            p = safe_fetch(f"bls/{sid}.json", f"https://api.bls.gov/publicAPI/v2/timeseries/data/{sid}?startyear=2021&endyear=2026", "BLS_CPS_NATIVITY_MONTHLY", "U.S. federal public-use data")
            if p is None:
                continue
            response = json.loads(p.read_text())
            if response.get("status") != "REQUEST_SUCCEEDED":
                errors.append(dict(file=str(p), error=response))
                continue
            for series in response["Results"]["series"]:
                assert series["seriesID"] == sid
                for row in series["data"]:
                    if row["period"] == "M13":
                        continue
                    labor.append(dict(series_id=sid, date=row["year"] + "-" + row["period"][1:], nativity=nativity, metric=metric, unit=unit, value=row["value"], footnotes=json.dumps(row["footnotes"])))
    if labor:
        out = STAGE / "bls/nativity_monthly.csv"
        with out.open("w") as file:
            w = csv.DictWriter(file, fieldnames=list(labor[0])); w.writeheader(); w.writerows(labor)
        audits["BLS"] = dict(rows=len(labor), series=len({r["series_id"] for r in labor}), dates=sorted({r["date"] for r in labor}),
                             duplicates=len(labor)-len({(r["series_id"], r["date"]) for r in labor}),
                             missing_values=sum(r["value"] in ("", "-", ".") for r in labor), latest=[r for r in labor if r["date"] == max(x["date"] for x in labor)])
    # The source's two header rows are kept in raw. These names reflect those headers.
    bps_header = ["survey_date", "state_fips", "county_fips", "region", "division", "county_name"]
    for group in ("one", "two", "three_four", "five_plus", "one_reported", "two_reported", "three_four_reported", "five_plus_reported"):
        bps_header.extend(f"{group}_{x}" for x in ("buildings", "units", "value"))
    for filename in ("co2025a.txt", "co2605c.txt", "co2606c.txt", "co2607c.txt"):
        p = safe_fetch(f"bps/{filename}", f"https://www2.census.gov/econ/bps/County/{filename}", "CENSUS_BPS_COUNTY_2025_2026", "U.S. federal public-use data")
        if p:
            audit, rows = csv_audit(p, skip=3, header=bps_header)
            audit["duplicate_county_keys"] = len(rows) - len({(r[0], r[1], r[2]) for r in rows})
            audit["total_authorized_units"] = sum(sum(int(r[i]) for i in (7, 10, 13, 16)) for r in rows)
            audit["total_reported_units"] = sum(sum(int(r[i]) for i in (19, 22, 25, 28)) for r in rows)
            audits[filename] = audit
    safe_fetch("bps/documentation-index.html", "https://www2.census.gov/econ/bps/Documentation/", "CENSUS_BPS_COUNTY_2025_2026", "U.S. federal documentation", "codebook")
    commit = tree["sha"]
    candidates = [r for r in tree["tree"] if r["path"].startswith("raw/FY2026/") and r["path"].lower().endswith(".xlsx") and r["path"].split("/")[-1][:10] <= "2026-09-05"]
    if not candidates:
        raise ValueError("No verified ICE FY2026 raw paths in GitHub tree")
    selected = sorted(candidates, key=lambda r: r["path"])[-1]
    assert selected["size"] < LIMIT
    base = f"https://raw.githubusercontent.com/vera-institute/ice-fytd-stats/{commit}/"
    for path, kind in ((selected["path"], "data"), ("processed_data/monthly_detention_summary.csv", "data"), ("README.md", "codebook"), ("License.md", "license")):
        p = safe_fetch("ice/" + Path(path).name, base + quote(path), "ICE_DETENTION_VERA_ARCHIVE", "ICE federal raw data; Vera processed series governed by archived License.md", kind)
        if p and p.suffix == ".csv":
            audit, rows = csv_audit(p)
            audits["ICE_processed"] = audit
        if p and p.suffix == ".xlsx":
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            audits["ICE_raw"] = dict(source_git_blob=selected["sha"], archive_tree=commit, sheets={})
            for ws in wb:
                rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
                audits["ICE_raw"]["sheets"][ws.title] = dict(nonempty_rows=len(rows), columns=ws.max_column, preview=rows[:35])
            wb.close()
    # Reuse the existing BEA principal file instead of downloading the same vintage.
    bea = Path("/Volumes/2TBPNY/corpus/bea_data/SAINC/SAINC35__ALL_AREAS_1929_2024.csv")
    if bea.exists():
        audit, rows = csv_audit(bea)
        audit["path"] = str(bea); audit["sha256"] = sha(bea); audit["bytes"] = bea.stat().st_size
        audit["2024_US_components"] = [dict(zip(audit["columns"], r)) for r in rows if r[0].strip().strip('"') == "00000"]
        audits["BEA_existing"] = audit
    else:
        errors.append(dict(file=str(bea), error="Required existing BEA source is unavailable; principal checks cannot complete"))
    if not errors and bea.exists():
        audits["principal_checks"] = principal_checks(labor, bea)
    (STAGE / "manifest.json").write_text(json.dumps(dict(files=manifest, errors=errors), indent=2))
    (STAGE / "inspection.json").write_text(json.dumps(audits, indent=2, default=str))
    print(json.dumps(dict(downloaded=len(manifest), bytes=sum(r["bytes"] for r in manifest), errors=errors,
                         labor=audits.get("BLS"), bps={k:{x:v[x] for x in ("rows", "duplicate_county_keys", "total_authorized_units", "total_reported_units")} for k,v in audits.items() if k.startswith("co")},
                         ice_csv=audits.get("ICE_processed"), ice_raw=selected["path"], bea_existing_rows=audits.get("BEA_existing",{}).get("rows")), indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
