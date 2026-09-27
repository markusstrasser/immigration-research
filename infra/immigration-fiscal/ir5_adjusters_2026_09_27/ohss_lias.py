#!/usr/bin/env python3
"""Adjustments of status against new arrivals, FY2019 and FY2022-FY2025, from the OHSS Legal
Immigration and Adjustment of Status Report (fourth-quarter workbooks, cached in _cache/ohss_lias/).

Every workbook holds the same six tables: 1A (nationality x type of admission), 1B (major class x
type), 2 (refugee arrivals), 3 (naturalizations), 4A and 4B (nonimmigrant admissions). None crosses
nationality with class, and none records an adjuster's prior status or last nonimmigrant class; the
script checks the table titles and stops if that changes. It extracts two marginal series: parents
of US citizens of all nationalities (Table 1B) and Mexican nationals of all classes (Table 1A).
Nationality here, not country of birth as in the OHSS yearbook tables the tail lane uses.

Run from the repository root:
  OPENBLAS_NUM_THREADS=1 uv run --no-project python3 infra/immigration-fiscal/ir5_adjusters_2026_09_27/ohss_lias.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE / "derived"
CACHE = HERE / "_cache/ohss_lias"
# URLs as listed in research/immigration-admission-work-access-2026-09-05.md
URLS = {
    2019: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2019_q4_d_final.xlsx",
    2022: "https://ohss.dhs.gov/sites/default/files/2023-12/"
          "2023_0308_plcy_legal_immigration_adjustment_of_status_report_fy_2022q4_final_d_0.xlsx",
    2023: "https://ohss.dhs.gov/sites/default/files/2024-06/"
          "2024_0507_ohss_legal-immigration-adjustment-of-status-fy-2023q4.xlsx",
    2024: "https://ohss.dhs.gov/sites/default/files/2025-06/"
          "2025_0624_ohss_legal-immigration-adjustment-of-status-fy-2024q4_0.xlsx",
    2025: "https://ohss.dhs.gov/system/files/2026-06/"
          "2026_0604_ohss_legal-immigration-adjustment-of-status-fy-2025q4.xlsx",
}
TITLES = {
    "Table 1A": "PERSONS OBTAINING LAWFUL PERMANENT RESIDENT STATUS BY TYPE OF ADMISSION AND REGION AND COUNTRY OF "
                "NATIONALITY",
    "Table 1B": "PERSONS OBTAINING LAWFUL PERMANENT RESIDENT STATUS BY TYPE AND MAJOR CLASS OF ADMISSION",
    "Table 2": "REFUGEE ARRIVALS BY REGION AND COUNTRY OF NATIONALITY",
    "Table 3": "PERSONS NATURALIZED BY REGION AND COUNTRY OF NATIONALITY",
    "Table 4A": "NONIMMIGRANT ADMISSIONS",
    "Table 4B": "NONIMMIGRANT ADMISSIONS BY CLASS OF ADMISSION",
}
# row label -> (table, pattern); columns: label, total and four quarters, adjustments and four
# quarters, new arrivals and four quarters
SERIES = {"parents_of_us_citizens_all_nationalities": ("Table 1B", r"Parents\d*"),
          "mexico_nationals_all_classes": ("Table 1A", r"Mexico")}


def fetch(fy: int) -> Path:
    path = CACHE / f"lias_fy{fy}_q4.xlsx"
    if not path.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        subprocess.run(["curl", "-sSL", "--fail", "--max-time", "120", "-A", "Mozilla/5.0", "-o", str(path),
                        URLS[fy]], check=True)
    return path


def main() -> None:
    OUT.mkdir(exist_ok=True)
    rows, prov = [], {}
    for fy in URLS:
        path = fetch(fy)
        prov[str(fy)] = {"url": URLS[fy], "file": str(path.relative_to(HERE)),
                         "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        tables = [s for s in wb.sheetnames if s != "TOC"]
        assert tables == list(TITLES), (fy, wb.sheetnames)
        for sheet, title in TITLES.items():
            head = " ".join(str(c) for r in wb[sheet].iter_rows(max_row=4, values_only=True) for c in r if c)
            assert title in head.upper(), (fy, sheet, head[:200])
        for series, (sheet, pattern) in SERIES.items():
            hit = [r for r in wb[sheet].iter_rows(values_only=True)
                   if isinstance(r[0], str) and re.fullmatch(pattern, r[0].strip())]
            assert len(hit) == 1, (fy, series, len(hit))
            total, adjust, new = (int(hit[0][i]) for i in (1, 6, 11))
            assert abs(total - adjust - new) <= 20, (fy, series, total, adjust, new)  # cells rounded to 10 from FY2023
            rows.append({"fy": fy, "series": series, "source_table": sheet, "total": total,
                         "adjustments": adjust, "new_arrivals": new,
                         "adjustment_share": round(adjust / (adjust + new), 4)})
    with (OUT / "ohss_lias_adjust_shares.csv").open("w", newline="") as fh:
        wr = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        wr.writeheader()
        wr.writerows(rows)
    prov["tables_checked"] = TITLES
    (OUT / "ohss_lias_provenance.json").write_text(json.dumps(prov, indent=1, sort_keys=True) + "\n")
    for r in rows:
        print(r)


if __name__ == "__main__":
    main()
