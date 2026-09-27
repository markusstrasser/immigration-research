"""Part 2: deconfound the European native PISA decline (BRIEF_2.md, 1a82c45).

Tests: (1) covariates on the cross-country slope; (2) exclusion bounds; (3) regional panel with country FE on
first differences (= region FE + country x cycle FE in a two-period panel); (4) 2006/2009 pre-trends;
(5) timing of exposure (age at arrival; cohort-matched grade-4 share when derived/grade4_exposure.csv exists).
Inputs: part-1 derived/crosscountry.csv and pisa2022_annex_long.csv; helper outputs derived/exclusion_coverage.csv,
regional_panel.csv, pretrend_inputs.csv, grade4_exposure.csv; _cache/unesco_duration_school_closures.xlsx
(UNESCO UIS, days closed 16/02/2020-30/04/2022); _cache/wb_gdppc_ppp_kd.json (World Bank NY.GDP.PCAP.PP.KD).
Output: derived/deconfound_covariates.csv, derived/deconfound_slopes.csv, derived/exclusion_bounds.csv.
"""
import csv, json, math, os
import numpy as np
import openpyxl
from decompose import ols, cell, share, change, write
from parse_tables import parse_sheet

EUROPE = {"Albania", "Austria", "Belgium", "Bulgaria", "Croatia", "Cyprus", "Czech Republic", "Denmark", "Estonia",
          "Finland", "France", "Germany", "Greece", "Hungary", "Iceland", "Ireland", "Italy", "Latvia", "Lithuania",
          "Luxembourg", "Malta", "Moldova", "Montenegro", "Netherlands", "North Macedonia", "Norway", "Poland",
          "Portugal", "Romania", "Serbia", "Slovak Republic", "Slovenia", "Spain", "Sweden", "Switzerland",
          "Türkiye", "United Kingdom", "Kosovo"}
UNESCO_ALIAS = {"Republic of Korea": "Korea", "Czechia": "Czech Republic", "Slovakia": "Slovak Republic",
                "United Kingdom of Great Britain and Northern Ireland": "United Kingdom",
                "United States of America": "United States", "Republic of Moldova": "Moldova", "Turkey": "Türkiye",
                "Netherlands (Kingdom of the)": "Netherlands"}
ISO3 = {"Australia": "AUS", "Austria": "AUT", "Belgium": "BEL", "Canada": "CAN", "Chile": "CHL", "Colombia": "COL",
        "Costa Rica": "CRI", "Czech Republic": "CZE", "Denmark": "DNK", "Estonia": "EST", "Finland": "FIN",
        "France": "FRA", "Germany": "DEU", "Greece": "GRC", "Hungary": "HUN", "Iceland": "ISL", "Ireland": "IRL",
        "Israel": "ISR", "Italy": "ITA", "Japan": "JPN", "Korea": "KOR", "Latvia": "LVA", "Lithuania": "LTU",
        "Luxembourg": "LUX", "Mexico": "MEX", "Netherlands": "NLD", "New Zealand": "NZL", "Norway": "NOR",
        "Poland": "POL", "Portugal": "PRT", "Slovak Republic": "SVK", "Slovenia": "SVN", "Spain": "ESP",
        "Sweden": "SWE", "Switzerland": "CHE", "Türkiye": "TUR", "United Kingdom": "GBR", "United States": "USA",
        "Albania": "ALB", "Bulgaria": "BGR", "Croatia": "HRV", "Cyprus": "CYP", "Malta": "MLT", "Moldova": "MDA",
        "Montenegro": "MNE", "North Macedonia": "MKD", "Romania": "ROU", "Serbia": "SRB", "Kosovo": "XKX"}
clean = lambda n: n.replace("*", "").strip()
f = lambda v: float(v) if v not in ("", None) else float("nan")


def covariates():
    rows = {clean(r["country"]): r for r in csv.DictReader(open("derived/crosscountry.csv"))}
    wb = openpyxl.load_workbook("_cache/unesco_duration_school_closures.xlsx", read_only=True, data_only=True)
    clos = {}
    for r in list(wb["database"].iter_rows(values_only=True))[1:]:
        if r and r[0]:
            clos[UNESCO_ALIAS.get(r[0], r[0])] = (r[2] / 7, r[3] / 7, r[4] / 7)
    wbj = json.load(open("_cache/wb_gdppc_ppp_kd.json"))[1]
    gdp = {(d["countryiso3code"], d["date"]): d["value"] for d in wbj if d["value"] is not None}
    wq = openpyxl.load_workbook("_cache/statlink_qmuad8.xlsx", read_only=True, data_only=True)
    escs = {}
    for n, _, c, v in parse_sheet(wq["Table I.B1.7.8"]):
        if c.startswith("Difference between PISA 2012 and PISA 2022 / Students' socio-economic status") and \
                c.endswith("/ Non-immigrant students / Dif.") and isinstance(v, (int, float)):
            escs[clean(n)] = float(v)
    out = []
    for c, r in sorted(rows.items()):
        o = dict(country=c, section=r["section"], europe=int(c in EUROPE),
                 dshare_12_22=f(r["dshare_12_22"]), dshare_18_22=f(r["dshare_18_22"]),
                 nat_math_2012=f(r.get("nat_math_2012")), dnat_math_12_22=f(r.get("dnat_math_12_22")),
                 dnat_math_18_22=f(r.get("dnat_math_18_22")), dnat_reading_12_22=f(r.get("dnat_reading_12_22")),
                 dnat_science_12_22=f(r.get("dnat_science_12_22")))
        cw = clos.get(c)
        o["closure_full_wk"], o["closure_partial_wk"], o["closure_total_wk"] = cw if cw else (float("nan"),) * 3
        iso = ISO3.get(c)
        g12, g22 = gdp.get((iso, "2012")), gdp.get((iso, "2022"))
        o["dlog_gdppc_12_22"] = math.log(g22 / g12) if g12 and g22 else float("nan")
        o["dnat_escs_12_22"] = escs.get(c, float("nan"))
        out.append(o)
    return out


def fit(rows, y, xs, label, sample, test):
    use = [r for r in rows if all(not math.isnan(r[k]) for k in [y] + xs)]
    if len(use) < len(xs) + 3:
        return None
    b, se, r2, e, _ = ols([r[y] for r in use], np.array([[r[k] for k in xs] for r in use]))
    res = dict(test=test, sample=sample, outcome=y, regressors="+".join(xs), label=label, n=len(use),
               slope_per10pp=round(10 * b[1], 2), se_per10pp=round(10 * se[1], 2), r2=round(r2, 3),
               intercept=round(b[0], 2))
    for k, bk, sk in zip(xs[1:], b[2:], se[2:]):
        res[f"coef_{k}"] = round(bk, 3); res[f"se_{k}"] = round(sk, 3)
    de = next((i for i, r in enumerate(use) if r["country"] == "Germany"), None)
    res["germany_resid"] = round(e[de], 2) if de is not None else ""
    loo = []
    for i in range(len(use)):
        keep = [j for j in range(len(use)) if j != i]
        bk = ols([use[j][y] for j in keep], np.array([[use[j][k] for k in xs] for j in keep]))[0]
        loo.append((10 * bk[1], use[i]["country"]))
    res["loo_min_per10"] = f"{min(loo)[0]:.2f} (drop {min(loo)[1]})"
    res["loo_max_per10"] = f"{max(loo)[0]:.2f} (drop {max(loo)[1]})"
    return res


def test1(cov):
    out = []
    specs = [("share only", ["dshare_12_22"]), ("+2012 level", ["dshare_12_22", "nat_math_2012"]),
             ("+closure total wk", ["dshare_12_22", "closure_total_wk"]),
             ("+closure full wk", ["dshare_12_22", "closure_full_wk"]),
             ("+dlog GDP pc", ["dshare_12_22", "dlog_gdppc_12_22"]),
             ("+native ESCS change", ["dshare_12_22", "dnat_escs_12_22"]),
             ("joint", ["dshare_12_22", "nat_math_2012", "closure_total_wk", "dlog_gdppc_12_22", "dnat_escs_12_22"]),
             ("joint without GDP", ["dshare_12_22", "nat_math_2012", "closure_total_wk", "dnat_escs_12_22"])]
    for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1)):
        rows = [r for r in cov if keep(r)]
        for lab, xs in specs:
            for y in ("dnat_math_12_22",) + (("dnat_reading_12_22", "dnat_science_12_22") if lab in ("share only", "joint") else ()):
                x = fit(rows, y, xs, lab, sample, "1 covariates")
                if x: out.append(x)
        for lab, xs in (("2018-22 share only", ["dshare_18_22"]), ("2018-22 +closure total wk", ["dshare_18_22", "closure_total_wk"])):
            x = fit(rows, "dnat_math_18_22", xs, lab, sample, "1 covariates")
            if x: out.append(x)
    return out


def natl(table, country, cyc, unit):
    wb = natl.wb
    for n, _, c, v in natl.cache.setdefault(table, parse_sheet(wb[table])):
        if clean(n) == country and c.startswith(f"PISA {cyc} /") and c.endswith(unit) and isinstance(v, (int, float)):
            return float(v)
    return float("nan")


def test2(cov):
    ex = {}
    for r in csv.DictReader(open("derived/exclusion_coverage.csv")):
        ex[(clean(r["country"]), int(r["cycle"]))] = r
    natl.wb = openpyxl.load_workbook("_cache/statlink_wh9d4z.xlsx", read_only=True, data_only=True); natl.cache = {}
    xs = {clean(r["country"]): r for r in csv.DictReader(open("derived/crosscountry.csv"))}
    bounds, slopes = [], []
    for c, r in sorted(xs.items()):
        o = dict(country=c, europe=int(c in EUROPE), section=r["section"])
        for cyc in (2012, 2018, 2022):
            e = ex.get((c, cyc))
            o[f"excl_{cyc}"] = f(e["overall_excl_rate_pct"]) if e else float("nan")
            o[f"within_excl_{cyc}"] = f(e["within_school_excl_rate_pct"]) if e else float("nan")
            o[f"ci3_{cyc}"] = f(e["ci3"]) if e else float("nan")
            o[f"sd_{cyc}"] = natl("Table I.B1.5.10", c, cyc, "Standard deviation / S.D.")
        nat = {2012: f(r.get("nat_math_2012")), 2022: f(r.get("nat_math_2012")) + f(r.get("dnat_math_12_22")),
               2018: f(r.get("nat_math_2012")) + f(r.get("dnat_math_12_18"))}
        for z, lab in ((1.6449, "p5"), (1.2816, "p10")):
            # excluded pupils score at national percentile ~ national mean - z*SD; natives' level correction
            # = -e * (native mean - (native mean - z*SD)) is an upper bound (all excluded treated as natives)
            corr = {cyc: -o[f"excl_{cyc}"] / 100 * z * o[f"sd_{cyc}"] for cyc in (2012, 2018, 2022)}
            o[f"dnat_12_22_adj_{lab}"] = round(f(r.get("dnat_math_12_22")) + corr[2022] - corr[2012], 2)
            o[f"dnat_12_18_adj_{lab}"] = round(f(r.get("dnat_math_12_18")) + corr[2018] - corr[2012], 2)
            o[f"dnat_18_22_adj_{lab}"] = round(f(r.get("dnat_math_18_22")) + corr[2022] - corr[2018], 2)
        o["dnat_math_12_22"] = f(r.get("dnat_math_12_22")); o["dnat_math_12_18"] = f(r.get("dnat_math_12_18"))
        o["dnat_math_18_22"] = f(r.get("dnat_math_18_22"))
        o["dexcl_12_22"] = o["excl_2022"] - o["excl_2012"]; o["dwithin_12_22"] = o["within_excl_2022"] - o["within_excl_2012"]
        o["dci3_12_22"] = o["ci3_2022"] - o["ci3_2012"]; o["dshare_12_22"] = f(r["dshare_12_22"])
        bounds.append(o)
    for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1)):
        rows = [r for r in bounds if keep(r)]
        for y in ("dexcl_12_22", "dwithin_12_22", "dci3_12_22"):
            x = fit(rows, y, ["dshare_12_22"], f"{y} on share change", sample, "2 exclusion")
            if x: slopes.append(x)
        for y in ("dnat_math_12_22", "dnat_12_22_adj_p10", "dnat_12_22_adj_p5"):
            x = fit(rows, y, ["dshare_12_22"], f"native change ({y})", sample, "2 exclusion")
            if x: slopes.append(x)
    return bounds, slopes


def test3():
    P = {}
    for r in csv.DictReader(open("derived/regional_panel.csv")):
        P[(r["country"], r["region"], int(r["cycle"]))] = r
    regions = sorted({(c, g) for c, g, _ in P})
    out = []
    for subj, c0, c1 in (("math", 2012, 2022), ("read", 2018, 2022), ("read", 2012, 2022)):
        rows = []
        for c, g in regions:
            a, b = P.get((c, g, c0)), P.get((c, g, c1))
            if not a or not b or "" in (a[f"nat_{subj}"], b[f"nat_{subj}"], a["imm_share"], b["imm_share"]):
                continue
            rows.append(dict(country=c, region=g, dnat=float(b[f"nat_{subj}"]) - float(a[f"nat_{subj}"]),
                             dshare=float(b["imm_share"]) - float(a["imm_share"])))
        ctry = sorted({r["country"] for r in rows if sum(1 for q in rows if q["country"] == r["country"]) >= 2})
        rows_fe = [r for r in rows if r["country"] in ctry]
        for lab, rs, fe in (("pooled, no FE", rows, False), ("country FE", rows_fe, True)):
            if len(rs) < 5:
                continue
            X = [[r["dshare"]] + ([1.0 if r["country"] == k else 0.0 for k in ctry[1:]] if fe else []) for r in rs]
            b, se, r2, e, _ = ols([r["dnat"] for r in rs], np.array(X))
            loo = []
            for i in range(len(rs)):
                keep = [j for j in range(len(rs)) if j != i]
                Xk = np.array([X[j] for j in keep])
                Xk = Xk[:, [0] + [m for m in range(1, Xk.shape[1]) if Xk[:, m].any()]]
                loo.append((ols([rs[j]["dnat"] for j in keep], Xk)[0][1], rs[i]["region"]))
            lo, hi = min(loo), max(loo)
            out.append(dict(test="3 regional", sample=f"{subj} {c0}-{c1}", outcome=f"dnat_{subj}", regressors="dshare",
                            label=lab, n=len(rs), slope_per10pp=round(10 * b[1], 2), se_per10pp=round(10 * se[1], 2),
                            r2=round(r2, 3), intercept=round(b[0], 2), countries=";".join(sorted({r['country'] for r in rs})),
                            loo_min_per10=f"{10 * lo[0]:.2f} (drop {lo[1]})", loo_max_per10=f"{10 * hi[0]:.2f} (drop {hi[1]})"))
        for c in ctry:
            rs = [r for r in rows_fe if r["country"] == c]
            if len(rs) >= 4:
                b, se, r2, e, _ = ols([r["dnat"] for r in rs], np.array([[r["dshare"]] for r in rs]))
                out.append(dict(test="3 regional", sample=f"{subj} {c0}-{c1}", outcome=f"dnat_{subj}", regressors="dshare",
                                label=f"{c} only", n=len(rs), slope_per10pp=round(10 * b[1], 2),
                                se_per10pp=round(10 * se[1], 2), r2=round(r2, 3), intercept=round(b[0], 2), countries=c))
        write(f"derived/regional_changes_{subj}_{c0}_{c1}.csv", rows)
    return out


def test4(cov):
    pre, preshare = {}, {}
    for r in csv.DictReader(open("derived/pretrend_inputs.csv")):
        pre[(clean(r["country"]), int(r["cycle"]), r["subject"])] = f(r["nat_mean"])
        if r["subject"] == "math":
            preshare[(clean(r["country"]), int(r["cycle"]))] = f(r["imm_share"])
    xs = {clean(r["country"]): r for r in csv.DictReader(open("derived/crosscountry.csv"))}
    rows = []
    for c, r in sorted(xs.items()):
        o = dict(country=c, section=r["section"], europe=int(c in EUROPE), dshare_12_22=f(r["dshare_12_22"]),
                 dnat_math_12_22=f(r.get("dnat_math_12_22")))
        for s in ("math", "reading", "science"):
            o[f"pre_{s}_06_12"] = f(r.get(f"nat_{s}_2012")) - pre.get((c, 2006, s), float("nan"))
        o["pre_reading_09_12"] = f(r.get("nat_reading_2012")) - pre.get((c, 2009, "reading"), float("nan"))
        o["pre_dshare_06_12"] = f(r.get("share_imm_2012")) - preshare.get((c, 2006), float("nan"))
        rows.append(o)
    out = []
    for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1)):
        rs = [r for r in rows if keep(r)]
        for y in ("pre_math_06_12", "pre_reading_06_12", "pre_science_06_12", "pre_reading_09_12"):
            x = fit(rs, y, ["dshare_12_22"], f"pre-trend {y} on later share change", sample, "4 pretrend")
            if x: out.append(x)
        x = fit(rs, "dnat_math_12_22", ["dshare_12_22", "pre_math_06_12"], "2012-22 slope controlling pre-trend", sample, "4 pretrend")
        if x: out.append(x)
        for y in ("pre_math_06_12", "pre_science_06_12", "pre_reading_06_12"):
            x = fit(rs, y, ["pre_dshare_06_12"], f"same-period slope 2006-12 ({y})", sample, "4 pretrend")
            if x: out.append(x)
    write("derived/pretrend_changes.csv", rows)
    return out


def test5(cov):
    wq = openpyxl.load_workbook("_cache/statlink_qmuad8.xlsx", read_only=True, data_only=True)
    arr = {}
    for sh, cyc_tok in (("Table I.B1.7.13", {2022: ""}), ("Table I.B1.7.14", {2012: "/ 2012 /"})):
        for n, _, c, v in parse_sheet(wq[sh]):
            for cyc, tok in cyc_tok.items():
                if tok in c and "After age 12" in c and c.endswith("%") is False and isinstance(v, (int, float)):
                    pass
        # layout: 7.13 cols '<band> / %' ; 7.14 cols '<title> / <cycle> / <band truncated>' alternating value/SE
        vals = {}
        for n, _, c, v in parse_sheet(wq[sh]):
            vals.setdefault(clean(n), []).append((c, v))
        for n, cv in vals.items():
            if sh.endswith("7.13"):
                late = next((v for c, v in cv if c.startswith("After age 12") and c.endswith("/ %")), None)
                mid = next((v for c, v in cv if c.startswith("At ages 6-11") and c.endswith("/ %")), None)
                if isinstance(late, (int, float)): arr[(n, 2022)] = (float(mid), float(late))
            else:
                seq = [v for c, v in cv if "/ 2012 /" in c]
                if len(seq) >= 6 and all(isinstance(x, (int, float)) for x in (seq[2], seq[4])):
                    arr[(n, 2012)] = (float(seq[2]), float(seq[4]))  # value cells precede their SE cells
    rows = []
    for r in cov:
        c = r["country"]
        try:
            g1 = {y: share(c if c not in STAR else c + "*", y, "g1") * 100 for y in (2012, 2022)}
            imm = {y: share(c if c not in STAR else c + "*", y, "imm") * 100 for y in (2012, 2022)}
        except KeyError:
            continue
        if (c, 2012) not in arr or (c, 2022) not in arr:
            continue
        rec = {y: g1[y] * arr[(c, y)][1] / 100 for y in (2012, 2022)}
        mid = {y: g1[y] * arr[(c, y)][0] / 100 for y in (2012, 2022)}
        rows.append(dict(country=c, section=r["section"], europe=r["europe"], dnat_math_12_22=r["dnat_math_12_22"],
                         d_recent_12plus=rec[2022] - rec[2012], d_mid_6_11=mid[2022] - mid[2012],
                         d_settled=(imm[2022] - rec[2022]) - (imm[2012] - rec[2012]),
                         d_g1=g1[2022] - g1[2012], d_g2=(imm[2022] - g1[2022]) - (imm[2012] - g1[2012]),
                         recent_2012=rec[2012], recent_2022=rec[2022], mid_2012=mid[2012], mid_2022=mid[2022]))
    out = []
    for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1)):
        rs = [r for r in rows if keep(r)]
        for lab, xs in (("recent (arrived >12) vs settled", ["d_recent_12plus", "d_settled"]),
                        ("settled vs recent", ["d_settled", "d_recent_12plus"]),
                        ("first vs second generation", ["d_g1", "d_g2"]), ("second vs first generation", ["d_g2", "d_g1"]),
                        ("arrived 6-11 only", ["d_mid_6_11"])):
            x = fit(rs, "dnat_math_12_22", xs, lab, sample, "5 timing")
            if x: out.append(x)
    if os.path.exists("derived/grade4_exposure.csv"):
        g4 = {}
        for r in csv.DictReader(open("derived/grade4_exposure.csv")):
            # TIMSS 2015 grade 4, both parents born abroad, pupil report; national samples only, England stands
            # in for the United Kingdom and the grade-4 benchmark sample for Norway
            name = {"Korea, Rep. of": "Korea", "Turkey": "Türkiye", "Norway (4)": "Norway", "England": "United Kingdom",
                    "Norway": None}.get(r["country"], r["country"])
            if name and r["study"] == "TIMSS" and r["cycle"] == "2015" and r["definition"] == "both_parents_born_abroad_student" \
                    and r["share"] not in ("", None):
                g4[name] = float(r["share"])
        rs = []
        for r in cov:
            c = r["country"]
            if c in g4:
                try:
                    s12 = share(c if c not in STAR else c + "*", 2012, "imm") * 100
                    s22 = share(c if c not in STAR else c + "*", 2022, "imm") * 100
                except KeyError:
                    continue
                rs.append(dict(country=c, section=r["section"], europe=r["europe"], dnat_math_12_22=r["dnat_math_12_22"],
                               g4_minus_pisa12=g4[c] - s12, avg_g4_pisa22_minus_pisa12=(g4[c] + s22) / 2 - s12,
                               dshare_12_22=s22 - s12, g4_2015=g4[c]))
        for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1), ("all", lambda r: True)):
            q = [r for r in rs if keep(r)]
            for lab, xs in (("grade-4 2015 share minus PISA 2012 share", ["g4_minus_pisa12"]),
                            ("avg(grade-4 2015, PISA 2022) minus PISA 2012", ["avg_g4_pisa22_minus_pisa12"]),
                            ("contemporaneous share change, same countries", ["dshare_12_22"])):
                x = fit(q, "dnat_math_12_22", xs, lab, sample, "5 timing grade4")
                if x: out.append(x)
        write("derived/grade4_matched.csv", rs)
    write("derived/arrival_timing.csv", rows)
    if os.path.exists("derived/arrival_within5.csv"):
        # microdata split (arrival_microdata.py): recent = arrived at age >= 10 (within ~5-6 years), 2015 -> 2022
        inv = {v: k for k, v in ISO3.items()}; inv["KSV"] = "Kosovo"
        mic = {}
        for r in csv.DictReader(open("derived/arrival_within5.csv")):
            mic[(inv.get(r["cnt"]), int(r["cycle"]))] = r
        rs = []
        for r in cov:
            c = r["country"]
            a, b = mic.get((c, 2015)), mic.get((c, 2022))
            if not a or not b:
                continue
            try:
                d = change(c if c not in STAR else c + "*", "math", 2015, "nat")
            except KeyError:
                continue
            rs.append(dict(country=c, section=r["section"], europe=r["europe"], dnat_math_15_22=d,
                           d_recent10=f(b["recent10_share"]) - f(a["recent10_share"]),
                           d_recent12=f(b["recent12_share"]) - f(a["recent12_share"]),
                           d_settled10=f(b["settled10_share"]) - f(a["settled10_share"]),
                           d_imm=f(b["imm_share"]) - f(a["imm_share"])))
        for sample, keep in (("OECD", lambda r: r["section"] == "OECD"), ("Europe", lambda r: r["europe"] == 1)):
            q = [r for r in rs if keep(r)]
            for lab, xs in (("microdata: recent (arrived >=10) vs settled, 2015-22", ["d_recent10", "d_settled10"]),
                            ("microdata: settled vs recent (arrived >=10), 2015-22", ["d_settled10", "d_recent10"]),
                            ("microdata: recent (arrived >=12) only, 2015-22", ["d_recent12"]),
                            ("microdata: total share change, 2015-22", ["d_imm"])):
                x = fit(q, "dnat_math_15_22", xs, lab, sample, "5 timing microdata")
                if x: out.append(x)
        write("derived/arrival_within5_matched.csv", rs)
    return out


def quote_inputs(cov):
    """Verbatim source rows for the covariates, European and OECD countries (reads/deconfound_inputs.md)."""
    keep = sorted(r["country"] for r in cov if r["section"] == "OECD" or r["europe"])
    wb = openpyxl.load_workbook("_cache/unesco_duration_school_closures.xlsx", read_only=True, data_only=True)
    un = {UNESCO_ALIAS.get(r[0], r[0]): r for r in list(wb["database"].iter_rows(values_only=True))[1:] if r and r[0]}
    wbj = json.load(open("_cache/wb_gdppc_ppp_kd.json"))[1]
    gdp = {(d["countryiso3code"], d["date"]): d["value"] for d in wbj}
    wq = openpyxl.load_workbook("_cache/statlink_qmuad8.xlsx", read_only=True, data_only=True)
    cells = {}
    for sh in ("Table I.B1.7.8", "Table I.B1.7.13", "Table I.B1.7.14"):
        for n, _, c, v in parse_sheet(wq[sh]):
            if clean(n) in keep and (sh != "Table I.B1.7.8" or ("2012 and PISA 2022" in c and "ESCS" in c and "Non-immigrant" in c)):
                cells.setdefault(clean(n), []).append(f"{sh} | {c} | {v}")
    L = ["# Part-2 covariate inputs, verbatim", "",
         "UNESCO: `_cache/unesco_duration_school_closures.xlsx` (https://covid19.uis.unesco.org/wp-content/uploads/sites/11/2022/09/"
         "SDG-duration-of-school-closures-by-country.xlsx), sheet `database`: Country | SDG Region | Days fully closed | Days"
         " partially closed | Total (16/02/2020-30/04/2022, per sheet `codebook`).",
         "World Bank: `_cache/wb_gdppc_ppp_kd.json` (https://api.worldbank.org/v2/country/all/indicator/NY.GDP.PCAP.PP.KD?date=2011:2022),"
         " GDP per capita, PPP (constant international $).",
         "OECD: `_cache/statlink_qmuad8.xlsx` (https://stat.link/qmuad8) Tables I.B1.7.8 (ESCS change), I.B1.7.13/7.14 (age at arrival).", ""]
    for c in keep:
        u = un.get(c)
        L.append(f"## {c}")
        L.append(f"- UNESCO row: {' | '.join(str(x) for x in u[:5]) if u else 'not in file'}")
        iso = ISO3.get(c)
        L.append(f"- World Bank {iso}: 2012 = {gdp.get((iso, '2012'))}; 2022 = {gdp.get((iso, '2022'))}")
        for x in cells.get(c, []):
            L.append(f"- {x}")
        L.append("")
    open("reads/deconfound_inputs.md", "w").write("\n".join(L))


STAR = set()
if __name__ == "__main__":
    STAR |= {clean(r["country"]) for r in csv.DictReader(open("derived/crosscountry.csv")) if "*" in r["country"]}
    cov = covariates()
    write("derived/deconfound_covariates.csv", cov)
    quote_inputs(cov)
    miss = [r["country"] for r in cov if (r["section"] == "OECD" or r["europe"]) and not math.isnan(r["dnat_math_12_22"])
            and (math.isnan(r["closure_total_wk"]) or math.isnan(r["dlog_gdppc_12_22"]) or math.isnan(r["dnat_escs_12_22"]))]
    print("covariate misses:", miss)
    res = test1(cov)
    bounds, s2 = test2(cov); write("derived/exclusion_bounds.csv", bounds); res += s2
    res += test3(); res += test4(cov); res += test5(cov)
    write("derived/deconfound_slopes.csv", res)
    for r in res:
        extra = " ".join(f"{k[5:]}={r[k]}({r['se_' + k[5:]]})" for k in r if k.startswith("coef_"))
        print(f"{r['test'][:9]:9} {r['sample'][:14]:14} {r['outcome'][:20]:20} {r['label'][:42]:42} n={r['n']:<3} "
              f"b10={r['slope_per10pp']:>7} se={r['se_per10pp']:>6} r2={r['r2']:<6} {extra}")
