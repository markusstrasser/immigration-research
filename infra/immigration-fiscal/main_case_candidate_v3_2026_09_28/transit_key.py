#!/usr/bin/env python3
"""The Mexican-origin group's relative use of public transit (RU), so the transit deficit can be keyed by riders.

    OPENBLAS_NUM_THREADS=1 uv run --no-project python3 \
        infra/immigration-fiscal/main_case_candidate_v3_2026_09_28/transit_key.py

The consumer applies: riders' key = (the case's population key) x RU, with RU = acs.national.ru (central) or one of
the variants (acs.variants, state_weighted, the NHTS all-trip adjustments).

Group: the ACS proxy of the case's Mexican-origin population, HISP == 02 (Mexican) or POBP == 303 (born in Mexico),
as in receipt_side_long_run_2026_09_28/housing.py (lines 4 and 140). Transit commuter: JWTRNS 02 bus, 03 subway or
elevated rail, 04 long-distance train or commuter rail, 05 light rail/streetcar/trolley, 06 ferryboat, the ACS
"public transportation (excluding taxicab)"; the labels are read from the 2024 PUMS data dictionary at run time.

RU = (group transit commuters / group persons) / (all transit commuters / all persons), all persons including group
quarters, weight PWGTP; equivalently the group's share of transit commuters over its share of persons. Standard
errors use the ACS successive-difference replicate formula SE = sqrt(4/80 * sum_r (x_r - x)^2) over PWGTP1-80, each
ratio recomputed on every replicate. All weight sums are exact integer sums.

Checks and extensions:
- published control: PUMS transit mode shares against ACS 2024 1-year S0201 (POPGROUP 4015 Mexican, 001 total),
  B08301 and B08105I (Hispanic); the tolerances in TOL were fixed before any published figure was fetched;
- NHTS 2022 (no replicate weights are published: stratified household bootstrap, Rao-Wu with n_h - 1 draws, 1,000
  draws, fixed seed) and NHTS 2017 (98 replicate weights; SE = sqrt(6/7 * sum), which reproduces the User's Guide
  example exactly): transit trips per person, Hispanic against all, for all trips and for home-based-work trips;
  their ratio is the all-trip adjustment of a commute-based RU;
- state deficit weighting: D_s = transit current operations (LF0216) - transit utility revenue (LF0073), Census
  Annual Survey of State and Local Government Finances 2024 (state and local); RU_state_weighted =
  (sum_s D_s g_s / sum_s D_s) / the group's national population share, g_s = the group's share of the transit
  commuters living in state s. The replicate SE holds D_s fixed.

Reads only local files (run transit_acquire.py first) and verifies every hash. Any failed gate prints
[BLOCKED] gate <name>: <detail> for each failure, writes nothing and exits 1. Outputs: derived/transit_key.json
(sorted keys, floats to 12 significant digits, repo-relative paths, no timestamps) and derived/transit_key_states.csv.
"""
from __future__ import annotations

import os

os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")

import csv  # noqa: E402
import json  # noqa: E402
import math  # noqa: E402
import re  # noqa: E402
import sys  # noqa: E402
import zipfile  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
import pypdf  # noqa: E402

LANE = Path(__file__).resolve().parent
ROOT = LANE.parents[2]
sys.path.insert(0, str(LANE))
sys.dont_write_bytecode = True  # no __pycache__ in the shared lane directory
import transit_acquire as ACQ  # noqa: E402  (the cache files' URLs and pinned sha256 are defined there, once)

CACHE = ACQ.CACHE
OUT_JSON = LANE / "derived" / "transit_key.json"
OUT_CSV = LANE / "derived" / "transit_key_states.csv"
EXT = ROOT / "sources" / "immigration-fiscal" / "data" / "external"
LOCAL = {
    "pums": dict(
        path=EXT / "acs_pums_2024_1yr" / "csv_pus.zip",
        sha256="afdc6d90c6e2f0bab365ed32d95ba4c4d8ac651162f46ac7861295b2dc469894",
        url="https://www2.census.gov/programs-surveys/acs/data/pums/2024/1-Year/csv_pus.zip",
        note="ACS 2024 1-year person PUMS (psam_pusa.csv, psam_pusb.csv); size equals census.gov's 602,847,146 bytes"),
    "dictionary": dict(
        path=EXT / "acs_pums_2024_1yr" / "PUMS_Data_Dictionary_2024.pdf",
        sha256="929c2752995b0af1c16d5c64de8cdc43b4aa7d388ee2d45b4b4df90fecce1dff",
        url="https://www2.census.gov/programs-surveys/acs/tech_docs/pums/data_dict/PUMS_Data_Dictionary_2024.pdf",
        note="2024 ACS PUMS data dictionary: JWTRNS pp.38-39, HISP p.63, POBP pp.99-102"),
    "nipa": dict(
        path=EXT / "bea_nipa" / "Section3All_xls.xlsx",
        sha256="69b5c7aefb38675324887ce31d6feb4fcde7c903ab952db7328da0813096615e",
        url="https://apps.bea.gov/national/Release/XLS/Survey/Section3All_xls.xlsx",
        note="NIPA Table 3.8 (T30800-A) line 14, public transit current surplus (L31224), data published 2025-09-26"),
}

GROUP_DEF = "HISP == 02 (Mexican) or POBP == 303 (born in Mexico)"
GROUP_CITE = "infra/immigration-fiscal/receipt_side_long_run_2026_09_28/housing.py lines 4 and 140"
TRANSIT = (2, 3, 4, 5, 6)
EXPECTED_JWTRNS = {
    1: "Car, truck, or van", 2: "Bus", 3: "Subway or elevated rail", 4: "Long-distance train or commuter rail",
    5: "Light rail, streetcar, or trolley", 6: "Ferryboat", 7: "Taxi or ride-hailing services", 8: "Motorcycle",
    9: "Bicycle", 10: "Walked", 11: "Worked from home", 12: "Other method"}
VARIANTS = {
    "without_04_commuter_rail": (2, 3, 5, 6),
    "bus_only_02": (2,),
    "rail_only_03_05": (3, 5),
}
# person categories: 0 neither Hispanic nor born in Mexico; 1 Hispanic (not Mexican origin), not born in Mexico;
# 2 born in Mexico, not Hispanic; 3 born in Mexico, Hispanic other than Mexican; 4 Mexican origin (HISP 02)
POPS = {"all": (0, 1, 2, 3, 4), "group": (2, 3, 4), "hispanic_all": (1, 3, 4), "mexican_origin_hisp02": (4,)}
WCOLS = ["PWGTP"] + [f"PWGTP{i}" for i in range(1, 81)]
PUMS_MEMBERS = ["psam_pusa.csv", "psam_pusb.csv"]
CHUNK = 250_000
NY = 36

# Share and count tolerances were written to the lane log on 2026-09-28 before any published figure was fetched;
# hispanic_ru_rel before the 2017/2022/2024 B03003/B08105I tables were fetched.
TOL = dict(group_share_pp=0.30, total_share_pp=0.15, hispanic_share_pp=0.20, transit_commuters_rel=0.03,
           persons_vs_b01003_rel=0.005, persons_vs_340m_rel=0.02, state_sum_rel=1e-6, nhts_person_total_rel=0.01,
           hispanic_ru_rel=0.03)
PERSONS_REF = 340e6
PARENT_CHAT = {"mexican_pct": 2.8, "total_pct": 3.7}

NHTS22_PERSON_TOTAL = 305_560_925  # 2022 Weighting Report, Table 15 (p.32): Final Person 16,997 / 305,560,925
NHTS17_PERSON_TOTAL = 301_599_169  # 2017 User's Guide, Table 7-1: Person 264,234 / 301,599,169
NHTS17_EXAMPLE = (9_444_506_727, 212_640_677)  # 2017 User's Guide 7.12: total transit trips and its standard error
NHTS17_FACTOR = 6 / 7  # 2017 User's Guide 7.12: SE = sqrt(sum_{i=1..98} (6/7) (REP_i - x)^2)
BOOT_B, BOOT_SEED = 1000, 20260928

GATES: list[dict] = []


def blocked(msg: str) -> None:
    print(f"[BLOCKED] {msg}")
    sys.exit(1)


def gate(name: str, ok: bool, detail: str) -> None:
    GATES.append({"gate": name, "pass": bool(ok), "detail": detail})


def rel(p: Path) -> str:
    """Repo-relative path; a path outside the repository (never an output) is shown by name only."""
    return p.relative_to(ROOT).as_posix() if p.is_relative_to(ROOT) else p.name


def sig(x):
    """Round floats to 12 significant digits, recursively; ints and strings pass through."""
    if isinstance(x, (bool, np.bool_)):
        return bool(x)
    if isinstance(x, (int, np.integer)):
        return int(x)
    if isinstance(x, (float, np.floating)):
        x = float(x)
        if not math.isfinite(x):
            raise ValueError(f"non-finite value {x}")
        return float(f"{x:.12g}") if x else 0.0
    if isinstance(x, dict):
        return {str(k): sig(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [sig(v) for v in x]
    return x


def se_acs(x: np.ndarray) -> float:
    return math.sqrt(4 / 80 * float(np.sum((x[1:] - x[0]) ** 2)))


def est(x: np.ndarray) -> dict:
    return {"value": float(x[0]), "se": se_acs(x)}


def fmt(x: float) -> str:
    return f"{x:.12g}"


# ----------------------------------------------------------------------------------------------------------------
# inputs


def verify_inputs() -> list[dict]:
    sources = []
    for spec in LOCAL.values():
        p = spec["path"]
        if not p.is_file():
            blocked(f"missing input {rel(p)}")
        got = ACQ.sha256(p)
        if got != spec["sha256"]:
            blocked(f"gate input_hashes: {rel(p)} sha256 {got} != pinned {spec['sha256']}")
        sources.append(dict(path=rel(p), sha256=got, url=spec["url"], note=spec["note"]))
    for name in sorted(ACQ.FILES):
        p, pin = CACHE / name, ACQ.PINS.get(name)
        if not pin:
            blocked(f"gate input_hashes: {name} has no pinned sha256 in transit_acquire.PINS")
        if not p.is_file():
            blocked(f"missing input {rel(p)}; run transit_acquire.py")
        got = ACQ.sha256(p)
        if got != pin:
            blocked(f"gate input_hashes: {rel(p)} sha256 {got} != pinned {pin}")
        sources.append(dict(path=rel(p), sha256=got, url=ACQ.FILES[name]["url"], note=ACQ.FILES[name]["note"]))
    gate("input_hashes", True, f"{len(sources)} input files match their pinned sha256")
    return sources


def read_dictionary(pdf: Path) -> tuple[dict[int, str], list[dict]]:
    lines = []
    for page_no, page in enumerate(pypdf.PdfReader(str(pdf)).pages, 1):
        for raw in (page.extract_text() or "").splitlines():
            t = " ".join(raw.split())
            if t:
                lines.append((page_no, t))
    header_re = re.compile(r"^[A-Z][A-Z0-9_]* (Character|Numeric) \d+$")

    def header(var: str, width: int) -> int:
        hits = [i for i, (_, t) in enumerate(lines) if t == f"{var} Character {width}"]
        if len(hits) != 1:
            blocked(f"gate acs_codes_match_dictionary: header '{var} Character {width}' found {len(hits)} times")
        return hits[0]

    used: list[dict] = []
    j = header("JWTRNS", 2)
    used += [dict(page=lines[j][0], line=lines[j][1]), dict(page=lines[j + 1][0], line=lines[j + 1][1])]
    codes: dict[int, str] = {}
    for page_no, t in lines[j + 1:]:
        if header_re.match(t):
            break
        m = re.match(r"^(\d{2}) \.(.+)$", t)
        if m:
            codes[int(m.group(1))] = m.group(2).strip()
            used.append(dict(page=page_no, line=t))
    j = header("HISP", 2)
    hisp = [(p, t) for p, t in lines[j:j + 6] if t.startswith("02 .")]
    used += [dict(page=lines[j][0], line=lines[j][1])] + [dict(page=p, line=t) for p, t in hisp]
    j = header("POBP", 3)
    pobp = []
    for page_no, t in lines[j + 1:]:
        if header_re.match(t):
            break
        if t.startswith("303 ."):
            pobp.append((page_no, t))
    used += [dict(page=lines[j][0], line=lines[j][1])] + [dict(page=p, line=t) for p, t in pobp]
    ok = (codes == EXPECTED_JWTRNS and [t for _, t in hisp] == ["02 .Mexican"]
          and [t for _, t in pobp] == ["303 .Mexico"])
    detail = (f"JWTRNS 01-12 labels {'match' if codes == EXPECTED_JWTRNS else 'DIFFER: ' + repr(codes)}; "
              f"HISP 02 = {[t for _, t in hisp]}; POBP 303 = {[t for _, t in pobp]}")
    gate("acs_codes_match_dictionary", ok, detail)
    return codes, used


# ----------------------------------------------------------------------------------------------------------------
# ACS PUMS


class Tab:
    """Exact weighted totals by (state, JWTRNS code 0-12, person category 0-4) for PWGTP and PWGTP1-80."""

    def __init__(self, table: pd.DataFrame):
        keys = table.index.to_numpy()
        self.states = sorted({int(k) // 1000 for k in keys})
        self.sidx = {s: i for i, s in enumerate(self.states)}
        self.w = np.zeros((len(self.states), 13, 5, 81), dtype=np.int64)
        self.n = np.zeros((len(self.states), 13, 5), dtype=np.int64)
        vals = table[list(range(81))].to_numpy(np.int64)
        cnt = table["n"].to_numpy(np.int64)
        for k, v, c in zip(keys, vals, cnt):
            s, jw, cat = int(k) // 1000, (int(k) % 1000) // 10, int(k) % 10
            self.w[self.sidx[s], jw, cat] = v
            self.n[self.sidx[s], jw, cat] = c

    def _sel(self, pop: str, codes, states):
        s = list(range(len(self.states))) if states is None else [self.sidx[x] for x in states]
        jw = list(range(13)) if codes is None else list(codes)
        return np.ix_(s, jw, list(POPS[pop]))

    def tot(self, pop: str, codes=None, states=None) -> np.ndarray:
        return self.w[self._sel(pop, codes, states)].sum(axis=(0, 1, 2))

    def cnt(self, pop: str, codes=None, states=None) -> int:
        return int(self.n[self._sel(pop, codes, states)].sum())


def pums_aggregate(path: Path) -> tuple[Tab, dict, dict]:
    usecols = ["STATE", "AGEP", "HISP", "POBP", "JWTRNS"] + WCOLS
    dtypes = {c: "int64" for c in WCOLS} | {"STATE": "int64", "AGEP": "int64", "HISP": "int64", "POBP": "int64",
                                              "JWTRNS": "float64"}
    parts = []
    direct = {k: np.zeros(81, np.int64) for k in ("persons_all", "persons_group", "transit_all", "transit_group")}
    chk = dict(rows=0, bad_jwtrns=0, bad_hisp=0, workers_under_16=0)
    with zipfile.ZipFile(path) as z:
        members = sorted(n for n in z.namelist() if n.lower().endswith(".csv"))
        if members != PUMS_MEMBERS:
            blocked(f"gate pums_structure: csv members {members} != {PUMS_MEMBERS}")
        for m in members:
            with z.open(m) as fh:
                for i, ch in enumerate(pd.read_csv(fh, usecols=usecols, dtype=dtypes, chunksize=CHUNK), 1):
                    jwf = ch["JWTRNS"].to_numpy()
                    jw = np.where(np.isnan(jwf), 0, jwf).astype(np.int64)
                    chk["bad_jwtrns"] += int(np.sum((jw < 0) | (jw > 12) | (~np.isnan(jwf) & (jwf != jw))))
                    hisp, pobp, agep = (ch[c].to_numpy() for c in ("HISP", "POBP", "AGEP"))
                    chk["bad_hisp"] += int(np.sum((hisp < 1) | (hisp > 24)))
                    chk["workers_under_16"] += int(np.sum((jw > 0) & (agep < 16)))
                    mex, mxb, hsp = hisp == 2, pobp == 303, hisp != 1
                    cat = np.select([mex, mxb & hsp, mxb, hsp], [4, 3, 2, 1], default=0)
                    st = ch["STATE"].to_numpy()
                    w = ch[WCOLS].to_numpy(np.int64)
                    grp, tr = cat >= 2, np.isin(jw, TRANSIT)
                    direct["persons_all"] += w.sum(axis=0)
                    direct["persons_group"] += w[grp].sum(axis=0)
                    direct["transit_all"] += w[tr].sum(axis=0)
                    direct["transit_group"] += w[tr & grp].sum(axis=0)
                    df = pd.DataFrame(w, columns=range(81))
                    df["n"] = 1
                    parts.append(df.groupby(st * 1000 + jw * 10 + cat, sort=True).sum())
                    chk["rows"] += len(ch)
                    print(f"[pums] {m} chunk {i}: {chk['rows']:,} persons read", flush=True)
    table = pd.concat(parts).groupby(level=0, sort=True).sum()
    return Tab(table), direct, chk


def acs_block(tab: Tab, labels: dict[int, str]) -> tuple[dict, dict]:
    n_all, n_grp = tab.tot("all"), tab.tot("group")
    t_all, t_grp = tab.tot("all", TRANSIT), tab.tot("group", TRANSIT)
    ru = (t_grp / n_grp) / (t_all / n_all)
    work = {p: tab.tot(p, range(1, 13)) for p in POPS}
    mode = {}
    for p in ("group", "all", "hispanic_all", "mexican_origin_hisp02"):
        shares = {f"{c:02d}": {"label": labels[c], **est(tab.tot(p, (c,)) / work[p])} for c in range(1, 13)}
        shares["public_transportation_02_06"] = est(tab.tot(p, TRANSIT) / work[p])
        mode[p] = shares
    national = {
        "ru": est(ru),
        "direct_share_of_transit_commuters": est(t_grp / t_all),
        "group_population_share": est(n_grp / n_all),
        "transit_commuters_per_person": {"group": est(t_grp / n_grp), "all": est(t_all / n_all)},
        "persons": {p: int(tab.tot(p)[0]) for p in POPS},
        "transit_commuters": {p: int(tab.tot(p, TRANSIT)[0]) for p in POPS},
        "workers_with_jwtrns": {p: int(work[p][0]) for p in POPS},
        "sample_records": {"persons_all": tab.cnt("all"), "persons_group": tab.cnt("group"),
                           "transit_commuters_all": tab.cnt("all", TRANSIT),
                           "transit_commuters_group": tab.cnt("group", TRANSIT)},
        "mode_shares_among_workers": mode,
        "definition": "RU = (group transit commuters / group persons) / (all transit commuters / all persons)",
    }

    def variant(pop: str, codes, states=None, note: str = "") -> dict:
        tg, ng = tab.tot(pop, codes, states), tab.tot(pop, None, states)
        ta, na = tab.tot("all", codes, states), tab.tot("all", None, states)
        out = {"jwtrns_codes": [f"{c:02d}" for c in codes], "ru": est((tg / ng) / (ta / na)),
               "direct_share_of_transit_commuters": est(tg / ta), "population_share": est(ng / na),
               "transit_commuters": {"population": int(tg[0]), "all": int(ta[0])},
               "sample_transit_commuters": {"population": tab.cnt(pop, codes, states),
                                            "all": tab.cnt("all", codes, states)}}
        if note:
            out["note"] = note
        return out

    variants = {name: variant("group", codes) for name, codes in VARIANTS.items()}
    variants["mexican_origin_hisp02_only"] = variant(
        "mexican_origin_hisp02", TRANSIT,
        note="HISP == 02 alone, without the Mexican-born addition; like-for-like with S0201")
    ex_ny = [s for s in tab.states if s != NY]
    variants["excluding_new_york"] = variant(
        "group", TRANSIT, ex_ny, note="residents of New York State dropped from numerator and denominator; sensitivity")
    return national, variants


# ----------------------------------------------------------------------------------------------------------------
# published control


def census_row(name: str) -> dict:
    rows = json.loads((CACHE / name).read_text())
    return dict(zip(rows[0], rows[1]))


def published_hispanic_ru() -> dict:
    out = {}
    for y in ("2017", "2022", "2024"):
        d = census_row(f"census_acs1_{y}_hispanic_commute_us.json")
        v = {k: int(d[k]) for k in d if k.startswith("B")}
        ru = (v["B08105I_004E"] / v["B03003_003E"]) / (v["B08301_010E"] / v["B01003_001E"])
        out[y] = {"ru": ru, **v}
    return out


def published_control(tab: Tab, national: dict, hisp_ru: dict) -> dict:
    s = {pg: census_row(f"census_acs1_2024_s0201_popgroup_{pg}.json") for pg in ("001", "400", "4015")}
    want = {"001": "Total population", "400": "Hispanic or Latino (of any race)", "4015": "Mexican"}
    labels_ok = all(s[pg]["POPGROUP_LABEL"] == want[pg] and s[pg]["POPGROUP"] == pg for pg in want)
    d = census_row("census_acs1_2024_detail_us.json")
    pct = {pg: float(s[pg]["S0201_171E"]) for pg in s}
    work = {pg: int(s[pg]["S0201_168E"]) for pg in s}
    popn = {pg: int(s[pg]["S0201_001E"]) for pg in s}
    b_total, b_transit = int(d["B08301_001E"]), int(d["B08301_010E"])
    bi_total, bi_transit = int(d["B08105I_001E"]), int(d["B08105I_004E"])
    persons_pub = int(d["B01003_001E"])
    mode = national["mode_shares_among_workers"]
    pums_pct = {p: 100 * mode[p]["public_transportation_02_06"]["value"] for p in mode}
    comps = []

    def comp(name: str, pums: float, pub: float, tol: float, unit: str, gated: bool = True) -> None:
        diff = pums - pub if unit == "pp" else pums / pub - 1
        ok = abs(diff) <= tol
        comps.append({"name": name, "pums": pums, "published": pub, "difference": diff, "unit": unit,
                      "tolerance": tol if gated else None, "pass": ok if gated else None})
        if gated:
            gate(f"published_{name}", ok, f"PUMS {fmt(pums)} vs published {fmt(pub)}: "
                 f"{'diff' if unit == 'pp' else 'rel diff'} {diff:+.4f} {unit}, tolerance {tol}")

    comp("group_transit_share_vs_s0201_mexican_4015", pums_pct["group"], pct["4015"], TOL["group_share_pp"], "pp")
    comp("total_transit_share_vs_s0201_001", pums_pct["all"], pct["001"], TOL["total_share_pp"], "pp")
    comp("total_transit_share_vs_b08301", pums_pct["all"], 100 * b_transit / b_total, TOL["total_share_pp"], "pp")
    comp("hispanic_transit_share_vs_b08105i", pums_pct["hispanic_all"], 100 * bi_transit / bi_total,
         TOL["hispanic_share_pp"], "pp")
    comp("hispanic_transit_share_vs_s0201_400", pums_pct["hispanic_all"], pct["400"], TOL["hispanic_share_pp"], "pp")
    comp("mexican_hisp02_transit_share_vs_s0201_4015", pums_pct["mexican_origin_hisp02"], pct["4015"],
         TOL["group_share_pp"], "pp", gated=False)
    comp("transit_commuters_vs_b08301_010", national["transit_commuters"]["all"], b_transit,
         TOL["transit_commuters_rel"], "rel")
    comp("persons_vs_b01003", national["persons"]["all"], persons_pub, TOL["persons_vs_b01003_rel"], "rel")
    comp("persons_vs_340m", national["persons"]["all"], PERSONS_REF, TOL["persons_vs_340m_rel"], "rel")
    comp("group_persons_vs_s0201_mexican_population", national["persons"]["group"], popn["4015"], 0, "rel", gated=False)
    comp("mexican_hisp02_persons_vs_s0201_mexican_population", national["persons"]["mexican_origin_hisp02"],
         popn["4015"], 0, "rel", gated=False)
    comp("workers_vs_b08301_001", national["workers_with_jwtrns"]["all"], b_total, 0, "rel", gated=False)
    for code in (2, 3, 4, 5, 6):
        var = f"B08301_{code + 9:03d}E"  # 011 bus, 012 subway, 013 commuter rail, 014 light rail, 015 ferry
        comp(f"jwtrns_{code:02d}_commuters_vs_{var[:-1].lower()}", int(tab.tot("all", (code,))[0]), int(d[var]), 0,
             "rel", gated=False)
    hru = published_hispanic_ru()
    comp("hispanic_commute_ru_vs_published_tables", hisp_ru["value"], hru["2024"]["ru"], TOL["hispanic_ru_rel"], "rel")
    gate("published_labels", labels_ok, f"S0201 POPGROUP labels {[s[pg]['POPGROUP_LABEL'] for pg in s]}")
    ru_pub = (pct["4015"] * work["4015"] / popn["4015"]) / (pct["001"] * work["001"] / popn["001"])
    return {
        "tolerances_fixed_before_comparison": {k: TOL[k] for k in ("group_share_pp", "total_share_pp",
                                                                   "hispanic_share_pp", "transit_commuters_rel",
                                                                   "persons_vs_b01003_rel", "persons_vs_340m_rel",
                                                                   "hispanic_ru_rel")},
        "comparisons": comps,
        "s0201": {pg: {"label": s[pg]["POPGROUP_LABEL"], "population": popn[pg], "workers_16_plus": work[pg],
                       "public_transportation_pct": pct[pg], "worked_from_home_pct": float(s[pg]["S0201_174E"])}
                  for pg in s},
        "s0201_popgroup_note": "POPGROUP 401 (Mexican) returns HTTP 204 for 2024; 4015 is the 2024 code labelled "
                               "Mexican",
        "b08301": {k: int(d[k]) for k in sorted(d) if k.startswith("B08301")},
        "b08105i": {k: int(d[k]) for k in sorted(d) if k.startswith("B08105I")},
        "b01003_001e": persons_pub,
        "published_ru_mexican_from_s0201": {
            "value": ru_pub,
            "formula": "(pct_4015 x workers_4015 / population_4015) / (pct_001 x workers_001 / population_001)",
            "note": "S0201 percentages are rounded to 0.1, so this carries about +-2% rounding error"},
        "published_hispanic_commute_ru": {
            "by_year": hru, "formula": "(B08105I_004E / B03003_003E) / (B08301_010E / B01003_001E)",
            "note": "ACS 1-year published tables; the 2024 value checks the PUMS Hispanic RU, 2017 and 2022 pair with "
                    "the NHTS waves"},
        "parent_chat_estimate": {**PARENT_CHAT, "api_mexican_pct": pct["4015"], "api_total_pct": pct["001"],
                                 "matches_api": PARENT_CHAT["mexican_pct"] == pct["4015"]
                                 and PARENT_CHAT["total_pct"] == pct["001"]},
    }


# ----------------------------------------------------------------------------------------------------------------
# NHTS


def codebook(path: Path, sheet: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    out, cur = {}, None
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        if row[0]:
            cur = str(row[0]).strip()
            out[cur] = {"label": row[1], "codes": {}}
        if cur and row[4] is not None:
            code, _, lab = str(row[4]).partition("=")
            out[cur]["codes"][code.strip()] = (lab.strip(), row[5], row[6])
    wb.close()
    return out


def codebook_check(cb: dict, expect: dict[tuple[str, str], str]) -> tuple[bool, list[str]]:
    lines, ok = [], True
    for (var, code), label in expect.items():
        got = cb.get(var, {}).get("codes", {}).get(code, (None,))[0]
        ok &= got == label
        lines.append(f"{var} {code}={got}")
    return ok, lines


def ratio_stats(p_all, p_h, t_all, t_h):
    return (t_h / p_h) / (t_all / p_all)


def summarize_draws(x: np.ndarray, full: float) -> dict:
    ok = np.isfinite(x)
    v = x[ok]
    return {"value": full, "se": float(np.std(v, ddof=1)), "ci95": [float(q) for q in np.percentile(v, [2.5, 97.5])],
            "invalid_draws": int((~ok).sum())}


def rao_wu(strata: np.ndarray, b: int, seed: int) -> np.ndarray:
    rng = np.random.default_rng(seed)
    groups = [np.flatnonzero(strata == s) for s in np.unique(strata)]
    mult = np.zeros((b, len(strata)))
    for r in range(b):
        for idx in groups:
            n = len(idx)
            mult[r, idx] = np.bincount(rng.integers(0, n, size=n - 1), minlength=n) * (n / (n - 1))
    return mult


WITHIN = "ACS 2024 group RU x NHTS Hispanic (all-trip RU / HBW commute RU); independent errors on the log scale"
CROSS = ("ACS 2024 group RU x (NHTS Hispanic all-trip RU / same-year published ACS Hispanic commute RU); log-scale "
         "errors of the ACS group RU and the NHTS RU (the published ACS RU's own error, about 1%, is left out)")


def adjusted(ru_group: dict, adj_value: float, se_log_adj: float, formula: str = WITHIN) -> dict:
    v = ru_group["value"] * adj_value
    se_log = math.sqrt((ru_group["se"] / ru_group["value"]) ** 2 + se_log_adj ** 2)
    return {"value": v, "se_log": se_log, "ci95": [v * math.exp(-1.96 * se_log), v * math.exp(1.96 * se_log)],
            "formula": formula}


def nhts_2022(acs_hisp_ru: dict, ru_group: dict, acs_same_year: float) -> dict:
    z = zipfile.ZipFile(CACHE / "nhts2022_csv.zip")
    per = pd.read_csv(z.open("perv2pub.csv"), usecols=["HOUSEID", "PERSONID", "WTPERFIN", "R_HISP", "R_AGE",
                                                       "STRATUMID"], dtype=str)
    trp = pd.read_csv(z.open("tripv2pub.csv"), usecols=["HOUSEID", "PERSONID", "TRPTRANS", "PUBTRANS", "TRIPPURP",
                                                        "WHYFROM", "WHYTO", "WTTRDFIN"], dtype=str)
    per["w"], trp["w"] = per["WTPERFIN"].astype(float), trp["WTTRDFIN"].astype(float)
    cb_p = codebook(CACHE / "nhts2022_codebook.xlsx", "Person")
    cb_t = codebook(CACHE / "nhts2022_codebook.xlsx", "Trip")
    ok_p, lp = codebook_check(cb_p, {("R_HISP", "01"): "Hispanic", ("R_HISP", "02"): "Not Hispanic"})
    ok_t, lt = codebook_check(cb_t, {("TRPTRANS", "08"): "Public or commuter bus",
                                     ("TRPTRANS", "10"): "Street car or trolley car",
                                     ("TRPTRANS", "11"): "Subway or elevated rail", ("TRPTRANS", "12"): "Commuter rail",
                                     ("PUBTRANS", "01"): "Used public transit",
                                     ("TRIPPURP", "01"): "Home-based work (HBW)"})
    pub_re = trp["TRPTRANS"].isin(["08", "10", "11", "12"])
    home, work = ["01", "02"], ["03", "04", "05"]
    hbw_re = ((trp["WHYFROM"].isin(home) & trp["WHYTO"].isin(work))
              | (trp["WHYFROM"].isin(work) & trp["WHYTO"].isin(home)))
    transit, hbw = trp["PUBTRANS"].eq("01"), trp["TRIPPURP"].eq("01")
    mism_pub, mism_hbw = int((pub_re != transit).sum()), int((hbw_re != hbw).sum())
    gate("nhts2022_codebook_and_derivations", ok_p and ok_t and mism_pub == 0 and mism_hbw == 0,
         f"codebook {lp + lt}; PUBTRANS vs TRPTRANS in (08,10,11,12) mismatches {mism_pub}; "
         f"TRIPPURP 01 vs WHYFROM/WHYTO home(01,02)<->work(03,04,05) mismatches {mism_hbw}")
    cb_total = sum(v[2] for v in cb_p["WTPERFIN"]["codes"].values())
    p_sum = float(per["w"].sum())
    gate("nhts2022_person_total", abs(p_sum / NHTS22_PERSON_TOTAL - 1) <= TOL["nhts_person_total_rel"],
         f"sum WTPERFIN {p_sum:,.1f} vs published 305,560,925 (Weighting Report Table 15; codebook {cb_total:,.1f})")
    if per.duplicated(["HOUSEID", "PERSONID"]).any():
        blocked("gate nhts2022_structure: duplicate person ids")
    trp = trp.merge(per[["HOUSEID", "PERSONID", "R_HISP"]].rename(columns={"R_HISP": "P_HISP"}),
                    on=["HOUSEID", "PERSONID"], how="left", validate="many_to_one")
    if trp["P_HISP"].isna().any():
        blocked("gate nhts2022_structure: trips without a person record")
    hh = per.groupby("HOUSEID")["STRATUMID"].agg(["nunique", "first"])
    if (hh["nunique"] != 1).any():
        blocked("gate nhts2022_structure: a household spans strata")
    hh_ids = hh.index.to_numpy()
    strata = hh["first"].to_numpy()
    n_h = pd.Series(strata).value_counts()
    gate("nhts2022_strata_bootstrap", int(n_h.min()) >= 2,
         f"{len(n_h)} strata (STRATUMID), households per stratum min {int(n_h.min())}, max {int(n_h.max())}")
    pos = {h: i for i, h in enumerate(hh_ids)}
    hisp_p, hisp_t = per["R_HISP"].eq("01"), trp["P_HISP"].eq("01")

    def hh_vec(frame: pd.DataFrame, mask: pd.Series) -> np.ndarray:
        s = frame.loc[mask].groupby("HOUSEID")["w"].sum()
        v = np.zeros(len(hh_ids))
        v[[pos[h] for h in s.index]] = s.to_numpy()
        return v

    vec = {"p_all": hh_vec(per, per["w"].notna()), "p_h": hh_vec(per, hisp_p),
           "t_all": hh_vec(trp, transit), "t_h": hh_vec(trp, transit & hisp_t),
           "c_all": hh_vec(trp, transit & hbw), "c_h": hh_vec(trp, transit & hbw & hisp_t)}
    full = {k: float(v.sum()) for k, v in vec.items()}
    ru_all = ratio_stats(full["p_all"], full["p_h"], full["t_all"], full["t_h"])
    ru_c = ratio_stats(full["p_all"], full["p_h"], full["c_all"], full["c_h"])
    mult = rao_wu(strata, BOOT_B, BOOT_SEED)
    rep = {k: (mult * v[None, :]).sum(axis=1) for k, v in vec.items()}
    with np.errstate(divide="ignore", invalid="ignore"):
        b_all = ratio_stats(rep["p_all"], rep["p_h"], rep["t_all"], rep["t_h"])
        b_c = ratio_stats(rep["p_all"], rep["p_h"], rep["c_all"], rep["c_h"])
        b_adj = b_all / b_c
        b_log_adj = np.log(b_adj)
    s_all, s_c = summarize_draws(b_all, ru_all), summarize_draws(b_c, ru_c)
    s_adj = summarize_draws(b_adj, ru_all / ru_c)
    se_log_adj = float(np.std(b_log_adj[np.isfinite(b_log_adj)], ddof=1))
    se_log_all = float(np.std(np.log(b_all[np.isfinite(b_all) & (b_all > 0)]), ddof=1))
    gate("nhts2022_bootstrap_finite", all(math.isfinite(x["se"]) and x["se"] > 0 for x in (s_all, s_c, s_adj)),
         f"B={BOOT_B}, seed {BOOT_SEED}; invalid draws all-trip {s_all['invalid_draws']}, "
         f"commute {s_c['invalid_draws']}, "
         f"adjustment {s_adj['invalid_draws']}")
    hisp_tr = trp.loc[transit & hisp_t]
    return {
        "source": "2022 NextGen NHTS public use V2.1 (perv2pub.csv, tripv2pub.csv); persons 5 and older",
        "definitions": {
            "hispanic": "person R_HISP == 01 (Hispanic)",
            "transit_trip": "PUBTRANS == 01, i.e. TRPTRANS in 08 public or commuter bus, 10 street car or trolley, "
                            "11 subway or elevated rail, 12 commuter rail (Amtrak 13, paratransit 17 excluded)",
            "commute_trip": "TRIPPURP == 01 home-based work: WHYFROM in (01,02) and WHYTO in (03,04,05), or reverse",
            "rate": "annual transit trips (WTTRDFIN) per person (WTPERFIN)",
        },
        "codebook_lines": lp + lt,
        "sample": {"households": int(len(hh_ids)), "persons": int(len(per)), "persons_hispanic": int(hisp_p.sum()),
                   "trips": int(len(trp)), "transit_trips": int(transit.sum()),
                   "transit_trips_hispanic": int((transit & hisp_t).sum()),
                   "hbw_transit_trips": int((transit & hbw).sum()),
                   "hbw_transit_trips_hispanic": int((transit & hbw & hisp_t).sum()),
                   "households_with_hispanic_transit_trips": int(hisp_tr["HOUSEID"].nunique()),
                   "households_with_hispanic_hbw_transit_trips":
                       int(trp.loc[transit & hbw & hisp_t, "HOUSEID"].nunique())},
        "weighted": {"persons_all": full["p_all"], "persons_hispanic": full["p_h"],
                     "transit_trips_all": full["t_all"], "transit_trips_hispanic": full["t_h"],
                     "hbw_transit_trips_all": full["c_all"], "hbw_transit_trips_hispanic": full["c_h"]},
        "transit_trips_per_person_year": {"all": full["t_all"] / full["p_all"], "hispanic": full["t_h"] / full["p_h"]},
        "hbw_transit_trips_per_person_year": {"all": full["c_all"] / full["p_all"],
                                              "hispanic": full["c_h"] / full["p_h"]},
        "ru_all_trips_hispanic": s_all,
        "ru_commute_hbw_hispanic": s_c,
        "adjustment_all_trips_over_commute": {**s_adj, "se_log": se_log_adj},
        "acs_hispanic_commute_ru_2024": acs_hisp_ru,
        "nhts_commute_ru_over_acs_hispanic_commute_ru": ru_c / acs_hisp_ru["value"],
        "group_ru_all_trip_adjusted": adjusted(ru_group, ru_all / ru_c, se_log_adj),
        "acs_hispanic_commute_ru_same_year_published": {"year": 2022, "value": acs_same_year},
        "nhts_commute_ru_over_same_year_acs": ru_c / acs_same_year,
        "cross_survey_adjustment_all_trips_over_same_year_acs_commute": {
            "value": ru_all / acs_same_year, "se_log": se_log_all},
        "group_ru_all_trip_adjusted_cross_survey": adjusted(ru_group, ru_all / acs_same_year, se_log_all, CROSS),
        "variance": {"method": "stratified household-cluster bootstrap (Rao-Wu, n_h - 1 draws per stratum, "
                               "multiplier n_h/(n_h - 1)); strata STRATUMID, clusters HOUSEID, as the User's Guide "
                               "specifies for Taylor series; the NextGen NHTS publishes no replicate weights",
                     "draws": BOOT_B, "seed": BOOT_SEED, "ci": "percentile 2.5-97.5"},
        "person_total_check": {"sum_wtperfin": p_sum, "published": NHTS22_PERSON_TOTAL, "codebook": cb_total},
    }


def nhts_2017(acs_hisp_ru: dict, ru_group: dict, acs_same_year: float) -> dict:
    z = zipfile.ZipFile(CACHE / "nhts2017_csv.zip")
    per = pd.read_csv(z.open("perpub.csv"), usecols=["HOUSEID", "PERSONID", "WTPERFIN", "R_HISP", "R_AGE"],
                      dtype={"HOUSEID": str, "PERSONID": str, "R_HISP": str})
    trp = pd.read_csv(z.open("trippub.csv"), usecols=["HOUSEID", "PERSONID", "TRPTRANS", "PUBTRANS", "TRIPPURP",
                                                       "WTTRDFIN"],
                      dtype={"HOUSEID": str, "PERSONID": str, "TRPTRANS": str, "PUBTRANS": str, "TRIPPURP": str})
    rcols = ["WTPERFIN"] + [f"WTPERFIN{i}" for i in range(1, 99)]
    rep = pd.read_csv(zipfile.ZipFile(CACHE / "nhts2017_replicates_csv.zip").open("perwgt.csv"),
                      usecols=["HOUSEID", "PERSONID"] + rcols, dtype={"HOUSEID": str, "PERSONID": str})
    cb_p = codebook(CACHE / "nhts2017_codebook_v1.2.xlsx", "CODEBOOK_PER")
    cb_t = codebook(CACHE / "nhts2017_codebook_v1.2.xlsx", "CODEBOOK_TRIP")
    ok_p, lp = codebook_check(cb_p, {("R_HISP", "01"): "Yes, Hispanic or Latino",
                                     ("R_HISP", "02"): "No, Not Hispanic or Latino"})
    ok_t, lt = codebook_check(cb_t, {("PUBTRANS", "01"): "Yes", ("TRPTRANS", "11"): "Public or commuter bus",
                                     ("TRPTRANS", "15"): "Amtrak / Commuter rail",
                                     ("TRPTRANS", "16"): "Subway / elevated / light rail / street car",
                                     ("TRIPPURP", "HBW"): "Home-based trip (work)"})
    transit, hbw = trp["PUBTRANS"].eq("01"), trp["TRIPPURP"].eq("HBW")
    mism_pub = int((trp["TRPTRANS"].isin(["11", "15", "16"]) != transit).sum())
    m = per.merge(rep, on=["HOUSEID", "PERSONID"], how="left", validate="one_to_one", suffixes=("", "_rep"))
    if m["WTPERFIN_rep"].isna().any():
        blocked("gate nhts2017_structure: persons without replicate weights")
    w_diff = float((m["WTPERFIN"] - m["WTPERFIN_rep"]).abs().max())
    tw = trp.merge(per[["HOUSEID", "PERSONID", "WTPERFIN", "R_HISP"]], on=["HOUSEID", "PERSONID"], how="left",
                   validate="many_to_one")
    if tw["WTPERFIN"].isna().any():
        blocked("gate nhts2017_structure: trips without a person record")
    tw_rel = float((tw["WTTRDFIN"] / (365 * tw["WTPERFIN"]) - 1).abs().max())
    gate("nhts2017_codebook_and_derivations", ok_p and ok_t and mism_pub == 0 and w_diff < 1e-6 and tw_rel < 1e-9,
         f"codebook {lp + lt}; PUBTRANS vs TRPTRANS in (11,15,16) mismatches {mism_pub}; replicate-file WTPERFIN max "
         f"abs diff {w_diff:.2e}; WTTRDFIN = 365 x WTPERFIN max rel diff {tw_rel:.2e}")
    kt = tw.loc[transit].groupby(["HOUSEID", "PERSONID"]).size().rename("kt")
    kc = tw.loc[transit & hbw].groupby(["HOUSEID", "PERSONID"]).size().rename("kc")
    m = m.join(kt, on=["HOUSEID", "PERSONID"]).join(kc, on=["HOUSEID", "PERSONID"])
    m[["kt", "kc"]] = m[["kt", "kc"]].fillna(0)
    W = m[[c if c != "WTPERFIN" else "WTPERFIN_rep" for c in rcols]].to_numpy(float)
    hisp = m["R_HISP"].eq("01").to_numpy()
    kt_, kc_ = m["kt"].to_numpy(), m["kc"].to_numpy()
    tot = {"p_all": W.sum(axis=0), "p_h": W[hisp].sum(axis=0),
           "t_all": 365 * (W[kt_ > 0] * kt_[kt_ > 0, None]).sum(axis=0),
           "t_h": 365 * (W[hisp & (kt_ > 0)] * kt_[hisp & (kt_ > 0), None]).sum(axis=0),
           "c_all": 365 * (W[kc_ > 0] * kc_[kc_ > 0, None]).sum(axis=0),
           "c_h": 365 * (W[hisp & (kc_ > 0)] * kc_[hisp & (kc_ > 0), None]).sum(axis=0)}

    def se_jk(x: np.ndarray) -> float:
        return math.sqrt(NHTS17_FACTOR * float(np.sum((x[1:] - x[0]) ** 2)))

    t_trip = float(tw.loc[transit, "WTTRDFIN"].sum())
    ex_ok = (abs(tot["t_all"][0] / NHTS17_EXAMPLE[0] - 1) < 1e-9 and abs(se_jk(tot["t_all"]) / NHTS17_EXAMPLE[1] - 1)
             < 1e-6 and abs(t_trip / tot["t_all"][0] - 1) < 1e-9)
    gate("nhts2017_documented_example", ex_ok,
         f"total transit trips {tot['t_all'][0]:,.1f} (trip file {t_trip:,.1f}) SE {se_jk(tot['t_all']):,.1f} vs "
         f"User's Guide 7.12: 9,444,506,727 SE 212,640,677")
    p_sum = float(tot["p_all"][0])
    gate("nhts2017_person_total", abs(p_sum / NHTS17_PERSON_TOTAL - 1) <= TOL["nhts_person_total_rel"],
         f"sum WTPERFIN {p_sum:,.1f} vs published 301,599,169 (User's Guide Table 7-1)")
    ru_all = ratio_stats(tot["p_all"], tot["p_h"], tot["t_all"], tot["t_h"])
    ru_c = ratio_stats(tot["p_all"], tot["p_h"], tot["c_all"], tot["c_h"])
    adj = ru_all / ru_c
    se_log_adj = se_jk(np.log(adj))
    out_est = lambda x: {"value": float(x[0]), "se": se_jk(x)}  # noqa: E731
    return {
        "source": "2017 NHTS public use v1.2 (perpub.csv, trippub.csv) with the 98 person replicate weights "
                  "(perwgt.csv); persons 5 and older",
        "definitions": {
            "hispanic": "R_HISP == 01 (Yes, Hispanic or Latino); -7/-8 counted in all persons only",
            "transit_trip": "PUBTRANS == 01, i.e. TRPTRANS 11 public or commuter bus, 15 Amtrak/commuter rail, "
                            "16 subway/elevated/light rail/street car",
            "commute_trip": "TRIPPURP == HBW (home-based work)",
            "trip_replicates": "WTTRDFIN = 365 x WTPERFIN on every trip, so trip replicate r = 365 x WTPERFIN_r",
        },
        "codebook_lines": lp + lt,
        "sample": {"persons": int(len(per)), "persons_hispanic": int(hisp.sum()), "trips": int(len(trp)),
                   "transit_trips": int(transit.sum()),
                   "transit_trips_hispanic": int((transit & tw["R_HISP"].eq("01")).sum()),
                   "hbw_transit_trips": int((transit & hbw).sum()),
                   "hbw_transit_trips_hispanic": int((transit & hbw & tw["R_HISP"].eq("01")).sum()),
                   "households_with_hispanic_transit_trips": int(tw.loc[transit & tw["R_HISP"].eq("01"),
                                                                        "HOUSEID"].nunique())},
        "weighted": {k: float(v[0]) for k, v in tot.items()},
        "transit_trips_per_person_year": {"all": float(tot["t_all"][0] / tot["p_all"][0]),
                                          "hispanic": float(tot["t_h"][0] / tot["p_h"][0])},
        "ru_all_trips_hispanic": out_est(ru_all),
        "ru_commute_hbw_hispanic": out_est(ru_c),
        "adjustment_all_trips_over_commute": {**out_est(adj), "se_log": se_log_adj},
        "acs_hispanic_commute_ru_2024": acs_hisp_ru,
        "nhts_commute_ru_over_acs_hispanic_commute_ru": float(ru_c[0]) / acs_hisp_ru["value"],
        "group_ru_all_trip_adjusted": adjusted(ru_group, float(adj[0]), se_log_adj),
        "acs_hispanic_commute_ru_same_year_published": {"year": 2017, "value": acs_same_year},
        "nhts_commute_ru_over_same_year_acs": float(ru_c[0]) / acs_same_year,
        "cross_survey_adjustment_all_trips_over_same_year_acs_commute": {
            "value": float(ru_all[0]) / acs_same_year, "se_log": se_jk(np.log(ru_all))},
        "group_ru_all_trip_adjusted_cross_survey": adjusted(ru_group, float(ru_all[0]) / acs_same_year,
                                                            se_jk(np.log(ru_all)), CROSS),
        "variance": {"method": "98 jackknife replicates, SE = sqrt((6/7) sum (REP_i - x)^2) (User's Guide 7.12); "
                               "reproduces the guide's transit-trip SE"},
        "documented_example_check": {"transit_trips": float(tot["t_all"][0]), "se": se_jk(tot["t_all"]),
                                     "published": list(NHTS17_EXAMPLE)},
    }


# ----------------------------------------------------------------------------------------------------------------
# state deficit weighting


def govs(code: str, tag: str) -> dict[int, tuple[str, int]]:
    rows = json.loads((CACHE / f"census_govslocalfin_2024_{code}_{tag}.json").read_text())
    h = rows[0]
    out = {}
    for r in rows[1:]:
        d = dict(zip(h, r))
        if d["GOVTYPE"] != "001" or d["YEAR"] != "2024" or d["AGG_DESC"] != code:
            blocked(f"gate govs_structure: unexpected row {d['AGG_DESC']} {d['GOVTYPE']} {d['YEAR']}")
        out[int(d["state"]) if tag == "states" else 0] = (d["NAME"], int(d["AMOUNT"]))
    return out


def nipa_transit() -> float:
    wb = openpyxl.load_workbook(LOCAL["nipa"]["path"], read_only=True)
    rows = list(wb["T30800-A"].iter_rows(values_only=True))
    wb.close()
    head = next(r for r in rows if r and r[0] == "Line")
    col = list(head).index("2024")
    line = [r for r in rows if r and r[0] == "14"]
    if len(line) != 1 or line[0][1].strip() != "Public transit" or line[0][2] != "L31224":
        blocked("gate nipa_line: T30800-A line 14 is not Public transit (L31224)")
    return float(line[0][col]) / 1000


def state_block(tab: Tab, direct: dict) -> tuple[dict, list[list]]:
    rev, ops, cap = govs("LF0073", "states"), govs("LF0216", "states"), govs("LF0217", "states")
    us = {c: govs(c, "us")[0][1] for c in ("LF0073", "LF0216", "LF0217")}
    sums = {"LF0073": sum(v for _, v in rev.values()), "LF0216": sum(v for _, v in ops.values()),
            "LF0217": sum(v for _, v in cap.values())}
    gate("govs_states_sum_to_us", all(abs(sums[c] - us[c]) <= 1e-5 * us[c] for c in us),
         "; ".join(f"{c} states {sums[c]:,} vs U.S. {us[c]:,} ($1,000)" for c in us))
    same = set(rev) == set(ops) == set(cap) == set(tab.states)
    gate("govs_states_match_pums_states", same, f"{len(ops)} finance states, {len(tab.states)} PUMS states")
    if not same:
        return {}, []
    D = {s: ops[s][1] - rev[s][1] for s in tab.states}
    K = {s: cap[s][1] for s in tab.states}
    t_all_s = {s: tab.tot("all", TRANSIT, [s]) for s in tab.states}
    t_grp_s = {s: tab.tot("group", TRANSIT, [s]) for s in tab.states}
    n_all_s = {s: tab.tot("all", None, [s]) for s in tab.states}
    n_grp_s = {s: tab.tot("group", None, [s]) for s in tab.states}
    pos_ok = all((t_all_s[s] > 0).all() for s in tab.states)
    gate("state_replicate_transit_positive", pos_ok,
         "every state's transit-commuter total is positive in all 81 weights")
    g = {s: t_grp_s[s] / t_all_s[s] for s in tab.states}
    n_all, n_grp = tab.tot("all"), tab.tot("group")
    pop_share = n_grp / n_all
    t_all, t_grp = tab.tot("all", TRANSIT), tab.tot("group", TRANSIT)
    sum_d, sum_k = sum(D.values()), sum(K.values())
    dw = sum(D[s] * g[s] for s in tab.states) / sum_d
    kw = sum(K[s] * g[s] for s in tab.states) / sum_k
    cw = sum(t_all_s[s] * g[s] for s in tab.states) / t_all
    rows = []
    for s in tab.states:
        rows.append([f"{s:02d}", ops[s][0], int(n_all_s[s][0]), int(n_grp_s[s][0]), n_grp_s[s][0] / n_all_s[s][0],
                     int(t_all_s[s][0]), int(t_grp_s[s][0]), float(g[s][0]), se_acs(g[s]),
                     t_all_s[s][0] / t_all[0], tab.cnt("all", TRANSIT, [s]), tab.cnt("group", TRANSIT, [s]),
                     ops[s][1], rev[s][1], D[s], D[s] / sum_d, K[s], K[s] / sum_k])
    col = {k: sum(r[i] for r in rows) for k, i in (("n_all", 2), ("n_grp", 3), ("t_all", 5), ("t_grp", 6))}
    ref = {"n_all": direct["persons_all"][0], "n_grp": direct["persons_group"][0], "t_all": direct["transit_all"][0],
           "t_grp": direct["transit_group"][0]}
    dev = max(abs(col[k] / ref[k] - 1) for k in col)
    rep_dev = max(float(np.max(np.abs(a / b - 1))) for a, b in (
        (n_all, direct["persons_all"]), (n_grp, direct["persons_group"]), (t_all, direct["transit_all"]),
        (t_grp, direct["transit_group"])))
    gate("state_rows_sum_to_national", dev <= TOL["state_sum_rel"] and rep_dev <= TOL["state_sum_rel"],
         f"state rows vs chunk-level national totals (independent path): max rel dev {dev:.2e} (full weight), "
         f"{rep_dev:.2e} (all 81 weights)")
    nipa = nipa_transit()
    top = sorted(tab.states, key=lambda s: -D[s])[:10]
    ny = next(r for r in rows if r[0] == f"{NY:02d}")
    block = {
        "deficit_source": "Census Bureau, Annual Survey of State and Local Government Finances 2024 (API "
                          "timeseries/govslocalfin, GOVTYPE 001 state and local): D_s = LF0216 transit current "
                          "operations - LF0073 transit utility revenue; K_s = LF0217 transit capital outlay",
        "deficit_year": "2024 (state and local fiscal years ending in 2024, mostly June 30)",
        "sum_opex_bn": sums["LF0216"] / 1e6, "sum_revenue_bn": sums["LF0073"] / 1e6, "sum_D_bn": sum_d / 1e6,
        "sum_capital_outlay_bn": sum_k / 1e6,
        "nipa_3_8_line_14_2024_bn": nipa,
        "ratio_sum_D_to_nipa_deficit": (sum_d / 1e6) / abs(nipa),
        "nipa_note": "NIPA 3.8 line 14 is calendar 2024 current surplus of S&L public transit (BEA concept); "
                     "Census D_s is fiscal-year opex less transit utility revenue, so no hard tolerance",
        "group_population_share": est(pop_share),
        "deficit_weighted_group_share": est(dw),
        "ru_state_weighted": est(dw / pop_share),
        "capital_outlay_weighted_group_share": est(kw),
        "ru_capital_outlay_weighted": est(kw / pop_share),
        "commuter_weighted_identity_check": {"ru": float(cw[0] / pop_share[0]),
                                             "note": "weights T_s reproduce acs.national.ru"},
        "top10_by_deficit": [{"state_fips": f"{s:02d}", "state": ops[s][0], "D_bn": D[s] / 1e6, "D_share": D[s] / sum_d,
                              "transit_commuter_share": float(t_all_s[s][0] / t_all[0]), "g_s": float(g[s][0]),
                              "g_s_se": se_acs(g[s]), "group_population_share_in_state": float(n_grp_s[s][0] /
                                                                                              n_all_s[s][0])}
                             for s in top],
        "new_york": {"D_share": ny[15], "transit_commuter_share": ny[9], "g_s": ny[7], "g_s_se": ny[8],
                     "group_population_share_in_state": ny[4], "capital_outlay_share": ny[17]},
        "method": "RU_state_weighted = (sum_s D_s g_s / sum_s D_s) / (group persons / all persons); g_s by state of "
                  "residence; D_s held fixed across the 80 replicates",
        "limits": "D_s is booked to the operating government's state (WMATA to DC, the MTA to NY) while g_s is by "
                  "state of residence; cross-state commuters (NJ, CT, MD, VA) ride in another state's system",
    }
    return block, rows


# ----------------------------------------------------------------------------------------------------------------


def candidates(national: dict, variants: dict, sw: dict, n17: dict, n22: dict) -> dict:
    """The RU values a consumer may choose between, with where each lives in this file."""
    def c(path: str, x: dict, note: str) -> dict:
        out = {"path": path, "value": x["value"], "note": note}
        out |= {"se": x["se"]} if "se" in x else {"ci95": x["ci95"]}
        return out

    return {
        "central_commuters_national": c("acs.national.ru", national["ru"],
                                        "ACS 2024 PUMS transit commuters per person, group vs all"),
        "state_deficit_weighted": c("state_weighted.ru_state_weighted", sw["ru_state_weighted"],
                                    "g_s weighted by 2024 Census state transit deficits instead of transit commuters"),
        "all_trip_nhts2017_within": c("nhts_2017.group_ru_all_trip_adjusted", n17["group_ru_all_trip_adjusted"],
                                      "central x NHTS 2017 Hispanic all-trip RU / HBW RU"),
        "all_trip_nhts2017_cross_survey": c("nhts_2017.group_ru_all_trip_adjusted_cross_survey",
                                            n17["group_ru_all_trip_adjusted_cross_survey"],
                                            "central x NHTS 2017 Hispanic all-trip RU / ACS 2017 Hispanic commute RU"),
        "all_trip_nhts2022_within": c("nhts_2022.group_ru_all_trip_adjusted", n22["group_ru_all_trip_adjusted"],
                                      "thin: 22 Hispanic HBW transit trips in 12 households"),
        "without_04_commuter_rail": c("acs.variants.without_04_commuter_rail.ru",
                                      variants["without_04_commuter_rail"]["ru"], "drops JWTRNS 04"),
        "excluding_new_york": c("acs.variants.excluding_new_york.ru", variants["excluding_new_york"]["ru"],
                                "sensitivity only"),
    }


def main() -> int:
    print("[transit_key] verifying input hashes", flush=True)
    sources = verify_inputs()
    labels, dict_lines = read_dictionary(LOCAL["dictionary"]["path"])
    print("[transit_key] reading ACS 2024 PUMS persons", flush=True)
    tab, direct, chk = pums_aggregate(LOCAL["pums"]["path"])
    gate("pums_structure", chk["bad_jwtrns"] == 0 and chk["bad_hisp"] == 0 and chk["workers_under_16"] == 0
         and len(tab.states) == 51,
         f"{chk['rows']:,} person records, {len(tab.states)} states; JWTRNS outside 01-12 {chk['bad_jwtrns']}, "
         f"HISP outside 01-24 {chk['bad_hisp']}, JWTRNS set with AGEP < 16 {chk['workers_under_16']}")
    national, variants = acs_block(tab, labels)
    hisp = {
        "definition": "HISP != 01 (all Hispanic or Latino origins)",
        "ru": est((tab.tot("hispanic_all", TRANSIT) / tab.tot("hispanic_all"))
                  / (tab.tot("all", TRANSIT) / tab.tot("all"))),
        "direct_share_of_transit_commuters": est(tab.tot("hispanic_all", TRANSIT) / tab.tot("all", TRANSIT)),
        "population_share": est(tab.tot("hispanic_all") / tab.tot("all")),
        "mode_share_public_transportation": national["mode_shares_among_workers"]["hispanic_all"][
            "public_transportation_02_06"],
    }
    ses = [national["ru"]["se"], national["direct_share_of_transit_commuters"]["se"],
           national["group_population_share"]["se"], hisp["ru"]["se"]] + [v["ru"]["se"] for v in variants.values()]
    gate("replicate_se_finite_positive", all(math.isfinite(x) and x > 0 for x in ses),
         f"{len(ses)} replicate SEs, min {min(ses):.3g}, max {max(ses):.3g}")
    print(f"[transit_key] RU central {national['ru']['value']:.4f} (SE {national['ru']['se']:.4f})", flush=True)
    control = published_control(tab, national, hisp["ru"])
    pub_h = control["published_hispanic_commute_ru"]["by_year"]
    print("[transit_key] NHTS 2022", flush=True)
    n22 = nhts_2022(hisp["ru"], national["ru"], pub_h["2022"]["ru"])
    print("[transit_key] NHTS 2017", flush=True)
    n17 = nhts_2017(hisp["ru"], national["ru"], pub_h["2017"]["ru"])
    print("[transit_key] state deficit weighting", flush=True)
    sw, rows = state_block(tab, direct)
    failed = [g for g in GATES if not g["pass"]]
    nyb = sw.get("new_york", {"transit_commuter_share": 0.0, "D_share": 0.0, "g_s": 0.0})
    if failed:
        for g in failed:
            print(f"[BLOCKED] gate {g['gate']}: {g['detail']}")
        return 1
    out = {
        "meta": {
            "sources": sources,
            "group": {"definition": GROUP_DEF, "citation": GROUP_CITE,
                      "universe": "all persons including group quarters, ACS 2024 1-year person PUMS, weight PWGTP"},
            "transit_codes": {f"{c:02d}": labels[c] for c in TRANSIT},
            "dictionary_lines": dict_lines,
            "use": "riders' key = (the case's population key) x RU; central RU = acs.national.ru.value",
            "se_formula": "ACS: sqrt(4/80 * sum_r (x_r - x)^2), PWGTP1-80; NHTS 2017: sqrt(6/7 * sum_i (x_i - x)^2), "
                          "98 replicates; NHTS 2022: stratified household bootstrap",
            "script": rel(Path(__file__).resolve()),
            "ru_candidates": candidates(national, variants, sw, n17, n22),
            "limits": [
                "Commuters are not riders: ACS records the usual mode of workers 16+ at work in the reference week; "
                "riders also include children, students, non-work trips and part-week riders. The NHTS adjustments "
                "address this, but disagree in sign (nhts_2017 within-survey vs cross-survey; nhts_2022 is thin)",
                f"New York: its residents are {100 * nyb['transit_commuter_share']:.1f}% of U.S. transit commuters "
                f"while its systems carry {100 * nyb['D_share']:.1f}% of the 2024 S&L transit deficit (Census D_s); "
                f"the group is {100 * nyb['g_s']:.1f}% of New York's transit commuters against "
                f"{100 * national['group_population_share']['value']:.1f}% of U.S. persons, so the commuter-weighted "
                "national RU leans on New York more than the deficit does (see state_weighted)",
                "State weighting assigns D_s to the operating government's state and g_s to riders' state of "
                "residence, and misses within-state differences (for example Los Angeles against the Bay Area)",
                "The group proxy adds Mexican-born persons of other or no Hispanic origin to HISP 02 "
                f"(+{100 * (national['persons']['group'] / national['persons']['mexican_origin_hisp02'] - 1):.1f}% "
                "persons)"],
        },
        "acs": {"national": national, "variants": variants, "hispanic_all": hisp},
        "published_control": control,
        "nhts_2022": n22,
        "nhts_2017": n17,
        "state_weighted": sw,
        "gates": GATES,
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(sig(out), indent=1, sort_keys=True) + "\n")
    header = ["state_fips", "state", "persons_all", "persons_group", "group_population_share",
              "transit_commuters_all", "transit_commuters_group", "g_s_group_share_of_state_transit_commuters",
              "g_s_se", "state_share_of_national_transit_commuters", "sample_transit_commuters_all",
              "sample_transit_commuters_group", "transit_current_operations_k", "transit_utility_revenue_k",
              "transit_deficit_D_k", "D_share", "transit_capital_outlay_k", "capital_outlay_share"]
    with open(OUT_CSV, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([fmt(v) if isinstance(v, float) else v for v in r])
    print(f"[transit_key] wrote {rel(OUT_JSON)} and {rel(OUT_CSV)} ({len(rows)} states); {len(GATES)} gates pass")
    return 0


if __name__ == "__main__":
    sys.exit(main())
