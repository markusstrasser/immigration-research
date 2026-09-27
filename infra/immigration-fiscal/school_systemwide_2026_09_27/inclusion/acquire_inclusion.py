#!/usr/bin/env python3
"""NAEP state inclusion/exclusion panel (reading and mathematics, grades 4 and 8).

Builds a state x year x grade x subject panel of the percentages of public school students
identified as English learners (EL), students with disabilities (SD) and SD and/or EL, and
the percentages excluded from NAEP, from NCES primary sources:

* Report-card technical appendices (XLSX, the machine-readable twin of the PDF tables):
  2024 "Appendix Tables for 2024 Mathematics/Reading Report Card" (state_district), state
  trend tables "... identified as English learners excluded and assessed ... when
  accommodations were permitted, by state/jurisdiction: Various years, 2000-24" (math) and
  "1998-2024" (reading).  These are the primary source for identified/excluded/assessed as a
  percentage of all students.  The 2022 and 2019 appendices carry the same tables through
  their own year and are kept as cross-check vintages.
* Excluded as a percentage of identified students: 2017 state appendix trend tables
  ("... ELL excluded ..., as a percentage of identified ELL students ...: Various years,
  1992-2017") and the per-year tables of the 2017, 2019, 2022 and 2024 appendices.
* NAEP Technical Documentation "Weighted student response and exclusion rates" pages
  (one or two decimals; not censored at 0.5 the way report-card "#" cells are).
* Validation: the 2024 state_district PDFs (integers) and the 2024 national appendix PDFs.
* Primary-text quotes: 20 U.S.C. 6311(b)(3)(A) (law.cornell.edu) and the NAGB 2010 policy.

Run from the repository root:
    OPENBLAS_NUM_THREADS=1 uv run --no-project python3 \
        infra/immigration-fiscal/school_systemwide_2026_09_27/inclusion/acquire_inclusion.py

Raw downloads go to inclusion/_cache/ (git-ignored) and are fetched only when absent.
All values are parsed with code (openpyxl, BeautifulSoup, pdftotext -layout); none are
typed in by hand.  Outputs are deterministic (sorted rows, "\n" line endings).
"""
from __future__ import annotations

import csv
import hashlib
import random
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
import requests
from bs4 import BeautifulSoup

LANE = Path(__file__).resolve().parent
CACHE = LANE / "_cache"
DERIVED = LANE / "derived"
UA = "Mozilla/5.0 (X11; Linux x86_64) research-data-fetch"

NRC = "https://www.nationsreportcard.gov"
TDW = "https://nces.ed.gov/nationsreportcard/tdw/sample_design"

# ---------------------------------------------------------------------------------------
# Sources
# ---------------------------------------------------------------------------------------
TA = {  # report-card technical appendices: (vintage, subject) -> url
    (2024, "mathematics"): f"{NRC}/reports/mathematics/2024/g4_8/supporting-files/2024_technical_appendix_math_state_district.xlsx",
    (2024, "reading"): f"{NRC}/reports/reading/2024/g4_8/supporting-files/2024_technical_appendix_reading_state_district.xlsx",
    (2022, "mathematics"): f"{NRC}/mathematics/supportive_files/2022_technical_appendix_math.xlsx",
    (2022, "reading"): f"{NRC}/reading/supportive_files/2022_technical_appendix_reading.xlsx",
    (2019, "mathematics"): f"{NRC}/mathematics/supportive_files/2019_Technical_Appendix_Math.xlsx",
    (2019, "reading"): f"{NRC}/reading/supportive_files/2019_Technical_Appendix_Reading.xlsx",
    (2017, "mathematics"): f"{NRC}/math_2017/files/2017_Technical_Appendix_Math_State.xlsx",
    (2017, "reading"): f"{NRC}/reading_2017/files/2017_Technical_Appendix_Reading_State.xlsx",
}
TA_PDF = {  # PDFs used for validation of the XLSX values
    (2024, "mathematics"): f"{NRC}/reports/mathematics/2024/g4_8/supporting-files/2024_technical_appendix_math_state_district.pdf",
    (2024, "reading"): f"{NRC}/reports/reading/2024/g4_8/supporting-files/2024_technical_appendix_reading_state_district.pdf",
    (2017, "mathematics"): f"{NRC}/math_2017/files/2017_Technical_Appendix_Math_State.pdf",
    (2017, "reading"): f"{NRC}/reading_2017/files/2017_Technical_Appendix_Reading_State.pdf",
}
NATIONAL_PDF = {
    "mathematics": f"{NRC}/reports/mathematics/2024/g4_8/supporting-files/2024_technical_appendix_math_national.pdf",
    "reading": f"{NRC}/reports/reading/2024/g4_8/supporting-files/2024_technical_appendix_reading_national.pdf",
}
TDW_PAGES = {  # (year, subject) -> url; located by crawling each year's sample-design tree
    (2002, "reading"): f"{TDW}/2002_2003/sampdsgn_2002_state_studresp_table1.aspx",
    (2003, "reading"): f"{TDW}/2002_2003/sampdsgn_2003_state_studresp_table1.aspx",
    (2003, "mathematics"): f"{TDW}/2002_2003/sampdsgn_2003_state_studresp_table2.aspx",
    (2005, "reading"): f"{TDW}/2004_2005/sampdsgn_2005_state_studresp_table1.aspx",
    (2005, "mathematics"): f"{TDW}/2004_2005/sampdsgn_2005_state_studresp_table2.aspx",
    (2007, "reading"): f"{TDW}/2007/sampdsgn_2007_state_studresp_table1.aspx",
    (2007, "mathematics"): f"{TDW}/2007/sampdsgn_2007_state_studresp_table2.aspx",
    (2009, "reading"): f"{TDW}/2009/2009_sampdsgn_state_studresp_reading.aspx",
    (2009, "mathematics"): f"{TDW}/2009/2009_sampdsgn_state_studresp_math.aspx",
    (2011, "reading"): f"{TDW}/2011/2011_sampdsgn_state_studresp_reading.aspx",
    (2011, "mathematics"): f"{TDW}/2011/2011_sampdsgn_state_studresp_math.aspx",
}
for _y in (2013, 2015, 2017, 2019, 2022, 2024):
    for _s in ("reading", "mathematics"):
        TDW_PAGES[(_y, _s)] = (f"{TDW}/{_y}/weighted_student_response_and_exclusion_rates_for_the_"
                               f"{_y}_state_{_s}_assessment.aspx")
ESSA_URL = "https://www.law.cornell.edu/uscode/text/20/6311"
NAGB_URL = ("https://www.nagb.gov/content/dam/nagb/en/documents/policies/"
            "naep_testandreport_studentswithdisabilities.pdf")
NCES_HISTORY_URL = "https://nces.ed.gov/nationsreportcard/about/history_inclusion.aspx"

ASSESSMENT_YEARS = {
    "mathematics": [2000, 2003, 2005, 2007, 2009, 2011, 2013, 2015, 2017, 2019, 2022, 2024],
    "reading": [1998, 2002, 2003, 2005, 2007, 2009, 2011, 2013, 2015, 2017, 2019, 2022, 2024],
}

STATES = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR", "California": "CA",
    "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE", "Florida": "FL", "Georgia": "GA",
    "Hawaii": "HI", "Idaho": "ID", "Illinois": "IL", "Indiana": "IN", "Iowa": "IA",
    "Kansas": "KS", "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS",
    "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV", "New Hampshire": "NH",
    "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY", "North Carolina": "NC",
    "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR", "Pennsylvania": "PA",
    "Rhode Island": "RI", "South Carolina": "SC", "South Dakota": "SD", "Tennessee": "TN",
    "Texas": "TX", "Utah": "UT", "Vermont": "VT", "Virginia": "VA", "Washington": "WA",
    "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
}
TYPE_RANK = {"nation_public": 0, "state": 1, "dc": 2, "dodea": 3, "territory": 4}
CATS = ("el", "sd", "sdel")
MEASURES = ("identified", "excluded", "assessed", "assessed_noacc", "assessed_acc")
CAT_LABEL = {"el": "EL", "sd": "SD", "sdel": "SD and/or EL"}


# ---------------------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------------------
def local_path(url: str) -> Path:
    name = re.sub(r"^https?://", "", url).replace("/", "__")
    return CACHE / name


def fetch(url: str) -> Path:
    """Download url into the cache once; later runs reuse the cached bytes."""
    path = local_path(url)
    if path.exists() and path.stat().st_size > 0:
        return path
    CACHE.mkdir(parents=True, exist_ok=True)
    last = None
    for attempt in range(4):
        try:
            r = requests.get(url, headers={"User-Agent": UA}, timeout=600)
            if r.status_code == 200 and r.content:
                tmp = path.with_name(path.name + ".part")
                tmp.write_bytes(r.content)
                tmp.rename(path)
                return path
            last = f"HTTP {r.status_code}"
        except requests.RequestException as exc:  # transport errors: retry, then fail loudly
            last = repr(exc)
        time.sleep(5 * (attempt + 1))
    raise SystemExit(f"[BLOCKED] download failed for {url}: {last}")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------------------
# Cell and label normalisation
# ---------------------------------------------------------------------------------------
FLAG_TOKENS = {
    "#": ("0", "rounds_to_zero"),
    "‡": ("", "reporting_standards_not_met"),
    "—": ("", "not_available"),
    "–": ("", "not_available"),
    "†": ("", "not_applicable"),
}
NUM_RE = re.compile(r"^-?\d+(?:\.\d+)?$")


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).replace(" ", " ").replace("​", "")
    return re.sub(r"\s+", " ", s).strip()


def parse_val(raw: str) -> tuple[str, str]:
    """Return (value, flag).  '#' -> ('0', rounds_to_zero); '‡' -> ('', reporting_standards_not_met)."""
    s = raw.strip()
    if s in FLAG_TOKENS:
        return FLAG_TOKENS[s]
    if s == "":
        return "", "blank"
    if NUM_RE.match(s):
        return s, ""
    if re.fullmatch(r"\d+,\d+", s):  # decimal comma typo in a source cell: keep, but flag it
        return s.replace(",", "."), "decimal_comma_in_source"
    raise ValueError(f"unparsed cell value {raw!r}")


def norm_name(raw) -> str:
    s = cell_str(raw).replace("’", "'")
    s = re.sub(r"[¹²³⁰-⁹]+$", "", s).strip()
    s = re.sub(r"\s+\d$", "", s)
    s = re.sub(r"(?<=[A-Za-z])\d$", "", s)
    return re.sub(r"\s+", " ", s).strip()


def juris(name: str):
    if name in STATES:
        return name, STATES[name], "state"
    if name == "Nation (public)":
        return name, "NP", "nation_public"
    if name == "District of Columbia":
        return name, "DC", "dc"
    if name in ("DoDEA", "Department of Defense Education Activity (DoDEA)",
                "Department of Defense Education Activity", "Department of Defense",
                "D epartment of Defense Education Activity (DoDEA)"):
        return "DoDEA", "DD", "dodea"
    if name == "Puerto Rico":
        return name, "PR", "territory"
    return None


def measure_of(label: str):
    k = re.sub(r"[^a-z]", "", label.lower())
    return {
        "identified": "identified", "excluded": "excluded", "assessed": "assessed",
        "assessedwithoutaccommodations": "assessed_noacc", "withoutaccommodations": "assessed_noacc",
        "assessedwithaccommodations": "assessed_acc", "withaccommodations": "assessed_acc",
    }.get(k)


def classify_title(t: str) -> dict:
    t0 = cell_str(t)
    if "fourth- and eighth-grade" in t0:
        grade = None
    else:
        grade = 4 if "fourth-grade" in t0 else 8 if "eighth-grade" in t0 else None
    subject = ("mathematics" if "NAEP mathematics" in t0 else
               "reading" if "NAEP reading" in t0 else None)
    if re.search(r"students with disabilities(?: \(SD\))? and/or English (?:language )?learners", t0):
        cat = "sdel"
    elif re.search(r"students with disabilities", t0):
        cat = "sd"
    elif re.search(r"English (?:language )?learners", t0):
        cat = "el"
    else:
        cat = None
    regime = ("not_permitted" if "accommodations were not permitted" in t0 else
              "permitted" if "accommodations were permitted" in t0 else None)
    kind = "pct_identified" if "as a percentage of identified" in t0 else "pct_all"
    m = re.search(r":\s*(\d{4})\s*$", t0)
    return dict(grade=grade, subject=subject, cat=cat, regime=regime, kind=kind,
                by_state="by state/jurisdiction" in t0, single_year=int(m.group(1)) if m else None,
                various="Various years" in t0, title=t0)


def sheet_title(rows) -> str:
    for r in rows[:9]:
        for v in r:
            s = cell_str(v)
            if s.startswith("Table A-") or s.startswith("Percentage of"):
                return s
    return ""


def find_year_row(rows, lo=0, hi=12):
    for i in range(lo, min(hi, len(rows))):
        yrs = [(j, v) for j, v in enumerate(rows[i])
               if j >= 1 and isinstance(v, int) and 1985 <= v <= 2030]
        if yrs:
            return i, yrs
    return None, []


# ---------------------------------------------------------------------------------------
# XLSX parsers
# ---------------------------------------------------------------------------------------
def parse_pct_all_sheet(rows, sheet, meta, vintage, url, unmatched):
    """State trend table: identified/excluded/assessed as a percentage of all students."""
    yi, yrs = find_year_row(rows)
    if yi is None:
        raise SystemExit(f"[BLOCKED] no year row in {vintage} {sheet}")
    hi = next((i for i in range(yi, yi + 4)
               if any(cell_str(v) == "Identified" for v in rows[i])), None)
    if hi is None:
        raise SystemExit(f"[BLOCKED] no header row in {vintage} {sheet}")
    ymap = dict(yrs)
    ncol = max(len(r) for r in rows)
    cur, colmap = None, {}
    for j in range(1, ncol):
        if j in ymap:
            cur = ymap[j]
        lab = cell_str(rows[hi][j]) if j < len(rows[hi]) else ""
        m = measure_of(lab) if lab else None
        if m and cur:
            colmap[j] = (cur, m)
    out = []
    for i in range(hi + 1, len(rows)):
        row = rows[i]
        name = norm_name(row[0]) if row else ""
        if not name:
            continue
        jj = juris(name)
        if jj is None:
            unmatched[(vintage, sheet)].add(name)
            continue
        for j, (year, m) in colmap.items():
            raw = cell_str(row[j]) if j < len(row) else ""
            val, flag = parse_val(raw)
            out.append(dict(vintage=vintage, subject=meta["subject"], grade=meta["grade"],
                            cat=meta["cat"], regime=meta["regime"], year=year,
                            jurisdiction=jj[0], abbr=jj[1], jtype=jj[2], measure=m,
                            raw=raw, value=val, flag=flag, sheet=sheet, url=url))
    return out


def parse_pct_ident_year_sheet(rows, sheet, meta, vintage, url, unmatched):
    """Per-year table: excluded/assessed as a percentage of identified SD and/or EL, SD, EL."""
    ci = next((i for i in range(0, 12) if any(
        cell_str(v) in ("SD and/or EL", "SD and/or ELL") for v in rows[i])), None)
    if ci is None:
        raise SystemExit(f"[BLOCKED] no category row in {vintage} {sheet}")
    li = next((i for i in range(ci, ci + 3) if any(cell_str(v) == "Excluded" for v in rows[i])), None)
    ncol = max(len(r) for r in rows)
    cur, colmap = None, {}
    for j in range(1, ncol):
        c = cell_str(rows[ci][j]) if j < len(rows[ci]) else ""
        if c:
            cur = {"SD and/or EL": "sdel", "SD and/or ELL": "sdel", "SD": "sd",
                   "EL": "el", "ELL": "el"}.get(c)
        lab = cell_str(rows[li][j]) if j < len(rows[li]) else ""
        m = measure_of(lab) if lab else None
        if cur and m:
            colmap[j] = (cur, m)
    out = []
    for i in range(li + 1, len(rows)):
        row = rows[i]
        name = norm_name(row[0]) if row else ""
        if not name:
            continue
        jj = juris(name)
        if jj is None:
            unmatched[(vintage, sheet)].add(name)
            continue
        for j, (cat, m) in colmap.items():
            raw = cell_str(row[j]) if j < len(row) else ""
            val, flag = parse_val(raw)
            out.append(dict(vintage=vintage, src="year", subject=meta["subject"], grade=meta["grade"], cat=cat,
                            regime="permitted", year=meta["single_year"], jurisdiction=jj[0],
                            abbr=jj[1], jtype=jj[2], measure=m + "_pct_identified", raw=raw,
                            value=val, flag=flag, sheet=sheet, url=url))
    return out


def parse_2017_ident_trend(rows, sheet, meta, url, unmatched):
    """2017 state trend: SD or ELL excluded as a percentage of identified, 1990/1992-2017."""
    yi, yrs = find_year_row(rows)
    if yi is None:
        raise SystemExit(f"[BLOCKED] no year row in 2017 {sheet}")
    notperm = set()
    for j, y in yrs:
        nxt = rows[yi][j + 1] if j + 1 < len(rows[yi]) else None
        if cell_str(nxt) == "1":
            notperm.add(y)
    out = []
    for i in range(yi + 1, len(rows)):
        row = rows[i]
        name = norm_name(row[0]) if row else ""
        if not name:
            continue
        jj = juris(name)
        if jj is None:
            unmatched[(2017, sheet)].add(name)
            continue
        for j, y in yrs:
            raw = cell_str(row[j]) if j < len(row) else ""
            val, flag = parse_val(raw)
            out.append(dict(vintage=2017, src="trend", subject=meta["subject"], grade=meta["grade"],
                            cat=meta["cat"], regime="not_permitted" if y in notperm else "permitted",
                            year=y, jurisdiction=jj[0], abbr=jj[1], jtype=jj[2],
                            measure="excluded_pct_identified", raw=raw, value=val, flag=flag,
                            sheet=sheet, url=url))
    return out


def load_appendix(vintage, subject, unmatched):
    url = TA[(vintage, subject)]
    wb = openpyxl.load_workbook(fetch(url), read_only=True, data_only=True)
    pct_all, pct_ident, national = [], [], []
    for sheet in wb.sheetnames:
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        if not rows:
            continue
        meta = classify_title(sheet_title(rows))
        if meta["subject"] != subject or meta["grade"] is None and not meta["title"]:
            continue
        if vintage == 2017:
            if meta["kind"] == "pct_identified" and meta["various"] and meta["cat"] in ("sd", "el"):
                pct_ident += parse_2017_ident_trend(rows, sheet, meta, url, unmatched)
            elif (meta["kind"] == "pct_identified" and meta["single_year"] == 2017
                  and meta["grade"] and "State_Iden" in sheet):
                pct_ident += parse_pct_ident_year_sheet(rows, sheet, meta, 2017, url, unmatched)
            continue
        if not meta["by_state"]:
            if (meta["kind"] == "pct_all" and meta["regime"] == "permitted" and meta["various"]
                    and meta["grade"] is None and vintage == 2024):
                national.append((sheet, rows))
            continue
        if meta["kind"] == "pct_all" and meta["regime"] and meta["various"] and meta["cat"]:
            pct_all += parse_pct_all_sheet(rows, sheet, meta, vintage, url, unmatched)
        elif meta["kind"] == "pct_all" and meta["regime"] and meta["cat"] and meta["grade"]:
            # 1992/1996/2000 (math) and 1992/1994/1998 (reading) not-permitted tables are titled
            # with explicit years rather than "Various years".
            pct_all += parse_pct_all_sheet(rows, sheet, meta, vintage, url, unmatched)
        elif meta["kind"] == "pct_identified" and meta["single_year"] and meta["grade"]:
            pct_ident += parse_pct_ident_year_sheet(rows, sheet, meta, vintage, url, unmatched)
    return pct_all, pct_ident, national


# ---------------------------------------------------------------------------------------
# TDW parser
# ---------------------------------------------------------------------------------------
def parse_tdw(year, subject, unmatched):
    url = TDW_PAGES[(year, subject)]
    soup = BeautifulSoup(fetch(url).read_text(encoding="utf-8", errors="replace"), "html.parser")
    out = []
    for tb in soup.find_all("table"):
        rows = [[cell_str(c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"])]
                for tr in tb.find_all("tr")]
        hi = next((i for i, r in enumerate(rows[:6])
                   if re.search(r"Fourth|Grade 4", " ".join(r)) and re.search(r"Eighth|Grade 8", " ".join(r))),
                  None)
        if hi is None:
            continue
        labels = rows[hi + 1]
        kinds = []
        for lab in labels:
            compact = re.sub(r"[^A-Za-z]", "", lab)  # some pages print "E L" for "EL"
            if "responserate" in compact.lower():
                kinds.append("response_rate")
            elif re.search(r"(?:are|were)SDand|SDstudents", compact):
                kinds.append("sd_excluded")
            elif re.search(r"(?:are|were)(?:ELL|EL)and|LEPstudents", compact):
                kinds.append("el_excluded")
            else:
                raise SystemExit(f"[BLOCKED] unknown TDW column {lab!r} in {url}")
        if len(kinds) != 6 or kinds[:3] != kinds[3:]:
            raise SystemExit(f"[BLOCKED] unexpected TDW column layout {kinds} in {url}")
        for r in rows[hi + 2:]:
            if len(r) != 1 + len(kinds):
                continue
            name = norm_name(r[0])
            if name == "Total":
                jj, label = ("Nation (public)", "NP", "nation_public"), "Total"
            else:
                jj, label = juris(name), name
            if jj is None:
                unmatched[("TDW", year, subject)].add(name)
                continue
            for k, (kind, raw) in enumerate(zip(kinds, r[1:])):
                grade = 4 if k < 3 else 8
                val, flag = parse_val(raw)
                out.append(dict(year=year, subject=subject, grade=grade, jurisdiction=jj[0],
                                abbr=jj[1], jtype=jj[2], kind=kind, raw=raw, value=val, flag=flag,
                                row_label=label, row_text=" | ".join(r), url=url))
        break
    if not out:
        raise SystemExit(f"[BLOCKED] no TDW table parsed from {url}")
    return out


# ---------------------------------------------------------------------------------------
# PDF validation parsers (pdftotext -layout)
# ---------------------------------------------------------------------------------------
TOK_RE = re.compile(r"^(?:\d+|#|—|‡|†)$")


def pdf_pages(path: Path) -> list[str]:
    res = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True)
    return res.stdout.decode("utf-8").split("\f")


def parse_state_pdf(subject):
    """Integer cells of the 2024 state trend tables (accommodations permitted), keyed like the xlsx."""
    url = TA_PDF[(2024, subject)]
    cells, anomalies = {}, []
    for pno, page in enumerate(pdf_pages(fetch(url)), start=1):
        lines = page.splitlines()
        head = cell_str(" ".join(lines[:8]))
        if "Table of Contents" in head:
            continue
        if not re.search(r"Percentage of (fourth|eighth)-grade public school students identified as", head):
            continue
        if "when accommodations were permitted" not in head or "by state/jurisdiction" not in head:
            continue
        meta = classify_title(head.split(" Various years")[0] + " NAEP " + subject)
        grade, cat = meta["grade"], meta["cat"]
        years = None
        start = None
        for i, l in enumerate(lines):
            if years is None and re.fullmatch(r"\s*\d{4}(?:\s+\d{4})*\s*", l):
                years = [int(x) for x in l.split()]
            if l.strip().startswith("State/jurisdiction"):
                start = i
                break
        if years is None or start is None:
            anomalies.append(f"p{pno}: no year/header line")
            continue
        ntok = 5 * len(years)
        prev_fragment, pending = None, None
        for l in lines[start + 1:]:
            st = l.strip()
            if not st:
                continue
            if st.startswith(("— Not", "# Rounds", "‡ Reporting", "See notes", "NOTE", "SOURCE", "¹")):
                break
            toks = st.split()
            k = next((n for n, t in enumerate(toks) if TOK_RE.match(t)), len(toks))
            name, vals = " ".join(toks[:k]), toks[k:]
            if not vals:
                if pending is not None:  # wrapped name: "District of" / values / "Columbia"
                    full = norm_name(pending[0] + " " + name)
                    cells_row(cells, subject, grade, cat, years, full, pending[1], pno, pending[2], anomalies)
                    pending = None
                else:
                    prev_fragment = name
                continue
            if not name and prev_fragment:
                pending = (prev_fragment, vals, st)
                prev_fragment = None
                continue
            if len(vals) != ntok:
                anomalies.append(f"p{pno}: {len(vals)} tokens for {name!r}")
                continue
            cells_row(cells, subject, grade, cat, years, norm_name(name), vals, pno, st, anomalies)
    return cells, anomalies


def cells_row(cells, subject, grade, cat, years, name, vals, pno, line, anomalies):
    jj = juris(name)
    if jj is None:
        if name not in ("Other jurisdictions",):
            anomalies.append(f"p{pno}: unmatched name {name!r}")
        return
    if len(vals) != 5 * len(years):
        anomalies.append(f"p{pno}: {len(vals)} tokens for {name!r}")
        return
    for yi, y in enumerate(years):
        for mi, m in enumerate(MEASURES):
            cells[(subject, grade, cat, y, jj[0], m)] = (vals[5 * yi + mi], pno, line)


def pdf_rows(lines, start, ntok):
    """Yield (name, tokens, line) for state rows after a header line; joins 'District of' / 'Columbia'."""
    prev_fragment, pending = None, None
    for l in lines[start + 1:]:
        st = l.strip()
        if not st:
            continue
        if st.startswith(("— Not", "# Rounds", "‡ Reporting", "See notes", "NOTE", "SOURCE", "¹", "1 Acc", "2 Dep")):
            break
        toks = st.split()
        k = next((n for n, t in enumerate(toks) if TOK_RE.match(t)), len(toks))
        name, vals = " ".join(toks[:k]), toks[k:]
        if not vals:
            if pending is not None:
                yield norm_name(pending[0] + " " + name), pending[1], pending[2]
                pending = None
            else:
                prev_fragment = name
            continue
        if not name and prev_fragment:
            pending = (prev_fragment, vals, st)
            prev_fragment = None
            continue
        if len(vals) == ntok:
            yield norm_name(name), vals, st


def parse_pct_identified_pdfs():
    """Integer cells of the %-of-identified tables: 2017 state trends and the 2024 per-year tables."""
    cells = {}
    for (vint, subject), url in sorted(TA_PDF.items()):
        for pno, page in enumerate(pdf_pages(fetch(url)), start=1):
            lines = page.splitlines()
            head = cell_str(" ".join(lines[:10]))
            if "Table of Contents" in head or "as a percentage of identified" not in head:
                continue
            m = re.search(r"Percentage of (fourth|eighth)-grade public (?:and nonpublic )?school", head)
            if not m:
                continue
            grade = 4 if m.group(1) == "fourth" else 8
            hdr = next((i for i, l in enumerate(lines) if l.strip().startswith("State/jurisdiction")), None)
            if hdr is None:
                continue
            if vint == 2017 and "Various years" in head:
                cat = "el" if "English language learners (ELL) excluded" in head else (
                    "sd" if "students with disabilities (SD) excluded" in head else None)
                if cat is None:
                    continue
                # years sit on the header line or the line above it; "1" markers are footnotes
                yl = lines[hdr] if re.search(r"(?:19|20)\d\d", lines[hdr]) else lines[hdr + 1]
                years = [int(t) for t in yl.split() if re.fullmatch(r"(?:19|20)\d\d", t)]
                start = hdr if yl is lines[hdr] else hdr + 1
                for name, vals, st in pdf_rows(lines, start, len(years)):
                    jj = juris(name)
                    if jj:
                        for y, tok in zip(years, vals):
                            cells[(2017, subject, grade, cat, y, jj[0])] = (tok, url, pno, st)
            elif vint == 2024 and ": 2024" in head and "by state/jurisdiction" in head:
                for name, vals, st in pdf_rows(lines, hdr, 12):
                    jj = juris(name)
                    if jj:
                        for cat, idx in (("sdel", 0), ("sd", 4), ("el", 8)):
                            cells[(2024, subject, grade, cat, 2024, jj[0])] = (vals[idx], url, pno, st)
    return cells


def parse_national_pdf(subject):
    """National (public and nonpublic) table of the 2024 national appendix: rows by category."""
    url = NATIONAL_PDF[subject]
    out = {}
    for pno, page in enumerate(pdf_pages(fetch(url)), start=1):
        lines = page.splitlines()
        head = cell_str(" ".join(lines[:10]))
        if "Table of Contents" in head:
            continue
        m = re.search(r"Percentage of (fourth|eighth)-grade students with disabilities \(SD\) and/or "
                      r"English learners \(EL\) identified, excluded, and assessed", head)
        if not m:
            continue
        grade = 4 if m.group(1) == "fourth" else 8
        cat = None
        for l in lines:
            st = l.strip()
            if st in ("SD and/or EL", "SD", "EL"):
                cat = {"SD and/or EL": "sdel", "SD": "sd", "EL": "el"}[st]
                continue
            mm = re.match(r"^(Identified|Excluded)\s+(.*)$", st)
            if mm and cat:
                out[(grade, cat, mm.group(1).lower())] = (mm.group(2).split(), pno, st)
    return out


# ---------------------------------------------------------------------------------------
# Quotes
# ---------------------------------------------------------------------------------------
def essa_quote() -> str:
    soup = BeautifulSoup(fetch(ESSA_URL).read_text(encoding="utf-8"), "html.parser")
    anchor = soup.find("a", attrs={"name": "b_3"})
    if anchor is None:
        raise SystemExit("[BLOCKED] 20 USC 6311(b)(3) anchor not found in the Cornell page")
    para = anchor.parent
    lines = []

    def norm(s: str) -> str:
        return re.sub(r"\s+", " ", s).strip()

    def own_parts(div):
        """(enumerator, heading, body) of one statutory unit, without its nested units."""
        num = head = ""
        body = []
        for ch in div.children:
            name = getattr(ch, "name", None)
            cls = set(ch.get("class", [])) if name else set()
            if name == "div" and not cls & {"content", "continuation"}:
                continue  # nested statutory unit, rendered on its own line
            if name == "a" and ch.get("name"):
                continue
            text = ch.get_text("") if name else str(ch)
            if "num" in cls:
                num = norm(text)
            elif "heading" in cls:
                head = norm(text)
            else:
                body.append(text)
        return num, head, norm(" ".join(body))

    for div in [para] + para.find_all("div"):
        a = div.find("a", attrs={"name": True}, recursive=False)
        if a is None or not a["name"].startswith("b_3"):
            continue
        if a["name"].startswith("b_3_B"):
            break
        pad = "    " * (a["name"].count("_") - 1)
        num, head, body = own_parts(div)
        if head:  # statutory heading on its own line, text below it
            lines.append(f"{pad}{num} {head}")
            if body:
                lines.append(f"{pad}{body}")
        else:
            lines.append(f"{pad}{num} {body}".rstrip())
    return "\n".join(lines)


def nces_history_quote() -> str:
    """NCES account of the pre-2010 EL exclusion guideline (in force from 1996)."""
    soup = BeautifulSoup(fetch(NCES_HISTORY_URL).read_text(encoding="utf-8"), "html.parser")
    lines = [re.sub(r"\s+", " ", l).strip() for l in soup.get_text("\n").splitlines()]
    lines = [l for l in lines if l]
    try:
        i = next(n for n, l in enumerate(lines) if l.startswith("A student who was identified as LEP or EL"))
        j = next(n for n, l in enumerate(lines) if l.startswith("The goal of all these activities"))
        k = next(n for n, l in enumerate(lines) if l.startswith("Beginning with the 2002 assessments"))
    except StopIteration:
        raise SystemExit("[BLOCKED] EL guideline passage not found on the NCES inclusion-history page")
    return "\n".join(lines[i:j] + ["", "[section \"NAEP in 2002\"]", lines[k]])


def nagb_quotes() -> list[tuple[str, str]]:
    path = fetch(NAGB_URL)
    res = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, check=True)
    t = re.sub(r"\s+", " ", res.stdout.decode("utf-8"))
    pats = [
        ("Header", r"Adopted: March 6, 2010.*?Updated: August 2, 2014"),
        ("Policy Principles, item 3", r"3\. The proportion of all students excluded from any NAEP sample.*?95 percent\."),
        ("Policy Principles, item 4", r"4\. Among students classified as either ELL or SD.*?NAEP reporting\."),
        ("Implementation Guidelines, For English Language Learners, item 1",
         r"1\. All English language learners selected for the NAEP sample.*?primary language\."),
        ("Implementation Guidelines, For English Language Learners, item 1 (continued after the page footnote)",
         r"One year or more shall be defined as one full academic year before the year of the assessment\."),
        ("Implementation Guidelines, For English Language Learners, item 3",
         r"3\. Bilingual versions of NAEP in Spanish and English.*?for these subjects\."),
    ]
    out = []
    for label, p in pats:
        m = re.search(p, t)
        if not m:
            raise SystemExit(f"[BLOCKED] NAGB passage not found: {label}")
        out.append((label, m.group(0)))
    return out


# ---------------------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------------------
def dec(s: str):
    return Decimal(s) if s not in ("", None) else None


def decimals(s: str) -> int:
    return len(s.split(".")[1]) if "." in s else 0


def round_half_up(s: str) -> int:
    return int(Decimal(s).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def fmt_ratio(num: str, den: str) -> str:
    q = (Decimal(num) / Decimal(den) * 100).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    return format(q, "f")


def write_csv(path: Path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([r.get(h, "") for h in header])


def sort_key(r):
    return (r["subject"], int(r["grade"]), int(r["year"]), TYPE_RANK[r["jurisdiction_type"]],
            r["jurisdiction"])


# ---------------------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------------------
def main() -> int:
    unmatched = defaultdict(set)
    pct_all, pct_ident, national_sheets = [], [], {}
    for (vintage, subject) in sorted(TA):
        a, b, nat = load_appendix(vintage, subject, unmatched)
        pct_all += a
        pct_ident += b
        if vintage == 2024:
            national_sheets[subject] = nat
    tdw = []
    for key in sorted(TDW_PAGES):
        tdw += parse_tdw(key[0], key[1], unmatched)

    # ---- index report-card cells ------------------------------------------------------
    A = {}  # (vintage, subject, grade, cat, regime, year, juris, measure) -> cell
    for c in pct_all:
        k = (c["vintage"], c["subject"], c["grade"], c["cat"], c["regime"], c["year"],
             c["jurisdiction"], c["measure"])
        if k in A and A[k]["raw"] != c["raw"]:
            raise SystemExit(f"[BLOCKED] conflicting duplicate cell {k}: {A[k]['raw']} vs {c['raw']}")
        A[k] = c
    I = {}  # (vintage, src, subject, grade, cat, regime, year, juris, measure) -> cell
    for c in pct_ident:
        k = (c["vintage"], c["src"], c["subject"], c["grade"], c["cat"], c["regime"], c["year"],
             c["jurisdiction"], c["measure"])
        if k in I and I[k]["raw"] != c["raw"]:
            raise SystemExit(f"[BLOCKED] conflicting duplicate %identified cell {k}")
        I[k] = c
    T = {}
    for c in tdw:
        T[(c["subject"], c["grade"], c["year"], c["jurisdiction"], c["kind"])] = c

    # ---- main panel: accommodations permitted, 2024 appendix as primary ----------------
    keys = sorted({(k[1], k[2], k[5], k[6]) for k in A if k[0] == 2024 and k[4] == "permitted"})
    juris_meta = {c["jurisdiction"]: (c["abbr"], c["jtype"]) for c in pct_all}
    panel, xcheck = [], []
    xsummary = Counter()
    xmax = defaultdict(Decimal)
    for (subject, grade, year, jn) in keys:
        abbr, jtype = juris_meta[jn]
        row = dict(subject=subject, grade=grade, year=year, jurisdiction=jn,
                   jurisdiction_abbr=abbr, jurisdiction_type=jtype, accommodations="permitted")
        sheets = []
        for cat in CATS:
            for m in ("identified", "excluded", "assessed"):
                c = A.get((2024, subject, grade, cat, "permitted", year, jn, m))
                f = f"{cat}_{m}_pct_all"
                row[f] = c["value"] if c else ""
                row[f + "_flag"] = c["flag"] if c else "not_in_source"
            c = A.get((2024, subject, grade, cat, "permitted", year, jn, "identified"))
            if c:
                sheets.append(f"{CAT_LABEL[cat]}: {c['sheet']}")
        row["pct_all_source_url"] = TA[(2024, subject)]
        row["pct_all_source_table"] = "; ".join(sheets)
        # published excluded as % of identified
        srcs = set()
        for cat in CATS:
            f = f"{cat}_excluded_pct_identified"
            c = None
            order = ([(2017, "trend"), (2017, "year")] if year <= 2017 else [(year, "year")])
            for vint, src in order:
                c = I.get((vint, src, subject, grade, cat, "permitted", year, jn, "excluded_pct_identified"))
                if c:
                    break
            row[f] = c["value"] if c else ""
            row[f + "_flag"] = c["flag"] if c else "not_in_source"
            if c:
                srcs.add((c["url"], f"{CAT_LABEL[cat]}: {c['sheet']}"))
            # derived ratio from the decimal %-of-all cells
            num = A.get((2024, subject, grade, cat, "permitted", year, jn, "excluded"))
            den = A.get((2024, subject, grade, cat, "permitted", year, jn, "identified"))
            ok = num and den and not num["flag"] and not den["flag"] and Decimal(den["value"]) > 0
            row[f + "_calc"] = fmt_ratio(num["value"], den["value"]) if ok else ""
            if c and not c["flag"] and ok:
                d = abs(Decimal(c["value"]) - Decimal(row[f + "_calc"]))
                tol = Decimal("0.5") if decimals(c["value"]) == 0 else Decimal("0.05")
                agree = d <= tol
                xsummary[("pctident_published_vs_calc", cat, agree)] += 1
                xmax[("pctident_published_vs_calc", cat)] = max(xmax[("pctident_published_vs_calc", cat)], d)
                if not agree:
                    xcheck.append(dict(check="pctident_published_vs_calc", subject=subject, grade=grade,
                                       year=year, jurisdiction=jn, field=f, primary=c["value"],
                                       other=row[f + "_calc"], abs_diff=str(d),
                                       note=f"published {c['sheet']} vs excluded/identified from {TA[(2024, subject)].rsplit('/', 1)[1]}"))
        row["pct_identified_source_url"] = "; ".join(sorted({s[0] for s in srcs}))
        row["pct_identified_source_table"] = "; ".join(sorted({s[1] for s in srcs}))
        # earlier vintages of the same %-of-all cells
        for vint in (2022, 2019):
            for cat in ("el", "sd"):
                for m in ("identified", "excluded"):
                    c = A.get((vint, subject, grade, cat, "permitted", year, jn, m))
                    row[f"{cat}_{m}_pct_all_ta{vint}"] = c["raw"] if c else ""
        # TDW
        for kind, f in (("el_excluded", "tdw_el_excluded_pct_all"), ("sd_excluded", "tdw_sd_excluded_pct_all"),
                        ("response_rate", "tdw_student_response_rate")):
            c = T.get((subject, grade, year, jn, kind))
            row[f] = c["value"] if c else ""
            row[f + "_flag"] = (c["flag"] if c else ("not_in_source" if (year, subject) in TDW_PAGES else "no_tdw_page"))
        c = T.get((subject, grade, year, jn, "el_excluded"))
        row["tdw_row_label"] = c["row_label"] if c else ""
        row["tdw_source_url"] = TDW_PAGES.get((year, subject), "")
        panel.append(row)

    # ---- published %-of-identified against excluded/identified from the older vintages -----
    # (locates vintage revisions: e.g. 2009 cells that the 2019 appendix printed differently)
    for row in panel:
        for cat in CATS:
            pub, pflag = row[f"{cat}_excluded_pct_identified"], row[f"{cat}_excluded_pct_identified_flag"]
            if not pub or pflag:
                continue
            for vint in (2022, 2019):
                num = A.get((vint, row["subject"], row["grade"], cat, "permitted", row["year"], row["jurisdiction"], "excluded"))
                den = A.get((vint, row["subject"], row["grade"], cat, "permitted", row["year"], row["jurisdiction"], "identified"))
                if not (num and den) or num["flag"] or den["flag"] or Decimal(den["value"]) <= 0:
                    continue
                if min(decimals(num["value"]), decimals(den["value"])) < 2:
                    continue  # integer-rounded vintage cells cannot resolve the ratio
                calc_v = fmt_ratio(num["value"], den["value"])
                d = abs(Decimal(pub) - Decimal(calc_v))
                tol = Decimal("0.5") if decimals(pub) == 0 else Decimal("0.05")
                agree = d <= tol
                key = f"pctident_published_vs_calc_ta{vint}"
                xsummary[(key, cat, agree)] += 1
                xmax[(key, cat)] = max(xmax[(key, cat)], d)
                if not agree:
                    xcheck.append(dict(check=key, subject=row["subject"], grade=row["grade"], year=row["year"],
                                       jurisdiction=row["jurisdiction"], field=f"{cat}_excluded_pct_identified",
                                       primary=pub, other=calc_v, abs_diff=str(d),
                                       note=f"published vs excluded/identified from the {vint} appendix ({num['sheet']})"))

    # ---- cross-vintage agreement (all measures, all categories) -------------------------
    for k, c in A.items():
        vint = k[0]
        if vint == 2024:
            continue
        p = A.get((2024,) + k[1:])
        if p is None:
            xsummary[(f"vintage_{vint}_vs_2024", k[3], "missing_in_2024")] += 1
            continue
        if p["flag"] or c["flag"]:
            agree = p["flag"] == c["flag"] and p["value"] == c["value"]
            d = Decimal(0)
        else:
            d = abs(Decimal(p["value"]) - Decimal(c["value"]))
            # one unit in the last place of the coarser cell: absorbs re-rounding (3.6255 vs 3.625449)
            tol = Decimal(1) / (Decimal(10) ** min(decimals(p["value"]), decimals(c["value"])))
            agree = d <= tol
        xsummary[(f"vintage_{vint}_vs_2024", k[3], agree)] += 1
        xmax[(f"vintage_{vint}_vs_2024", k[3])] = max(xmax[(f"vintage_{vint}_vs_2024", k[3])], d)
        if not agree:
            xcheck.append(dict(check=f"vintage_{vint}_vs_2024", subject=k[1], grade=k[2], year=k[5],
                               jurisdiction=k[6], field=f"{k[3]}_{k[7]}_pct_all ({k[4]})",
                               primary=p["raw"], other=c["raw"], abs_diff=str(d),
                               note=f"2024 {p['sheet']} vs {vint} {c['sheet']}"))

    # ---- 2017 %identified: trend table vs 2017 per-year table ---------------------------
    for k, c in I.items():
        if k[0] != 2017 or k[1] != "trend" or k[6] != 2017 or k[8] != "excluded_pct_identified":
            continue
        other = I.get((2017, "year") + k[2:])
        if other is None:
            continue
        if c["flag"] or other["flag"]:
            agree = c["flag"] == other["flag"]
            d = Decimal(0)
        else:
            d = abs(Decimal(c["value"]) - Decimal(other["value"]))
            agree = d <= Decimal("0.0001")
        xsummary[("pctident_2017_trend_vs_year_table", k[4], agree)] += 1
        xmax[("pctident_2017_trend_vs_year_table", k[4])] = max(xmax[("pctident_2017_trend_vs_year_table", k[4])], d)
        if not agree:
            xcheck.append(dict(check="pctident_2017_trend_vs_year_table", subject=k[2], grade=k[3], year=2017,
                               jurisdiction=k[7], field=f"{k[4]}_excluded_pct_identified",
                               primary=c["raw"], other=other["raw"], abs_diff=str(d),
                               note=f"{c['sheet']} vs {other['sheet']}"))

    # ---- TDW vs report card (EL and SD excluded, % of all) ----------------------------
    tdw_cmp = []
    for row in panel:
        for cat in ("el", "sd"):
            row[f"tdw_{cat}_agrees_with_reportcard"] = ""
            t, t_flag = row[f"tdw_{cat}_excluded_pct_all"], row[f"tdw_{cat}_excluded_pct_all_flag"]
            r, r_flag = row[f"{cat}_excluded_pct_all"], row[f"{cat}_excluded_pct_all_flag"]
            if t == "" or t_flag not in ("", "rounds_to_zero"):
                continue
            # a TDW "#" is below half a unit of the page's two-decimal display
            t_half = Decimal("0.005") if t_flag else Decimal(5) / (Decimal(10) ** (decimals(t) + 1))
            if r_flag == "rounds_to_zero":  # report-card "#": below 0.5
                agree = Decimal(t) < Decimal("0.5") + t_half
                d = max(Decimal(0), Decimal(t) - Decimal("0.5"))  # distance beyond the censoring bound
            elif r and not r_flag:
                d = abs(Decimal(t) - Decimal(r))
                agree = d <= Decimal("0.05") + t_half
            else:
                continue
            row[f"tdw_{cat}_agrees_with_reportcard"] = "1" if agree else "0"
            xsummary[("tdw_vs_reportcard", cat, agree)] += 1
            xmax[("tdw_vs_reportcard", cat)] = max(xmax[("tdw_vs_reportcard", cat)], d)
            tdw_cmp.append((row["subject"], row["grade"], row["year"], row["jurisdiction"], cat, t, r, r_flag, agree, d))
            if not agree:
                xcheck.append(dict(check="tdw_vs_reportcard", subject=row["subject"], grade=row["grade"],
                                   year=row["year"], jurisdiction=row["jurisdiction"],
                                   field=f"{cat}_excluded_pct_all",
                                   primary="#" if r_flag == "rounds_to_zero" else r, other=t,
                                   abs_diff=str(d), note=f"report card (2024 appendix) vs TDW row '{row['tdw_row_label']}'"))

    # ---- PDF validation of the 2024 xlsx ----------------------------------------------
    pdf_cells, pdf_anom = {}, []
    for subject in ("mathematics", "reading"):
        cells, anom = parse_state_pdf(subject)
        pdf_cells.update(cells)
        pdf_anom += [f"{subject} {a}" for a in anom]
    pdf_counts = Counter()
    pdf_mismatch = []
    for (subject, grade, cat, year, jn, m), (tok, pno, line) in sorted(pdf_cells.items(), key=lambda kv: str(kv[0])):
        c = A.get((2024, subject, grade, cat, "permitted", year, jn, m))
        if c is None:
            pdf_counts["pdf_cell_without_xlsx"] += 1
            pdf_mismatch.append((subject, grade, cat, year, jn, m, tok, "", pno, line))
            continue
        if tok == "#":
            ok = c["flag"] == "rounds_to_zero" or (not c["flag"] and Decimal(c["value"]) < Decimal("0.5"))
        elif tok in ("—", "‡", "†"):
            ok = c["raw"] == tok
        else:
            ok = (not c["flag"] and c["value"] != "" and
                  abs(Decimal(c["value"]) - Decimal(tok)) <= Decimal("0.5000001"))
        pdf_counts["match" if ok else "mismatch"] += 1
        if not ok:
            pdf_mismatch.append((subject, grade, cat, year, jn, m, tok, c["raw"], pno, line))
    # %-of-identified cells (2017 trend XLSX, 2024 per-year XLSX) against their PDFs
    pid_cells = parse_pct_identified_pdfs()
    pid_counts, pid_mismatch, pid_pairs = Counter(), [], []
    for key in sorted(pid_cells):
        vint, subject, grade, cat, year, jn = key
        tok, url, pno, line = pid_cells[key]
        src = "trend" if vint == 2017 else "year"
        c = (I.get((vint, src, subject, grade, cat, "permitted", year, jn, "excluded_pct_identified")) or
             I.get((vint, src, subject, grade, cat, "not_permitted", year, jn, "excluded_pct_identified")))
        if c is None:
            pid_counts[f"{vint}_pdf_cell_without_xlsx"] += 1
            continue
        if tok == "#":
            ok = c["flag"] == "rounds_to_zero" or (not c["flag"] and Decimal(c["value"]) < Decimal("0.5"))
        elif tok in ("—", "‡", "†"):
            ok = c["raw"] == tok
        else:
            ok = (not c["flag"] and c["value"] != "" and
                  abs(Decimal(c["value"]) - Decimal(tok)) <= Decimal("0.5000001"))
        pid_counts[f"{vint}_{'match' if ok else 'mismatch'}"] += 1
        if ok and not c["flag"]:
            pid_pairs.append((key, c))
        if not ok:
            pid_mismatch.append((key, tok, c["raw"], pno, line))
    xlsx_permitted_2024 = [k for k in A if k[0] == 2024 and k[4] == "permitted"]
    pdf_counts["xlsx_cells_permitted_2024"] = len(xlsx_permitted_2024)
    pdf_counts["xlsx_cells_without_pdf"] = sum(1 for k in xlsx_permitted_2024
                                               if (k[1], k[2], k[3], k[5], k[6], k[7]) not in pdf_cells)

    # ---- national check: A-15 (xlsx, all schools) vs national appendix PDF -------------
    national_lines = []
    for subject in ("mathematics", "reading"):
        nat_pdf = parse_national_pdf(subject)
        for sheet, rows in national_sheets[subject]:
            yi, yrs = find_year_row(rows)
            years = [y for _, y in yrs]
            grade = cat = None
            for r in rows[yi + 1:]:
                lab = cell_str(r[0])
                if lab in ("Grade 4", "Grade 8"):
                    grade = int(lab[-1])
                    continue
                if lab in ("SD and/or EL", "SD", "EL"):
                    cat = {"SD and/or EL": "sdel", "SD": "sd", "EL": "el"}[lab]
                    continue
                if lab in ("Identified", "Excluded") and grade and cat:
                    xv = [cell_str(v) for v in r[1:1 + len(years)]]
                    pv, pno, pline = nat_pdf.get((grade, cat, lab.lower()), ([], None, ""))
                    tail = pv[-len(years):] if pv else []
                    res = []
                    for y, x, p in zip(years, xv, tail):
                        if y < 2002:
                            # the PDF prints extra not-permitted/† columns before 2002 whose
                            # alignment with the XLSX columns is ambiguous; compare 2002+ only
                            continue
                        if not NUM_RE.match(x):
                            ok = x == p or (x in ("—", "†", "") and p in ("—", "†"))
                        else:
                            ok = NUM_RE.match(p) is not None and round_half_up(x) == int(p)
                        res.append((y, x, p, ok))
                    national_lines.append((subject, grade, cat, lab.lower(), sheet, pno, pline, res))

    # ---- not-permitted file -------------------------------------------------------------
    np_keys = sorted({(k[1], k[2], k[5], k[6]) for k in A if k[0] == 2024 and k[4] == "not_permitted"})
    np_rows = []
    for (subject, grade, year, jn) in np_keys:
        abbr, jtype = juris_meta[jn]
        row = dict(subject=subject, grade=grade, year=year, jurisdiction=jn, jurisdiction_abbr=abbr,
                   jurisdiction_type=jtype, accommodations="not_permitted")
        sheets = []
        for cat in CATS:
            for m in ("identified", "excluded", "assessed"):
                c = A.get((2024, subject, grade, cat, "not_permitted", year, jn, m))
                row[f"{cat}_{m}_pct_all"] = c["value"] if c else ""
                row[f"{cat}_{m}_pct_all_flag"] = c["flag"] if c else "not_in_source"
            c = A.get((2024, subject, grade, cat, "not_permitted", year, jn, "identified"))
            if c:
                sheets.append(f"{CAT_LABEL[cat]}: {c['sheet']}")
            c = I.get((2017, "trend", subject, grade, cat, "not_permitted", year, jn, "excluded_pct_identified"))
            row[f"{cat}_excluded_pct_identified"] = c["value"] if c else ""
            row[f"{cat}_excluded_pct_identified_flag"] = c["flag"] if c else "not_in_source"
        row["pct_all_source_url"] = TA[(2024, subject)]
        row["pct_all_source_table"] = "; ".join(sheets)
        row["pct_identified_source_url"] = TA[(2017, subject)]
        np_rows.append(row)

    # ---- write outputs ------------------------------------------------------------------
    base = ["subject", "grade", "year", "jurisdiction", "jurisdiction_abbr", "jurisdiction_type", "accommodations"]
    vals = []
    for cat in CATS:
        for m in ("identified", "excluded", "assessed"):
            vals += [f"{cat}_{m}_pct_all", f"{cat}_{m}_pct_all_flag"]
    ident = []
    for cat in CATS:
        ident += [f"{cat}_excluded_pct_identified", f"{cat}_excluded_pct_identified_flag"]
    calc = [f"{cat}_excluded_pct_identified_calc" for cat in CATS]
    vint_cols = [f"{cat}_{m}_pct_all_ta{v}" for v in (2022, 2019) for cat in ("el", "sd")
                 for m in ("identified", "excluded")]
    tdw_cols = ["tdw_el_excluded_pct_all", "tdw_el_excluded_pct_all_flag", "tdw_sd_excluded_pct_all",
                "tdw_sd_excluded_pct_all_flag", "tdw_student_response_rate", "tdw_student_response_rate_flag",
                "tdw_el_agrees_with_reportcard", "tdw_sd_agrees_with_reportcard", "tdw_row_label",
                "tdw_source_url"]
    header = (base + vals + ["pct_all_source_url", "pct_all_source_table"] + ident +
              ["pct_identified_source_url", "pct_identified_source_table"] + calc + vint_cols + tdw_cols)
    panel.sort(key=sort_key)
    write_csv(DERIVED / "naep_inclusion.csv", header, panel)
    np_header = (base + vals + ["pct_all_source_url", "pct_all_source_table"] + ident +
                 ["pct_identified_source_url"])
    np_rows.sort(key=sort_key)
    write_csv(DERIVED / "naep_inclusion_accom_not_permitted.csv", np_header, np_rows)
    xcheck.sort(key=lambda r: (r["check"], r["subject"], int(r["grade"]), int(r["year"]), r["jurisdiction"], r["field"]))
    write_csv(DERIVED / "naep_inclusion_crosscheck_disagreements.csv",
              ["check", "subject", "grade", "year", "jurisdiction", "field", "primary", "other", "abs_diff", "note"],
              xcheck)
    summ = []
    for (check, cat, agree), n in sorted(xsummary.items(), key=lambda kv: (kv[0][0], kv[0][1], str(kv[0][2]))):
        summ.append(dict(check=check, category=cat, outcome=str(agree), cells=n,
                         max_abs_diff=format(xmax.get((check, cat), Decimal(0)).normalize(), "f")))
    write_csv(DERIVED / "naep_inclusion_crosscheck_summary.csv",
              ["check", "category", "outcome", "cells", "max_abs_diff"], summ)

    essa_text, nagb, history_text = essa_quote(), nagb_quotes(), nces_history_quote()
    # sources manifest
    src_rows = []
    for (v, s), u in sorted(TA.items()):
        src_rows.append(dict(role=f"technical_appendix_xlsx_{v}", subject=s, url=u))
    for (v, s), u in sorted(TA_PDF.items()):
        src_rows.append(dict(role=f"technical_appendix_pdf_{v}", subject=s, url=u))
    for s, u in sorted(NATIONAL_PDF.items()):
        src_rows.append(dict(role="national_appendix_pdf_2024", subject=s, url=u))
    for (y, s), u in sorted(TDW_PAGES.items()):
        src_rows.append(dict(role=f"tdw_exclusion_{y}", subject=s, url=u))
    src_rows.append(dict(role="essa_20usc6311", subject="", url=ESSA_URL))
    src_rows.append(dict(role="nagb_2010_policy", subject="", url=NAGB_URL))
    src_rows.append(dict(role="nces_inclusion_history", subject="", url=NCES_HISTORY_URL))
    for r in src_rows:
        p = local_path(r["url"])
        r["cache_file"] = p.name
        r["bytes"] = p.stat().st_size
        r["sha256"] = sha256(p)
    write_csv(DERIVED / "sources.csv", ["role", "subject", "url", "cache_file", "bytes", "sha256"], src_rows)

    # quotes
    q = ["# Primary-text quotes: EL exclusion rules", "",
         "Generated by `acquire_inclusion.py` from cached primary sources; text is verbatim "
         "(whitespace normalised, line breaks at statutory enumerators).", "",
         "## (a) ESSA exception for recently arrived English learners, 20 U.S.C. 6311(b)(3)(A)", "",
         f"Source: {ESSA_URL} (Legal Information Institute rendering of the U.S. Code; section 1111(b)(3)(A) "
         "of the Elementary and Secondary Education Act as amended by the Every Student Succeeds Act). "
         "uscode.house.gov refused the connection during acquisition (port 443 timeout).", "",
         "```text", essa_text, "```", "",
         "## (b) NAGB policy: NAEP Testing and Reporting on Students with Disabilities and English Language Learners", "",
         f"Source: {NAGB_URL}", ""]
    for label, text in nagb:
        q += [f"{label}:", "", f"> {text}", ""]
    q += ["## (c) Context: the EL exclusion guideline before the 2010 policy (NCES)", "",
          f"Source: {NCES_HISTORY_URL}, sections \"NAEP in 1996\" and \"NAEP in 2002\". The page does "
          "not say when the three-year rule stopped applying; the NAGB policy in (b) was adopted on "
          "March 6, 2010.", "",
          "```text", history_text, "```", ""]
    (DERIVED / "quotes.md").write_text("\n".join(q), encoding="utf-8")

    # ---- validation report ------------------------------------------------------------
    rep = ["# Validation report (generated by acquire_inclusion.py)", ""]
    rep += ["## Panel coverage: jurisdictions with a numeric value, by subject, grade and year", "",
            "Counts cover the 50 states and DC (nation, DoDEA and Puerto Rico excluded). "
            "`el_id` = EL identified (% of all), `el_ex` = EL excluded (% of all, numeric or '#'), "
            "`el_exid` = EL excluded as % of identified (published, numeric or '#'), `tdw_el` = TDW EL "
            "excluded. Each count includes '#' cells (stored as 0).", "",
            "| subject | grade | year | el_id | el_ex | el_ex # | el_exid | el_exid ‡ | sd_id | sd_exid | tdw_el |",
            "|---|---|---|---|---|---|---|---|---|---|---|"]
    cov = defaultdict(Counter)
    for r in panel:
        if r["jurisdiction_type"] not in ("state", "dc"):
            continue
        k = (r["subject"], r["grade"], r["year"])
        cov[k]["el_id"] += r["el_identified_pct_all_flag"] in ("", "rounds_to_zero")
        cov[k]["el_ex"] += r["el_excluded_pct_all_flag"] in ("", "rounds_to_zero")
        cov[k]["el_ex_rtz"] += r["el_excluded_pct_all_flag"] == "rounds_to_zero"
        cov[k]["el_exid"] += r["el_excluded_pct_identified_flag"] in ("", "rounds_to_zero")
        cov[k]["el_exid_rsnm"] += r["el_excluded_pct_identified_flag"] == "reporting_standards_not_met"
        cov[k]["sd_id"] += r["sd_identified_pct_all_flag"] in ("", "rounds_to_zero")
        cov[k]["sd_exid"] += r["sd_excluded_pct_identified_flag"] in ("", "rounds_to_zero")
        cov[k]["tdw_el"] += r["tdw_el_excluded_pct_all_flag"] in ("", "rounds_to_zero") and r["tdw_el_excluded_pct_all"] != ""
    for k in sorted(cov):
        c = cov[k]
        rep.append(f"| {k[0]} | {k[1]} | {k[2]} | {c['el_id']} | {c['el_ex']} | {c['el_ex_rtz']} | {c['el_exid']} | "
                   f"{c['el_exid_rsnm']} | {c['sd_id']} | {c['sd_exid']} | {c['tdw_el']} |")
    rep += ["", "## Cross-check summary", "",
            "Tolerances: vintages agree when the difference is within one unit in the last place of the "
            "coarser cell (flags must match exactly); published vs calculated %-of-identified within 0.05 "
            "points (0.5 when the published cell is an integer); TDW vs report card within 0.05 points plus "
            "half a unit of the TDW display; a report-card '#' agrees with a TDW value below 0.5 (for a '#' "
            "cell, the disagreement file's abs_diff is the TDW value's distance above 0.5).", "",
            "| check | category | agree | cells | max abs diff |", "|---|---|---|---|---|"]
    for s in summ:
        rep.append(f"| {s['check']} | {s['category']} | {s['outcome']} | {s['cells']} | {s['max_abs_diff']} |")
    rep += ["", "## 2024 XLSX against the 2024 PDF (integers)", "",
            f"PDF cells parsed: {len(pdf_cells)}; " + "; ".join(f"{k}: {v}" for k, v in sorted(pdf_counts.items())), ""]
    if pdf_anom:
        rep += ["PDF parse anomalies:", ""] + [f"- {a}" for a in pdf_anom] + [""]
    if pdf_mismatch:
        rep += ["Mismatches:", ""] + [f"- {m}" for m in pdf_mismatch[:50]] + [""]
    no_pdf = Counter((k[1], k[2], k[3], k[5]) for k in xlsx_permitted_2024
                     if (k[1], k[2], k[3], k[5], k[6], k[7]) not in pdf_cells)
    if no_pdf:
        rep += ["XLSX cells with no PDF counterpart (subject, grade, category, year: cells):", ""]
        rep += [f"- {k}: {v}" for k, v in sorted(no_pdf.items())] + [""]
    rep += ["## Excluded as a percentage of identified: XLSX against PDF (integers)", "",
            "2017 = state trend tables (SD, ELL; 1990/1992–2017) of the 2017 appendix; 2024 = per-year "
            "tables A-29/A-30 of the 2024 appendix.", "",
            "; ".join(f"{k}: {v}" for k, v in sorted(pid_counts.items())), ""]
    if pid_mismatch:
        rep += ["Mismatches:", ""] + [f"- {m}" for m in pid_mismatch[:50]] + [""]
    rep += ["## National (public and nonpublic) rows: 2024 appendix XLSX Table A-15 vs 2024 national appendix PDF", ""]
    for subject, grade, cat, meas, sheet, pno, pline, res in national_lines:
        nbad = sum(1 for x in res if not x[3])
        rep.append(f"- {subject} grade {grade} {CAT_LABEL[cat]} {meas}: {len(res)} years compared, {nbad} mismatches "
                   f"(PDF p{pno}: `{pline}`)")
        if cat == "el":
            rep.append("  - " + ", ".join(f"{y}: {x}→{p}" for y, x, p, ok in res))
    rep += ["", "## Spot checks: random panel cells against the 2024 PDF text", "",
            "Seeded (random.Random(20260927)); each line quotes the PDF text row that carries the cell.", ""]
    rng = random.Random(20260927)
    cand = sorted(k for k in pdf_cells if k[2] in ("el", "sd") and k[5] in ("identified", "excluded")
                  and A.get((2024, k[0], k[1], k[2], "permitted", k[3], k[4], k[5])) is not None
                  and not A[(2024, k[0], k[1], k[2], "permitted", k[3], k[4], k[5])]["flag"])
    picks = rng.sample(cand, 10)
    for k in picks:
        tok, pno, line = pdf_cells[k]
        c = A[(2024, k[0], k[1], k[2], "permitted", k[3], k[4], k[5])]
        years_on_page = sorted({kk[3] for kk in pdf_cells if kk[:3] == k[:3] and pdf_cells[kk][1] == pno})
        rep.append(f"- {k[0]} grade {k[1]} {k[4]} {k[3]} {CAT_LABEL[k[2]]} {k[5]} (% of all): panel value "
                   f"{c['value']} [{c['sheet']}] → PDF p{pno} (years {years_on_page}) shows {tok}; line: `{line}`")
    rep += ["", "## Spot checks: random %-of-identified cells against the PDF text", ""]
    for key, c in rng.sample(pid_pairs, 6):
        tok, url, pno, line = pid_cells[key]
        rep.append(f"- {key[1]} grade {key[2]} {key[5]} {key[4]} {CAT_LABEL[key[3]]} excluded (% of identified): "
                   f"panel source value {c['value']} [{c['sheet']}] → {url.rsplit('/', 1)[1]} p{pno} shows {tok}; "
                   f"line: `{line}`")
    rep += ["", "## Spot checks: random TDW cells against the TDW HTML row text", ""]
    tcand = sorted((c["subject"], c["grade"], c["year"], c["jurisdiction"], c["kind"]) for c in tdw
                   if c["jtype"] in ("state", "dc") and c["kind"] != "response_rate" and not c["flag"])
    for k in rng.sample(tcand, 4):
        c = T[k]
        rep.append(f"- {k[0]} grade {k[1]} {k[2]} {k[3]} {k[4]}: panel value {c['value']} ← {c['url']} row: `{c['row_text']}`")
    rep += ["", "## Row labels not mapped to a jurisdiction (skipped)", "",
            "Urban districts (TUDA), BIE, the separate DDESS/DoDDS systems and outlying areas are out of "
            "scope; 'Nation' in the 2017 tables is public plus nonpublic and is skipped in favour of "
            "'Nation (public)'.", ""]
    noise = ("—", "#", "‡", "NOTE", "SOURCE", "See notes", "1 ", "2 ", "†", "¹", "²", "Other jurisdictions")
    by_source = defaultdict(set)
    for key, names in unmatched.items():
        src = f"TDW {key[1]} {key[2]}" if key[0] == "TDW" else f"{key[0]} appendix"
        by_source[src] |= {n for n in names if n and not n.startswith(noise)}
    for src in sorted(by_source):
        if by_source[src]:
            rep.append(f"- {src}: {', '.join(sorted(by_source[src]))}")
    (DERIVED / "validation_report.md").write_text("\n".join(rep) + "\n", encoding="utf-8")

    print(f"panel rows: {len(panel)}; not-permitted rows: {len(np_rows)}; TDW cells: {len(tdw)}; "
          f"disagreements: {len(xcheck)}; pdf: {dict(pdf_counts)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
