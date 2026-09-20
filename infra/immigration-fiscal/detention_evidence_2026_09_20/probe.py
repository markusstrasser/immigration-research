"""Reproduce bounded source probes; fail on corrupt files or inconsistent totals."""
from pathlib import Path
import csv, json, hashlib, calendar, math
import openpyxl
from source_contract import load_manifest, require, verify
ROOT=Path(__file__).resolve().parent
CACHE=ROOT/'_cache'
require(CACHE.is_dir(), 'Required _cache directory is missing')
DEST=ROOT/'derived'
DEST.mkdir(exist_ok=True)
manifest=load_manifest(ROOT)
for f in manifest:
    verify(ROOT/f['local_file'], f)
wb=openpyxl.load_workbook(CACHE/'ice_fy2024_yearend.xlsx',read_only=True,data_only=True)
ws=wb['Detention FY24']
require(ws['A99'].value.strip()=='Average', 'ICE annual row label changed')
adp=ws['N99'].value
days=[calendar.monthrange(y,m)[1] for y,m in [(2023,10),(2023,11),(2023,12)]+[(2024,m) for m in range(1,10)]]
mandays=sum(ws.cell(99,c).value*d for c,d in zip(range(2,14),days))
require(sum(days)==366, 'FY2024 must have 366 days')
require(math.isclose(mandays/366,adp,rel_tol=1e-12), 'ICE monthly/annual ADP mismatch')
require(math.isclose(sum(ws[f'N{r}'].value for r in [100,101,102]),adp,rel_tol=1e-12), 'ICE history categories do not sum to ADP')
require(ws['F20'].value==sum(ws[f'F{r}'].value for r in [21,22,23]), 'ICE snapshot categories do not sum')
ice=dict(period='2023-10-01/2024-09-30',adp=adp,midnight_detainee_days=round(mandays),yearend_count=ws['F20'].value,adp_by_history={ws[f'A{r}'].value:ws[f'N{r}'].value for r in [100,101,102]},initial_bookins=ws['E29'].value,source_sheet='Detention FY24',source_cells=['N99:N102','F20:F23','E29','B99:M99'])
rows=list(csv.DictReader((DEST/'scaap_fy2024_awards.csv').open()))
require(len(rows)==485, 'Expected 485 SCAAP application rows')
require(len({r['application'] for r in rows})==485, 'Duplicate SCAAP applications')
totals={k:sum(float(r[k]) for r in rows) for k in ['salary_usd','total_days','confirmed_days','unknown_days','award_usd']}
require(all(float(r['total_days'])>0 and float(r['award_usd'])>=0 for r in rows), 'Invalid SCAAP exposure or awards')
bjs=list(csv.reader((CACHE/'bjs_jails2023_table12.csv').open(encoding='cp1252')))
header=next(r for r in bjs if r and r[0]=='Year')
ice_cols=[i for i,name in enumerate(header) if name=='U.S. Immigration and Customs Enforcement']
require(len(ice_cols)==2, 'BJS ICE count/SE headers absent')
se_headers=[r for r in bjs if 'Standard errors' in r]
require(len(se_headers)==1, 'BJS standard-error header missing or duplicated')
boundary=se_headers[0].index('Standard errors')
require(ice_cols[0]<boundary<=ice_cols[1], 'BJS count/SE sections disagree')
row2023=next(r for r in bjs if r and r[0]=='2023*')
held_ice,ice_se=[int(row2023[i].replace(',','')) for i in ice_cols]
require(held_ice>0 and ice_se>=0, 'Invalid BJS count or SE')
results=dict(ice_fy2024=ice,bjs_midyear2023=dict(local_jail_held_for_ice=held_ice,standard_error=ice_se,source_columns_zero_based=ice_cols,scope='last weekday in June; rounded estimated count, not annual days'),scaap_fy2024=dict(rows=len(rows),reporting_period='2022-07-01/2023-06-30',totals=totals,limitation='Unknown days do not all become confirmed undocumented days; salary total covers all reported inmates, not exclusively undocumented inmates.'),verified_source_files=len(manifest))
(DEST/'probes.json').write_text(json.dumps(results,indent=2))
print(json.dumps(results,indent=2))
