#!/usr/bin/env python3
"""Pinned official incident proxies; default is offline and never silently repins.

Run with --download to fetch missing files against existing hashes. --initialize-pins
is a one-time explicit acquisition step and refuses to replace an existing manifest.
These are counts/location proxies, never measured spending or public-capital shares.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import resource
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import zipfile

import openpyxl
from bs4 import BeautifulSoup
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
PINS = ROOT / "inputs/incidents_pins.json"
OUTPUT = ROOT / "inputs/incident_shares.json"
SOURCES = {
    "fbi_location_2023": {"path": "_cache/incident_fbi_location_2023.zip", "cde_key": "nibrs/tables/2023/location.zip"},
    "fbi_location_2024": {"path": "_cache/incident_fbi_location_2024.zip", "cde_key": "nibrs/tables/2024/location.zip"},
    "fbi_contents_2023": {"path": "_cache/incident_fbi_tables_contents_2023.pdf", "cde_key": "nibrs/tables/2023/2023-nibrs-list-of-contents.pdf"},
    "usfa_runs_2020": {"path": "_cache/incident_usfa_run_profile_2020.pdf", "url": "https://www.usfa.fema.gov/downloads/pdf/statistics/v22i1-fire-department-run-profile.pdf"},
    "usfa_nonresidential_2023": {"path": "_cache/incident_usfa_nonresidential.html", "url": "https://www.usfa.fema.gov/statistics/nonresidential-fires/"},
    "usfa_allfires": {"path": "_cache/incident_usfa_statistics.html", "url": "https://www.usfa.fema.gov/statistics/"},
    "bjs_contract_trials_2005": {"path": "_cache/incident_bjs_contract_trials_2005.pdf", "url": "https://bjs.ojp.gov/content/pub/pdf/cbajtsc05.pdf"},
}

# Classification is an explicit analytic choice, not an FBI business-victim flag.
# Mixed residential/public/health/transport categories are excluded from narrow.
POLICE_NARROW = {
    "ATM Separate From Bank", "Auto Dealership New/Used", "Bank/Savings and Loan",
    "Bar/Nightclub", "Commercial/Office Building", "Construction Site",
    "Convenience Store", "Department/Discount Store", "Farm Facility",
    "Gambling Facility/Casino/Race Track", "Grocery/Supermarket", "Industrial Site",
    "Liquor Store", "Rental Storage Facility", "Restaurant", "Service/Gas Station",
    "Shopping Mall", "Specialty Store",
}
POLICE_BROAD_ADDITIONS = {
    "Air/Bus/Train Terminal", "Amusement Park", "Arena/Stadium/Fairgrounds/Coliseum",
    "Dock/Wharf/Freight/Modal Terminal", "Drug Store/Doctor's Office/Hospital",
    "Hotel/Motel/Etc.", "Parking/Drop Lot/Garage",
}
FIRE_NARROW = {"Stores/business", "Storage", "Manufacturing", "Industry"}
FIRE_NONRESIDENTIAL_STRUCTURE = FIRE_NARROW | {"Assembly", "Educational", "Health care"}


def require(condition, message):
    if not condition:
        raise RuntimeError("[BLOCKED] " + message)


def fetch(source):
    target = ROOT / source["path"]
    if "cde_key" in source:
        key = source["cde_key"]
        gateway = "https://cde.ucr.cjis.gov/LATEST/s3/signedurl?" + urllib.parse.urlencode({"key": key})
        with urllib.request.urlopen(gateway, timeout=60) as response:
            urls = json.load(response)
        require(key in urls, "FBI gateway returned no signed URL for " + key)
        url = urls[key]
    else:
        url = source["url"]
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (source verification)"})
    with urllib.request.urlopen(request, timeout=60) as response:
        data = response.read(20_000_001)
    require(0 < len(data) <= 20_000_000, "empty/oversized source " + source["path"])
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    # Strip expiring signature credentials from provenance; cde_key records retrieval.
    return url.split("?")[0] if "cde_key" in source else url


def verify_bytes(data, pin, name):
    require(hashlib.sha256(data).hexdigest() == pin["sha256"], "source hash mismatch: " + name)
    require(len(data) == pin["bytes"], "source length mismatch: " + name)


def initialize():
    require(not PINS.exists(), "manifest already exists; refusing to repin")
    pins = {}
    for name, source in SOURCES.items():
        target = ROOT / source["path"]
        url = source.get("url")
        if not target.exists():
            url = fetch(source)
        if "cde_key" in source:
            url = "https://cde-prd-data.s3.us-gov-east-1.amazonaws.com/" + source["cde_key"]
        data = target.read_bytes()
        pins[name] = {**source, "source_url": url, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}
    PINS.parent.mkdir(parents=True, exist_ok=True)
    PINS.write_text(json.dumps(pins, indent=2, sort_keys=True) + "\n")


def verified_sources(download):
    require(PINS.exists(), "source pins missing; acquire explicitly with --initialize-pins")
    pins = json.loads(PINS.read_text())
    require(set(pins) == set(SOURCES), "source manifest keys changed")
    for name, source in SOURCES.items():
        pin = pins[name]
        require(all(pin.get(k) == v for k, v in source.items()), "source declaration drift: " + name)
        target = ROOT / source["path"]
        if not target.exists() and download:
            fetch(source)
        require(target.exists(), "missing pinned source: " + source["path"])
        data = target.read_bytes()
        verify_bytes(data, pin, name)
    return pins


def police(year, path):
    tables = {}
    with zipfile.ZipFile(path) as archive:
        members = [n for n in archive.namelist() if n.endswith(".xlsx")]
        require(len(members) == 3, "expected three FBI location workbooks")
        for name in sorted(members):
            require(str(year) in name, "FBI workbook year differs")
            workbook = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
            rows = list(workbook.active.values)
            title = " ".join(str(r[0]) for r in rows[:5] if r[0] is not None)
            require(f"by Location, {year}" in title, "FBI title/year absent")
            require(any(r[0] == "Location" and "Total" in str(r[1]) for r in rows[:7]), "FBI location/total columns changed")
            totals = [r for r in rows if r[0] == "Total"]
            require(len(totals) == 1, "FBI total row nonunique")
            total = totals[0][1]
            values = {}
            for row in rows:
                if row[0] is None or row[0] == "Total" or not isinstance(row[1], (int, float)):
                    continue
                label = str(row[0]).strip()
                require(label not in values, "duplicate FBI location " + label)
                require(row[1] >= 0 and int(row[1]) == row[1], "FBI offense count invalid")
                detail = [v for v in row[2:] if isinstance(v, (int, float))]
                require(sum(detail) == row[1], "FBI row category closure failed: " + label)
                values[label] = int(row[1])
            require(sum(values.values()) == total, "FBI locations fail total closure")
            require(POLICE_NARROW | POLICE_BROAD_ADDITIONS <= set(values), "classified FBI labels missing")
            category = next(c for c in ["Persons", "Property", "Society"] if c in name)
            tables[category] = {"member": name, "total": int(total), "by_location": values}
            workbook.close()
    pooled = {}
    for table in tables.values():
        for label, count in table["by_location"].items():
            pooled[label] = pooled.get(label, 0) + count
    denominator = sum(t["total"] for t in tables.values())
    narrow = sum(pooled[label] for label in POLICE_NARROW)
    broad = sum(pooled[label] for label in POLICE_NARROW | POLICE_BROAD_ADDITIONS)
    return {
        "year": year, "unit": "reported Group A offenses, not unique incidents or victims",
        "denominator": denominator, "narrow_numerator": narrow, "broad_numerator": broad,
        "narrow_share": narrow / denominator, "broad_share": broad / denominator,
        "narrow_labels": sorted(POLICE_NARROW), "broad_added_labels": sorted(POLICE_BROAD_ADDITIONS),
        "tables": tables, "pooled_locations": pooled,
        "spending_share_identified": False,
        "limit": "Location is not victim type; persons/society offenses at businesses are included. Crime reporting, coverage, multi-offense counts, patrol/noncrime work, and incident resource intensity prevent identifying police spending/capital shares.",
    }


def fire(path):
    reader = PdfReader(path)
    page = reader.pages[6].extract_text(extraction_mode="layout")
    require("Figure 8" in page and "Stores/business" in page, "USFA Figure 8 missing")
    section = page.split("Figure 8", 1)[1].split("Percent of runs\n", 1)[0]
    labels = ["Residential", "Outside/special property", "Health care", "Stores/business", "Assembly", "None", "Other", "Storage", "Educational", "Manufacturing", "Industry", "Missing/blank", "Undetermined"]
    values = {}
    for label in labels:
        matches = re.findall(r"^\s*" + re.escape(label) + r"\s+(\d+\.\d+)\s*$", section, flags=re.M)
        require(len(matches) == 1, "USFA property label/value nonunique: " + label)
        values[label] = float(matches[0])
    # Source prints one decimal percent: exact total is not recoverable. Bound is
    # one half last decimal per printed category; do not silently renormalize.
    require(abs(sum(values.values()) - 100) <= len(values) * 0.05 + 1e-10, "USFA printed percentage closure exceeds rounding bound")
    footnotes = reader.pages[9].extract_text()
    count = re.findall(r"actual count of NFIRS runs used for the analyses in this report was ([\d,]+)", footnotes)
    require(len(count) == 1, "USFA actual run denominator missing")
    return {
        "year": 2020, "source_figure": "Figure 8, PDF page 7; count in note 3, PDF page 10",
        "unit": "fire-department collective responses (runs), including mutual aid",
        "denominator_runs": int(count[0].replace(",", "")), "printed_percent_by_property": values,
        "narrow_labels": sorted(FIRE_NARROW), "narrow_share": sum(values[x] for x in FIRE_NARROW) / 100,
        "nonresidential_structure_labels": sorted(FIRE_NONRESIDENTIAL_STRUCTURE),
        "nonresidential_structure_share": sum(values[x] for x in FIRE_NONRESIDENTIAL_STRUCTURE) / 100,
        "known_nonresidential_all_property_share": sum(values[x] for x in labels if x not in {"Residential", "Missing/blank", "Undetermined"}) / 100,
        "rounding_bound_pp": len(values) * 0.05,
        "spending_share_identified": False,
        "limit": "2020 voluntary NFIRS coverage, not national expenditure weights. Business-coded locations include medical/personal responses; storage need not be privately commercial. Broad nonresidential includes public education/healthcare and outdoor/public property. Hotels are classified residential. Counts and capital service requirements differ.",
    }


def fire_estimates(pins):
    parsed = {}
    for name in ["usfa_nonresidential_2023", "usfa_allfires"]:
        soup = BeautifulSoup((ROOT / pins[name]["path"]).read_text(), "html.parser")
        tables = soup.select("table#FiresData")
        require(len(tables) == 1, "USFA FiresData table missing/nonunique: " + name)
        table = tables[0]
        years = [int(cell.get_text(strip=True)) for cell in table.select("thead th")]
        cells = table.select("tbody td")
        require(len(years) == len(cells), "USFA fire-estimate columns differ")
        parsed[name] = dict(zip(years, [int(cell.get_text(strip=True).replace(",", "")) for cell in cells]))
        require(2023 in parsed[name], "USFA 2023 column absent")
    numerator = parsed["usfa_nonresidential_2023"][2023]
    denominator = parsed["usfa_allfires"][2023]
    require(0 < numerator < denominator, "USFA building numerator invalid")
    soup = BeautifulSoup((ROOT / pins["usfa_allfires"]["path"]).read_text(), "html.parser")
    table = soup.select_one("table#incidentTypeData")
    require(table is not None and "2024" in table.caption.get_text(), "USFA response-type table year absent")
    responses = {row.th.get_text(strip=True): float(row.td.get_text(strip=True)) for row in table.select("tbody tr")}
    require(abs(sum(responses.values()) - 100) <= len(responses) * 0.05 + 1e-10, "USFA response percentages exceed rounding closure")
    return {
        "year": 2023, "nonresidential_building_fires": numerator, "all_fires": denominator,
        "nonresidential_building_share_of_fires": numerator / denominator,
        "responses_2024_percent_by_type": responses,
        "limit": "USFA national fire estimates are not counts of all fire department runs. Never treat this fire-only share as an all-response share or multiply across these years to infer a business spending share.",
    }


def court_counterevidence(path):
    reader = PdfReader(path)
    text = reader.pages[2].extract_text()
    table = text.split("Table 3.", 1)[1].split("Table 4.", 1)[0]
    matches = re.findall(r"Totald\s+((?:\d+\.\d+%\s*){8})", table)
    require(len(matches) == 1, "BJS contract-trial Table 3 total row changed")
    values = [float(x) for x in re.findall(r"(\d+\.\d+)%", matches[0])]
    for half in [values[:4], values[4:]]:
        require(abs(sum(half) - 100) <= 4 * 0.05 + 1e-10, "BJS printed party percentages fail rounding closure")
    return {"year": 2005, "contract_trial_business_plaintiff_share": values[2] / 100,
            "contract_trial_business_defendant_share": values[6] / 100,
            "applied_to_courts_spending": False,
            "limit": "Conditional on contract trials in state general-jurisdiction courts; hierarchical party classification, not all cases/work/spending. A business defendant often faces an individual plaintiff. This measures litigants, not which activity makes service costs avoidable."}


def self_check(pins):
    # Failure catalog: corrupted pinned source; internally inconsistent source table.
    name = "fbi_location_2024"
    path = ROOT / pins[name]["path"]
    try:
        verify_bytes(path.read_bytes() + b"corrupt", pins[name], name)
    except RuntimeError as error:
        require("source hash mismatch" in str(error), "wrong corrupt-source failure")
    else:
        raise RuntimeError("[BLOCKED] corrupt-source negative control unexpectedly passed")
    with zipfile.ZipFile(path) as source:
        members = {n: source.read(n) for n in source.namelist()}
    name = sorted(members)[0]
    workbook = openpyxl.load_workbook(io.BytesIO(members[name]))
    worksheet = workbook.active
    row = next(r for r in worksheet.iter_rows() if r[0].value == "Total")
    row[1].value += 1
    changed = io.BytesIO()
    workbook.save(changed)
    members[name] = changed.getvalue()
    with tempfile.TemporaryDirectory(prefix="incident_negative_") as directory:
        broken = Path(directory) / "broken.zip"
        with zipfile.ZipFile(broken, "w") as archive:
            for member, data in members.items():
                archive.writestr(member, data)
        try:
            police(2024, broken)
        except RuntimeError as error:
            require("locations fail total closure" in str(error), "wrong count-closure failure")
        else:
            raise RuntimeError("[BLOCKED] count-closure negative control unexpectedly passed")
    print("PASS: failure catalog corrupt-source and count-closure negative controls")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--initialize-pins", action="store_true")
    parser.add_argument("--download", action="store_true")
    parser.add_argument("--self-check", action="store_true")
    args = parser.parse_args()
    start = time.perf_counter()
    if args.initialize_pins:
        initialize()
    pins = verified_sources(args.download)
    if args.self_check:
        self_check(pins)
    output = {
        "schema": "business-service-incident-proxies-v1",
        "source_sha256": {k: v["sha256"] for k, v in pins.items()},
        "police": {str(y): police(y, ROOT / pins[f"fbi_location_{y}"]["path"]) for y in [2023, 2024]},
        "fire": fire(ROOT / pins["usfa_runs_2020"]["path"]),
        "fire_only_crosscheck": fire_estimates(pins),
        "court_counterevidence": court_counterevidence(ROOT / pins["bjs_contract_trials_2005"]["path"]),
        "courts": {"business_share_used": 0, "status": "ASSUMPTION: no identified national expenditure-weighted share", "counterevidence": "BJS 2005 contract trials measure business parties but exclude nontrial civil, criminal, family, federal and limited-jurisdiction cases. US Courts C-2 tabulates federal civil nature of suit rather than business service expenditures."},
        "general_administration": {"business_share_used": 0, "status": "ASSUMPTION: no identified national expenditure-weighted share"},
        "central_measured_spending_share": None,
        "interpretation": "All incident fractions are sensitivity assumptions if applied to spending or public capital; none is a measured expenditure share.",
    }
    OUTPUT.write_text(json.dumps(output, indent=2, sort_keys=True) + "\n")
    rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    rss_mb = rss / (1024**2 if sys.platform == "darwin" else 1024)
    print(f"PASS: source hashes; police exact count closures; fire printed rounding closure; elapsed_s={time.perf_counter()-start:.3f}; peak_rss_mb={rss_mb:.2f}")
    print("police:", {y: {k: v[k] for k in ["denominator", "narrow_share", "broad_share"]} for y, v in output["police"].items()})
    print("fire:", {k: output["fire"][k] for k in ["narrow_share", "nonresidential_structure_share", "known_nonresidential_all_property_share"]})


if __name__ == "__main__":
    main()
