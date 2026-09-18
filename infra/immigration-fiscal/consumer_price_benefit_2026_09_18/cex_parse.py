#!/usr/bin/env python3
"""Parse BLS Consumer Expenditure Survey 2024 published tables.

Table 2500 (cu-all-detail-2024.xlsx) -> every detail line's mean for all consumer units.
Table 1101 (cu-income-quintiles-before-taxes-2024.xlsx) -> aggregate lines by quintile.

Writes derived/cex_detail_all.csv, derived/cex_quintile_parents.csv, derived/cex_audit.json.
Gate: the major Table 1101 components must sum to 'Average annual expenditures' per column.
"""
from __future__ import annotations
import hashlib, json
from pathlib import Path
import openpyxl, pandas as pd

HERE = Path(__file__).resolve().parent
OUT = HERE / "derived"; OUT.mkdir(exist_ok=True)
CEX = Path("/Users/alien/research-data/immigration-fiscal/data/external/cex_2024")
DETAIL = CEX / "cu-all-detail-2024.xlsx"
QUINT = CEX / "cu-income-quintiles-before-taxes-2024.xlsx"
STATS = {"Mean", "Share", "SE", "RSE", "Percent Reporting"}

MAJOR = ["Food", "Alcoholic beverages", "Housing", "Apparel and services", "Transportation",
         "Healthcare", "Entertainment", "Personal care products and services", "Reading",
         "Education", "Tobacco products and smoking supplies", "Miscellaneous",
         "Cash contributions", "Personal insurance and pensions"]


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def parse(path: Path, ncols: int) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows, item = [], None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None or lab == "":
            continue
        lab = str(lab).strip()
        if lab in STATS:
            if item is None:
                continue
            vals = [ws.cell(r, c).value for c in range(2, 2 + ncols)]
            rows.append(dict(row=r, item=item, stat=lab,
                             **{f"c{i}": v for i, v in enumerate(vals)}))
        else:
            vals = [ws.cell(r, c).value for c in range(2, 2 + ncols)]
            if any(v is not None for v in vals):
                # single-line item with values on the label row (characteristics)
                rows.append(dict(row=r, item=lab, stat="Mean",
                                 **{f"c{i}": v for i, v in enumerate(vals)}))
                item = lab
            else:
                item = lab
    return pd.DataFrame(rows)


def main() -> None:
    det = parse(DETAIL, 1)
    det = det[det.stat == "Mean"][["row", "item", "c0"]].rename(columns={"c0": "mean_all_cu"})
    det.to_csv(OUT / "cex_detail_all.csv", index=False)

    q = parse(QUINT, 6)
    cols = ["all_cu", "q1_lowest", "q2_second", "q3_third", "q4_fourth", "q5_highest"]
    q = q[q.stat == "Mean"].rename(columns={f"c{i}": c for i, c in enumerate(cols)})
    q = q[["row", "item"] + cols]
    q.to_csv(OUT / "cex_quintile_parents.csv", index=False)

    # ---- gate: major components sum to Average annual expenditures, per column ----
    def get(df, name, dup_first=True):
        s = df[df.item == name]
        if s.empty:
            raise SystemExit(f"[BLOCKED] missing CEX line: {name}")
        return s.iloc[0] if dup_first else s
    tot = get(q, "Average annual expenditures")
    gate = {}
    for c in cols:
        s = sum(float(get(q, m)[c]) for m in MAJOR)
        t = float(tot[c])
        gate[c] = dict(sum_major=s, published_total=t, abs_diff=abs(s - t),
                       rel_diff=abs(s - t) / t)
        if abs(s - t) / t > 0.005:
            raise SystemExit(f"[BLOCKED] CEX gate failed for {c}: {s} vs {t}")
    ncu = get(q, "Number of consumer units (in thousands) a/")
    audit = dict(detail_file=str(DETAIL), detail_sha256=sha256(DETAIL),
                 quintile_file=str(QUINT), quintile_sha256=sha256(QUINT),
                 source_urls=["https://www.bls.gov/cex/tables/calendar-year/mean/cu-all-detail-2024.xlsx",
                              "https://www.bls.gov/cex/tables/calendar-year/mean-item-share-average-standard-error/cu-income-quintiles-before-taxes-2024.xlsx"],
                 fetched="2026-09-18",
                 consumer_units_thousands={c: float(ncu[c]) for c in cols},
                 major_component_gate=gate,
                 n_detail_lines=int(len(det)), n_quintile_lines=int(len(q)))
    (OUT / "cex_audit.json").write_text(json.dumps(audit, indent=2))
    print(json.dumps(gate, indent=2))
    print(f"[ok] {len(det)} detail lines, {len(q)} quintile lines")


if __name__ == "__main__":
    main()
