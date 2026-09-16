import openpyxl, json, csv
wb=openpyxl.load_workbook('cmc_analysis_sexual_offences_by_nationality_2018_2024.xlsx',data_only=True)
ws=wb['SO Proceeded against']
hdr=[c for c in next(ws.iter_rows(min_row=7,max_row=7,values_only=True))]
cmc=[]
for r in ws.iter_rows(min_row=8,max_row=150,values_only=True):
    if not r[0] or r[8] in (None,''): continue
    if str(r[0]).strip() in ('Grand Total','Not Recorded'): continue
    num=lambda x: float(x) if isinstance(x,(int,float)) else None
    cmc.append(dict(nationality=str(r[0]).strip(), total_2018_2024=num(r[8]),
                    pop_aps=num(r[9]), rate_cmc=num(r[10])))
    if cmc[-1]['total_2018_2024'] is None: cmc.pop()
# census country of birth, London
d=json.load(open('ons_cob190_rgn.json'))
cob={}
for o in d['observations']:
    dm={x['dimension_id']:x for x in o['dimensions']}
    if dm['rgn']['option_id']=='E12000007':
        cob[dm['country_of_birth_190a']['option'].split(': ')[-1]]=o['observation']
ALIAS={'Moroccan':'Morocco','Congo (Democratic Republic)':'Congo (Democratic Republic)',
       'Trinidad And Tobago':'Trinidad and Tobago','St. Lucia':'Saint Lucia','Ivory Coast':"Côte d'Ivoire"}
uk_born=cob.get('England',0)+cob.get('Scotland',0)+cob.get('Wales',0)+cob.get('Northern Ireland',0)+cob.get('United Kingdom not otherwise specified',0)
rows=[]
for c in cmc:
    n=ALIAS.get(c['nationality'],c['nationality'])
    p=cob.get(n)
    if c['nationality'] in ('United Kingdom','British','England'): p=uk_born
    rows.append(dict(nationality=c['nationality'], proceeded_2018_2024=c['total_2018_2024'],
        pop_APS_nationality_2021=c['pop_aps'],
        rate_per10k_CMC=round(c['rate_cmc'],1) if c['rate_cmc'] else None,
        pop_census2021_country_of_birth=p,
        rate_per10k_census=round(10000*c['total_2018_2024']/p,1) if p else None))
rows.sort(key=lambda r:-(r['rate_per10k_CMC'] or 0))
with open('uk_london_sex_offence_rates.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
tot=sum(c['total_2018_2024'] for c in cmc)
print(f"CMC file: {len(cmc)} nationalities, {tot:,.0f} persons proceeded against 2018-2024")
print(f"UK-born London census 2021: {uk_born:,}")
print()
print(f"{'nationality':<22}{'persons':>9}{'APSpop':>10}{'rate_CMC':>10}{'CensusCoB':>11}{'rate_cens':>10}{'ratio':>7}")
for r in rows[:14]:
    rr = round(r['rate_per10k_CMC']/r['rate_per10k_census'],1) if (r['rate_per10k_census'] and r['rate_per10k_CMC']) else ''
    print(f"{r['nationality']:<22}{r['proceeded_2018_2024']:>9,.0f}{(r['pop_APS_nationality_2021'] or 0):>10,.0f}{str(r['rate_per10k_CMC']):>10}{(r['pop_census2021_country_of_birth'] or 0):>11,}{str(r['rate_per10k_census']):>10}{str(rr):>7}")
print("\n--- CMC summary block (rows 152-159 of the same sheet) ---")
for lab,n,p in [("British",4631,7128000),("Non-British",2809,1944000)]:
    print(f"  {lab:<12} {n:>6,} persons / APS pop {p:>10,} = {10000*n/p:>5.1f} per 10,000 (2018-2024 cumulative)")
print(f"  Non-British / British rate ratio = {(2809/1944000)/(4631/7128000):.2f}x")
