#!/usr/bin/env python3
"""K4: finer IRS tax patterns within frozen BEA cells; all-form and limited C-only arms.

OPENBLAS_NUM_THREADS=1 uv run --no-project python3 {lane}/k4_detail.py
Source-tax levels, between-cell allocation, mining and other lines are held fixed.
"""
from __future__ import annotations
import csv
import hashlib
import json
import resource
import re
import time
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
CACHE = HERE / '_cache'
INPUTS = HERE / 'inputs'
OUT = HERE / 'derived'
# Each bin is IRS Table 1 worksheet row(s) and public CPS code(s). Aggregate bins
# whenever either published code crosses IRS minor industries; never split such a code.
# Only cells with complete minor tax values enter the all-form arm.
BINS = {
 '311FT': [([35],[1070]),([36],[1080]),([37],[1090]),([38],[1170]),([39,41],[1180,1280]),([40],[1190,1270]),([43,44,45],[1370]),([46],[1390])],
 '322': [([57],[1870]),([58],[1880,1890])],
 '324': [([61],[2070]),([62],[2090])],
 '325': [([64,69],[2180,2290]),([65],[2170]),([66],[2190]),([67],[2270]),([68],[2280])],
 '326': [([71],[2370]),([72],[2380,2390])],
 '327': [([74],[2470,2480,2590]),([75],[2490]),([76],[2570])],
 '331': [([78],[2670]),([79],[2680,2690]),([80],[2770])],
 '332': [([82],[2780]),([83,87],[2790,2880,2970,2980]),([84,85],[2870]),([86],[2890])],
 '333': [([89],[3070,3080]),([90,92,95],[3291]),([91],[3095]),([93],[3170]),([94],[3180])],
 '334': [([97],[3365]),([98,99,100],[3370,3390]),([101],[3380])],
 '3364OT': [([108],[3580,3590]),([109],[3680]),([110],[3670,3690])],
 '339': [([113],[3960]),([114],[3970,3980])],
 '42': [([117],[4070]),([118],[4090]),([119],[4170]),([120],[4180]),([121],[4195]),([122],[4265]),([123],[4270]),([124],[4080,4280,4290]),([126],[4370]),([127,131,134],[4380,4570,4580]),([128],[4390]),([129],[4470]),([130],[4480]),([132],[4490]),([133],[4560]),([135],[4585])],
 '441': [([138],[4670]),([139],[4680,4690,4681,4691])],
 '445': [([146],[4971,4972,4973,4980,4981]),([147],[4990,4991])],
 '4A0': [([149],[4770,4771]),([150],[4780,4795,4796]),([141,143],[4870,4871]),([142],[4880,4881]),([144],[4890,4891]),([152],[5070,5080,5071,5081]),([153],[5090,5680]),([154],[5170,5180,5190,5171,5181,5191]),([155],[5275,5280,5295,5370,5470,5480,5490,5570,5580,5593,5670,5690,5276,5281,5296,5371,5471,5481,5491,5571,5581])],
 '511': [([171],[6470,6471]),([172,173,174],[6480,6481]),([175],[6490])],
 '512': [([168],[6570]),([169],[6590])],
 '514': [([178],[6695]),([179],[6672,6770,6780,6781])],
 '521CI': [([182],[6870,6871]),([183],[6880,6881]),([185,186,187,188],[6890])],
 '524': [([194,195],[6991]),([196,197],[6992])],
 '5412OP': [([214],[7280]),([215],[7290]),([216],[7370]),([218],[7390]),([219],[7460]),([220],[7470]),([221],[7480,7490])],
}
# C-only columns are explicit concordance to the Table 5.3 labels, checked below.
COLUMNS = {
35:26,36:27,37:28,38:29,39:30,40:31,41:32,43:33,44:34,45:35,46:36,
57:44,58:45,61:47,62:48,64:49,65:50,66:51,67:52,68:53,69:54,
71:55,72:56,74:57,75:58,76:59,78:60,79:61,80:62,82:63,83:64,84:65,85:66,86:67,87:68,
89:69,90:70,91:71,92:72,93:73,94:74,95:75,97:76,98:77,99:78,100:79,101:80,
108:85,109:86,110:87,113:89,114:90,
117:92,118:93,119:94,120:95,121:96,122:97,123:98,124:99,126:100,127:101,128:102,129:103,130:104,131:105,132:106,133:107,134:108,135:109,
138:111,139:112,141:113,142:114,143:115,144:116,146:117,147:118,149:119,150:120,152:122,153:123,154:124,155:125,
168:136,169:137,171:138,172:139,173:140,174:141,175:142,178:145,179:146,
182:148,183:149,185:150,186:151,187:152,188:153,194:157,195:158,196:159,197:160,
214:173,215:174,216:175,218:177,219:178,220:179,221:180,
}
FILES = {
 'k4_22co01ccr.xlsx': 'https://www.irs.gov/pub/irs-soi/22co01ccr.xlsx',
 'k4_22co53ccr.xlsx': 'https://www.irs.gov/pub/irs-soi/22co53ccr.xlsx',
 'k4_p16_ty2022.pdf': 'https://www.irs.gov/pub/irs-pdf/p16.pdf',
 'k4_bea_supply_use.zip': 'https://apps.bea.gov/industry/release/zip/SUPPLY-USE.zip',
 'k4_census_industry.xlsx': 'https://www2.census.gov/programs-surveys/demo/guidance/industry-occupation/2022-Census-Industry-Code-List-with-Crosswalk.xlsx',
}

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def main():
    started = time.perf_counter()
    INPUTS.mkdir(exist_ok=True)
    OUT.mkdir(exist_ok=True)
    pins_path = INPUTS / 'k4_pins.json'
    pins = json.loads(pins_path.read_text()) if pins_path.exists() else {}
    for name, url in FILES.items():
        path = CACHE / name
        actual = digest(path)
        if name in pins and actual != pins[name]['sha256']:
            raise SystemExit(f'[BLOCKED] changed K4 source: {name}')
        pins[name] = {'url':url, 'sha256':actual, 'bytes':path.stat().st_size}
    pins_path.write_text(json.dumps(pins,indent=1)+'\n')
    alltable = openpyxl.load_workbook(CACHE/'k4_22co01ccr.xlsx',data_only=True).worksheets[0]
    ctable = openpyxl.load_workbook(CACHE/'k4_22co53ccr.xlsx',data_only=True).worksheets[0]
    assert '2022' in alltable.cell(2,1).value and 'after' in alltable.cell(4,12).value
    assert '1120S' in ctable.cell(1,1).value and ctable.cell(80,1).value == 'Total income tax after credits'
    assert 'thousands' in alltable.cell(3,1).value and 'thousands' in ctable.cell(3,1).value
    core = np.load(CACHE/'core.npz')
    meta = json.loads((OUT/'core.json').read_text())
    cells = core['cells'].tolist()
    codes = core['industry_codes']
    groups = ['union','G1','G2','G3plus','white']
    pc = core['pair_cell_index']
    pi = core['pair_code_index']
    gates = []
    def gate(name, ok, detail=None):
        gates.append({'name':name,'passed':bool(ok),'detail':detail})
        if not ok:
            raise SystemExit(f'[BLOCKED] {name}: {detail}')
    gate('all frozen core gates passed',all(g['passed'] for g in meta['gates']))
    sectors=[9,13,19,24,33,115,136,156,166,180,201,212,222,225,231,232,240,243,246]
    national_k=alltable.cell(8,12).value
    sector_k=sum(alltable.cell(r,12).value for r in sectors)
    rounding_gap_k=sector_k-national_k
    gate('IRS national inside published sector rounding intervals before normalization',abs(rounding_gap_k)<=len(sectors)*0.5,rounding_gap_k)
    gate('IRS source documents rounding nonclosure','rounding' in alltable.cell(252,1).value)
    reconciled=np.array([alltable.cell(r,12).value for r in sectors],float)*national_k/sector_k
    gate('IRS reconciled national closure <=1e-6bn',abs(reconciled.sum()/1e6-national_k/1e6)<1e-6)
    reconstructed={}
    for g in ['national']+groups:
        sums=np.zeros((len(cells),161))
        np.add.at(sums,pc,core[f'pair_wages_{g}'])
        reconstructed[g]=sums
    for g in groups:
        pooled=[]
        for pool in core['cell_pool']:
            indices=[int(i) for i in str(pool).split(',')]
            numer=reconstructed[g][indices].sum(axis=0)
            denom=reconstructed['national'][indices].sum(axis=0)
            pooled.append(np.divide(numer,denom,out=np.zeros(161),where=denom>0))
        gap=float(np.max(abs(np.array(pooled)-core[f'shares_{g}'])))
        gate(f'{g} code-cell aggregate positive control reproduces frozen shares <=1e-9',gap<1e-9,gap)
    def norm(x):
        return re.sub('[^a-z0-9]','',str(x).lower().replace('&','and'))
    exceptions={
      122:('Hardware, plumbing, and heating equipment and supplies','Hardware, plumbing, heating equipment, and supplies'),
      124:('Furniture, sports, toys, recyclable, jewelry, and other durable goods','Furniture, sports, toys, recycle, jewelry, and other durable goods'),
      141:('Home centers; paint and wallpaper stores','Home centers; paint and wallpaper retailers'),
      143:('Other building material dealers','Other building materials dealers'),
      169:('Sound recording industries','Sound recoding industries'),
      188:('Activities related to credit intermediation (loan brokers, check clearing, and money transmitting)','Activities related to credit intermediation (including loan brokers, check clearing, and money transmitting)'),
      197:('Other insurance related activities (including third-party administration of insurance and pension funds)','Other insurance related activities (including third-party administrator of insurance and pension funds)'),
    }
    for r,c in COLUMNS.items():
        a,b=alltable.cell(r,1).value,ctable.cell(6,c).value
        same=norm(a)==norm(b)
        accepted=r in exceptions and (norm(a),norm(b))==tuple(norm(v) for v in exceptions[r])
        gate(f'IRS concordance exact normalized labels or named exception {r}/{c}',same or accepted)
    census=openpyxl.load_workbook(CACHE/'k4_census_industry.xlsx',data_only=True).worksheets[1]
    primary_codes={int(str(row[3])):(str(row[1]).strip(),str(row[4]).strip()) for row in census.values if str(row[3]).isdigit()}
    frozen_crosswalk=HERE.parent/'job_capital_taxes_2026_10_14/inputs/census_industry_crosswalk.csv'
    with frozen_crosswalk.open() as f:
        frozen_codes={int(r['census_code']):r for r in csv.DictReader(f)}
    code_trace=[]
    # Existing frozen crosswalk includes earlier NAICS spellings. These verified
    # 2022 changes preserve the matched bins, and are retained in the trace.
    naics_exceptions={2090:('32419','32412, 32419'),2480:('327120','32712'),
      2970:('332992, 332993, 332994, 332996','332992, 332993, 332994'),
      2980:('3325, 3326, 3329 exc. 332992, 332993, 332994, 332996','3325, 3326, 3329 exc. 332992, 332993, 332994'),
      5090:('447','4571'),5680:('454310','4572'),6490:('5112','5132'),6770:('51912','51921')}
    for cell,bins in BINS.items():
        for rr,cc in bins:
            for code in cc:
                if code not in codes:
                    continue
                gate(f'CPS live code {code} exists in primary 2022 list',code in primary_codes)
                desc,naics=primary_codes[code]
                frozen=frozen_codes[code]
                accepted=code in naics_exceptions and (norm(frozen['naics']),norm(naics))==tuple(norm(v) for v in naics_exceptions[code])
                gate(f'CPS live code {code} primary NAICS agrees or named 2022 change',norm(naics)==norm(frozen['naics']) or accepted)
                code_trace.append({'cell':cell,'cps_code':code,'primary_description':desc,'primary_naics':naics,'frozen_naics':frozen['naics'],'source_rows':';'.join(map(str,rr)),'irs_labels':';'.join(str(alltable.cell(r,1).value).strip() for r in rr),'interpretation':'matched published classes; broader not-specified classes are diagnosed separately'})
    traces=[]
    rows=[]
    outputs={}
    coverage={}
    for mode in ['allform','conly','strict_allform','strict_conly']:
        targets={g:core[f'shares_{g}'].T @ core['tax_corporate'] for g in groups}
        share_matrices={g:core[f'shares_{g}'].copy() for g in groups}
        used=[]
        skipped=[]
        for cell,bins in BINS.items():
            ci=cells.index(cell)
            vals=[]
            for source_rows, ccodes in bins:
                values=[alltable.cell(r,12).value if mode.endswith('allform') else ctable.cell(80,COLUMNS[r]).value for r in source_rows]
                vals.append(sum(values) if all(isinstance(v,(int,float)) for v in values) else None)
            if any(v is None for v in vals):
                skipped.append({'cell':cell,'reason':'IRS suppressed minor tax; retain frozen cell key'})
                continue
            native=(pc==ci)
            defined=set(c for _,cs in bins for c in cs)
            gate(f'{mode} {cell} bins disjoint',sum(len(cs) for _,cs in bins)==len(defined))
            residual=native & ~np.isin(codes[pi],list(defined))
            if mode.startswith('strict_') and (residual.any() or cell=='333'):
                skipped.append({'cell':cell,'reason':'public nonspecified industry can cross bins; strict arm retains whole frozen cell'})
                continue
            binmasks=[native & np.isin(codes[pi],cs) for _,cs in bins]
            nobs=np.stack([core['pair_wages_national'][b].sum(axis=0) for b in binmasks])
            gate(f'{mode} {cell} every bin has national wages',bool((nobs>0).all()))
            nres=core['pair_wages_national'][residual].sum(axis=0)
            nall=nobs.sum(axis=0)+nres
            records=np.array([core['pair_records'][b].sum() for b in binmasks])
            tax=np.array(vals,dtype=float)/sum(vals)*core['tax_corporate'][ci]
            gate(f'{mode} {cell} tax closes within 1e-6bn',abs(tax.sum()-core['tax_corporate'][ci])<1e-6)
            # Unknown/not-specified CPS wages remain an explicit separate term at
            # the frozen parent share. Tax is partitioned by identifiable national
            # wage coverage; the unknown share of each bin's tax is kept at parent.
            known_frac=nobs/(nobs+nres[None,:]*(nobs/nobs.sum(axis=0)))
            for g in groups:
                obs=np.stack([core[f'pair_wages_{g}'][b].sum(axis=0) for b in binmasks])
                old=core[f'shares_{g}'][ci]
                fine=obs/nobs
                fine[records<100]=old
                blend=fine*known_frac+old[None,:]*(1-known_frac)
                revised=tax @ blend
                share_matrices[g][ci]=revised/core['tax_corporate'][ci]
                targets[g]+=revised-core['tax_corporate'][ci]*old
                for bi,(rr,cs) in enumerate(bins):
                    traces.append({'mode':mode,'cell':cell,'bin':bi,'group':g,'source_rows':';'.join(map(str,rr)),'cps_codes':';'.join(map(str,cs)),
                      'records':int(records[bi]),'pool_at_parent':bool(records[bi]<100),'tax_bn':float(tax[bi]),'known_wage_fraction':float(known_frac[bi,0]),'share':float(blend[bi,0])})
            used.append(cell)
            rows.extend({'mode':mode,'cell':cell,'source_row':r,'allform_label':str(alltable.cell(r,1).value),'allform_tax_thousand':alltable.cell(r,12).value,'conly_column':COLUMNS[r],'conly_label':ctable.cell(6,COLUMNS[r]).value,'conly_tax_thousand':ctable.cell(80,COLUMNS[r]).value} for rs,_ in bins for r in rs)
            # Ignoring residual labels is never silent: preserve the exact codes.
            skipped.append({'cell':cell,'reason':'unmatched CPS code wages held at frozen share','codes':np.unique(codes[pi[residual]]).tolist(),'national_wage_fraction':float(nres[0]/nall[0])})
            if cell=='333':
                skipped.append({'cell':cell,'reason':'3291 includes part of 333, beyond its named n.e.c. classes; fuller arm treats it as 3332/3334/3339, strict arm excludes whole cell'})
        for g in groups:
            outputs[f'{mode}_{g}']=targets[g]
            outputs[f'fine_shares_{mode}_{g}']=share_matrices[g]
        coverage[mode]={'refined_cells':used,'refined_tax_bn':float(sum(core['tax_corporate'][cells.index(c)] for c in used)),'limitations':skipped}
    gate('fine targets generations add to union',all(np.max(abs(outputs[f'{a}_union']-sum(outputs[f'{a}_{g}'] for g in ['G1','G2','G3plus'])))<1e-9 for a in coverage))
    total=core['full_corporate_bn'].sum()
    gate('frozen national tax preserved <=1e-6bn',abs(total-663.685)<1e-6,float(total))
    def se(v):return float(np.sqrt(4/160*np.square(v[1:]-v[0]).sum()))
    results={}
    for mode in coverage:
        results[mode]={}
        for reading in ['low','high']:
            r=meta['responses'][reading]['corporate']
            old={g:core[f'shares_{g}'].T @ core['tax_corporate'] for g in groups}
            change=r*(old['union']-outputs[f'{mode}_union'])
            lineage_change=change+sum(meta['lineage_rates'][g]*r*(old[g]-outputs[f'{mode}_{g}']) for g in ['G3plus','white'])
            results[mode][reading]={'net_cost_change_union_bn':float(change[0]),'union_change_se_bn':se(change),'net_cost_change_lineage_bn':float(lineage_change[0]),'lineage_change_conditional_se_bn':se(lineage_change),'over_5bn':bool(abs(lineage_change[0])>=5)}
    np.savez(CACHE/'k4_targets.npz',**outputs,detail_union=outputs['allform_union'],detail_G3plus=outputs['allform_G3plus'],detail_white=outputs['allform_white'])
    for name,values in [('k4_source_rows.csv',rows),('k4_detail_bins.csv',traces),('k4_code_concordance.csv',code_trace)]:
        with (OUT/name).open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(values[0]),lineterminator='\n');w.writeheader();w.writerows(values)
    with zipfile.ZipFile(CACHE/'k4_bea_supply_use.zip') as z:
        detail=openpyxl.load_workbook(z.open('Use_SUT_Detail.xlsx'),data_only=True,read_only=True)
        detail_years=[s for s in detail.sheetnames if s.isdigit()]
        otop_rows=[i for i,row in enumerate(detail[detail_years[-1]].values,1) if 'T00OTOP' in row]
    gate('property-pattern analogue has detailed other-production-tax row',bool(otop_rows))
    summary={'sources':pins,'source_totals_bn':{'allform':alltable.cell(8,12).value/1e6,'conly':ctable.cell(80,2).value/1e6,'allform_sector_sum':sector_k/1e6,'raw_rounding_gap':rounding_gap_k/1e6},'coverage':coverage,'results':results,'gates':gates,'business_property_analogue':{'file':'Use_SUT_Detail.xlsx','years':detail_years,'otop_rows_latest':otop_rows,'limitation':'Detailed other taxes on production exist only through 2017; not a 2024 business-property-tax-only minor table. No property key changed.'},'method':'Within each frozen summary cell, IRS minor taxes partition the unchanged corporate amount. Public CPS codes aggregated to identifiable IRS bins; bins under 100 earners and unresolved code wages retain frozen parent share. Corporate response, other lines, national tax and lineage rule unchanged. All-form and C-only limited arms are separate.','uncertainty':meta['uncertainty']}
    (OUT/'k4_summary.json').write_text(json.dumps(summary,indent=1)+'\n')
    print(json.dumps(results,indent=1))
    print(f'elapsed_seconds={time.perf_counter()-started:.3f} maxrss_bytes={resource.getrusage(resource.RUSAGE_SELF).ru_maxrss}')

if __name__=='__main__':main()
