#!/usr/bin/env python3
"""K1 public enterprise-size legal-form payroll proxy, never a measured group C wage bill.

Run with --acquire once to stage missing source bytes and pin them. Ordinary runs
verify pins, parse primary workbooks, and use common core.npz (built by core.py).
The Census C-and-other-corporate category separates S corporations. Its finest
public industry detail is sector, not the frozen tax cells. Source probabilities
are applied under conditional exchangeability within sector and enterprise size.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import resource
import time
from pathlib import Path

import numpy as np
import openpyxl
import requests

HERE = Path(__file__).resolve().parent
FROZEN = HERE.parent / "job_capital_taxes_2026_10_14"
SOURCES = {
    "susb_2022.xlsx": "https://www2.census.gov/programs-surveys/susb/tables/2022/us_state_naicssector_lfo_2022.xlsx",
    "22co53ccr.xlsx": "https://www.irs.gov/pub/irs-soi/22co53ccr.xlsx",
    "22co61ccr.xlsx": "https://www.irs.gov/pub/irs-soi/22co61ccr.xlsx",
    "p16_2022.pdf": "https://www.irs.gov/pub/irs-pdf/p16.pdf",
    "susb_methodology.html": "https://www.census.gov/programs-surveys/susb/technical-documentation/methodology.html",
    "susb_glossary.html": "https://www.census.gov/programs-surveys/susb/about/glossary.html",
    "susb_lfo.html": "https://www.census.gov/programs-surveys/susb/technical-documentation/data-user-resources/legal-form-of-organization.html",
    "susb_tables.html": "https://www.census.gov/programs-surveys/susb/data/tables.html",
    "irs_tables.html": "https://www.irs.gov/statistics/soi-tax-stats-corporation-income-tax-returns-complete-report-publication-16",
    "ec2022_lfo_group.json": "https://api.census.gov/data/2022/ecnsize/groups/EC2200SIZELFO.json",
    "ec2022_firmsize_group.json": "https://api.census.gov/data/2022/ecnsize/groups/EC2200SIZEEMPFIRM.json",
    "abs2023_variables.json": "https://api.census.gov/data/2023/abscs/variables.json",
    "cbp2023_variables.json": "https://api.census.gov/data/2023/cbp/variables.json",
}
SECTOR_CODE = {
    "agriculture": "11", "mining": "21", "utilities": "22", "construction": "23",
    "durable": "31-33", "nondurable": "31-33", "wholesale": "42", "retail": "44-45",
    "transportation": "48-49", "information": "51", "finance": "52", "real_estate": "53",
    "professional": "54", "management": "55", "administrative": "56", "education": "61",
    "health": "62", "arts": "71", "accommodation_food": "72", "other_services": "81",
}
IRS_SECTORS = ["11", "21", "22", "23", "31-33", "42", "44-45", "48-49", "51", "52", "53", "54", "55", "56", "61", "62", "71", "72", "81"]
GATES = []


def gate(name, passed, value=None):
    GATES.append({"name": name, "passed": bool(passed), "value": value})
    if not passed:
        raise RuntimeError(f"[BLOCKED] {name}: {value}")


def dump(path, obj):
    path.write_text(json.dumps(obj, indent=2, allow_nan=False) + "\n")


def csv_out(path, rows):
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def acquire(do_acquire):
    (HERE / "inputs").mkdir(exist_ok=True)
    (HERE / "derived").mkdir(exist_ok=True)
    (HERE / "_cache").mkdir(exist_ok=True)
    pp = HERE / "inputs/k1_pins.json"
    pins = json.loads(pp.read_text()) if pp.exists() else {}
    for name, url in SOURCES.items():
        path = HERE / "_cache" / f"k1_{name}"
        if not path.exists():
            gate(f"acquisition authorized for missing {name}", do_acquire)
            response = requests.get(url, timeout=60)
            response.raise_for_status()
            path.write_bytes(response.content)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if name in pins:
            gate(f"pin {name}", pins[name]["sha256"] == digest)
        else:
            gate(f"pin creation authorized for {name}", do_acquire)
            pins[name] = {"url": url, "path": str(path.relative_to(HERE)), "sha256": digest, "bytes": path.stat().st_size}
    if do_acquire:
        dump(pp, pins)
    return pins


def parse_susb():
    wb = openpyxl.load_workbook(HERE / "_cache/k1_susb_2022.xlsx", read_only=True, data_only=True)
    sheet = wb["LFO"]
    rows = [r for r in sheet.values if r[0] == "00"]
    gate("primary workbook identifies enterprise employment size", "Enterprise Employment Size: 2022" in sheet.cell(1, 1).value)
    gate("S corporate category separated", {"2: Corporation", "3: S-Corporation"} <= {r[2] for r in rows})
    lookup = {(r[3], r[2].split(":")[0], r[5].split(":")[0]): r for r in rows}
    sectors = sorted({r[3] for r in rows})
    forms = sorted({r[2].split(":")[0] for r in rows})
    # Disjoint source bins: <5, 5-9, 10-19, 20-99, 100-499, 500+.
    bins = ["2", "3", "4", "6", "7", "9"]
    unpublished = []
    for sec in sectors:
        sec_forms = [form for form in forms if (sec, form, "1") in lookup]
        for form in sec_forms:
            total = lookup[sec, form, "1"][10]
            missing = [b for b in bins if (sec, form, b) not in lookup]
            gap = sum(lookup[sec, form, b][10] for b in bins if b not in missing) - total
            if missing:
                gate(f"SUSB used all/C profiles completely published {sec}/{form}", form not in ["1", "2"] or sec == "99", missing)
                unpublished.append({"sector": sec, "form": form, "unpublished_size_codes": ",".join(missing), "residual_payroll_bn": -gap / 1e6})
            else:
                gate(f"SUSB size payroll closure {sec}/{form} at 1e-6bn", abs(gap) <= 1, gap / 1e6)
        for b in ["1"]:
            gap = sum(lookup[sec, form, b][10] for form in sec_forms if form != "1") - lookup[sec, "1", b][10]
            gate(f"SUSB legal payroll closure {sec}/{b} at 1e-6bn", abs(gap) <= 1, gap / 1e6)
    if unpublished:
        csv_out(HERE / "derived/k1_unpublished_rows.csv", unpublished)
    for form in forms:
        gap = sum(lookup[sec, form, "1"][10] for sec in sectors if sec != "--" and (sec, form, "1") in lookup) - lookup["--", form, "1"][10]
        gate(f"SUSB national sector payroll closure form {form} at 1e-6bn", abs(gap) <= 1, gap / 1e6)
    output, probabilities = [], {}
    for sec in sectors:
        if sec == "99":
            # Unclassified industry's all-size payroll closes nationally, but its
            # omitted size cells are never imputed or used as a key profile.
            continue
        def fraction(size_codes):
            return sum(lookup[sec, "2", b][10] for b in size_codes) / sum(lookup[sec, "1", b][10] for b in size_codes)
        # Exact shared bins; no split of source 20-99 or CPS 10-24 is inferred.
        p = np.array([fraction(["1"]), fraction(["2", "3"]), fraction(["4", "6"]), fraction(["7"]), fraction(["9"])])
        probabilities[sec] = p
        gate(f"C corporate payroll probabilities in [0,1] {sec}", bool(np.all((p >= 0) & (p <= 1))))
        for label, codes, prop in zip(["total", "under10", "10-99", "100-499", "500plus"], [["1"], ["2", "3"], ["4", "6"], ["7"], ["9"]], p):
            all_pay = sum(lookup[sec, "1", b][10] for b in codes)
            c_pay = sum(lookup[sec, "2", b][10] for b in codes)
            output.append({"sector": sec, "size": label, "all_payroll_bn": all_pay / 1e6, "c_and_other_corporate_payroll_bn": c_pay / 1e6, "c_and_other_corporate_share": prop})
    csv_out(HERE / "derived/k1_payroll.csv", output)
    wb.close()
    return probabilities, lookup


def parse_irs(lookup):
    wb = openpyxl.load_workbook(HERE / "_cache/k1_22co53ccr.xlsx", read_only=True, data_only=True)
    rows = list(wb.worksheets[0].values)
    gate("IRS table excludes S, REIT and RIC", "OTHER THAN FORMS 1120S, 1120-REIT, AND 1120-RIC" in rows[0][0])
    cols = [(i, x) for i, x in enumerate(rows[4]) if x and i >= 2]
    gate("IRS 19 named sectors in primary header", len(cols) == len(IRS_SECTORS))
    salaries = next(r for r in rows if r[0] == "Salaries and wages")
    officers = next(r for r in rows if r[0] == "Compensation of officers")
    output = []
    for sec, (col, label) in zip(IRS_SECTORS, cols):
        c_payroll = lookup[sec, "2", "1"][10] / 1e6
        wages = salaries[col] / 1e6
        officer = officers[col] / 1e6
        output.append({"sector": sec, "irs_label": label, "susb_c_and_other_payroll_bn": c_payroll, "irs_c_ex_reit_ric_salary_deduction_bn": wages, "irs_c_officer_compensation_bn": officer, "irs_salary_plus_officer_over_susb_payroll": (wages + officer) / c_payroll})
    # National closure compares the table's sector totals with its national field,
    # never requires SOI's salary deduction to equal Census payroll.
    for row, label in [(salaries, "salaries"), (officers, "officers")]:
        gap = sum(row[col] for col, _ in cols) - row[1]
        gate(f"IRS {label} 19-sector closure at 1e-6bn", abs(gap) <= 1, gap / 1e6)
    output.append({"sector": "--", "irs_label": "All industries", "susb_c_and_other_payroll_bn": lookup["--", "2", "1"][10] / 1e6, "irs_c_ex_reit_ric_salary_deduction_bn": salaries[1] / 1e6, "irs_c_officer_compensation_bn": officers[1] / 1e6, "irs_salary_plus_officer_over_susb_payroll": (salaries[1] + officers[1]) / lookup["--", "2", "1"][10]})
    csv_out(HERE / "derived/k1_irs_crosscheck.csv", output)
    wb.close()
    national = output[-1].copy()
    swb = openpyxl.load_workbook(HERE / "_cache/k1_22co61ccr.xlsx", read_only=True, data_only=True)
    srows = list(swb.worksheets[0].values)
    gate("IRS S table is Form1120-S", "FORM 1120-S" in srows[0][0])
    ssalary = next(r for r in srows if r[0] == "Salaries and wages")[1]
    sofficer = next(r for r in srows if r[0] == "Compensation of officers")[1]
    national["irs_s_salary_plus_officer_bn"] = (ssalary + sofficer) / 1e6
    national["susb_s_payroll_bn"] = lookup["--", "3", "1"][10] / 1e6
    national["irs_c_share_of_c_and_s_salary_plus_officer"] = (salaries[1] + officers[1]) / (salaries[1] + officers[1] + ssalary + sofficer)
    national["susb_c_share_of_c_and_s_payroll"] = lookup["--", "2", "1"][10] / (lookup["--", "2", "1"][10] + lookup["--", "3", "1"][10])
    swb.close()
    return national


def coverage_metadata():
    def variables(name):
        return json.loads((HERE / f"_cache/k1_{name}.json").read_text())["variables"]
    lfo = variables("ec2022_lfo_group")
    size = variables("ec2022_firmsize_group")
    absvars = variables("abs2023_variables")
    cbp = variables("cbp2023_variables")
    gate("EC legal-form table has no employment-size dimension", "LFO" in lfo and not any(k.startswith("EMPSZ") for k in lfo))
    gate("EC firm-size table has no legal-form dimension", "EMPSZFF" in size and "LFO" not in size)
    gate("ABS public company summary lacks legal form", "LFO" not in absvars)
    gate("CBP legal-form size is establishment not enterprise", "LFO" in cbp and "establishments" in cbp["EMPSZES"]["label"])


def aggregate_bins(array):
    return np.stack([array[:, 0], array[:, 1], array[:, 2:4].sum(axis=1), array[:, 4], array[:, 5:7].sum(axis=1)], axis=1)


def se(v):
    return float(np.sqrt(4 / 160 * np.square(v[1:] - v[0]).sum()))


def calculate(probabilities):
    # Never create interpreter caches in the frozen source lane.
    import sys
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location("frozen_job_capital_k1", FROZEN / "job_capital.py")
    source = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(source)
    z = np.load(HERE / "_cache/core.npz")
    core = json.loads((HERE / "derived/core.json").read_text())
    cells = z["cells"].astype(str)
    pools = [np.fromstring(p, sep=",", dtype=int) for p in z["cell_pool"]]
    nat = aggregate_bins(z["size_wages_national"])
    gate("exact common-bin wage conservation", np.allclose(nat.sum(axis=1), z["size_wages_national"].sum(axis=1), rtol=1e-12, atol=1e-3))
    # Primary SUSB methodology excludes crop/animal production and railroads.
    # Do not extrapolate support-service or nonrail payroll profiles into them.
    covered = np.array([source.SECTOR[c] in SECTOR_CODE and c not in ["111CA", "482"] for c in cells])
    profiles = np.array([probabilities[SECTOR_CODE[source.SECTOR[c]]] if is_covered else np.ones(5) for c, is_covered in zip(cells, covered)])
    unsupported = [{"cell": c, "national_corporate_tax_bn": float(t), "treatment": "held at frozen all-employer key; no matching SUSB sector"} for c, t, is_covered in zip(cells, z["tax_corporate"], covered) if t > 0 and not is_covered]
    gate("uncovered corporate cells are exactly government, crop/animal and rail cells", {r["cell"] for r in unsupported} == {"GOVF", "GOVSL", "111CA", "482"})
    csv_out(HERE / "derived/k1_uncovered_cells.csv", unsupported)
    unknown = []
    gate("NOEMP0 contains no national wage dollars", float(nat[:, 0, :].sum()) == 0)
    targets, results = {}, {}
    for g in ["union", "G1", "G2", "G3plus", "white"]:
        grp = aggregate_bins(z[f"size_wages_{g}"])
        unknown.append({"group": g, "NOEMP0_wages_bn": float(grp[:, 0, 0].sum() / 1e9), "all_wages_bn": float(grp[:, :, 0].sum() / 1e9), "NOEMP0_share": float(grp[:, 0, 0].sum() / grp[:, :, 0].sum())})
        baseline = z["tax_corporate"] @ z[f"shares_{g}"]
        for variant, prop in [("neutral", np.ones_like(profiles)), ("size_only", np.tile(probabilities["--"], (len(cells), 1))), ("industry_size", profiles), ("industry_size_selection_min", profiles), ("industry_size_selection_max", profiles)]:
            prop = prop.copy()
            prop[~covered] = 1
            adjusted = np.zeros_like(z[f"shares_{g}"])
            for i, idx in enumerate(pools):
                n = nat[idx].sum(axis=0)
                w = grp[idx].sum(axis=0)
                raw = np.divide(w.sum(axis=0), n.sum(axis=0), out=np.zeros(161), where=n.sum(axis=0) > 0)
                # Pools within a fine sector carry its public broad-sector profile;
                # the one all-private pool (management) gets per-cell profiles.
                if variant == "industry_size_selection_min":
                    numerator = np.maximum(0, nat[idx] * prop[idx, :, None] - (nat[idx] - grp[idx])).sum(axis=(0, 1))
                elif variant == "industry_size_selection_max":
                    numerator = np.minimum(grp[idx], nat[idx] * prop[idx, :, None]).sum(axis=(0, 1))
                else:
                    numerator = (grp[idx] * prop[idx, :, None]).sum(axis=(0, 1))
                denominator = (nat[idx] * prop[idx, :, None]).sum(axis=(0, 1))
                legal_share = np.divide(numerator, denominator, out=np.zeros(161), where=denominator > 0)
                factor = np.divide(legal_share, raw, out=np.ones(161), where=raw > 0)
                adjusted[i] = z[f"shares_{g}"][i] * factor
            target = z["tax_corporate"] @ adjusted
            if variant == "neutral":
                gate(f"neutral legal probabilities recover frozen pooled corporate {g} in all161 columns", np.max(np.abs(target - baseline)) < 1e-9, float(np.max(np.abs(target - baseline))))
            else:
                targets[f"{variant}_{g}"] = target
                # Zero mining remains inside unchanged tax vector; national amounts
                # remain frozen, legal form redistributes only corporate receipts.
                targets[f"{variant}_cell_shares_{g}"] = adjusted
    for g in ["union", "G1", "G2", "G3plus", "white"]:
        lo = targets[f"industry_size_selection_min_{g}"]
        point = targets[f"industry_size_{g}"]
        hi = targets[f"industry_size_selection_max_{g}"]
        gate(f"selection envelopes contain exchangeable proxy {g} all161 columns", bool(np.all(lo <= point + 1e-9) and np.all(point <= hi + 1e-9)))
    for variant in ["size_only", "industry_size"]:
        gap = sum(targets[f"{variant}_{g}"] for g in ["G1", "G2", "G3plus"]) - targets[f"{variant}_union"]
        gate(f"corporate generation conservation {variant} all161 columns at1e-9bn", float(np.max(np.abs(gap))) < 1e-9, float(np.max(np.abs(gap))))
    csv_out(HERE / "derived/k1_unknown_size.csv", unknown)
    for variant in ["size_only", "industry_size", "industry_size_selection_min", "industry_size_selection_max"]:
        base_union = z["tax_corporate"] @ z["shares_union"]
        base_lineage = base_union + sum(core["lineage_rates"][g] * (z["tax_corporate"] @ z[f"shares_{g}"]) for g in ["G3plus", "white"])
        new_union = targets[f"{variant}_union"]
        new_lineage = new_union + sum(core["lineage_rates"][g] * targets[f"{variant}_{g}"] for g in ["G3plus", "white"])
        results[variant] = {}
        for end in ["low", "high"]:
            response = core["responses"][end]["corporate"]
            results[variant][end] = {}
            for pop, new, old in [("union", new_union, base_union), ("lineage", new_lineage, base_lineage)]:
                delta = -(new - old) * response
                results[variant][end][pop] = {"corporate_target_bn": float(new[0]), "delta_net_cost_bn": float(delta[0]), "delta_conditional_se_bn": se(delta), "item_net_cost_bn": -core["control"][end][f"{pop}_bn"] + float(delta[0]), "over_5bn_review_bar": bool(abs(delta[0]) >= 5)}
    np.savez_compressed(HERE / "_cache/k1_targets.npz", **targets)
    return results


def main():
    start = time.perf_counter()
    parser = argparse.ArgumentParser()
    parser.add_argument("--acquire", action="store_true")
    parser.add_argument("--sources-only", action="store_true")
    args = parser.parse_args()
    pins = acquire(args.acquire)
    probabilities, lookup = parse_susb()
    crosscheck = parse_irs(lookup)
    coverage_metadata()
    results = {} if args.sources_only else calculate(probabilities)
    dump(HERE / "derived/k1_summary.json", {"status": "conditional sector-enterprise-size C-and-other-corporate payroll proxy", "source_year": 2022, "results": results, "irs_national_crosscheck": crosscheck, "gates": GATES, "source_pins": pins, "assumptions": ["Conditional C-and-other-corporate payroll propensity of the group's wages equals the national propensity within industry sector and matched enterprise size.", "2022 SUSB propensity applies to 2024 wages.", "Fine industries share their broad sector legal-form propensity.", "NOEMP0 uses the source sector's all-size corporate payroll share, but this bin carries zero wage dollars in core.npz.", "Common core's frozen thin-cell pools and top-wage spread retained.", "Lineage added members inherit common core G3plus/white rates; ancestry-count uncertainty excluded.", "Government enterprise, crop/animal production, and rail cells held at the frozen key because SUSB cannot match them.", "SUSB excludes some funds inside merged523 and notaries inside5411: partial scope profile remains conditional.", "Selection_min/max remove conditional exchangeability only; retaining transported national C payroll marginals. Lineage bounds are relaxed outer envelopes rather than sharp jointly feasible bounds."], "irs_limit": "SOI salary deductions exclude cost-of-goods-sold labor and officer compensation; officers separately shown. SOI industry is consolidated return, Census establishment. SOI excludes REIT/RIC; Census category includes other corporate forms. No payroll-level parity claimed."})
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024**3
    gate("peak resident memory under2.5GB", peak < 2.5, peak)
    print(json.dumps({"elapsed_seconds": time.perf_counter() - start, "peak_rss_gb": peak, "gates": len(GATES), "results": results}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
