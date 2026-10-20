"""Unfunded public-employee pension legacy by group: stock entering 2024 and 2024 interest on it.

State-local and federal defined-benefit plans owe benefits already earned by past service that their
assets do not cover. BEA books the interest accruing on that shortfall inside government interest
payments; the main case holds that interest row (Table 3.1 line 28, $1,118.87bn) at zero response.
The federal legacy model capitalizes NIPA gaps, so its pension financing overlap is unreconciled.
This lane measures the shortfall and its 2024 interest from pinned primary files and attributes it
to equal-sized groups by their use of public services, retaining engine and matched union keys.

Run from the repository root (a worktree drops --no-project):
  uv run --no-project python3 infra/immigration-fiscal/pension_legacy_2026_09_30/pension_legacy.py
  uv run --no-project python3 infra/immigration-fiscal/pension_legacy_2026_09_30/pension_legacy.py --fetch
  uv run --no-project python3 infra/immigration-fiscal/pension_legacy_2026_09_30/pension_legacy.py --case oct16
The live default is oct16, after group_lines.py --case oct16; writes derived/oct16/.
Oct16 uses population_v9, separates federal VA service from civilian/state health,
and routes college operating changes separately from capital. The booked national
plan stocks and interest retain their imported valuations; item20's annual real
capital return does not revalue either. service_rules.json records the source
receipts and VA stock/annual-interest reductions in distinct units.
Earlier entries in CASES locate archived inputs and outputs, which the live run never regenerates.
`--fetch` downloads every input to `_cache/` (ignored; ASPEP needs CENSUS_API_KEY in the environment,
never written) and then runs. A default run reads `_cache/` and stops unless every file matches its pin.

Outputs (derived/, LF line endings):
  measured_2024.csv      every measured input with its source cell
  domestic_interest.csv  what the $1,118.87bn row holds: pension interest and the rest
  function_mix.csv       each plan's payroll split over the account's spending lines
  group_shares.csv       each group's share of each line and the main case's responses
  headcount_path.csv     the Mexican-origin population-share path and the service-year weights
  attribution.csv        every arm: group x plan x basis x time weighting x responses x end
  by_line.csv            the adopted arm's interest by plan and spending line
  summary.csv            the adopted arm and its companions, with the Mexican-origin minus white difference
  opeb_federal.csv       federal retiree health (beside; not in the NIPA interest row)
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import importlib.util
import json
import math
import os
import re
import subprocess
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
FISCAL = HERE.parent
_contract_spec = importlib.util.spec_from_file_location("pension_comparison_contract",
    FISCAL / "white_replacement_2026_09_28/comparison_contract.py")
COMPARISON = importlib.util.module_from_spec(_contract_spec)
_contract_spec.loader.exec_module(COMPARISON)
CACHE = HERE / "_cache"
UNION_ROW4 = 39_712_493.33          # dataset audit row 4: the people the account prices (per-member denominator)
UNION_PER_HEAD_KEY = 0.117175
RESIDENTS = 340_110_988              # the account's 2024 resident control
# The case the groups' lines come from (--case). oct05, the v5 case adopted 2026-10-05: group_lines.py --case oct05,
# whose union rows carry the 3,039,720 added people at the case lane's amounts and whose slices are on the lineage's
# 42,752,213; per member divides by that count, and the per-head key is the case's (0.117175 x the count's ratio).
# The headcount path stays the union's: the added people's population share is taken to move as the union's
# [ASSUMPTION; the back-cast instead carries them on the identified third-plus path].
# oct07, main case v6 of 2026-10-07: oct05 plus four items on the same 42,752,213. Its retiree-health item scales ten
# lines' national totals (engine scaleLine), three of them under a state-price overlay; a group's share of such a line
# keeps the overlay over the line's unscaled national, as if the item scaled the overlay with the line (it scales every
# cost on the line) [ASSUMPTION; dividing by the scaled national instead moves a share by at most 5.2e-5 and any
# group's or difference's interest by at most $0.0044bn in any arm, $0.0036bn in the adopted arm].
CASES = {"sept29": dict(lines="group_lines_sept29.csv", dump="engine_lines_sept29.json"),
         "oct05": dict(lines="group_lines_oct05.csv", dump="engine_lines_oct05.json",
                       lineage="main_case_2026_10_05/derived/corrections.json"),
         "oct07": dict(lines="group_lines_oct07.csv", dump="engine_lines_oct07.json",
                       lineage="main_case_2026_10_07/derived/corrections.json"),
         "oct15": dict(lines="group_lines_oct15.csv", dump="engine_lines_oct15.json",
                       lineage="main_case_2026_10_15/derived/corrections.json"),
         "oct16": dict(lines="group_lines_oct16.csv", dump="engine_lines_oct16.json",
                       lineage="main_case_2026_10_16/derived/corrections.json")}
CASE = "sept29"
POPULATION = UNION_ROW4              # the case's priced count (use_case())
SCALED = {}                          # line -> national total after the case's items' national-scale edits (use_case())
COMPLETION_CARRIERS = {}
COVERAGE = None
SERVICE_ATTRIBUTIONS = []
POPULATION_V8 = {}
COMPLETION_COUNTS = {}
DEFENSE_STOCK_BRIDGES = []
V9_META = {}
V9_SERVICE_PARTS = []
V9_SERVICE_RUNS = {}

PINS = {
    "Section3All_xls.xlsx": "69b5c7aefb38675324887ce31d6feb4fcde7c903ab952db7328da0813096615e",
    "Section6All_xls.xlsx": None,    # filled below from the fetched file; see PINS_FILE
    "Section7All_xls.xlsx": "ce107c8ce92393613c0abae38156afcfaef8ab571eedf0302dc09ca44738b9ef",
}
SOURCES = {
    "Section3All_xls.xlsx": "https://apps.bea.gov/national/Release/XLS/Survey/Section3All_xls.xlsx",
    "Section6All_xls.xlsx": "https://apps.bea.gov/national/Release/XLS/Survey/Section6All_xls.xlsx",
    "Section7All_xls.xlsx": "https://apps.bea.gov/national/Release/XLS/Survey/Section7All_xls.xlsx",
    "z1_csv_files.zip": "https://www.federalreserve.gov/releases/z1/current/z1_csv_files.zip",
    "c2010br-04.pdf": "https://www.census.gov/content/dam/Census/library/publications/2011/dec/c2010br-04.pdf",
    "c2kbr01-3.pdf": "https://www2.census.gov/library/publications/decennial/2000/briefs/c2kbr01-03.pdf",
    "we-02r.pdf": "https://www2.census.gov/library/publications/decennial/1990/we-the-americans/we-02r.pdf",
    "fr2025_note13.pdf": "https://fiscal.treasury.gov/system/files/2026-04/2025-note-13-financial-statements.pdf",
}
PINS_FILE = HERE / "source_pins.json"   # tracked: sha256 and bytes of every cached input

ASPEP_URL = "https://api.census.gov/data/timeseries/govsemp"


# ----------------------------------------------------------------------------------------------- fetch
def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fetch(env_file: Path | None = None) -> None:
    CACHE.mkdir(exist_ok=True)
    agent = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/120 Safari/537.36"}
    for name, url in SOURCES.items():
        if (CACHE / name).exists():          # keep a cached vintage; delete the file to refresh it
            continue
        body = urllib.request.urlopen(urllib.request.Request(url, headers=agent), timeout=180).read()
        if name.endswith(".pdf") and not body.startswith(b"%PDF"):
            raise ValueError(f"[BLOCKED] {url} did not return a PDF")
        (CACHE / name).write_bytes(body)
    if (CACHE / "aspep_2024.json").exists():
        write_pins()
        return
    key = os.environ.get("CENSUS_API_KEY")
    if not key and env_file:              # acquire/config.local.env, KEY=VALUE lines; the key is never printed
        found = re.search(r"^\s*(?:export\s+)?CENSUS_API_KEY=['\"]?([^'\"\s]+)", env_file.read_text(), re.M)
        key = found.group(1) if found else None
    if not key:
        raise ValueError("[BLOCKED] ASPEP needs CENSUS_API_KEY in the environment")
    q = {"get": "AGG_DESC,AGG_DESC_LABEL,GOVTYPE,GOVTYPE_LABEL,TOT_PAY,FTE", "for": "us:*", "time": "2024", "key": key}
    try:
        rows = json.load(urllib.request.urlopen(ASPEP_URL + "?" + urllib.parse.urlencode(q), timeout=120))
    except Exception as e:                                   # the URL carries the key: never print it
        raise ValueError("[BLOCKED] ASPEP request failed: " + str(e).replace(key, "<KEY>")) from None
    (CACHE / "aspep_2024.json").write_text(json.dumps(rows, indent=0) + "\n")
    write_pins()


def write_pins() -> None:
    pins = {p.name: {"sha256": sha(p), "bytes": p.stat().st_size}
            for p in sorted(CACHE.iterdir()) if p.is_file() and p.suffix in {".xlsx", ".zip", ".pdf", ".json"}}
    PINS_FILE.write_text(json.dumps(pins, indent=1, sort_keys=True) + "\n")


def pinned(name: str) -> Path:
    path = CACHE / name
    pins = json.loads(PINS_FILE.read_text())
    if not path.exists():
        raise ValueError(f"[BLOCKED] missing {path}; run with --fetch")
    if sha(path) != pins[name]["sha256"] or (PINS.get(name) and PINS[name] != pins[name]["sha256"]):
        raise ValueError(f"[BLOCKED] {name} is not the pinned vintage")
    return path


# ----------------------------------------------------------------------------------------------- BEA
def bea_table(book, sheet: str, year: int = 2024, units: str = "[Millions of dollars]") -> tuple[dict[int, dict], list[str]]:
    rows = list(book[sheet].values)
    headers = [r for r in rows if r[0] == "Line"]
    if len(headers) != 1 or (units and rows[1][0] != units):
        raise ValueError(f"[BLOCKED] {sheet}: header or units changed")
    cols = [i for i, x in enumerate(headers[0]) if str(x) == str(year)]
    if len(cols) != 1:
        raise ValueError(f"[BLOCKED] {sheet}: no single {year} column")
    j = cols[0]
    lines = {int(r[0]): dict(label=str(r[1]).strip(), code=r[2], value=r[j]) for r in rows if str(r[0]).isdigit()}
    notes = [str(r[0]) for r in rows if r[0] and not str(r[0]).isdigit() and r[0] != "Line"]
    return lines, notes


def bea_series(book, sheet: str, number: int) -> dict[int, float]:
    rows = list(book[sheet].values)
    header = [r for r in rows if r[0] == "Line"][0]
    row = [r for r in rows if str(r[0]) == str(number)][0]
    return {int(float(y)): float(v) for y, v in zip(header[3:], row[3:]) if isinstance(v, (int, float))}


def cell(lines: dict, number: int, label: str) -> float:
    got = lines[number]
    if not got["label"].replace("\\", "").startswith(label):
        raise ValueError(f"[BLOCKED] line {number} is {got['label']!r}, expected {label!r}")
    return got["value"] / 1000.0      # $bn


def measure_bea() -> tuple[dict, dict]:
    s3 = openpyxl.load_workbook(pinned("Section3All_xls.xlsx"), read_only=True, data_only=True)
    s7 = openpyxl.load_workbook(pinned("Section7All_xls.xlsx"), read_only=True, data_only=True)
    t31, n31 = bea_table(s3, "T30100-A")
    t32, n32 = bea_table(s3, "T30200-A")
    t33, n33 = bea_table(s3, "T30300-A")
    t3105, _ = bea_table(s3, "T31005-A")
    t3115, _ = bea_table(s3, "T31105-A")
    t317, _ = bea_table(s3, "T31700-A")
    t723, n723 = bea_table(s7, "T72300-A")
    t724, n724 = bea_table(s7, "T72400-A")
    t78, _ = bea_table(s7, "T70800-A")
    m = {
        "domestic_interest_total": (cell(t31, 28, "To persons and business"), "BEA T3.1 line 28"),
        "domestic_interest_federal": (cell(t32, 34, "To persons and business"), "BEA T3.2 line 34"),
        "domestic_interest_state_local": (cell(t33, 29, "To persons and business"), "BEA T3.3 line 29"),
        "sl_actual_employer_contributions": (cell(t724, 5, "Actual employer contributions"), "BEA T7.24 line 5"),
        "sl_imputed_employer_contributions": (cell(t724, 6, "Imputed employer contributions"), "BEA T7.24 line 6"),
        "sl_plan_monetary_interest": (cell(t724, 12, "Monetary interest"), "BEA T7.24 line 12"),
        "sl_imputed_interest": (cell(t724, 13, "Imputed interest on plans' claims on employers"), "BEA T7.24 line 13"),
        "sl_plan_dividends": (cell(t724, 14, "Dividends"), "BEA T7.24 line 14"),
        "sl_interest_accrued_on_entitlements": (cell(t724, 31, "Interest accrued on benefit entitlements"), "BEA T7.24 line 31"),
        "fed_actual_employer_contributions": (cell(t723, 5, "Actual employer contributions"), "BEA T7.23 line 5"),
        "fed_imputed_employer_contributions": (cell(t723, 8, "Imputed employer contributions"), "BEA T7.23 line 8"),
        "fed_plan_interest_total": (cell(t723, 17, "Interest"), "BEA T7.23 line 17"),
        "fed_plan_monetary_interest": (cell(t723, 18, "Monetary interest"), "BEA T7.23 line 18"),
        "fed_imputed_interest": (cell(t723, 19, "Imputed interest on plans' claims on employers"), "BEA T7.23 line 19"),
        "fed_interest_accrued_on_entitlements": (cell(t723, 41, "Interest accrued on benefit entitlements"), "BEA T7.23 line 41"),
        "supp_federal_civilian_pension": (cell(t78, 6, "Federal civilian pension plans"), "BEA T7.8 line 6"),
        "supp_federal_military_pension": (cell(t78, 7, "Federal military pension plans"), "BEA T7.8 line 7"),
        "supp_state_local_retirement": (cell(t78, 10, "State and local employee retirement"), "BEA T7.8 line 10"),
        "defense_comp_military": (cell(t3115, 6, "Military"), "BEA T3.11.5 line 6"),
        "defense_comp_civilian": (cell(t3115, 7, "Civilian"), "BEA T3.11.5 line 7"),
        "nondefense_comp": (cell(t3105, 37, "Compensation of general government employees"), "BEA T3.10.5 line 37"),
        "sl_comp": (cell(t3105, 50, "Compensation of general government employees"), "BEA T3.10.5 line 50"),
    }
    fed_fn = {"general_public_services": (12, "General public service"), "public_order_safety": (14, "Public order and safety"),
              "economic_affairs_services": (15, "Economic affairs"), "housing_community_services": (16, "Housing and community services"),
              "health_services": (17, "Health"), "recreation_culture": (18, "Recreation and culture"),
              "education_services": (19, "Education"), "income_security_services": (20, "Income security")}
    for line, (n, label) in fed_fn.items():
        m[f"fed_consumption_{line}"] = (cell(t317, n, label), f"BEA T3.17 line {n}")
    notes = {"T3.1": n31, "T3.2": n32, "T3.3": n33, "T7.23": n723, "T7.24": n724}
    population = bea_series(s7, "T70100-A", [k for k, v in bea_table(s7, "T70100-A", units="")[0].items()
                                             if v["label"].startswith("Population (midperiod")][0])
    s6 = openpyxl.load_workbook(pinned("Section6All_xls.xlsx"), read_only=True, data_only=True)
    t65, _ = bea_table(s6, "T60500D-A", units="")
    fte = {y: (bea_series(s6, "T60500D-A", 93)[y] + bea_series(s6, "T60500D-A", 96)[y]) for y in (1998, 2024)}
    if not (t65[93]["label"].startswith("General government") and t65[96]["label"].startswith("Government enterprises")):
        raise ValueError("[BLOCKED] T6.5D state-local lines moved")
    return m, dict(notes=notes, population=population, sl_fte=fte)


# ----------------------------------------------------------------------------------------------- Z.1
def z1_q4(archive: zipfile.ZipFile, table: str, series: list[str]) -> dict[int, dict[str, float]]:
    text = archive.read(f"csv/{table}.csv").decode()
    out = {}
    for row in csv.DictReader(io.StringIO(text)):
        if row["date"].endswith("Q4"):
            out[int(row["date"][:4])] = {s: float(row[s]) / 1000.0 for s in series}   # $bn
    return out


def measure_z1() -> dict[int, dict[str, float]]:
    with zipfile.ZipFile(pinned("z1_csv_files.zip")) as z:
        sl = z1_q4(z, "S129s1_3_s", ["FL223073045.Q", "FL224190043.Q", "FL224090045.Q"])
        fe = z1_q4(z, "S129s1_2_s", ["FL343073045.Q", "FL344190045.Q", "FL344090045.Q", "FL343069245.Q"])
        dictionary = z.read("data_dictionary/S129s1_2_s.txt").decode() + z.read("data_dictionary/S129s1_3_s.txt").decode()
    for code, words in [("FL223073045", "State and local government employee defined benefit pension funds; claims of pension fund on sponsor"),
                        ("FL343073045", "Federal government defined benefit pension funds; claims of pension funds on sponsor"),
                        ("FL344190045", "Federal government defined benefit pension funds; pension entitlements (total liabilities)"),
                        ("FL224190043", "State and local government employee defined benefit pension funds; pension entitlements (total liabilities)"),
                        ("FL343069245", "Federal government defined benefit pension funds; nonmarketable Treasury securities"),
                        ("FL344090045", "Federal government defined benefit pension funds; total financial assets")]:
        line = [l for l in dictionary.splitlines() if l.startswith(code)]
        if not line or words not in line[0]:
            raise ValueError(f"[BLOCKED] Z.1 series {code} label changed")
    return {y: {**sl.get(y, {}), **fe.get(y, {})} for y in sorted(set(sl) | set(fe))}


# ----------------------------------------------------------------------------------------------- text sources
def pdf_text(name: str) -> str:
    return subprocess.run(["pdftotext", "-layout", str(pinned(name)), "-"], capture_output=True, text=True, check=True).stdout


def num(s: str) -> float:
    return float(s.replace(",", "").replace("(", "-").replace(")", ""))


def measure_fr() -> dict:
    """FY 2025 Financial Report Note 13, FY 2024 columns (Civilian 2025, 2024, Military 2025, 2024, Total 2025, 2024)."""
    t = pdf_text("fr2025_note13.pdf")
    def row(pattern: str, start: int = 0) -> tuple[list[float], int]:
        m = re.compile(pattern + r"\s+" + r"\s+".join([r"\(?([\d,]+\.\d)\)?"] * 6)).search(t, start)
        if not m:
            raise ValueError(f"[BLOCKED] Note 13 row {pattern!r} not found")
        return [num(g) for g in m.groups()], m.end()
    pension, _ = row(r"Pension benefits")
    opeb, _ = row(r"Post-retirement health benefits")
    begin, pos = row(r"Actuarial accrued pension liability, beginning of\s+fiscal year")
    interest, _ = row(r"Interest on liability", pos)
    obegin, opos = row(r"Actuarial accrued post-retirement health benefits\s+liability, beginning of fiscal year")
    ointerest, _ = row(r"Interest on liability", opos)
    return dict(pension_end_fy24=dict(civilian=pension[1], military=pension[3], total=pension[5]),
                pension_begin_fy24=dict(civilian=begin[1], military=begin[3], total=begin[5]),
                pension_interest_fy24=dict(civilian=interest[1], military=interest[3], total=interest[5]),
                opeb_end_fy24=dict(civilian=opeb[1], military=opeb[3], total=opeb[5]),
                opeb_begin_fy24=dict(civilian=obegin[1], military=obegin[3], total=obegin[5]),
                opeb_interest_fy24=dict(civilian=ointerest[1], military=ointerest[3], total=ointerest[5]))


def measure_census_counts() -> dict[int, float]:
    """Mexican-origin counts at the decennial censuses, millions. 2000 from the 2010 brief's table 1; 1990 from the
    2000 brief's growth rate; 1980 and 1970 from the 1993 brief's growth chart (1980-90 and 1970-80)."""
    t10 = pdf_text("c2010br-04.pdf")
    m = re.search(r"\nMexican[ .]+([\d,]{9,})\s+[\d .]+?\s+([\d,]{9,})", t10)
    c2000, c2010 = num(m.group(1)), num(m.group(2))
    t00 = pdf_text("c2kbr01-3.pdf")
    g90 = float(re.search(r"Mexicans increased by (\d+\.\d) percent", t00).group(1))
    if not re.search(r"from 13\.5 million to 20\.6 million", t00):
        raise ValueError("[BLOCKED] 2000 brief: 1990 Mexican count text changed")
    t93 = pdf_text("we-02r.pdf")
    block = t93[t93.index("The Mexican population nearly doubled"):t93.index("Puerto Rican", t93.index("The Mexican population nearly doubled"))]
    if "54.4" not in t93[t93.index("The Mexican population nearly doubled") - 400:t93.index("The Mexican population nearly doubled")] or "92.8" not in block:
        raise ValueError("[BLOCKED] 1993 brief: Mexican growth chart values moved")
    c1990 = c2000 / (1 + g90 / 100)
    c1980 = c1990 / 1.544      # WE-2R chart: Mexican growth 1980-90, 54.4 percent
    c1970 = c1980 / 1.928      # WE-2R chart: 1970-80, 92.8 percent ("nearly doubled")
    if abs(c1990 / 1e6 - 13.5) > 0.05:
        raise ValueError("[BLOCKED] derived 1990 count does not round to the brief's 13.5 million")
    return {1970: c1970 / 1e6, 1980: c1980 / 1e6, 1990: c1990 / 1e6, 2000: c2000 / 1e6, 2010: c2010 / 1e6}


def measure_aspep() -> dict[str, float]:
    rows = json.loads(pinned("aspep_2024.json").read_text())
    head = rows[0]
    out = {}
    for r in rows[1:]:
        d = dict(zip(head, r))
        if d["GOVTYPE"] == "001" and d["time"] == "2024":
            out[d["AGG_DESC_LABEL"]] = float(d["TOT_PAY"]) / 1e9      # March 2024 monthly payroll, $bn
    return out


# ----------------------------------------------------------------------------------------------- mapping
# ASPEP function -> the account's spending line. Enterprises (utilities, transit, liquor stores) are priced to users;
# they take the per-head key of the enterprise surplus line, as the case keys that line.
ASPEP_MAP = {
    "Financial Administration": "general_public_services",
    "Other Government Administration": "general_public_services",
    "All other and unallocable": "general_public_services",
    "Judicial and Legal": "public_order_safety",
    "Police Protection Total": "public_order_safety",
    "Fire Protection Total": "public_order_safety",
    "Corrections": "public_order_safety",
    "Highways": "highways",
    "Air Transportation": "economic_affairs_services",
    "Sea and Inland Port Facilities": "economic_affairs_services",
    "Natural Resources": "economic_affairs_services",
    "Public Welfare": "income_security_services",
    "Social Insurance Administration": "income_security_services",
    "Health": "health_services",
    "Hospitals": "health_services",
    "Solid Waste Management": "housing_community_services",
    "Sewerage": "housing_community_services",
    "Housing and Community Development": "housing_community_services",
    "Parks and Recreation": "recreation_culture",
    "Education Total": "education_services",
    "Libraries": "education_services",          # BEA puts libraries in education (T3.16 lines 33-34)
    "Water Supply": "enterprises",
    "Electric Power": "enterprises",
    "Gas Supply": "enterprises",
    "Transit": "enterprises",
    "State liquor stores": "enterprises",
}
ASPEP_PARTS = {"Police Protection - Persons with Power of Arrest", "Police Protection - Other", "Fire Protection - Firefighters",
               "Fire Protection - Other", "Education - Elementary and Secondary Total", "Education - Elementary and Secondary Instructional",
               "Education - Elementary and Secondary Other", "Education - Higher Education Total", "Education - Higher Education Instructional",
               "Education - Higher Education Other", "Education - Other", "Total - All Government Employment Functions"}
LINES = ["general_public_services", "defense", "public_order_safety", "economic_affairs_services", "highways",
         "housing_community_services", "health_services", "recreation_culture", "education_services",
         "income_security_services", "enterprises"]


def function_mixes(aspep: dict, m: dict) -> dict[str, dict[str, float]]:
    extra = set(aspep) - set(ASPEP_MAP) - ASPEP_PARTS
    if extra:
        raise ValueError(f"[BLOCKED] unmapped ASPEP functions: {sorted(extra)}")
    total = aspep["Total - All Government Employment Functions"]
    mapped = sum(aspep[k] for k in ASPEP_MAP)
    if abs(mapped - total) > 1e-6 * total:
        raise ValueError(f"[BLOCKED] ASPEP functions sum to {mapped} against total {total}")
    sl = {l: 0.0 for l in LINES}
    for k, line in ASPEP_MAP.items():
        sl[line] += aspep[k] / total
    # Federal civilian: defense civilians by T3.11.5; nondefense compensation spread by nondefense consumption by function.
    civ_def = m["defense_comp_civilian"][0]
    nondef = m["nondefense_comp"][0]
    fed_fn = {k.replace("fed_consumption_", ""): v for k, (v, _) in m.items() if k.startswith("fed_consumption_")}
    fn_total = sum(fed_fn.values())
    civ = {l: 0.0 for l in LINES}
    civ["defense"] = civ_def / (civ_def + nondef)
    for line, v in fed_fn.items():
        civ[line] += nondef / (civ_def + nondef) * v / fn_total
    mil = {l: 0.0 for l in LINES}
    mil["defense"] = 1.0
    return {"state_local": sl, "federal_civilian": civ, "federal_military": mil}


# ----------------------------------------------------------------------------------------------- group shares
GROUP_SOURCES = {"mexican_origin": "mexican_origin_engine", "mexican_origin_rough": "mexican_origin_rough",
                 "A1_third_plus_nh_white": "A1_third_plus_nh_white", "all_residents_slice": "all_residents_slice"}
GROUPS = list(GROUP_SOURCES)
GROUP_LINES = FISCAL / "legacy_comparators_2026_09_30/derived/group_lines_sept29.csv"
DUMP_FILE = FISCAL / "white_replacement_2026_09_28/derived/engine_lines_sept29.json"


def use_case(case: str) -> None:
    """Point the group shares at a case's line export, engine dump, priced count and items' national-scale edits."""
    global CASE, POPULATION, GROUP_LINES, DUMP_FILE, SCALED, COMPLETION_CARRIERS, COVERAGE, SERVICE_ATTRIBUTIONS
    global POPULATION_V8, COMPLETION_COUNTS, DEFENSE_STOCK_BRIDGES, V9_META, V9_SERVICE_PARTS, V9_SERVICE_RUNS
    conf = CASES[case]
    CASE = case
    GROUP_LINES = FISCAL / "legacy_comparators_2026_09_30/derived" / conf["lines"]
    DUMP_FILE = FISCAL / "white_replacement_2026_09_28/derived" / conf["dump"]
    POPULATION, SCALED, COMPLETION_CARRIERS, COVERAGE = UNION_ROW4, {}, {}, None
    SERVICE_ATTRIBUTIONS = []
    POPULATION_V8, COMPLETION_COUNTS, DEFENSE_STOCK_BRIDGES = {}, {}, []
    V9_META, V9_SERVICE_PARTS = {}, []
    V9_SERVICE_RUNS = {}
    if "lineage" in conf:
        payload = json.loads((FISCAL / conf["lineage"]).read_text())
        POPULATION = payload["meta"]["lineage"]["counts"]["lineage_population"]
        if case in ("oct15", "oct16"):
            population_key = "population_v9" if case == "oct16" else "population_v8"
            POPULATION = payload["meta"][population_key]["with_coverage"]
            COMPLETION_CARRIERS = payload["meta"]["union_completion"]["carriers"]
            COVERAGE = payload["meta"].get("coverage")
            POPULATION_V8 = payload["meta"]["population_v8"]
            COMPLETION_COUNTS = payload["meta"]["union_completion"]["counts"]
            if not math.isfinite(POPULATION) or POPULATION <= 0:
                raise ValueError(f"[BLOCKED] {case} {population_key}.with_coverage is not a positive finite count")
            if case == "oct16":
                V9_META = payload["meta"]
        for item in payload["meta"].get("items", []):
            if item["kind"] != "edit_set" or not item["applied"]:
                continue
            first = item["edits"]["first"]
            for e in payload["edits"][first:first + item["edits"]["count"]]:
                if "national_bn" in e:
                    if e["line"] in SCALED or "by" in e:
                        raise ValueError(f"[BLOCKED] {item['id']} edit on {e['line']} is not one national-scale edit")
                    SCALED[e["line"]] = e["national_bn"]
OVERLAY = {"public_order_safety": ("state_price_public_order_safety", 519.153),
           "health_services": ("state_price_health_services", 306.539),
           "recreation_culture": ("state_price_recreation_culture", 54.331)}
HIGHWAYS_NATIONAL = 201.005        # S&L highways inside economic affairs, the national of the case's road key


def service_carrier_amount(parent, dump_lines, amounts, carriers):
    """Use only the selected parent key; inactive carriers can also have response zero."""
    selected = dump_lines[parent]
    parent_side = "receipt" if selected["side"] == "receipts" else selected["side"]
    parent_key = selected.get("scenario") if parent_side == "receipt" else selected.get("key")
    candidates = [source for source in carriers.values() if source.get("kind") == "line"
                  and source.get("side") == parent_side and source.get("parent_line") == parent]
    if candidates and (not isinstance(parent_key, str) or parent_key not in {s.get("parent_key") for s in candidates}):
        raise ValueError(f"[BLOCKED] selected service parent key/scenario is absent from the declared carrier map: {parent}")
    total = 0.0
    for carrier, source in carriers.items():
        if (source.get("kind") != "line" or source.get("side") != parent_side
                or source.get("parent_line") != parent or source.get("parent_key") != parent_key):
            continue
        if carrier not in dump_lines or carrier not in amounts:
            raise ValueError(f"[BLOCKED] missing selected service carrier {carrier}")
        line = dump_lines[carrier]
        if line["national_bn"] != 0 or abs(line["response"] - selected["response"]) > 1e-9:
            raise ValueError(f"[BLOCKED] service carrier {carrier} differs from its selected parent response")
        amount = amounts[carrier]
        if not math.isfinite(amount):
            raise ValueError(f"[BLOCKED] nonfinite selected service carrier {carrier}")
        # Receipt amounts live on signed spending carriers. Recover their economic
        # receipt amount before measuring the parent's service-use share.
        total += amount if source["side"] == "spending" else -amount
    return total


def coverage_service_amounts(coverage, end, dump_lines, required, source_root=FISCAL):
    """Coverage people use the same generation/age household prices as coverage.

    The pinned source's line_parts retain the exact missing-person-weighted service
    composition. Zero-response per-head lines retain their household raw amount
    per person; net-cost division cannot recover those amounts.
    """
    if not coverage or coverage["net_missing_people"] == 0:
        return {line: 0.0 for line in required}
    if coverage.get("stub") or coverage.get("which") != "set":
        raise ValueError("[BLOCKED] pension coverage requires the live explicit set profile")
    source = source_root / coverage["file"]
    if not source.exists() or hashlib.sha256(source.read_bytes()).hexdigest() != coverage["sha256"]:
        raise ValueError("[BLOCKED] pension coverage source is absent or stale")
    raw = json.loads(source.read_text())
    if raw.get("key") != "oct15":
        raise ValueError("[BLOCKED] pension coverage profile is not oct15")
    scenario = raw["sets"]["set"]["scenarios"][coverage["construction"]]
    entry = scenario["ends"][end]
    decomposition = entry["decomposition"]
    parts = decomposition["line_parts"]
    target = coverage["by_end"][end]["total_bn"]
    if (not all(math.isfinite(row["coverage_bn"]) for row in parts)
            or abs(sum(row["coverage_bn"] for row in parts) - target) > 1e-6
            or abs(decomposition["total_cost_bn"] - target) > 1e-6):
        raise ValueError("[BLOCKED] pension coverage typed parts do not close to the payload")
    age_prices = scenario["by_generation_age_band"]
    composition = 0.0
    for generation, bands in age_prices.items():
        for band, price in bands.items():
            count = scenario["implied_net_missing_people_by_age_band"][generation][band]
            quoted = price["by_end"][end]
            if abs(price["net_missing_people"] - count) > 1e-6:
                raise ValueError("[BLOCKED] pension coverage age-price population mismatch")
            amount = count * quoted["per_member_usd"] / 1e9 if count else 0.0
            if not math.isfinite(amount) or abs(amount - quoted["coverage_bn"]) > 1e-6:
                raise ValueError("[BLOCKED] pension coverage missing-person age-price mismatch")
            composition += amount
    if abs(composition - target) > 1e-6:
        raise ValueError("[BLOCKED] pension coverage generation/age service composition does not close (1e-6bn)")
    # Coverage prices the frozen, coverage-off household export, not a recursively
    # coverage-on amount. Verify that principal input's receipt before raw pricing.
    base = source.parent.parent / "inputs/oct15_base"
    receipt = json.loads((base / "manifest.json").read_text())
    pin = next((p for p in receipt["files"] if p["name"] == "lines.json"), None)
    household_path = base / "lines.json"
    if not pin or not household_path.exists() or sha(household_path) != pin["sha256"]:
        raise ValueError("[BLOCKED] pension coverage household prices are absent or stale")
    household = json.loads(household_path.read_text())
    if household["meta"].get("case") != "main_case_2026_10_15" or household["meta"].get("basis") != "set":
        raise ValueError("[BLOCKED] pension coverage household price identity differs")
    population_path = base / "generation_results.csv"
    population_pin = next((p for p in receipt["files"] if p["name"] == population_path.name), None)
    if not population_pin or not population_path.exists() or sha(population_path) != population_pin["sha256"]:
        raise ValueError("[BLOCKED] pension coverage household population is absent or stale")
    with population_path.open() as stream:
        populations = {r["generation"]: float(r["population"]) for r in csv.DictReader(stream)
                       if r["convention"] == "a" and r["band_end"] == end}
    costs = {line: 0.0 for line in required}
    for row in parts:
        line = row.get("source_line")
        if line not in costs or row["economic_class"] != "fiscal_line":
            continue
        economic_side = "receipt" if dump_lines[line]["side"] == "receipts" else dump_lines[line]["side"]
        if row["source_side"] != economic_side:
            raise ValueError(f"[BLOCKED] pension coverage source side differs for {line}")
        costs[line] += row["coverage_bn"]
    amounts = {}
    for line, cost in costs.items():
        response = dump_lines[line]["response"]
        sign = 1 if dump_lines[line]["side"] == "spending" else -1
        if response != 0:
            amount = sign * cost / response
        else:
            amount = 0.0
            for generation in ("G1", "G2", "G3plus"):
                group = household["generations"][generation][end]
                economic_side = "receipt" if dump_lines[line]["side"] == "receipts" else dump_lines[line]["side"]
                declarations = [r for r in group["rows"] if r["id"] == line and r["side"] == economic_side]
                if len(declarations) != 1 or declarations[0]["key"] != "population":
                    raise ValueError(f"[BLOCKED] pension coverage zero-response {line} lacks a per-head household price")
                people = populations[generation]
                if not math.isfinite(people) or people <= 0:
                    raise ValueError("[BLOCKED] pension coverage service-price population is invalid")
                missing = sum(scenario["implied_net_missing_people_by_age_band"][generation].values())
                amount += declarations[0]["amount_bn"] * missing / people
        if not math.isfinite(amount) or abs(sign * amount * response - cost) > 1e-9:
            raise ValueError(f"[BLOCKED] pension coverage raw service amount does not reproduce {line}")
        amounts[line] = amount
    return amounts


def completion_attribution(end, dump_lines):
    """Attribute every carrier to its declared parent, retaining its economic nature.

    Capital and production carriers are resource terms, not employee service-use
    amounts. Tax/transfer and unlocated constants are outside this pension kernel.
    """
    capital_services = {"k12": "education_services", "college": "education_services",
        "k12_user_fees": "education_services", "college_user_fees": "education_services",
        "pos_sl": "public_order_safety", "pos_fed": "public_order_safety",
        "health_sl": "health_services", "health_fed": "health_services",
        "health_sl_user_fees": "health_services", "health_fed_user_fees": "health_services",
        "gps_sl": "general_public_services", "gps_fed": "general_public_services",
        "hwy_sl": "highways", "hwy_fed": "highways", "air_fed": "economic_affairs_services",
        "rec_sl": "recreation_culture", "rec_fed": "recreation_culture"}
    for carrier, source in COMPLETION_CARRIERS.items():
        parent, kind = source.get("parent_line"), source.get("kind")
        if not parent or kind not in ("line", "capital", "production") or carrier not in dump_lines:
            raise ValueError(f"[BLOCKED] item10 carrier lacks its declared parent/kind/amount: {carrier}")
        if kind == "capital":
            service = "enterprises" if parent.startswith("ent_") else capital_services.get(parent)
            if service is None:
                raise ValueError(f"[BLOCKED] item10 capital parent has no service mapping: {parent}")
        elif kind == "production":
            service = "private_production"
        else:
            if parent != "lane_constants" and parent not in dump_lines:
                raise ValueError(f"[BLOCKED] item10 fiscal parent absent: {parent}")
            service = next((name for name, (overlay, _) in OVERLAY.items() if parent == overlay),
                           "highways" if parent.startswith("roads_vmt_") else
                           "education_services" if parent in ("school_reprice", "college_rekey") else
                           "enterprises" if parent == "enterprise_surplus" else
                           parent if parent in LINES else "outside_public_service")
        line = dump_lines[carrier]
        amount = line["amount_bn"]
        if not math.isfinite(amount):
            raise ValueError(f"[BLOCKED] item10 carrier amount nonfinite: {carrier}")
        if kind == "line" and parent in dump_lines:
            selected = dump_lines[parent]
            selected_key = selected.get("scenario") if source.get("side") == "receipt" else selected.get("key")
            active = selected_key == source.get("parent_key")
        else:
            active = source.get("end") == end if kind in ("capital", "production") else False
        included = kind == "line" and service in LINES and active
        SERVICE_ATTRIBUTIONS.append([end, carrier, kind, parent, service, amount, line["response"], active, included])
    if len([r for r in SERVICE_ATTRIBUTIONS if r[0] == end]) != len(COMPLETION_CARRIERS):
        raise ValueError("[BLOCKED] incomplete item10 service attribution")


def defense_stock_bridge(counts, population, union_raw, baseline_raw, full_raw, native_raw):
    """Stock use at union raw defense/n4; subtract native amounts already in full raw.

    The baseline lineage's raw amount is retained, including its historical tiny
    construction-rounding remainder. This does not alter any removal-response key.
    """
    n4 = counts["baseline_account_union"]
    completed = counts["reporters"] + counts["nonreporters"]
    covered = population["coverage"]
    values = (n4, completed, covered, union_raw, baseline_raw, full_raw, native_raw)
    if not all(math.isfinite(v) for v in values) or n4 <= 0:
        raise ValueError("[BLOCKED] defense stock source values are invalid")
    if (abs(completed - population["completion"]) > 1e-6
            or abs(counts["baseline_lineage_population"] + completed - population["without_coverage"]) > 1e-6
            or abs(population["without_coverage"] + covered - population["with_coverage"]) > 1e-6):
        raise ValueError("[BLOCKED] defense stock population bridge does not close (1e-6 persons)")
    if abs(full_raw - baseline_raw - native_raw) > 1e-9:
        raise ValueError("[BLOCKED] defense raw baseline/native bridge does not close (1e-9bn)")
    price = union_raw / n4
    completion_raw, coverage_raw = price * completed, price * covered
    additional = completion_raw + coverage_raw - native_raw
    desired = baseline_raw + completion_raw + coverage_raw
    if abs(full_raw + additional - desired) > 1e-9:
        raise ValueError("[BLOCKED] defense stock addition double assigns native defense")
    return dict(union_raw_bn=union_raw, n4=n4, raw_price_bn_per_person=price,
        baseline_raw_bn=baseline_raw, full_raw_bn=full_raw, native_already_included_bn=native_raw,
        reporters=counts["reporters"], nonreporters=counts["nonreporters"], coverage_people=covered,
        completed_people=completed, lineage_population=population["with_coverage"],
        completion_raw_bn=completion_raw, coverage_raw_bn=coverage_raw,
        additional_stock_raw_bn=additional, stock_raw_bn=desired,
        amount_basis="annual raw defense service-use amount for legacy-stock attribution; not a pension stock or fiscal-removal amount",
        rule="baseline lineage raw + (baseline union raw defense/n4)*(reporters+nonreporters+coverage); subtract already included native raw once")


def defense_stock_amount(end, dump_line):
    """Trace independently preserved v7 raw and the coverage's frozen household base."""
    white = FISCAL / "white_replacement_2026_09_28/derived"
    union = json.loads((white / "engine_lines_oct14_union.json").read_text())[end]
    baseline = json.loads((white / "engine_lines_oct14.json").read_text())[end]
    row = lambda block: next(r for r in block["lines"] if r["id"] == "defense")
    base = FISCAL / "coverage_bound_2026_10_14/inputs/oct15_base"
    receipt = json.loads((base / "manifest.json").read_text())
    pin = next(p for p in receipt["files"] if p["name"] == "lines.json")
    if sha(base / "lines.json") != pin["sha256"]:
        raise ValueError("[BLOCKED] defense stock household source is stale")
    household = json.loads((base / "lines.json").read_text())
    if household["meta"].get("case") != "main_case_2026_10_15" or household["meta"].get("basis") != "set":
        raise ValueError("[BLOCKED] defense stock household source identity differs")
    amount = lambda group: next(r["amount_bn"] for r in group["rows"] if r["id"] == "defense" and r["side"] == "spending")
    before, full = 0.0, 0.0
    for generation in ("G1", "G2", "G3plus"):
        group = household["generations"][generation][end]
        before += amount(group.get("before_native_children") or group)
        full += amount(group)
    native = household["meta"]["child_pricing"]["records"]
    native_people = sum(r["people"] for r in native)
    native_raw = sum(r["people"] * next(p["amount_bn"] for p in r["ends"][end]["rows"]
        if p["id"] == "defense" and p["side"] == "spending") for r in native)
    if (abs(before - row(baseline)["amount_bn"]) > 1e-9
            or abs(native_people - COMPLETION_COUNTS["nonreporters"]) > 1e-6
            or abs(full - dump_line["amount_bn"]) > 1e-9
            or abs(row(union)["national_bn"] - dump_line["national_bn"]) > 1e-9
            or dump_line["response"] != 0):
        raise ValueError("[BLOCKED] defense stock raw/native/source response trace does not close")
    bridge = defense_stock_bridge(COMPLETION_COUNTS, POPULATION_V8, row(union)["amount_bn"],
                                  before, full, native_raw)
    DEFENSE_STOCK_BRIDGES.append(dict(case=CASE, end=end, removal_defense_response=0,
        mortality_convention="not_applicable_service_share", **bridge))
    return bridge["stock_raw_bn"]


def v9_service_sources():
    """Require the live producer's raw snapshot, rather than undoing unknown edits."""
    path = GROUP_LINES.parent / "service_parts_oct16.json"
    if not path.exists():
        raise ValueError("[BLOCKED] oct16 pension source needs the raw-service/typed-part export")
    source = json.loads(path.read_text())
    if source.get("case") != "oct16":
        raise ValueError("[BLOCKED] pension service source case differs")
    for name, expected in source["source_sha256"].items():
        if sha(FISCAL / name) != expected:
            raise ValueError(f"[BLOCKED] stale pension service source: {name}")
    runs = {(r["group"], r["end"]): r for r in source["runs"] if r["basis"] == "accrual"}
    expected = {(group, end) for group in GROUP_SOURCES.values() for end in ("low", "high")}
    if set(runs) != expected:
        raise ValueError("[BLOCKED] pension service source lacks eight accrual group/end records")
    return runs


def v9_military_fraction():
    bridge = json.loads((FISCAL / "spending_keys_r5_2026_10_15/derived/item13.json").read_text())["extension"]["health_bridge"]
    fraction = bridge["central_national_bn"] / bridge["federal_health_consumption"]
    if not 0 < fraction < 1:
        raise ValueError("[BLOCKED] pension federal VA service fraction is invalid")
    if V9_META:
        declared = [part["national_component_bn"] for part in
            V9_META["veterans_compensation"]["by_end"]["low"]["parts"] if part.get("part") == "va_medical"]
        if not declared or any(abs(value-bridge["central_national_bn"]) > 1e-9 for value in declared):
            raise ValueError("[BLOCKED] pension VA source differs from the canonical live payload")
    return bridge, fraction


def v9_raw_health(cost, line, lines):
    """Undo only source13's aggregate price factor; VA is routed by plan below."""
    if line != "health_services":
        return cost
    bridge, _ = v9_military_fraction()
    factor = 1 - bridge["central_national_bn"] / lines[line]["national_bn"]
    if not 0 < factor < 1:
        raise ValueError("[BLOCKED] pension health source military factor is invalid")
    return cost / factor


def v9_coverage_service_amounts(end, lines, needed):
    source = FISCAL / COVERAGE["file"]
    if COVERAGE.get("stub") or COVERAGE.get("which") != "set" or sha(source) != COVERAGE["sha256"]:
        raise ValueError("[BLOCKED] pension oct16 coverage is stub, wrong basis or stale")
    raw = json.loads(source.read_text())
    if raw.get("key") != "oct16":
        raise ValueError("[BLOCKED] pension coverage is not oct16")
    entry = raw["sets"]["set"]["scenarios"][COVERAGE["construction"]]["ends"][end]
    parts = entry["decomposition"]["line_parts"]
    if abs(sum(r["coverage_bn"] for r in parts) - COVERAGE["by_end"][end]["total_bn"]) > 1e-9:
        raise ValueError("[BLOCKED] oct16 pension coverage parts do not close (1e-9bn)")
    # A zero-response raw defense key has no signed cost to invert. Retain its
    # separately pinned v8 household raw price, then calibrate people explicitly.
    old = json.loads((FISCAL / "main_case_2026_10_15/derived/corrections.json").read_text())["meta"]["coverage"]
    zero = [line for line in needed if lines[line]["response"] == 0]
    amounts = coverage_service_amounts(old, end, lines, zero) if zero else {}
    for line in needed:
        if line in amounts:
            continue
        cost = sum(r["coverage_bn"] for r in parts if r.get("source_side") in ("receipt", "spending")
                   and r.get("source_line") == line and r.get("economic_class") == "fiscal_line"
                   and not r["line"].startswith("item21:"))
        sign = -1 if lines[line]["side"] == "receipts" else 1
        amounts[line] = sign * v9_raw_health(cost, line, lines) / lines[line]["response"]
    return amounts


def v9_service_adjustments(group, end, lines, shares):
    """Apply operating service changes, never capital, to employee attribution."""
    run = V9_SERVICE_RUNS[(GROUP_SOURCES[group], end)]
    own = group in ("mexican_origin", "mexican_origin_rough")
    parts = ([dict(part, item=number) for number, key in
              ((14, "income_security_administration"), (15, "public_college_response"))
              for part in V9_META[key]["by_end"][end]["parts"]] if own else run["parts"])
    for line, selected in (("education_services", "college_operating"),
                           ("income_security_services", "income_security_administration")):
        delta = sum(p["delta_bn"] for p in parts if p.get("population") != "coverage" and
                    (p.get("kind") == selected or str(p.get("item")) == "14" or p.get("item") == selected)
                    and p.get("source_line") == line)
        shares[line] += delta / lines[line]["national_bn"]
        V9_SERVICE_PARTS.append(dict(group=group, end=end, source_line=line, kind=selected,
            delta_service_bn=delta, source="canonical source parts" if own else "own-group comparator parts"))
    V9_SERVICE_PARTS.append(dict(group=group, end=end, kind="college_capital",
        delta_capital_excluded_bn=sum(p["delta_bn"] for p in parts if p.get("kind") == "college_capital"),
        employee_service=False))
    changes = []
    if own:
        calibration, tenure = V9_META["native_minor_calibration"], V9_META["coverage_tenure"]
        if (calibration["parts_file"] != tenure["parts_file"] or
                tenure.get("subtract_parts_scenario") != calibration["parts_scenario"]):
            raise ValueError("[BLOCKED] pension count parts lack an explicit joint-minus-calibration declaration")
        path = FISCAL / calibration["parts_file"]
        with path.open() as stream:
            changes = [r for r in csv.DictReader(stream) if r["scenario"] == tenure["parts_scenario"] and r["end"] == end]
        if abs(sum(float(r["change_bn"]) for r in changes) - sum(
                V9_META[key]["by_end"][end]["total_bn"] for key in ("native_minor_calibration", "coverage_tenure"))) > 1e-9:
            raise ValueError("[BLOCKED] pension joint calibrated/coverage service parts do not close (1e-9bn)")
        for row in changes:
            if row["line"].startswith("item21:"):
                continue  # The new payer key never changes VA's old service key.
            line = row["source_line"]
            if row["economic_class"] != "fiscal_line" or line not in lines or row["source_side"] not in ("receipt", "spending"):
                continue
            cost = v9_raw_health(float(row["change_bn"]), line, lines)
            if lines[line]["response"] == 0:
                continue
            amount = cost / lines[line]["response"] * (-1 if row["source_side"] == "receipt" else 1)
            target = next((name for name, (overlay, _) in OVERLAY.items() if line == overlay), line)
            national = next((nat for _, (overlay, nat) in OVERLAY.items() if line == overlay), lines[line]["national_bn"])
            if target in shares:
                shares[target] += amount / national
                if line == "state_price_health_services":
                    shares["_health_state_price"] += amount / national
                if line == "health_services":
                    shares["health_services_va"] += amount / national
                if line == "economic_affairs_services":
                    shares["highways"] += amount / national
            elif line == "roads_vmt_sl":
                shares["highways"] += amount / HIGHWAYS_NATIONAL
            elif line == "enterprise_surplus":
                shares["enterprises"] += amount / national
        # Both are declared per-person source keys; no whole-account ratio is used.
        shares["defense"] = UNION_PER_HEAD_KEY * POPULATION / UNION_ROW4
    _, fraction = v9_military_fraction()
    state_price = shares.pop("_health_state_price", 0.0)
    shares["health_services_federal_other"] = (shares["health_services"] - state_price
        - fraction * shares["health_services_va"]) / (1-fraction)
    if own:
        v9_health_residual_adjustment(group, end, lines, shares, changes)


def v9_health_residual_adjustment(group, end, lines, shares, count_parts):
    """The residual key serves civilian/state health; VA and capital stay separate."""
    if "health_residual_key" not in V9_META:
        raise ValueError("[BLOCKED] oct16 pension needs the declared item21 health source")
    block = V9_META["health_residual_key"]["by_end"][end]
    parts = block["parts"]
    if abs(sum(p["delta_bn"] for p in parts)-block["total_bn"]) > 1e-9:
        raise ValueError("[BLOCKED] pension item21 source parts do not close (1e-9bn)")
    bridge, _ = v9_military_fraction()
    residual = lines["health_services"]["national_bn"]-bridge["central_national_bn"]
    if residual <= 0:
        raise ValueError("[BLOCKED] pension item21 residual national amount is invalid")
    changes = [p for p in count_parts if p["line"].startswith("item21:")]
    interactions = V9_META["health_residual_key"]["count_interactions"]
    target = sum(interactions[k]["by_end"][end]["total_bn"]
                 for k in ("native_minor_calibration", "coverage_tenure"))
    # The metadata keeps source2%/3% price coordinates; count parts carry20's
    # live rate. Undo that coordinate change only for the source-total gate.
    source_total = 0.0
    for part in changes:
        change = float(part["change_bn"])
        if part.get("economic_class") == "capital_return":
            source_rate, live_rate = float(part["source_capital_rate"]), float(part["capital_rate"])
            if source_rate <= 0 or live_rate <= 0:
                raise ValueError("[BLOCKED] pension item21 count capital lacks its rate bridge")
            change *= source_rate/live_rate
        source_total += change
    if abs(source_total-target) > 1e-9:
        raise ValueError("[BLOCKED] pension item21 calibrated/coverage parts do not close (1e-9bn)")
    deltas = {}
    for line in ("health_services", "state_price_health_services"):
        deltas[line] = (sum(p["delta_bn"] for p in parts if p["source_line"] == line)
            + sum(float(p["change_bn"]) for p in changes if p["source_line"] == line))
    residual_share = deltas["health_services"] / residual
    overlay_share = deltas["state_price_health_services"] / OVERLAY["health_services"][1]
    shares["health_services"] += residual_share + overlay_share
    shares["health_services_federal_other"] += residual_share
    V9_SERVICE_PARTS.append(dict(group=group, end=end, kind="item21_non_va_health",
        residual_consumption_delta_bn=deltas["health_services"], state_price_delta_bn=deltas["state_price_health_services"],
        residual_national_bn=residual, residual_share_change=residual_share, state_price_share_change=overlay_share,
        va_key_changed=False, capital_employee_service=False))


def v9_function_mixes(mixes):
    """The source VA national subblock applies to federal civilian service only."""
    _, fraction = v9_military_fraction()
    for plan, mix in mixes.items():
        health = mix["health_services"] if plan == "federal_civilian" else 0.0
        weight = health * fraction
        mix["health_services"] -= health
        mix["health_services_va"] = weight
        mix["health_services_federal_other"] = health-weight
        if abs(sum(mix.values()) - 1) > 1e-12:
            raise ValueError("[BLOCKED] VA/non-VA pension function mix does not close")
    return mixes


def group_shares() -> tuple[dict, dict]:
    """Read the white lane's gated sept29 line export on audit row 4, including each group's overlays.

    Keep the engine union as an anchor; use the rough union for comparisons on identical keys.
    The export's producer runs rekey_sept29.setup(), checks every cost, and records each slice's count.
    Reject a different population, line definition or engine amount before attributing any liability.
    """
    global V9_SERVICE_RUNS
    SERVICE_ATTRIBUTIONS.clear()
    DEFENSE_STOCK_BRIDGES.clear()
    V9_SERVICE_PARTS.clear()
    for path in (DUMP_FILE, GROUP_LINES):
        if not path.exists():
            raise ValueError(f"[BLOCKED] pension service-share input absent: {path}; build the matching group-line export")
    dump = json.loads(DUMP_FILE.read_text())
    with GROUP_LINES.open() as f:
        rows = [r for r in csv.DictReader(f) if r["basis"] == "accrual" and r["group"] in GROUP_SOURCES.values()]
    if CASE == "oct16":
        V9_SERVICE_RUNS = v9_service_sources()
        raw = {}
        for (group, end), run in V9_SERVICE_RUNS.items():
            for side, line, national, amount, response in run["raw_lines"]:
                raw[group, end, "receipt" if side == "receipts" else "spending", line] = (national, amount, response)
        rows = [dict(r) for r in rows]
        for row in rows:
            key = row["group"], row["end"], row["side"], row["line"]
            if key in raw:
                row.update(zip(("national_bn", "amount_bn", "response"), raw[key]))
    if CASE in ("oct15", "oct16") and any(r.get("engine") != COMPARISON.engine_for(r["group"]) for r in rows):
        raise ValueError("[BLOCKED] pension group lines have missing or incorrect engine tags")
    source = {(r["group"], r["end"], r["side"], r["line"]): r for r in rows}
    if len(source) != len(rows):
        raise ValueError("[BLOCKED] duplicate group line in row-4 export")

    def value(group, end, line, field):
        side = "scalar" if line == "population" else "receipt" if line == "enterprise_surplus" else "spending"
        try:
            v = float(source[(GROUP_SOURCES[group], end, side, line)][field])
        except (KeyError, ValueError) as e:
            raise ValueError(f"[BLOCKED] missing or invalid {group}/{end}/{line}/{field}") from e
        if not math.isfinite(v):
            raise ValueError(f"[BLOCKED] nonfinite {group}/{end}/{line}/{field}")
        return v

    base = ["general_public_services", "defense", "public_order_safety", "economic_affairs_services", "housing_community_services",
            "health_services", "recreation_culture", "education_services", "income_security_services"]
    share, response = {g: {} for g in GROUPS}, {}
    for end in ("low", "high"):
        lines = {l["id"]: l for l in dump[end]["lines"]}
        if CASE in ("oct15", "oct16"):
            completion_attribution(end, lines)
            defense_stock = defense_stock_amount(end, lines["defense"])
        if lines["domestic_interest"]["response"] != 0 or abs(lines["domestic_interest"]["national_bn"] - 1118.87) > 1e-9:
            raise ValueError("[BLOCKED] the case no longer holds domestic interest ($1,118.87bn) at zero response")
        resp = {l: lines[l]["response"] for l in base}
        resp["highways"] = lines["roads_vmt_sl"]["response"]
        resp["enterprises"] = lines["enterprise_surplus"]["response"]
        if CASE == "oct16":
            resp["health_services_va"] = 0.0
            resp["health_services_federal_other"] = resp["health_services"]
        response[end] = resp
        needed = base + [v[0] for v in OVERLAY.values()] + ["roads_vmt_sl", "enterprise_surplus"]
        coverage_amounts = (v9_coverage_service_amounts(end, lines, needed) if CASE == "oct16" else
                            coverage_service_amounts(COVERAGE, end, lines, needed) if CASE == "oct15" else {})
        for g in GROUPS:
            if abs(value(g, end, "population", "amount_bn") - POPULATION) > 0.01:
                raise ValueError(f"[BLOCKED] {g}/{end} population is not the case's priced count ({POPULATION:,.2f})")
            for l in needed:
                for field in ("national_bn", "response"):
                    if abs(value(g, end, l, field) - lines[l][field]) > 1e-8:
                        raise ValueError(f"[BLOCKED] {g}/{end}/{l} {field} differs from the case")
                if g == "mexican_origin" and abs(value(g, end, l, "amount_bn") - lines[l]["amount_bn"]) > 1e-8:
                    raise ValueError(f"[BLOCKED] {end}/{l} differs from the engine dump")
            s = {l: value(g, end, l, "amount_bn") / lines[l]["national_bn"] for l in base}
            if CASE in ("oct15", "oct16"):
                amounts = {r["line"]: float(r["amount_bn"]) for r in rows
                           if r["group"] == GROUP_SOURCES[g] and r["end"] == end}
                completion_amounts = {line: service_carrier_amount(line, lines, amounts, COMPLETION_CARRIERS)
                                      for line in needed}
                for line in base:
                    s[line] += completion_amounts[line] / lines[line]["national_bn"]
            if g in ("mexican_origin", "mexican_origin_rough"):
                for line in base:
                    s[line] += coverage_amounts.get(line, 0.0) / lines[line]["national_bn"]
                if CASE in ("oct15", "oct16"):
                    # A stock-only service attribution. The same per-head price/count
                    # applies to both lineage key variants; defense removal stays zero.
                    s["defense"] = defense_stock / lines["defense"]["national_bn"]
                if CASE == "oct16":
                    # The traced raw defense bridge runs on v8's counts (its household source is the oct15
                    # base). v9_service_adjustments then sets the per-head key at v9's count, so the 5e-7 key
                    # gate below would compare that key with itself. Test the trace here, on its own count.
                    v8_key = UNION_PER_HEAD_KEY * POPULATION_V8["with_coverage"] / UNION_ROW4
                    if abs(s["defense"] - v8_key) > 5e-7:
                        raise ValueError(f"[BLOCKED] {g}/{end} traced defense stock is not the per-head key "
                                         f"{v8_key:.6f} at its own v8 count")
            if CASE == "oct16":
                s["health_services_va"] = s["health_services"]
                s["_health_state_price"] = 0.0
            for line, (ov, nat) in OVERLAY.items():
                if abs(lines[line]["national_bn"] - SCALED.get(line, nat)) > 1e-9:
                    raise ValueError(f"[BLOCKED] {line} national changed")
                s[line] += value(g, end, ov, "amount_bn") / nat      # a scaled line's overlay stays on nat (CASES)
                if CASE in ("oct15", "oct16"):
                    s[line] += completion_amounts[ov] / nat
                if g in ("mexican_origin", "mexican_origin_rough"):
                    s[line] += coverage_amounts.get(ov, 0.0) / nat
                if CASE == "oct16" and line == "health_services":
                    s["_health_state_price"] = (value(g, end, ov, "amount_bn")
                        + completion_amounts[ov]
                        + (coverage_amounts.get(ov, 0.0) if g in ("mexican_origin", "mexican_origin_rough") else 0.0)) / nat
            s["highways"] = s["economic_affairs_services"] + value(g, end, "roads_vmt_sl", "amount_bn") / HIGHWAYS_NATIONAL
            s["enterprises"] = value(g, end, "enterprise_surplus", "amount_bn") / lines["enterprise_surplus"]["national_bn"]
            if CASE in ("oct15", "oct16"):
                s["highways"] += completion_amounts["roads_vmt_sl"] / HIGHWAYS_NATIONAL
                s["enterprises"] += completion_amounts["enterprise_surplus"] / lines["enterprise_surplus"]["national_bn"]
            if g in ("mexican_origin", "mexican_origin_rough"):
                s["highways"] += coverage_amounts.get("roads_vmt_sl", 0.0) / HIGHWAYS_NATIONAL
                s["enterprises"] += coverage_amounts.get("enterprise_surplus", 0.0) / lines["enterprise_surplus"]["national_bn"]
            if CASE == "oct16":
                v9_service_adjustments(g, end, lines, s)
            share[g][end] = s
    key = UNION_PER_HEAD_KEY * POPULATION / UNION_ROW4
    for g in GROUPS:
        for end in ("low", "high"):
            if abs(share[g][end]["defense"] - key) > 5e-7 or abs(share[g][end]["enterprises"] - key) > 5e-7:
                raise ValueError(f"[BLOCKED] {g} per-head key is not the case's {key:.6f}")
    return share, response


# ----------------------------------------------------------------------------------------------- headcount path
def headcount_path(population: dict[int, float], census: dict[int, float]) -> tuple[dict[int, float], dict[int, float]]:
    """The Mexican-origin population share relative to 2024, as the back-cast builds it: ACS 1-year counts 2005-2024
    (2020 interpolated) over BEA midperiod population; the decennial counts before, spliced onto the ACS level by
    the 2010 ratio (ACS 2010 / census 2010); log-linear between anchors; held at the 1970 share before 1970."""
    acs = {int(r["year"]): r["acs_mexican_origin"] for r in csv.DictReader(open(FISCAL / "historical_backcast_2026_09_20/inputs/acs_mexican_origin.csv"))}
    share = {y: float(v) / (population[y] * 1e3) for y, v in acs.items() if v}
    splice = float(acs[2010]) / (census[2010] * 1e6)
    if not 1.0 < splice < 1.06:
        raise ValueError(f"[BLOCKED] ACS/census 2010 splice {splice} out of range")
    for y, c in census.items():
        if y < 2005:
            share[y] = c * 1e6 * splice / (population[y] * 1e3)
    anchors = sorted(share)
    full = {}
    for y in range(1946, 2025):
        if y in share:
            full[y] = share[y]
        elif y < anchors[0]:
            full[y] = share[anchors[0]]
        else:
            a = max(k for k in anchors if k < y)
            b = min(k for k in anchors if k > y)
            w = (y - a) / (b - a)
            full[y] = math.exp((1 - w) * math.log(share[a]) + w * math.log(share[b]))
    return {y: v / full[2024] for y, v in full.items()}, full


def kernel(age: int, plateau: int, zero_at: int) -> float:
    """Weight of service rendered `age` years before 2024 in the end-2023 accrued liability: benefits not yet in payment
    are held whole for `plateau` years, then run off linearly to nothing at `zero_at` years."""
    if age <= plateau:
        return 1.0
    return max(0.0, (zero_at - age) / (zero_at - plateau))


VINTAGES = {"vintage_central": (15, 45), "vintage_short": (10, 35), "vintage_long": (20, 60)}


def time_factors(rel: dict, population: dict, z1: dict) -> tuple[dict, dict]:
    """F[plan][arm] = sum_v w(v) rel(v) / sum_v w(v) over service years 1946-2023; 'simple' is 1 (2024 shares)."""
    years = range(1946, 2024)
    weights = {}
    for arm, (p, z) in VINTAGES.items():
        weights[("state_local", arm)] = {v: kernel(2024 - v, p, z) * population[v] for v in years}
        for plan in ("federal_civilian", "federal_military"):
            weights[(plan, arm)] = {v: kernel(2024 - v, p, z) for v in years}
    sl_unf = {y: z1[y]["FL223073045.Q"] for y in z1}
    fed_unf = {y: z1[y]["FL344190045.Q"] - (z1[y]["FL344090045.Q"] - z1[y]["FL343073045.Q"] - z1[y]["FL343069245.Q"]) for y in z1}
    weights[("state_local", "increments")] = {v: sl_unf[v] - sl_unf[v - 1] for v in years}
    for plan in ("federal_civilian", "federal_military"):
        weights[(plan, "increments")] = {v: fed_unf[v] - fed_unf[v - 1] for v in years}
    factors = {}
    for (plan, arm), w in weights.items():
        factors.setdefault(plan, {"simple": 1.0})[arm] = sum(w[v] * rel[v] for v in years) / sum(w.values())
    return factors, weights


# ----------------------------------------------------------------------------------------------- output
def comparison_tags(group):
    """Live comparisons use rough keys; preserve case-vs-rough as a bridge."""
    if group == "case_minus_rough_bridge":
        return ["bridge", "case", "rough"]
    if group == "mexican_origin_rough_minus_A1_white":
        left, right = COMPARISON.engine_for("mexican_origin_rough"), COMPARISON.engine_for("A1_third_plus_nh_white")
        COMPARISON.require_same_engine({"engine": left}, {"engine": right})
        return ["comparison", left, right]
    if group not in GROUPS:
        raise ValueError(f"[BLOCKED] unrecognized live pension comparison group: {group}")
    return ["absolute", COMPARISON.engine_for("mexican_origin_engine" if group == "mexican_origin" else group), ""]


def write(path: Path, header: list[str], rows: list[list]) -> None:
    if CASE in ("oct15", "oct16") and "group" in header:
        group_index = header.index("group")
        header = header + ["case", "mortality_convention", "row_type", "lineage_engine", "comparator_engine"]
        rows = [row + [CASE, "not_applicable_service_share"] + comparison_tags(row[group_index]) for row in rows]
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow([f"{x:.6f}" if isinstance(x, float) else x for x in r])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fetch", action="store_true", help="download every input to _cache/ and rewrite source_pins.json")
    ap.add_argument("--census-env", type=Path, default=FISCAL / "acquire/config.local.env",
                    help="file holding CENSUS_API_KEY=... when the variable is not set (only --fetch reads it)")
    ap.add_argument("--out-dir", type=Path, default=None, help="default derived/ (sept29) or derived/<case>/")
    ap.add_argument("--case", default="oct16", choices=list(CASES),
                    help="oct16 (live v9); earlier keyed cases are archived")
    args = ap.parse_args()
    if args.fetch:
        fetch(args.census_env if args.census_env.exists() else None)
    use_case(args.case)
    out = args.out_dir or (HERE / "derived" if args.case == "sept29" else HERE / "derived" / args.case)
    out.mkdir(parents=True, exist_ok=True)

    m, extra = measure_bea()
    z1 = measure_z1()
    fr = measure_fr()
    census = measure_census_counts()
    aspep = measure_aspep()
    mix = function_mixes(aspep, m)
    attribution_lines = LINES
    if CASE == "oct16":
        mix = v9_function_mixes(mix)
        attribution_lines = LINES + ["health_services_va", "health_services_federal_other"]
    share, response = group_shares()
    if CASE in ("oct15", "oct16"):
        write(out / "completion_service_attribution.csv", ["end", "carrier", "kind", "parent_line", "service",
              "signed_carrier_amount_bn", "response", "selected_parent_or_end", "employee_service_amount"], SERVICE_ATTRIBUTIONS)
        (out / "defense_stock_bridge.json").write_text(json.dumps(DEFENSE_STOCK_BRIDGES, indent=2) + "\n")
    rel, abs_share = headcount_path(extra["population"], census)
    factors, weights = time_factors(rel, extra["population"], z1)

    v = {k: x for k, (x, _) in m.items()}
    # --- gates on the measured structure
    if abs(v["domestic_interest_federal"] + v["domestic_interest_state_local"] - v["domestic_interest_total"]) > 1e-9:
        raise ValueError("[BLOCKED] T3.2 + T3.3 domestic interest do not add to T3.1")
    for side, parts, whole in [("sl", ["sl_plan_monetary_interest", "sl_imputed_interest", "sl_plan_dividends"], None),
                               ("fed", ["fed_plan_monetary_interest", "fed_imputed_interest"], "fed_plan_interest_total")]:
        if whole and abs(sum(v[p] for p in parts) - v[whole]) > 1e-9:
            raise ValueError(f"[BLOCKED] {side} interest parts do not add")
    holding = v["sl_interest_accrued_on_entitlements"] - (v["sl_plan_monetary_interest"] + v["sl_imputed_interest"] + v["sl_plan_dividends"])
    if not all("actuarial liabilities" in " ".join(extra["notes"][t]) for t in ("T3.1", "T3.2", "T3.3")):
        raise ValueError("[BLOCKED] the interest footnote no longer says it includes interest on actuarial liabilities")
    nc_sl = v["sl_actual_employer_contributions"] + v["sl_imputed_employer_contributions"]
    t723 = openpyxl.load_workbook(pinned("Section7All_xls.xlsx"), read_only=True, data_only=True)
    t723l, _ = bea_table(t723, "T72300-A")
    nc_civ = cell(t723l, 6, "Civilian") + cell(t723l, 9, "Civilian")
    nc_mil = cell(t723l, 7, "Military") + cell(t723l, 10, "Military")

    y0, y1 = 2023, 2024
    sl_stock = z1[y0]["FL223073045.Q"]
    fed_claims = z1[y0]["FL343073045.Q"]
    fed_entitle = z1[y0]["FL344190045.Q"]
    fed_other_assets = z1[y0]["FL344090045.Q"] - z1[y0]["FL343073045.Q"] - z1[y0]["FL343069245.Q"]
    fed_consolidated = fed_entitle - fed_other_assets
    civ_share_int = fr["pension_interest_fy24"]["civilian"] / fr["pension_interest_fy24"]["total"]
    civ_share_stock = fr["pension_begin_fy24"]["civilian"] / fr["pension_begin_fy24"]["total"]

    plans = {  # plan -> basis -> (stock end-2023, 2024 interest)
        "state_local": {"narrow": (sl_stock, v["sl_imputed_interest"])},
        "federal_civilian": {"narrow": (fed_claims * civ_share_stock, v["fed_imputed_interest"] * civ_share_int),
                             "consolidated": (fed_consolidated * civ_share_stock, v["fed_plan_interest_total"] * civ_share_int)},
        "federal_military": {"narrow": (fed_claims * (1 - civ_share_stock), v["fed_imputed_interest"] * (1 - civ_share_int)),
                             "consolidated": (fed_consolidated * (1 - civ_share_stock), v["fed_plan_interest_total"] * (1 - civ_share_int))},
    }
    plans["state_local"]["consolidated"] = plans["state_local"]["narrow"]    # S&L plan assets are outside claims

    # --- measured_2024.csv
    rows = [[k, x, src] for k, (x, src) in m.items()]
    rows += [["sl_employer_pension_compensation", nc_sl, "BEA T7.24 lines 5+6, including service charges"],
             ["fed_civilian_employer_pension_compensation", nc_civ, "BEA T7.23 lines 6+9, including service charges"],
             ["fed_military_employer_pension_compensation", nc_mil, "BEA T7.23 lines 7+10, including service charges"],
             ["sl_implied_funding_from_holding_gains", holding, "BEA T7.24 line 31 - lines 12,13,14"]]
    for y in (y0, y1):
        rows += [[f"z1_sl_claims_on_sponsor_{y}q4", z1[y]["FL223073045.Q"], "Z.1 FL223073045.Q (S129s1.3.s line 19)"],
                 [f"z1_sl_entitlements_{y}q4", z1[y]["FL224190043.Q"], "Z.1 FL224190043.Q (S129s1.3.s line 21)"],
                 [f"z1_fed_claims_on_sponsor_{y}q4", z1[y]["FL343073045.Q"], "Z.1 FL343073045.Q (S129s1.2.s line 10)"],
                 [f"z1_fed_entitlements_{y}q4", z1[y]["FL344190045.Q"], "Z.1 FL344190045.Q (S129s1.2.s line 11)"],
                 [f"z1_fed_treasury_nonmarketable_{y}q4", z1[y]["FL343069245.Q"], "Z.1 FL343069245.Q (S129s1.2.s)"],
                 [f"z1_fed_total_assets_{y}q4", z1[y]["FL344090045.Q"], "Z.1 FL344090045.Q (S129s1.2.s line 1)"]]
    rows += [["fed_consolidated_unfunded_2023q4", fed_consolidated, "Z.1 entitlements less assets other than sponsor claims and nonmarketable Treasury securities"]]
    for k, d in fr.items():
        for part, x in d.items():
            rows.append([f"fr_{k}_{part}", x, "Financial Report FY2025 Note 13, FY2024 column"])
    for y, c in census.items():
        rows.append([f"census_mexican_origin_{y}_millions", c, "Census briefs C2010BR-04 / C2KBR/01-3 / WE-2R"])
    for y, f in extra["sl_fte"].items():
        rows.append([f"sl_fte_per_resident_{y}", f / extra["population"][y], "BEA T6.5D lines 93+96 over T7.1 population"])
    acs2010 = [float(r["acs_mexican_origin"]) for r in csv.DictReader(open(FISCAL / "historical_backcast_2026_09_20/inputs/acs_mexican_origin.csv"))
               if r["year"] == "2010"][0]
    rows.append(["acs_census_2010_splice_ratio", acs2010 / (census[2010] * 1e6), "ACS B03001_004E 2010 (back-cast input) / census 2010"])
    rows.append(["aspep_sl_march_payroll_2024", aspep["Total - All Government Employment Functions"], "Census ASPEP 2024, state and local, total payroll ($bn, March)"])
    write(out / "measured_2024.csv", ["item", "value_bn", "source"], rows)

    # --- domestic_interest.csv
    di = v["domestic_interest_total"]
    parts = [["state_local", "imputed interest on plans' claims on employers", v["sl_imputed_interest"], "T7.24 line 13"],
             ["federal", "imputed interest on plans' claims on employers", v["fed_imputed_interest"], "T7.23 line 19"],
             ["federal", "interest on Treasury securities held by the employee plans", v["fed_plan_monetary_interest"], "T7.23 line 18; in T3.2 line 34 per its footnote 4 (interest accrued = T7.23 line 41)"],
             ["state_local", "other interest to persons and business", v["domestic_interest_state_local"] - v["sl_imputed_interest"], "T3.3 line 29 less the above"],
             ["federal", "other interest to persons and business", v["domestic_interest_federal"] - v["fed_imputed_interest"] - v["fed_plan_monetary_interest"], "T3.2 line 34 less the above"]]
    if abs(sum(p[2] for p in parts) - di) > 1e-9:
        raise ValueError("[BLOCKED] domestic interest decomposition does not add")
    write(out / "domestic_interest.csv", ["government", "component", "bn_2024", "source"],
          [p[:2] + [p[2], p[3]] for p in parts] + [["all", "total (T3.1 line 28)", di, "T3.1 line 28"]])

    # --- function_mix.csv / group_shares.csv
    write(out / "function_mix.csv", ["plan", "line", "weight"], [[p, l, w] for p, d in mix.items() for l, w in d.items()])
    write(out / "group_shares.csv", ["group", "end", "line", "share", "response"],
          [[g, e, l, share[g][e][l], response[e][l]] for g in GROUPS for e in ("low", "high") for l in attribution_lines])

    # --- headcount_path.csv
    years = range(1946, 2025)
    write(out / "headcount_path.csv", ["year", "mexican_origin_population_share", "relative_to_2024", "bea_population_thousands"]
          + [f"w_{p}_{a}" for p in mix for a in ("vintage_central", "increments")],
          [[y, abs_share[y], rel[y], extra["population"][y]]
           + [weights[(p, a)].get(y, 0.0) / sum(weights[(p, a)].values()) for p in mix for a in ("vintage_central", "increments")]
           for y in years])

    # --- attribution.csv
    resp_arms = {"main_case": lambda e, l: response[e][l], "average_cost": lambda e, l: 1.0}
    rows, res = [], {}
    for g in GROUPS:
        for plan, bases in plans.items():
            for basis, (stock, interest) in bases.items():
                for tarm, F in factors[plan].items():
                    for rarm, rf in resp_arms.items():
                        for e in ("low", "high"):
                            phi = F * sum(mix[plan][l] * share[g][e][l] * rf(e, l) for l in attribution_lines)
                            key = (g, plan, basis, tarm, rarm, e)
                            res[key] = (phi, phi * stock, phi * interest)
                            rows.append([g, plan, basis, tarm, rarm, e, phi, phi * stock, phi * interest, phi * interest * 1e9 / POPULATION])
    write(out / "attribution.csv", ["group", "plan", "basis", "time_weighting", "responses", "end", "fraction",
                                    "stock_end2023_bn", "interest_2024_bn", "interest_per_member"], rows)

    # --- by_line.csv: the adopted arm's 2024 interest by plan and spending line (the parts of each group's figure)
    brows = []
    for g in GROUPS:
        for e in ("low", "high"):
            for plan, bases in plans.items():
                interest = bases["consolidated"][1]
                for l in attribution_lines:
                    if mix[plan][l]:
                        x = factors[plan]["vintage_central"] * mix[plan][l] * share[g][e][l] * response[e][l] * interest
                        brows.append([g, e, plan, l, x])
    write(out / "by_line.csv", ["group", "end", "plan", "line", "interest_2024_bn"], brows)
    if CASE in ("oct15", "oct16"):
        paired = {(r[0], r[1], r[2], r[3]): r[4] for r in brows}
        bridge_rows = [["case_minus_rough_bridge", e, p, l,
            paired[("mexican_origin", e, p, l)] - paired[("mexican_origin_rough", e, p, l)]]
            for g, e, p, l, _ in brows if g == "mexican_origin"]
        for e in ("low", "high"):
            target = sum(res[("mexican_origin", p, "consolidated", "vintage_central", "main_case", e)][2]
                         - res[("mexican_origin_rough", p, "consolidated", "vintage_central", "main_case", e)][2] for p in plans)
            if abs(sum(r[-1] for r in bridge_rows if r[1] == e) - target) > 1e-9:
                raise ValueError("[BLOCKED] itemized pension case/rough bridge does not close (1e-9bn)")
        write(out / "case_rough_bridge_by_line.csv", ["group", "end", "plan", "line", "interest_2024_bn"], bridge_rows)

    # --- summary.csv: totals over plans for named arms
    arms = [("adopted", "consolidated", "vintage_central", "main_case"),
            ("simple_2024_shares", "consolidated", "simple", "main_case"),
            ("increments", "consolidated", "vintage_central", "main_case"),
            ("vintage_short", "consolidated", "vintage_short", "main_case"),
            ("vintage_long", "consolidated", "vintage_long", "main_case"),
            ("average_cost", "consolidated", "vintage_central", "average_cost"),
            ("average_cost_simple", "consolidated", "simple", "average_cost"),
            ("narrow_federal", "narrow", "vintage_central", "main_case")]
    arms[2] = ("increments", "consolidated", "increments", "main_case")
    srows = []
    for name, basis, tarm, rarm in arms:
        for e in ("low", "high"):
            tot = {}
            for g in GROUPS:
                parts = {p: res[(g, p, basis, tarm, rarm, e)] for p in plans}
                stock = sum(x[1] for x in parts.values())
                inter = sum(x[2] for x in parts.values())
                tot[g] = (stock, inter)
                srows.append([name, basis, tarm, rarm, e, g, parts["state_local"][1], parts["federal_civilian"][1] + parts["federal_military"][1],
                              stock, parts["state_local"][2], parts["federal_civilian"][2] + parts["federal_military"][2], inter,
                              inter * 1e9 / POPULATION])
            for g in (("mexican_origin_rough",) if CASE in ("oct15", "oct16") else ("mexican_origin", "mexican_origin_rough")):
                d = [tot[g][i] - tot["A1_third_plus_nh_white"][i] for i in (0, 1)]
                srows.append([name, basis, tarm, rarm, e, f"{g}_minus_A1_white", "", "", d[0], "", "", d[1], d[1] * 1e9 / POPULATION])
            if CASE in ("oct15", "oct16"):
                d = [tot["mexican_origin"][i] - tot["mexican_origin_rough"][i] for i in (0, 1)]
                srows.append([name, basis, tarm, rarm, e, "case_minus_rough_bridge", "", "", d[0], "", "", d[1], d[1] * 1e9 / POPULATION])
    write(out / "summary.csv", ["arm", "federal_basis", "time_weighting", "responses", "end", "group", "stock_sl_bn", "stock_federal_bn",
                                "stock_total_bn", "interest_sl_bn", "interest_federal_bn", "interest_total_bn", "interest_per_member"], srows)

    # --- opeb_federal.csv (beside: NIPA records retiree health outside the interest row)
    orows = []
    for g in GROUPS:
        for e in ("low", "high"):
            for rarm, rf in resp_arms.items():
                for plan, part in (("federal_civilian", "civilian"), ("federal_military", "military")):
                    phi = factors[plan]["vintage_central"] * sum(mix[plan][l] * share[g][e][l] * rf(e, l) for l in attribution_lines)
                    orows.append([g, e, rarm, plan, fr["opeb_begin_fy24"][part], fr["opeb_interest_fy24"][part], phi,
                                  phi * fr["opeb_begin_fy24"][part], phi * fr["opeb_interest_fy24"][part]])
    write(out / "opeb_federal.csv", ["group", "end", "responses", "plan", "liability_begin_fy24_bn", "interest_fy24_bn", "fraction",
                                     "stock_bn", "interest_bn"], orows)
    if CASE == "oct16":
        bridge, fraction = v9_military_fraction()
        va = []
        for group in GROUPS:
            for end in ("low", "high"):
                phi = factors["federal_civilian"]["vintage_central"] * mix["federal_civilian"]["health_services_va"] * share[group][end]["health_services_va"]
                stock, interest = plans["federal_civilian"]["consolidated"]
                va.append(dict(group=group, end=end, service_share=share[group][end]["health_services_va"],
                    annual_va_service_use_2024_bn=share[group][end]["health_services_va"]*bridge["central_national_bn"],
                    fraction_of_federal_health=fraction, stock_end2023_removed_bn=phi*stock,
                    nominal_interest_2024_removed_bn=phi*interest, response=0))
        source_paths = [FISCAL / CASES[CASE]["lineage"], GROUP_LINES, GROUP_LINES.parent / "service_parts_oct16.json",
                        FISCAL / COVERAGE["file"], FISCAL / V9_META["native_minor_calibration"]["parts_file"],
                        FISCAL / "spending_keys_r5_2026_10_15/derived/item13.json"]
        (out / "service_rules.json").write_text(json.dumps(dict(case=CASE, population=POPULATION,
            va_medical=va, va_national_bridge=bridge, operating_adjustments=V9_SERVICE_PARTS,
            excluded_capital_rule="Item15 college_capital, item20 real capital return and item21 health capital never enter employee-service stock attribution.",
            health_key_rule="VA keeps the source13 old key; source21 residual consumption uses its residual national mass for civilian/state service, with state-price changes confined to state/local service. Own comparators already use OFD+STL without pooled ethnic translation.",
            plan_valuations="Pinned BEA/Z.1 booked stocks and interest unchanged; VA payroll-function split is the declared national medical proxy, not measured historical VA payroll.",
            alternative_scope="Average functional-response arms retain the oct16 operating key changes; average_cost restores the separate VA block's response to one.",
            source_sha256={str(p.relative_to(FISCAL)): sha(p) for p in source_paths}), indent=2) + "\n")

    for (plan, arm), w in sorted(weights.items()):
        print(f"{plan:18s} {arm:16s} F={factors[plan][arm]:.4f}")
    for r in srows:
        if r[0] in ("adopted", "simple_2024_shares", "average_cost"):
            print(r[0], r[4], r[5], f"stock {r[8]:.1f}bn", f"interest {r[11]:.2f}bn", f"{r[12]:.0f}/member")


if __name__ == "__main__":
    main()
    import sys
    sys.path.insert(0, str(FISCAL / "scripts"))
    from memory_budget import require_within_budget
    print(f"[memory] peak physical footprint {require_within_budget('Pension legacy') / 1e9:.3f} GB")
